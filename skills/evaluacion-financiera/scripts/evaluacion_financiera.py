#!/usr/bin/env python3
"""evaluacion-financiera · cómo financiar un activo y cuánto cambia para el dueño.

Compara alternativas de financiamiento de un mismo activo (contado, crédito con cuota francesa o
alemana, leasing con opción de compra, arriendo) después de impuesto y en valor presente, y arma el
flujo del inversionista sumando al flujo del proyecto puro lo que cambia cada alternativa.

Uso:
    python evaluacion_financiera.py --caso caso.json [--proyecto flujo.csv] [--perfil empresa]
                                    [--tasa 0.12] [--impuesto 0.27] [--salida carpeta]
    python evaluacion_financiera.py --probar

Solo biblioteca estándar + openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

RAIZ_SKILL = Path(__file__).resolve().parents[1]
RAIZ_REPO = Path(__file__).resolve().parents[3]
PLANTILLA_PERFIL = RAIZ_REPO / "perfiles" / "_plantilla.yaml"
TIPOS = ("contado", "credito", "leasing", "arriendo")


class ErrorEntrada(Exception):
    """Falta un dato o viene mal: el script no inventa, se detiene y dice cuál."""


# ----------------------------------------------------------------------------------------------
# Lectura: perfil YAML (subconjunto simple), caso JSON, flujo del proyecto
# ----------------------------------------------------------------------------------------------

def _valor_yaml(txt: str):
    txt = txt.strip()
    if txt in ("", "null", "~", "Null", "NULL"):
        return None
    if txt in ("[]", "{}"):
        return [] if txt == "[]" else {}
    if txt in ("true", "True"):
        return True
    if txt in ("false", "False"):
        return False
    if (txt[0] == txt[-1]) and txt[0] in "\"'":
        return txt[1:-1]
    try:
        return int(txt)
    except ValueError:
        pass
    try:
        return float(txt)
    except ValueError:
        return txt


def leer_yaml_simple(ruta: Path) -> dict:
    """Lee el perfil: mapas anidados por indentación, `clave: valor  # comentario`.
    Las listas se ignoran (este método no las usa)."""
    raiz: dict = {}
    pila = [(-1, raiz)]
    for linea in ruta.read_text(encoding="utf-8-sig").splitlines():
        sin_com = linea.split(" #", 1)[0].rstrip() if not linea.lstrip().startswith("#") else ""
        if not sin_com.strip() or sin_com.lstrip().startswith("-"):
            continue
        sangria = len(sin_com) - len(sin_com.lstrip())
        if ":" not in sin_com:
            continue
        clave, _, resto = sin_com.strip().partition(":")
        while pila and pila[-1][0] >= sangria:
            pila.pop()
        padre = pila[-1][1]
        if resto.strip() == "":
            nuevo: dict = {}
            padre[clave.strip()] = nuevo
            pila.append((sangria, nuevo))
        else:
            padre[clave.strip()] = _valor_yaml(resto)
    return raiz


def buscar_perfil(nombre: str | None) -> tuple[Path, str]:
    """Orden: ruta directa · $INGENIERO_DATOS/perfiles · ../ingeniero-agentico-datos/perfiles · plantilla."""
    if nombre:
        directo = Path(nombre)
        if directo.is_file():
            return directo, "ruta indicada en --perfil"
        archivo = nombre if nombre.endswith((".yaml", ".yml")) else nombre + ".yaml"
        candidatos = []
        if os.environ.get("INGENIERO_DATOS"):
            candidatos.append((Path(os.environ["INGENIERO_DATOS"]) / "perfiles" / archivo,
                               "variable INGENIERO_DATOS"))
        candidatos.append((RAIZ_REPO.parent / "ingeniero-agentico-datos" / "perfiles" / archivo,
                           "carpeta hermana ingeniero-agentico-datos"))
        for ruta, origen in candidatos:
            if ruta.is_file():
                return ruta, origen
        raise ErrorEntrada(f"No encontré el perfil «{nombre}» (busqué en: "
                           + "; ".join(str(c[0]) for c in candidatos) + ").")
    return PLANTILLA_PERFIL, "plantilla (no hay perfil de empresa)"


def _num(txt: str) -> float:
    t = str(txt).strip().replace("$", "").replace(" ", "").replace("−", "-")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    return float(t)


def leer_flujo_proyecto(ruta: Path) -> dict[int, float]:
    """Flujo del proyecto puro, anual. Acepta:
    - CSV (separador , ; o tabulador) con columnas año/anio/periodo y flujo/flujo_neto/flujo_proyecto;
    - JSON: lista de {"anio": n, "flujo": x}, o {"anios": [...], "flujos": [...]}, o {"flujo": <lista>}.
    Los números van sin separador de miles."""
    if not ruta.is_file():
        raise ErrorEntrada(f"No existe el archivo del flujo del proyecto: {ruta}")
    nombres_anio = ("anio", "año", "ano", "periodo", "year", "t")
    nombres_flujo = ("flujo", "flujo_neto", "flujo_proyecto", "flujo_caja", "fcn", "flujo_neto_caja")
    if ruta.suffix.lower() == ".json":
        datos = json.loads(ruta.read_text(encoding="utf-8-sig"))
        # La salida de evaluacion-economica trae {"flujo": {"anios": [...], "flujo_de_caja": [...]}}
        if isinstance(datos, dict) and isinstance(datos.get("flujo"), dict):
            datos = datos["flujo"]
        if isinstance(datos, dict) and "anios" in datos:
            serie = next((datos[k] for k in ("flujos", "flujo_de_caja", "flujo", "flujo_proyecto")
                          if isinstance(datos.get(k), list)), None)
            if serie is None:
                raise ErrorEntrada(f"{ruta.name}: trae «anios» pero no la serie del flujo "
                                   "(flujos · flujo_de_caja · flujo).")
            return {int(a): float(f) for a, f in zip(datos["anios"], serie)}
        if isinstance(datos, dict):
            for k in ("flujo", "flujo_proyecto", "filas"):
                if k in datos:
                    datos = datos[k]
                    break
        res = {}
        for fila in datos:
            fa = {k.lower(): v for k, v in fila.items()}
            a = next(fa[k] for k in nombres_anio if k in fa)
            f = next(fa[k] for k in nombres_flujo if k in fa)
            res[int(a)] = float(f)
        return res
    texto = ruta.read_text(encoding="utf-8-sig")
    sep = max((";", ",", "\t"), key=lambda s: texto.splitlines()[0].count(s))
    lector = csv.DictReader(texto.splitlines(), delimiter=sep)
    cols = {c.strip().lower(): c for c in (lector.fieldnames or [])}
    ca = next((cols[k] for k in nombres_anio if k in cols), None)
    cf = next((cols[k] for k in nombres_flujo if k in cols), None)
    if not ca or not cf:
        raise ErrorEntrada(f"El flujo {ruta.name} necesita columnas año y flujo; trae {list(cols)}")
    res = {}
    for fila in lector:
        if not (fila.get(ca) or "").strip():
            continue
        res[int(_num(fila[ca]))] = _num(fila[cf])
    return res


# ----------------------------------------------------------------------------------------------
# Registro de supuestos (trazabilidad: de dónde vino cada número)
# ----------------------------------------------------------------------------------------------

_FALTA = object()


class Registro:
    def __init__(self) -> None:
        self.filas: list[dict] = []
        self.avisos: list[str] = []

    def add(self, parametro, valor, unidad, origen, supuesto, nota=""):
        self.filas.append(dict(parametro=parametro, valor=valor, unidad=unidad, origen=origen,
                               marca="SUPUESTO" if supuesto else "dato", nota=nota or ""))

    def leer(self, d: dict, clave: str, ruta: str, unidad: str = "", defecto=_FALTA, nota_def=""):
        if clave in d and d[clave] is not None:
            v = d[clave]
            if isinstance(v, dict) and "valor" in v:
                self.add(f"{ruta}.{clave}", v["valor"], unidad, f"caso: {ruta}.{clave}",
                         bool(v.get("supuesto")), v.get("fuente", ""))
                return v["valor"]
            self.add(f"{ruta}.{clave}", v, unidad, f"caso: {ruta}.{clave}", False)
            return v
        if defecto is _FALTA:
            raise ErrorEntrada(f"Falta «{ruta}.{clave}» en el caso.")
        self.add(f"{ruta}.{clave}", defecto, unidad, "valor por defecto del script", True, nota_def)
        return defecto


# ----------------------------------------------------------------------------------------------
# Aritmética financiera
# ----------------------------------------------------------------------------------------------

def red(x: float, dec: int) -> float:
    q = Decimal(1).scaleb(-dec)
    return float(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))


def factor_anualidad(i: float, n: int) -> float:
    """a(n, i) = (1 - (1+i)^-n) / i : valor presente de n pagos unitarios vencidos."""
    return n if i == 0 else (1 - (1 + i) ** -n) / i


def cuota_francesa(monto: float, i: float, n: int) -> float:
    return monto / factor_anualidad(i, n)


def tabla_francesa(monto, i, n, dec):
    c = red(cuota_francesa(monto, i, n), dec)
    saldo, filas = monto, []
    for m in range(1, n + 1):
        interes = red(saldo * i, dec)
        amort = saldo if m == n else red(c - interes, dec)
        filas.append(dict(mes=m, saldo_inicial=saldo, cuota=red(amort + interes, dec), interes=interes,
                          amortizacion=amort, saldo_final=red(saldo - amort, dec)))
        saldo = filas[-1]["saldo_final"]
    return filas, c


def tabla_alemana(monto, i, n, dec):
    a = red(monto / n, dec)
    saldo, filas = monto, []
    for m in range(1, n + 1):
        interes = red(saldo * i, dec)
        amort = saldo if m == n else a
        filas.append(dict(mes=m, saldo_inicial=saldo, cuota=red(amort + interes, dec), interes=interes,
                          amortizacion=amort, saldo_final=red(saldo - amort, dec)))
        saldo = filas[-1]["saldo_final"]
    return filas, filas[0]["cuota"]


def mensual_desde_anual(r: float) -> float:
    return (1 + r) ** (1 / 12) - 1


def vp(flujos: list[float], r_anual: float) -> float:
    """Valor presente de un flujo mensual (índice = mes) a una tasa anual efectiva."""
    return sum(f / (1 + r_anual) ** (m / 12) for m, f in enumerate(flujos) if f)


def tir_raices(flujos: list[float], lo=-0.03, hi=0.15, paso=0.0005) -> list[float]:
    """Todas las TIR mensuales entre lo y hi (−31 % a +435 % anual), por barrido + bisección.
    Más de una raíz = el flujo cambia de signo varias veces y la TIR no es única."""
    def van(r):
        return sum(f / (1 + r) ** m for m, f in enumerate(flujos) if f)
    raices, previo = [], None
    for k in range(int(round((hi - lo) / paso)) + 1):
        r = lo + k * paso
        v = van(r)
        if previo is not None and (previo[1] < 0) != (v < 0):
            a, b, fa = previo[0], r, previo[1]
            for _ in range(100):
                mid = (a + b) / 2
                fm = van(mid)
                if (fa < 0) == (fm < 0):
                    a, fa = mid, fm
                else:
                    b = mid
            raices.append((a + b) / 2)
        previo = (r, v)
    return raices


def tir_mensual(flujos: list[float]) -> float | None:
    """TIR mensual de un flujo con un solo cambio de signo (CAE, tasa implícita). None si no hay."""
    raices = tir_raices(flujos)
    return raices[0] if raices else None


# ----------------------------------------------------------------------------------------------
# Modelo: flujos del activo por alternativa
# ----------------------------------------------------------------------------------------------

def anio_de_mes(m: int) -> int:
    return max(1, math.ceil(m / 12))


def depreciacion(base, metodo, vida, vida_acel, anio_ini):
    if base <= 0:
        return {}
    if metodo == "instantanea":
        return {anio_ini: base}
    v = vida_acel if metodo == "acelerada" else vida
    if not v or v <= 0:
        raise ErrorEntrada(f"Depreciación {metodo} sin vida útil válida.")
    return {anio_ini + k: base / v for k in range(int(v))}


def construir_alternativa(alt: dict, P: dict, t: float) -> dict:
    """Flujo mensual del activo para una alternativa (dueño, después de impuesto).
    Devuelve los componentes por mes y la tabla de amortización si corresponde."""
    H = P["meses"]
    dec = P["dec"]
    fiva = P["factor_iva"]
    pagos = [0.0] * (H + 1)
    residual = [0.0] * (H + 1)
    ded_fin: dict[int, float] = {}     # intereses, gastos financieros, seguros del crédito
    ded_activo: dict[int, float] = {}  # depreciación, cuotas de leasing, renta de arriendo
    ingresos = [0.0] * (H + 1)         # garantía devuelta, etc.

    def suma(d, anio, x):
        d[anio] = d.get(anio, 0.0) + x

    tipo = alt["tipo"]
    precio = alt["precio"] * fiva
    tabla, cuota, extra = None, None, {}
    base_dep, anio_compra, mes_compra = 0.0, 1, 0

    if tipo == "contado":
        pagos[0] -= precio
        base_dep, anio_compra = precio, 1
    elif tipo == "credito":
        pie = alt["pie"]
        monto = precio - pie
        n, i = alt["plazo_meses"], alt["tasa_mensual"]
        imp_cred = red(monto * min(alt["imp_credito_tasa"] * n, alt["imp_credito_tope"]), dec) \
            if alt["imp_credito_tasa"] else 0.0
        gastos = alt["gastos_iniciales"] + imp_cred
        if alt["gastos_financiados"]:
            monto += gastos
            pagos[0] -= pie
        else:
            pagos[0] -= pie + gastos
        suma(ded_fin, 1, gastos)
        tabla, cuota = (tabla_francesa if alt["sistema"] == "francesa" else tabla_alemana)(monto, i, n, dec)
        for f in tabla:
            m = f["mes"]
            if m > H:
                break
            pagos[m] -= f["cuota"] + alt["seguro_mensual"]
            suma(ded_fin, anio_de_mes(m), f["interes"] + alt["seguro_mensual"])
        if n > H:
            saldo = tabla[H - 1]["saldo_final"]
            pagos[H] -= saldo
            extra["saldo_pagado_al_horizonte"] = saldo
        base_dep, anio_compra = precio, 1
        # CAE: lo que de verdad llega al deudor (monto − gastos) contra todo lo que paga
        flujo_cae = [monto - gastos] + [-(f["cuota"] + alt["seguro_mensual"]) for f in tabla]
        tm = tir_mensual(flujo_cae)
        extra.update(monto_financiado=monto, gastos_iniciales=gastos, impuesto_credito=imp_cred,
                     cae=None if tm is None else tm * 12,
                     tasa_efectiva_anual_credito=None if tm is None else (1 + tm) ** 12 - 1)
    elif tipo == "leasing":
        n = alt["plazo_meses"]
        mes_op = alt["mes_opcion"] or n + 1
        if mes_op > H:
            raise ErrorEntrada(f"{alt['nombre']}: la opción de compra (mes {mes_op}) cae después del "
                               f"horizonte ({H} meses). Alarga el horizonte o acorta el leasing.")
        i = alt["tasa_mensual"]
        if alt["cuota_dada"] is not None:
            c = alt["cuota_dada"] * fiva
        else:
            # P = pie + C·a(n,i) + opción·v^mes_op ; pie y opción pueden ser «una_cuota» (k·C)
            k_pie = 1.0 if alt["cuota_inicial"] == "una_cuota" else 0.0
            k_op = 1.0 if alt["opcion_compra"] == "una_cuota" else 0.0
            fijo_pie = 0.0 if k_pie else float(alt["cuota_inicial"] or 0) * fiva
            fijo_op = 0.0 if k_op else float(alt["opcion_compra"] or 0) * fiva
            v_op = (1 + i) ** -mes_op
            c = (precio - fijo_pie - fijo_op * v_op) / (k_pie + factor_anualidad(i, n) + k_op * v_op)
            c = red(c, dec)
        pie = c if alt["cuota_inicial"] == "una_cuota" else float(alt["cuota_inicial"] or 0) * fiva
        opcion = c if alt["opcion_compra"] == "una_cuota" else float(alt["opcion_compra"] or 0) * fiva
        if alt["cuota_dada"] is not None:
            flujo_impl = [precio - pie] + [-c] * n + [0.0] * (mes_op - n)
            flujo_impl[mes_op] -= opcion
            tm = tir_mensual(flujo_impl)
            if tm is None:
                raise ErrorEntrada(f"{alt['nombre']}: no pude despejar la tasa implícita del leasing.")
            i = tm
        pagos[0] -= pie
        suma(ded_activo, 1, pie)
        saldo, tabla = precio - pie, []
        for m in range(1, n + 1):
            interes = red(saldo * i, dec)
            amort = red(c - interes, dec)
            tabla.append(dict(mes=m, saldo_inicial=saldo, cuota=c, interes=interes, amortizacion=amort,
                              saldo_final=red(saldo - amort, dec)))
            saldo = tabla[-1]["saldo_final"]
            pagos[m] -= c
            suma(ded_activo, anio_de_mes(m), c)
        meses_espera = mes_op - n
        interes_op = red(saldo * ((1 + i) ** meses_espera - 1), dec)
        tabla.append(dict(mes=mes_op, saldo_inicial=saldo, cuota=opcion, interes=opcion - saldo,
                          amortizacion=saldo, saldo_final=0.0, es_opcion=True,
                          descuadre_redondeo=red(opcion - saldo - interes_op, dec)))
        pagos[mes_op] -= opcion
        cuota = c
        base_dep, anio_compra, mes_compra = opcion, anio_de_mes(mes_op), mes_op
        flujo_cae = [precio - pie] + [-c] * n + [0.0] * meses_espera
        flujo_cae[mes_op] -= opcion
        tm = tir_mensual(flujo_cae)
        extra.update(monto_financiado=precio - pie, cuota_inicial=pie, opcion_compra=opcion,
                     mes_opcion=mes_op, tasa_implicita_mensual=i,
                     cae=None if tm is None else tm * 12,
                     tasa_efectiva_anual_credito=None if tm is None else (1 + tm) ** 12 - 1)
    elif tipo == "arriendo":
        n = alt["plazo_meses"]
        renta = alt["renta_mensual"] * fiva
        garantia = alt["garantia"]
        pagos[0] -= garantia
        for m in range(1, min(n, H) + 1):
            pagos[m] -= renta
            suma(ded_activo, anio_de_mes(m), renta)
        ingresos[min(n, H)] += garantia
        cuota = renta
        if n != H:
            extra["aviso"] = f"el arriendo dura {n} meses y el horizonte {H}: no cubre el mismo periodo"
    else:
        raise ErrorEntrada(f"Tipo de alternativa desconocido: {tipo} (usa {', '.join(TIPOS)}).")

    # Depreciación y valor residual del dueño (contado, crédito y leasing con opción ejercida)
    dep = depreciacion(base_dep, P["metodo_dep"], P["vida"], P["vida_acel"], anio_compra) if base_dep else {}
    anios_h = P["horizonte"]
    for a, d in dep.items():
        if a <= anios_h:
            suma(ded_activo, a, d)
    if tipo != "arriendo":
        dep_acum = sum(d for a, d in dep.items() if a <= anios_h)
        libro = base_dep - dep_acum
        vm = P["valor_mercado"]
        impuesto_venta = t * (vm - libro)
        residual[H] += vm - impuesto_venta
        extra.update(valor_libro_horizonte=libro, impuesto_venta=impuesto_venta,
                     base_depreciacion=base_dep, anio_inicio_depreciacion=anio_compra)

    ahorro_fin = [0.0] * (H + 1)
    ahorro_act = [0.0] * (H + 1)
    for a in range(1, anios_h + 1):
        ahorro_fin[12 * a] += t * ded_fin.get(a, 0.0)
        ahorro_act[12 * a] += t * ded_activo.get(a, 0.0)
    F = [pagos[m] + ingresos[m] + ahorro_fin[m] + ahorro_act[m] + residual[m] for m in range(H + 1)]
    return dict(nombre=alt["nombre"], tipo=tipo, F=F, pagos=pagos, ingresos=ingresos, ahorro_fin=ahorro_fin,
                ahorro_act=ahorro_act, residual=residual, ded_fin=ded_fin, ded_activo=ded_activo,
                tabla=tabla, cuota=cuota, extra=extra, precio=precio, notas=alt.get("notas", []),
                sistema=alt.get("sistema"))


# ----------------------------------------------------------------------------------------------
# Preparación de parámetros
# ----------------------------------------------------------------------------------------------

def preparar(caso: dict, ruta_caso: Path, perfil_nombre, tasa_arg, imp_arg, proyecto_arg, reg: Registro):
    ruta_perfil, origen_perfil = buscar_perfil(perfil_nombre or caso.get("perfil"))
    perfil = leer_yaml_simple(ruta_perfil)
    es_plantilla = ruta_perfil.resolve() == PLANTILLA_PERFIL.resolve()
    emp = perfil.get("empresa", {}) or {}
    fin = perfil.get("finanzas", {}) or {}
    reg.add("perfil", ruta_perfil.name, "", origen_perfil, es_plantilla,
            "la plantilla no describe a ninguna empresa: sus valores son supuestos" if es_plantilla else "")

    def del_perfil(clave, seccion, arg, nombre_arg, unidad):
        v = seccion.get(clave)
        if arg is not None:
            reg.add(clave, arg, unidad, f"argumento {nombre_arg}", True,
                    "reemplaza el valor del perfil" if v is not None else "el perfil lo trae vacío")
            return arg
        if v is None:
            raise ErrorEntrada(f"El perfil {ruta_perfil.name} no trae «{clave}». Pásalo con {nombre_arg} "
                               f"(queda marcado SUPUESTO en la salida).")
        reg.add(clave, v, unidad, f"perfil {ruta_perfil.name}", es_plantilla)
        return v

    tasa = float(del_perfil("tasa_descuento", fin, tasa_arg, "--tasa", "anual efectiva"))
    t = float(del_perfil("impuesto_renta", emp, imp_arg, "--impuesto", "fracción"))
    moneda = emp.get("moneda") or "CLP"
    dec = int(emp.get("decimales_moneda") if emp.get("decimales_moneda") is not None else 0)
    iva = float(emp.get("iva") if emp.get("iva") is not None else 0.0)
    reg.add("moneda", moneda, "", f"perfil {ruta_perfil.name}", es_plantilla)
    def del_caso_o_perfil(d, clave, ruta, unidad, v_perfil, defecto):
        if d.get(clave) is not None or v_perfil is None:
            return reg.leer(d, clave, ruta, unidad, defecto)
        reg.add(f"{ruta}.{clave}", v_perfil, unidad, f"perfil {ruta_perfil.name}", es_plantilla)
        return v_perfil

    horizonte = int(del_caso_o_perfil(caso, "horizonte_anios", "caso", "años", fin.get("horizonte_anios"), 5))
    iva_rec = bool(reg.leer(caso, "iva_recuperable", "caso", "sí/no", True,
                            "se supone que la empresa recupera el IVA como crédito fiscal"))
    if not iva_rec:
        reg.add("iva", iva, "fracción", f"perfil {ruta_perfil.name}", es_plantilla,
                "IVA no recuperable: se suma a precios, cuotas de leasing y rentas")
    a = caso.get("activo") or {}
    precio = float(reg.leer(a, "precio_neto", "activo", moneda))
    metodo = del_caso_o_perfil(a, "depreciacion", "activo", "", fin.get("depreciacion"), "lineal")
    if metodo not in ("lineal", "acelerada", "instantanea"):
        raise ErrorEntrada(f"Depreciación «{metodo}» no reconocida (lineal · acelerada · instantanea).")
    vida = reg.leer(a, "vida_util_anios", "activo", "años") if metodo == "lineal" else a.get("vida_util_anios")
    vida_acel = reg.leer(a, "vida_util_acelerada_anios", "activo", "años") if metodo == "acelerada" else None
    vm = float(reg.leer(a, "valor_mercado_al_horizonte", "activo", moneda))
    P = dict(meses=12 * horizonte, horizonte=horizonte, dec=dec, factor_iva=1.0 if iva_rec else 1 + iva,
             metodo_dep=metodo, vida=vida, vida_acel=vida_acel, valor_mercado=vm, precio=precio,
             moneda=moneda, tasa=tasa, t=t, perfil=ruta_perfil.name, descripcion=a.get("descripcion", "activo"))

    alts = []
    nombres = set()
    for k, x in enumerate(caso.get("alternativas") or []):
        nombre = x.get("nombre") or f"alternativa {k + 1}"
        if nombre in nombres:
            raise ErrorEntrada(f"Hay dos alternativas llamadas «{nombre}».")
        nombres.add(nombre)
        ruta = f"alternativas[{nombre}]"
        tipo = x.get("tipo")
        if tipo not in TIPOS:
            raise ErrorEntrada(f"{ruta}: tipo «{tipo}» no reconocido (usa {', '.join(TIPOS)}).")
        alt = dict(nombre=nombre, tipo=tipo, notas=list(x.get("notas_no_comparables") or []))
        # Precio propio solo si la alternativa lo cambia (p. ej., bono por financiar con la marca)
        alt["precio"] = float(reg.leer(x, "precio_neto", ruta, moneda)) if "precio_neto" in x else precio
        if tipo == "credito":
            alt["sistema"] = reg.leer(x, "sistema", ruta, "", "francesa", "cuota fija, la más común")
            if alt["sistema"] not in ("francesa", "alemana"):
                raise ErrorEntrada(f"{ruta}: sistema «{alt['sistema']}» (usa francesa o alemana).")
            if "pie_fraccion" in x:
                alt["pie"] = red(alt["precio"] * P["factor_iva"] *
                                 float(reg.leer(x, "pie_fraccion", ruta, "fracción del precio")), dec)
            else:
                alt["pie"] = float(reg.leer(x, "pie", ruta, moneda, 0.0, "sin pie"))
            alt["plazo_meses"] = int(reg.leer(x, "plazo_meses", ruta, "meses"))
            if "tasa_mensual" in x:
                alt["tasa_mensual"] = float(reg.leer(x, "tasa_mensual", ruta, "mensual"))
            else:
                alt["tasa_mensual"] = mensual_desde_anual(float(reg.leer(x, "tasa_anual", ruta, "anual efectiva")))
            alt["gastos_iniciales"] = float(reg.leer(x, "gastos_iniciales", ruta, moneda, 0.0, "sin gastos"))
            alt["gastos_financiados"] = bool(reg.leer(x, "gastos_financiados", ruta, "sí/no", False,
                                                      "los gastos se pagan al contado"))
            alt["seguro_mensual"] = float(reg.leer(x, "seguro_mensual", ruta, moneda, 0.0, "sin seguro"))
            ic = x.get("impuesto_credito") or {}
            alt["imp_credito_tasa"] = float(reg.leer(ic, "tasa_mensual", ruta + ".impuesto_credito",
                                                     "por mes de plazo", 0.0, "sin impuesto al crédito")) if ic else 0.0
            alt["imp_credito_tope"] = float(reg.leer(ic, "tope", ruta + ".impuesto_credito",
                                                     "fracción del monto", 1.0, "sin tope")) if ic else 1.0
        elif tipo == "leasing":
            alt["plazo_meses"] = int(reg.leer(x, "plazo_meses", ruta, "meses"))
            alt["cuota_dada"] = None
            if x.get("cuota_mensual") is not None:
                alt["cuota_dada"] = float(reg.leer(x, "cuota_mensual", ruta, moneda))
                alt["tasa_mensual"] = None
            elif "tasa_mensual" in x:
                alt["tasa_mensual"] = float(reg.leer(x, "tasa_mensual", ruta, "mensual"))
            else:
                alt["tasa_mensual"] = mensual_desde_anual(float(reg.leer(x, "tasa_anual", ruta, "anual efectiva")))
            alt["cuota_inicial"] = reg.leer(x, "cuota_inicial", ruta, f"{moneda} o una_cuota", "una_cuota",
                                            "lo usual en leasing: una cuota al firmar")
            alt["opcion_compra"] = reg.leer(x, "opcion_compra", ruta, f"{moneda} o una_cuota", "una_cuota",
                                            "lo usual en leasing: opción igual a una cuota")
            alt["mes_opcion"] = int(reg.leer(x, "mes_opcion", ruta, "mes", alt["plazo_meses"] + 1,
                                             "un mes después de la última cuota"))
        elif tipo == "arriendo":
            alt["renta_mensual"] = float(reg.leer(x, "renta_mensual", ruta, moneda))
            alt["plazo_meses"] = int(reg.leer(x, "plazo_meses", ruta, "meses", 12 * horizonte,
                                              "todo el horizonte"))
            alt["garantia"] = float(reg.leer(x, "garantia", ruta, moneda, 0.0, "sin garantía"))
        alts.append(alt)
    if len(alts) < 2:
        raise ErrorEntrada("El caso necesita al menos dos alternativas para comparar.")

    flujo_proy = None
    origen_proy = proyecto_arg or caso.get("flujo_proyecto")
    if origen_proy:
        rp = Path(origen_proy)
        if not rp.is_absolute():
            rp = (Path.cwd() / rp) if proyecto_arg else (ruta_caso.parent / rp)
        anual = leer_flujo_proyecto(rp)
        reg.add("flujo_proyecto", rp.name, "anual", "argumento --proyecto" if proyecto_arg else
                "caso: flujo_proyecto", False, f"{len(anual)} periodos (años {min(anual)} a {max(anual)})")
        if max(anual) > horizonte:
            reg.avisos.append(f"El flujo del proyecto llega al año {max(anual)} y el horizonte es "
                              f"{horizonte}: los años siguientes no se consideran.")
        if max(anual) < horizonte:
            reg.avisos.append(f"El flujo del proyecto termina en el año {max(anual)} y el horizonte es "
                              f"{horizonte}: revisa que sean el mismo periodo.")
        f0 = anual.get(0, 0.0)
        if abs(f0 + precio * P["factor_iva"]) > 0.25 * precio:
            reg.avisos.append(f"El año 0 del proyecto ({f0:,.0f}) no se parece a −precio "
                              f"({-precio:,.0f}): revisa que el proyecto compre este mismo activo.")
        flujo_proy = [0.0] * (P["meses"] + 1)
        for an, f in anual.items():
            if 0 <= an <= horizonte:
                flujo_proy[12 * an] += f
    else:
        reg.avisos.append("Sin flujo del proyecto: se compara el costo de financiar; el VAN del "
                          "inversionista no se calcula.")
    return P, alts, flujo_proy


# ----------------------------------------------------------------------------------------------
# Cálculo completo
# ----------------------------------------------------------------------------------------------

def evaluar(P, alts, flujo_proy, t=None, tasa=None, con_tir=True):
    t = P["t"] if t is None else t
    tasa = P["tasa"] if tasa is None else tasa
    ref_alt = dict(nombre="contado de referencia", tipo="contado", precio=P["precio"], notas=[])
    ref = construir_alternativa(ref_alt, P, t)
    vp_ref = vp(ref["F"], tasa)
    van_proy = vp(flujo_proy, tasa) if flujo_proy else None
    res = []
    for a in alts:
        r = construir_alternativa(a, P, t)
        delta = [r["F"][m] - ref["F"][m] for m in range(P["meses"] + 1)]
        r["delta"] = delta
        r["costo_vp"] = -vp(r["F"], tasa)
        r["ventaja_vs_contado"] = vp(r["F"], tasa) - vp_ref
        if flujo_proy:
            r["inversionista"] = [flujo_proy[m] + delta[m] for m in range(P["meses"] + 1)]
            r["van_inversionista"] = vp(r["inversionista"], tasa)
        raices = tir_raices(delta) if con_tir and any(abs(d) > 0.5 for d in delta) else []
        r["tir_multiple"] = len(raices) > 1
        tm = min(raices, key=lambda x: abs(x - mensual_desde_anual(tasa))) if raices else None
        r["tasa_efectiva_despues_impuesto"] = None if tm is None else (1 + tm) ** 12 - 1
        if r["tir_multiple"]:
            r["notas"] = list(r["notas"]) + [
                "el flujo frente al contado cambia de signo varias veces: la tasa efectiva no es única "
                f"({', '.join(f'{(1 + x) ** 12 - 1:.2%}' for x in raices)}); decide por el costo en VP"]
        elif con_tir and tm is None and any(abs(d) > 0.5 for d in delta):
            r["notas"] = list(r["notas"]) + [
                "sin tasa efectiva: el flujo frente al contado no cambia de signo (domina o es dominada "
                "a cualquier tasa); decide por el costo en VP"]
        res.append(r)
    orden = sorted(res, key=lambda r: r["costo_vp"])
    return dict(ref=ref, res=res, orden=orden, van_proyecto=van_proy, tasa=tasa, t=t)


def tasa_indiferencia(P, alts, flujo_proy, a_nombre, b_nombre):
    """Tasa de descuento (anual) a la que dos alternativas cuestan lo mismo en valor presente."""
    def dif(r):
        ev = evaluar(P, alts, flujo_proy, tasa=r, con_tir=False)
        c = {x["nombre"]: x["costo_vp"] for x in ev["res"]}
        return c[a_nombre] - c[b_nombre]
    raices = []
    previo = (0.0, dif(0.0))
    for k in range(1, 121):
        r = k * 0.005
        v = dif(r)
        if (previo[1] < 0) != (v < 0):
            a, b, fa = previo[0], r, previo[1]
            for _ in range(40):
                mid = (a + b) / 2
                fm = dif(mid)
                if (fa < 0) == (fm < 0):
                    a, fa = mid, fm
                else:
                    b = mid
            raices.append((a + b) / 2)
        previo = (r, v)
    return raices


def sensibilidad(P, alts, flujo_proy, impuestos_extra):
    tasas = sorted({round(max(0.0, P["tasa"] + d), 4) for d in (-0.06, -0.03, 0.0, 0.03, 0.06)})
    imps = []
    for x in [P["t"]] + list(impuestos_extra or []):
        if round(float(x), 4) not in imps:
            imps.append(round(float(x), 4))
    filas = []
    for ti in imps:
        for r in tasas:
            ev = evaluar(P, alts, flujo_proy, t=ti, tasa=r, con_tir=False)
            filas.append(dict(impuesto=ti, tasa=r, costos={x["nombre"]: x["costo_vp"] for x in ev["res"]},
                              ganadora=ev["orden"][0]["nombre"]))
    return filas


# ----------------------------------------------------------------------------------------------
# Salidas: resultado.json y Excel
# ----------------------------------------------------------------------------------------------

def resumen_json(P, ev, indif, sens, reg, alts_in):
    g, s = ev["orden"][0], ev["orden"][1]
    out = dict(
        moneda=P["moneda"], tasa_descuento=P["tasa"], impuesto_renta=P["t"], horizonte_anios=P["horizonte"],
        van_proyecto=None if ev["van_proyecto"] is None else round(ev["van_proyecto"], 2),
        ganadora=g["nombre"], segunda=s["nombre"], diferencia_vp=round(s["costo_vp"] - g["costo_vp"], 2),
        tasa_indiferencia_con_segunda=[round(x, 6) for x in indif],
        alternativas={},
        supuestos=[f for f in reg.filas if f["marca"] == "SUPUESTO"],
        avisos=reg.avisos,
        sensibilidad=[dict(impuesto=f["impuesto"], tasa=f["tasa"], ganadora=f["ganadora"]) for f in sens],
    )
    for r in ev["res"]:
        e = r["extra"]
        out["alternativas"][r["nombre"]] = dict(
            tipo=r["tipo"],
            cuota=r["cuota"],
            costo_vp=round(r["costo_vp"], 2),
            ventaja_vs_contado=round(r["ventaja_vs_contado"], 2),
            sobre_la_ganadora=round(r["costo_vp"] - g["costo_vp"], 2),
            van_inversionista=None if "van_inversionista" not in r else round(r["van_inversionista"], 2),
            tasa_efectiva_despues_impuesto=None if r["tasa_efectiva_despues_impuesto"] is None
            else round(r["tasa_efectiva_despues_impuesto"], 6),
            cae=None if e.get("cae") is None else round(e["cae"], 6),
            total_pagado=round(-sum(r["pagos"]), 2),
            total_intereses=None if not r["tabla"] else round(sum(f["interes"] for f in r["tabla"]), 2),
            ahorro_tributario_total=round(sum(r["ahorro_fin"]) + sum(r["ahorro_act"]), 2),
            valor_residual_neto=round(sum(r["residual"]), 2),
        )
    return out


def escribir_excel(ruta: Path, P, ev, indif, sens, reg, titulo):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fmt_m = "#,##0" if P["dec"] == 0 else "#,##0.00"
    fmt_p = "0.00%"
    neg = Font(bold=True)
    enc_fill = PatternFill("solid", fgColor="1F3A5F")
    enc_font = Font(bold=True, color="FFFFFF")
    sup_fill = PatternFill("solid", fgColor="FCE4B6")
    gan_fill = PatternFill("solid", fgColor="D9EAD3")

    wb = Workbook()

    def encabezado(ws, fila, textos):
        for j, tx in enumerate(textos, 1):
            c = ws.cell(row=fila, column=j, value=tx)
            c.font, c.fill = enc_font, enc_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")

    def anchos(ws, anchos_):
        for j, w in enumerate(anchos_, 1):
            ws.column_dimensions[get_column_letter(j)].width = w

    # SUPUESTOS ------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "SUPUESTOS"
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Método de cálculo, no asesoría financiera. Montos en " + P["moneda"] +
                ", netos de IVA salvo que el caso diga lo contrario. Fila naranja = SUPUESTO.")
    encabezado(ws, 4, ["Parámetro", "Valor", "Unidad", "De dónde vino", "Marca", "Fuente o nota"])
    f = 5
    for x in reg.filas:
        vals = [x["parametro"], x["valor"], x["unidad"], x["origen"], x["marca"], x["nota"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v if not isinstance(v, (list, dict)) else json.dumps(v))
            if x["marca"] == "SUPUESTO":
                c.fill = sup_fill
        f += 1
    fijos = [
        ("Tributación", "el ahorro tributario de cada año (intereses, depreciación, cuotas de leasing, renta) "
                        "se recibe al cierre de ese año; se supone que la empresa tiene utilidades para usarlo"),
        ("Depreciación", "año completo desde el año de compra; en leasing parte cuando se ejerce la opción"),
        ("Valor residual", "se vende al valor de mercado al horizonte y se paga impuesto sobre la diferencia con el valor libro"),
        ("Leasing", "las cuotas y la cuota inicial son gasto tributario del arrendatario (regla chilena); "
                    "otro país puede tratarlo como activo"),
        ("Descuento", "flujos mensuales descontados con la tasa anual efectiva: v = (1 + r)^(−mes/12)"),
        ("Flujo del inversionista", "flujo del proyecto + (flujo del activo con la alternativa − flujo del "
                                    "activo comprado al contado); así no se cuenta dos veces la compra"),
    ]
    f += 1
    ws.cell(row=f, column=1, value="Reglas del modelo").font = neg
    f += 1
    for a, b in fijos:
        ws.cell(row=f, column=1, value=a)
        ws.cell(row=f, column=2, value=b)
        f += 1
    if reg.avisos:
        f += 1
        ws.cell(row=f, column=1, value="Avisos").font = neg
        f += 1
        for a in reg.avisos:
            ws.cell(row=f, column=1, value=a)
            f += 1
    anchos(ws, [42, 22, 18, 36, 12, 70])

    # AMORTIZACION ---------------------------------------------------------------------------
    ws = wb.create_sheet("AMORTIZACION")
    f = 1
    for r in ev["res"]:
        if not r["tabla"]:
            continue
        e = r["extra"]
        ws.cell(row=f, column=1, value=f"{r['nombre']} · {r['tipo']}"
                + (f" ({r['sistema']})" if r["tipo"] == "credito" else "")).font = Font(bold=True, size=12)
        f += 1
        info = [("Monto financiado", e.get("monto_financiado"), fmt_m)]
        if r["tipo"] == "credito":
            info += [("Gastos iniciales (incluye impuesto al crédito)", e.get("gastos_iniciales"), fmt_m),
                     ("Impuesto al crédito", e.get("impuesto_credito"), fmt_m),
                     ("CAE (TIR mensual × 12, con gastos)", e.get("cae"), fmt_p),
                     ("Costo efectivo anual antes de impuesto", e.get("tasa_efectiva_anual_credito"), fmt_p)]
        else:
            info += [("Cuota inicial", e.get("cuota_inicial"), fmt_m),
                     ("Opción de compra", e.get("opcion_compra"), fmt_m),
                     ("Tasa implícita mensual", e.get("tasa_implicita_mensual"), "0.0000%"),
                     ("Tasa implícita anual (TIR mensual × 12)", e.get("cae"), fmt_p)]
        for et, v, fm in info:
            ws.cell(row=f, column=1, value=et)
            c = ws.cell(row=f, column=2, value=v)
            c.number_format = fm
            f += 1
        encabezado(ws, f, ["Mes", "Saldo inicial", "Cuota", "Interés", "Amortización", "Saldo final"])
        f += 1
        for x in r["tabla"]:
            vals = [x["mes"], x["saldo_inicial"], x["cuota"], x["interes"], x["amortizacion"], x["saldo_final"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=f, column=j, value=v)
                if j > 1:
                    c.number_format = fmt_m
            if x.get("es_opcion"):
                ws.cell(row=f, column=7, value="opción de compra")
            f += 1
        tot = ["Total", None, sum(x["cuota"] for x in r["tabla"]), sum(x["interes"] for x in r["tabla"]),
               sum(x["amortizacion"] for x in r["tabla"]), None]
        for j, v in enumerate(tot, 1):
            c = ws.cell(row=f, column=j, value=v)
            c.font = neg
            if j > 1 and v is not None:
                c.number_format = fmt_m
        f += 1
        cuadra = abs(sum(x["amortizacion"] for x in r["tabla"]) - e["monto_financiado"]) < 0.5
        ws.cell(row=f, column=1, value="Suma de amortizaciones = monto financiado")
        ws.cell(row=f, column=2, value="OK" if cuadra else "NO CUADRA")
        f += 3
    anchos(ws, [44, 18, 16, 16, 16, 16, 18])

    # FLUJO-INVERSIONISTA ---------------------------------------------------------------------
    ws = wb.create_sheet("FLUJO-INVERSIONISTA")
    H, anios = P["meses"], P["horizonte"]
    proy = ev["res"][0].get("inversionista") is not None

    def por_anio(serie):
        out = [0.0] * (anios + 1)
        for m, v in enumerate(serie):
            out[0 if m == 0 else anio_de_mes(m)] += v
        return out

    ws.cell(row=1, column=1, value="Flujo del inversionista por alternativa (montos nominales por año)").font = \
        Font(bold=True, size=12)
    ws.cell(row=2, column=1, value="Inversionista = proyecto + precio al contado que no se paga + pagos de la "
                                   "alternativa + su ahorro tributario + su valor residual − (ahorro tributario y "
                                   "residual del contado de referencia). El VAN se calcula mes a mes (tabla al final).")
    f = 4
    ref = ev["ref"]
    ref_aho = por_anio([ref["ahorro_fin"][m] + ref["ahorro_act"][m] for m in range(H + 1)])
    ref_res = por_anio(ref["residual"])
    for r in ev["res"]:
        ws.cell(row=f, column=1, value=r["nombre"]).font = Font(bold=True, size=11)
        f += 1
        cols = ["Año", "Flujo del proyecto", "+ Precio al contado que no se paga", "− Pagos de la alternativa",
                "+ Devoluciones (garantía)", "+ Ahorro tributario intereses y gastos financieros",
                "+ Ahorro tributario depreciación / cuotas / renta", "+ Valor residual neto",
                "− Ahorro tributario del contado", "− Residual del contado", "= Flujo del inversionista"]
        encabezado(ws, f, cols)
        f += 1
        series = [
            por_anio(ev["proy_mensual"]) if proy else [None] * (anios + 1),
            por_anio([ref["precio"] if m == 0 else 0.0 for m in range(H + 1)]),
            por_anio(r["pagos"]), por_anio(r["ingresos"]), por_anio(r["ahorro_fin"]),
            por_anio(r["ahorro_act"]), por_anio(r["residual"]),
            [-x for x in ref_aho], [-x for x in ref_res],
        ]
        for a in range(anios + 1):
            vals = [s[a] for s in series]
            total = sum(v for v in vals if v is not None)
            fila = [a] + vals + [total if proy else None]
            for j, v in enumerate(fila, 1):
                c = ws.cell(row=f, column=j, value=v)
                if j > 1:
                    c.number_format = fmt_m
            f += 1
        if proy:
            ws.cell(row=f, column=1, value="VAN del inversionista").font = neg
            c = ws.cell(row=f, column=11, value=r["van_inversionista"])
            c.number_format, c.font = fmt_m, neg
            f += 1
        f += 2
    if proy:
        ws.cell(row=f, column=1, value="Tabla mensual (con esta se calcula el VAN)").font = Font(bold=True, size=11)
        f += 1
        cab = ["Mes", "Factor de descuento", "Flujo del proyecto"] + [f"Inversionista · {r['nombre']}"
                                                                       for r in ev["res"]]
        encabezado(ws, f, cab)
        f += 1
        for m in range(H + 1):
            fac = (1 + ev["tasa"]) ** (-m / 12)
            vals = [m, fac, ev["proy_mensual"][m]] + [r["inversionista"][m] for r in ev["res"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=f, column=j, value=v)
                c.number_format = "0.000000" if j == 2 else (fmt_m if j > 2 else "0")
            f += 1
        vals = ["VAN", None, ev["van_proyecto"]] + [r["van_inversionista"] for r in ev["res"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v)
            c.font = neg
            if j > 2:
                c.number_format = fmt_m
    anchos(ws, [24, 16] + [18] * 9)

    # COMPARACION ------------------------------------------------------------------------------
    ws = wb.create_sheet("COMPARACION")
    g, s = ev["orden"][0], ev["orden"][1]
    ws["A1"] = "Comparación de alternativas de financiamiento"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"Tasa de descuento {ev['tasa']:.2%} anual · impuesto {ev['t']:.2%} · horizonte "
                f"{anios} años · {P['descripcion']}")
    cab = ["Posición", "Alternativa", "Tipo", "Cuota", "Total pagado (nominal)", "Ahorro tributario total",
           "Valor residual neto", "Costo en valor presente después de impuesto", "Ventaja frente al contado (VP)",
           "Costo efectivo anual después de impuesto", "CAE o tasa implícita anual", "VAN del inversionista",
           "Cuánto más cuesta que la ganadora (VP)"]
    encabezado(ws, 4, cab)
    f = 5
    for k, r in enumerate(ev["orden"], 1):
        e = r["extra"]
        vals = [k, r["nombre"], r["tipo"] + (f" ({r['sistema']})" if r["tipo"] == "credito" else ""),
                r["cuota"], -sum(r["pagos"]), sum(r["ahorro_fin"]) + sum(r["ahorro_act"]), sum(r["residual"]),
                r["costo_vp"], r["ventaja_vs_contado"], r["tasa_efectiva_despues_impuesto"], e.get("cae"),
                r.get("van_inversionista"), r["costo_vp"] - g["costo_vp"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v)
            if j in (4, 5, 6, 7, 8, 9, 12, 13):
                c.number_format = fmt_m
            if j in (10, 11):
                c.number_format = fmt_p
            if k == 1:
                c.fill = gan_fill
        f += 1
    f += 1
    ws.cell(row=f, column=1, value="Ganadora").font = neg
    ws.cell(row=f, column=2, value=(f"{g['nombre']}: cuesta {s['costo_vp'] - g['costo_vp']:,.0f} {P['moneda']} "
                                    f"menos en valor presente que {s['nombre']} (segunda)."))
    f += 1
    ws.cell(row=f, column=1, value="Tasa que cambia la decisión").font = neg
    ws.cell(row=f, column=2, value=("; ".join(f"{x:.2%}" for x in indif) +
                                    f" anual: a esa tasa de descuento {g['nombre']} y {s['nombre']} empatan")
            if indif else f"ninguna entre 0 % y 60 %: {g['nombre']} gana a {s['nombre']} en todo ese rango")
    f += 1
    ws.cell(row=f, column=1, value="Cómo leer el costo efectivo").font = neg
    ws.cell(row=f, column=2, value=("Costo efectivo anual después de impuesto = TIR de (flujo de la alternativa − "
                                    "flujo al contado). Si es menor que la tasa de descuento, financiar así conviene "
                                    "más que pagar al contado."))
    notas = [(r["nombre"], n) for r in ev["res"] for n in r["notas"]] + \
            [(r["nombre"], r["extra"]["aviso"]) for r in ev["res"] if r["extra"].get("aviso")]
    if notas:
        f += 2
        ws.cell(row=f, column=1, value="No comparable (se declara, no se promedia)").font = neg
        f += 1
        for a, n in notas:
            ws.cell(row=f, column=1, value=a)
            ws.cell(row=f, column=2, value=n)
            f += 1
    f += 2
    ws.cell(row=f, column=1, value="Sensibilidad: costo en VP después de impuesto").font = Font(bold=True, size=11)
    f += 1
    nombres = [r["nombre"] for r in ev["res"]]
    encabezado(ws, f, ["Impuesto", "Tasa de descuento"] + nombres + ["Ganadora"])
    f += 1
    for x in sens:
        vals = [x["impuesto"], x["tasa"]] + [x["costos"][n] for n in nombres] + [x["ganadora"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v)
            c.number_format = fmt_p if j <= 2 else fmt_m
            if x["impuesto"] == round(ev["t"], 4) and x["tasa"] == round(ev["tasa"], 4):
                c.font = neg
        f += 1
    anchos(ws, [22, 30, 18, 14, 18, 18, 18, 22, 20, 20, 18, 20, 20])
    wb.save(ruta)


# ----------------------------------------------------------------------------------------------
# Orquestación
# ----------------------------------------------------------------------------------------------

def correr(ruta_caso: Path, perfil=None, tasa=None, impuesto=None, proyecto=None, salida=None, silencioso=False):
    caso = json.loads(ruta_caso.read_text(encoding="utf-8-sig"))
    reg = Registro()
    P, alts, flujo_proy = preparar(caso, ruta_caso, perfil, tasa, impuesto, proyecto, reg)
    ev = evaluar(P, alts, flujo_proy)
    ev["proy_mensual"] = flujo_proy
    g, s = ev["orden"][0], ev["orden"][1]
    indif = tasa_indiferencia(P, alts, flujo_proy, g["nombre"], s["nombre"])
    sens = sensibilidad(P, alts, flujo_proy, (caso.get("sensibilidad") or {}).get("impuestos"))
    out = resumen_json(P, ev, indif, sens, reg, alts)
    carpeta = Path(salida) if salida else ruta_caso.parent / "salidas"
    carpeta.mkdir(parents=True, exist_ok=True)
    base = caso.get("caso") or ruta_caso.stem
    (carpeta / f"{base}-resultado.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    titulo = f"Evaluación financiera · {caso.get('caso', ruta_caso.stem)}"
    escribir_excel(carpeta / f"{base}-evaluacion-financiera.xlsx", P, ev, indif, sens, reg, titulo)
    if not silencioso:
        imprimir(P, ev, indif, reg, carpeta, base)
    return out, ev, P, alts, reg


def imprimir(P, ev, indif, reg, carpeta, base):
    mon = P["moneda"]
    print(f"Tasa de descuento {ev['tasa']:.2%} · impuesto {ev['t']:.2%} · horizonte {P['horizonte']} años")
    sup = [f["parametro"] for f in reg.filas if f["marca"] == "SUPUESTO"]
    if sup:
        print("SUPUESTO: " + ", ".join(sup))
    for a in reg.avisos:
        print("AVISO: " + a)
    if ev["van_proyecto"] is not None:
        print(f"VAN del proyecto puro: {ev['van_proyecto']:,.0f} {mon}")
    print(f"{'Alternativa':34} {'Costo VP d/i':>16} {'Costo ef. d/i':>13} {'CAE/impl.':>10} {'VAN inversionista':>18}")
    for r in ev["orden"]:
        te = r["tasa_efectiva_despues_impuesto"]
        cae = r["extra"].get("cae")
        van = r.get("van_inversionista")
        print(f"{r['nombre'][:34]:34} {r['costo_vp']:>16,.0f} {('—' if te is None else f'{te:.2%}'):>13} "
              f"{('—' if cae is None else f'{cae:.2%}'):>10} {('—' if van is None else f'{van:,.0f}'):>18}")
    g, s = ev["orden"][0], ev["orden"][1]
    print(f"GANA {g['nombre']}: {s['costo_vp'] - g['costo_vp']:,.0f} {mon} menos en VP que {s['nombre']}.")
    print("Tasa de indiferencia con la segunda: " + (", ".join(f"{x:.2%}" for x in indif) if indif else "ninguna en 0-60 %"))
    print(f"Salida: {carpeta / (base + '-evaluacion-financiera.xlsx')}")


# ----------------------------------------------------------------------------------------------
# --probar: el caso ficticio, comparado con esperado.json y con un segundo camino
# ----------------------------------------------------------------------------------------------

def probar() -> int:
    pr = RAIZ_SKILL / "pruebas"
    caso = pr / "caso-ficticio.json"
    esperado = json.loads((pr / "esperado.json").read_text(encoding="utf-8"))
    fallas: list[str] = []

    def chequear(cond, msg):
        print(("  ok   " if cond else "  FALLA ") + msg)
        if not cond:
            fallas.append(msg)

    with tempfile.TemporaryDirectory() as tmp:
        out, ev, P, alts, reg = correr(caso, perfil=str(pr / "perfil-ficticio.yaml"), salida=tmp, silencioso=True)
        print("1. Comparación con esperado.json (tolerancia 1 peso; tasas 1e-6)")
        for nombre, esp in esperado["alternativas"].items():
            got = out["alternativas"].get(nombre)
            if got is None:
                chequear(False, f"falta la alternativa {nombre}")
                continue
            for k, v in esp.items():
                g = got.get(k)
                if isinstance(v, float) and abs(v) < 5:
                    chequear(g is not None and abs(g - v) <= 1e-6, f"{nombre} · {k}: {g} vs {v}")
                elif isinstance(v, (int, float)):
                    chequear(g is not None and abs(g - v) <= 1.0, f"{nombre} · {k}: {g} vs {v}")
                else:
                    chequear(g == v, f"{nombre} · {k}: {g} vs {v}")
        chequear(out["ganadora"] == esperado["ganadora"], f"ganadora {out['ganadora']} vs {esperado['ganadora']}")
        chequear(abs(out["diferencia_vp"] - esperado["diferencia_vp"]) <= 1.0,
                 f"diferencia {out['diferencia_vp']} vs {esperado['diferencia_vp']}")
        chequear(out["van_proyecto"] is not None and abs(out["van_proyecto"] - esperado["van_proyecto"]) <= 1.0,
                 f"VAN del proyecto {out['van_proyecto']} vs {esperado['van_proyecto']}")

        print("2. Segundo camino: fórmula de anualidad contra tabla sumada, al peso")
        for r in ev["res"]:
            if not r["tabla"]:
                continue
            e, tab = r["extra"], r["tabla"]
            suma_am = sum(x["amortizacion"] for x in tab)
            chequear(abs(suma_am - e["monto_financiado"]) < 0.5,
                     f"{r['nombre']}: suma de amortizaciones {suma_am:,.0f} = monto {e['monto_financiado']:,.0f}")
            chequear(abs(sum(x["cuota"] for x in tab) - sum(x["interes"] for x in tab) - suma_am) < 0.5,
                     f"{r['nombre']}: cuotas − intereses = amortizaciones")
            if r["tipo"] == "credito" and r["sistema"] == "francesa":
                alt = next(a for a in alts if a["nombre"] == r["nombre"])
                i, n, L = alt["tasa_mensual"], alt["plazo_meses"], e["monto_financiado"]
                c_formula = L * i * (1 + i) ** n / ((1 + i) ** n - 1)
                chequear(abs(round(c_formula) - tab[0]["cuota"]) <= 0,
                         f"{r['nombre']}: cuota por fórmula {c_formula:,.2f} → {round(c_formula):,} = tabla {tab[0]['cuota']:,.0f}")
                chequear(abs(tab[-1]["cuota"] - tab[0]["cuota"]) <= n,
                         f"{r['nombre']}: última cuota ajusta el redondeo ({tab[-1]['cuota'] - tab[0]['cuota']:+,.0f})")
                vp_cuotas = sum(x["cuota"] / (1 + i) ** x["mes"] for x in tab)
                chequear(abs(vp_cuotas - L) <= 0.5 * n,
                         f"{r['nombre']}: VP de las cuotas a la tasa = monto ({vp_cuotas - L:+.2f}; "
                         f"tolerancia medio peso por cuota redondeada)")
            if r["tipo"] == "credito" and r["sistema"] == "alemana":
                alt = next(a for a in alts if a["nombre"] == r["nombre"])
                i, n, L = alt["tasa_mensual"], alt["plazo_meses"], e["monto_financiado"]
                int_formula = i * L * (n + 1) / 2
                chequear(abs(int_formula - sum(x["interes"] for x in tab)) <= n,
                         f"{r['nombre']}: intereses por fórmula i·L·(n+1)/2 = {int_formula:,.0f} ≈ tabla")
            if r["tipo"] == "leasing":
                chequear(abs(tab[-1].get("descuadre_redondeo", 0)) <= len(tab),
                         f"{r['nombre']}: la opción cierra el saldo (descuadre {tab[-1].get('descuadre_redondeo', 0):+,.0f})")
        print("3. Identidades del modelo")
        for r in ev["res"]:
            van_id = ev["van_proyecto"] + r["ventaja_vs_contado"]
            chequear(abs(van_id - r["van_inversionista"]) <= 1.0,
                     f"{r['nombre']}: VAN inversionista = VAN proyecto + ventaja frente al contado")
        contado = next((r for r in ev["res"] if r["tipo"] == "contado"), None)
        if contado:
            chequear(abs(contado["van_inversionista"] - ev["van_proyecto"]) <= 1.0,
                     "contado: VAN del inversionista = VAN del proyecto puro")
        from openpyxl import load_workbook
        wb = load_workbook(Path(tmp) / f"{json.loads(caso.read_text(encoding='utf-8'))['caso']}-evaluacion-financiera.xlsx")
        chequear(wb.sheetnames == ["SUPUESTOS", "AMORTIZACION", "FLUJO-INVERSIONISTA", "COMPARACION"],
                 f"hojas del Excel: {wb.sheetnames}")

        print("4. Perfil sin tasa: el script la exige y la marca SUPUESTO")
        try:
            correr(caso, perfil=str(PLANTILLA_PERFIL), salida=tmp, silencioso=True)
            chequear(False, "con la plantilla y sin --tasa debía detenerse")
        except ErrorEntrada as ex:
            chequear("--tasa" in str(ex), "sin --tasa se detiene y pide el argumento")
        o2, *_ = correr(caso, perfil=str(PLANTILLA_PERFIL), tasa=0.12, salida=tmp, silencioso=True)
        chequear(any(s["parametro"] == "tasa_descuento" and "--tasa" in s["origen"] for s in o2["supuestos"]),
                 "con --tasa la tasa sale marcada SUPUESTO")

    print("OK" if not fallas else f"FALLA ({len(fallas)})")
    return 0 if not fallas else 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Evaluación financiera: contado, crédito, leasing o arriendo.")
    ap.add_argument("--caso", help="JSON del caso (ver plantillas/caso.json)")
    ap.add_argument("--proyecto", help="flujo del proyecto puro: CSV o JSON con año y flujo")
    ap.add_argument("--perfil", help="nombre del perfil (busca en INGENIERO_DATOS y ../ingeniero-agentico-datos) o ruta")
    ap.add_argument("--tasa", type=float, help="tasa de descuento anual (obligatoria si el perfil no la trae)")
    ap.add_argument("--impuesto", type=float, help="impuesto a la renta (obligatorio si el perfil no lo trae)")
    ap.add_argument("--salida", help="carpeta de salida (por defecto <carpeta del caso>/salidas)")
    ap.add_argument("--probar", action="store_true", help="corre el caso ficticio y verifica que cuadre")
    a = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if a.probar:
        return probar()
    if not a.caso:
        ap.error("falta --caso (o usa --probar)")
    try:
        correr(Path(a.caso), a.perfil, a.tasa, a.impuesto, a.proyecto, a.salida)
    except ErrorEntrada as ex:
        print(f"ERROR: {ex}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
