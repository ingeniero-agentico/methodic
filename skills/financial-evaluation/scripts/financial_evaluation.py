#!/usr/bin/env python3
"""financial-evaluation · how to pay for an asset and how much it changes for the owner.

Compares financing alternatives for the same asset (cash, loan with French or German
amortization, financial leasing with purchase option, rental) after tax and in present value,
and builds the equity (investor) cash flow by adding to the pure project flow what each
alternative changes. Sheet names, labels, the assumption marker and the console text follow the
profile `language` (es by default: SUPUESTOS · AMORTIZACION · FLUJO-INVERSIONISTA · COMPARACION;
en: ASSUMPTIONS · AMORTIZATION · INVESTOR-CASH-FLOW · COMPARISON).

Usage:
    python financial_evaluation.py --case case.json [--project flow.csv|output.json]
                                   [--profile company] [--rate 0.12] [--tax 0.27] [--out folder]
    python financial_evaluation.py --test

Spanish flags still work as hidden aliases (--caso --proyecto --perfil --tasa --impuesto
--salida --probar), and Spanish case and profile keys are accepted through an alias map.

Standard library + openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import sys
import tempfile
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

SKILL_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = Path(__file__).resolve().parents[3]
TEMPLATES = (REPO_ROOT / "profiles" / "_template.yaml", REPO_ROOT / "perfiles" / "_plantilla.yaml")
TYPES = ("cash", "loan", "leasing", "rental")

# ----------------------------------------------------------------------------------------------
# Output language (profile `language`, default es)
# ----------------------------------------------------------------------------------------------

_LANG = "es"


def tr(es: str, en: str) -> str:
    return en if _LANG == "en" else es


def mark(assumption: bool) -> str:
    if assumption:
        return tr("SUPUESTO", "ASSUMPTION")
    return tr("dato", "data")


class InputError(Exception):
    """A datum is missing or wrong: the script does not invent it, it stops and says which."""


# ----------------------------------------------------------------------------------------------
# Spanish → English alias maps (case, profile, values)
# ----------------------------------------------------------------------------------------------

CASE_KEYS = {
    "caso": "case", "descripcion": "description", "perfil": "profile", "fecha": "date",
    "horizonte_anios": "horizon_years", "iva_recuperable": "vat_recoverable",
    "flujo_proyecto": "project_flow", "activo": "asset", "precio_neto": "net_price",
    "depreciacion": "depreciation", "vida_util_anios": "useful_life_years",
    "vida_util_acelerada_anios": "accelerated_life_years",
    "valor_mercado_al_horizonte": "market_value_at_horizon", "alternativas": "alternatives",
    "nombre": "name", "tipo": "type", "sistema": "system", "pie_fraccion": "down_payment_fraction",
    "pie": "down_payment", "plazo_meses": "term_months", "tasa_mensual": "monthly_rate",
    "tasa_anual": "annual_rate", "gastos_iniciales": "upfront_fees", "gastos_financiados": "fees_financed",
    "seguro_mensual": "monthly_insurance", "impuesto_credito": "loan_tax", "tope": "cap",
    "cuota_mensual": "monthly_installment", "cuota_inicial": "initial_installment",
    "opcion_compra": "purchase_option", "mes_opcion": "option_month", "renta_mensual": "monthly_rent",
    "garantia": "deposit", "notas_no_comparables": "non_comparable_notes", "sensibilidad": "sensitivity",
    "impuestos": "tax_rates", "valor": "value", "fuente": "source", "supuesto": "assumption",
    "investigacion": "research",
}
VALUE_KEYS = {"type", "system", "depreciation", "initial_installment", "purchase_option"}
VALUES = {"contado": "cash", "credito": "loan", "crédito": "loan", "arriendo": "rental",
          "francesa": "french", "alemana": "german", "una_cuota": "one_installment",
          "lineal": "straight_line", "acelerada": "accelerated", "instantanea": "instant",
          "instantánea": "instant"}

PROFILE_SECTIONS = {"empresa": "company", "compras": "purchasing", "precios": "pricing",
                    "finanzas": "finance", "idioma": "language"}
PROFILE_KEYS = {"nombre": "name", "rut": "tax_id", "pais": "country", "moneda": "currency",
                "decimales_moneda": "currency_decimals", "iva": "vat", "impuesto_renta": "income_tax",
                "proveedores": "suppliers", "plazo_entrega_dias": "lead_time_days",
                "desempate": "tie_break", "recargo_por_defecto": "default_markup",
                "gastos_generales": "overhead", "tasa_descuento": "discount_rate",
                "horizonte_anios": "horizon_years", "depreciacion": "depreciation"}
PROFILE_VALUES = {"mayor_disponibilidad": "highest_availability", "lineal": "straight_line",
                  "acelerada": "accelerated", "instantanea": "instant"}

ES_KEY = {v: k for k, v in CASE_KEYS.items()}
ES_KEY.update({v: k for k, v in PROFILE_KEYS.items()})
ES_KEY.update({"currency": "moneda", "vat": "iva", "profile": "perfil", "project_flow": "flujo_proyecto",
               "project_flow_regime": "regimen_flujo_proyecto"})
ES_VALUE = {"cash": "contado", "loan": "credito", "rental": "arriendo", "french": "francesa",
            "german": "alemana", "one_installment": "una_cuota", "straight_line": "lineal",
            "accelerated": "acelerada", "instant": "instantanea"}


def _code(v):
    return VALUES.get(v, v) if isinstance(v, str) else v


def normalize_case(obj):
    """Spanish keys (and code values, also inside {value, source, assumption}) → English."""
    if isinstance(obj, list):
        return [normalize_case(x) for x in obj]
    if not isinstance(obj, dict):
        return obj
    out: dict = {}
    for k in sorted(obj, key=lambda k: k not in CASE_KEYS):
        v = obj[k]
        nk = k if k.startswith("_") else CASE_KEYS.get(k, k)
        if nk in VALUE_KEYS:
            if isinstance(v, dict):
                v = normalize_case(v)
                v["value"] = _code(v.get("value"))
            else:
                v = _code(v)
        elif nk in ("non_comparable_notes",):
            pass
        elif isinstance(v, (dict, list)):
            v = normalize_case(v)
        out[nk] = v
    return out


def normalize_profile(raw: dict) -> dict:
    out: dict = {}
    for k in sorted(raw, key=lambda k: k not in PROFILE_SECTIONS):
        v = raw[k]
        nk = PROFILE_SECTIONS.get(k, k)
        if isinstance(v, dict):
            dest = out.setdefault(nk, {})
            for ik in sorted(v, key=lambda k: k not in PROFILE_KEYS):
                iv = v[ik]
                nik = PROFILE_KEYS.get(ik, ik)
                if nik in ("tie_break", "depreciation") and isinstance(iv, str):
                    iv = PROFILE_VALUES.get(iv, iv)
                dest[nik] = iv
        else:
            out[nk] = v
    return out


def loc_path(path: str) -> str:
    """'alternatives[X].monthly_rate' → 'alternativas[X].tasa_mensual' in es (names outside brackets)."""
    if _LANG != "es":
        return path
    parts = []
    for seg in path.split("."):
        base, br, rest = seg.partition("[")
        parts.append(ES_KEY.get(base, base) + (br + rest if br else ""))
    return ".".join(parts)


def loc_value(v):
    if _LANG == "es" and isinstance(v, str):
        return ES_VALUE.get(v, v)
    return v


def set_language(profile: dict):
    global _LANG
    lang = str(profile.get("language") or "es").strip().lower()
    if lang not in ("es", "en"):
        raise InputError(f"language «{lang}»: use es or en.")
    _LANG = lang
    return lang


# ----------------------------------------------------------------------------------------------
# Reading: YAML profile (simple subset), case JSON, project flow
# ----------------------------------------------------------------------------------------------

def _scalar(txt: str):
    t = txt.strip()
    if t in ("", "~", "null", "Null", "NULL"):
        return None
    if t in ("true", "True"):
        return True
    if t in ("false", "False"):
        return False
    if t == "[]":
        return []
    if t == "{}":
        return {}
    if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
        return t[1:-1]
    try:
        return int(t)
    except ValueError:
        pass
    try:
        return float(t)
    except ValueError:
        return t


def _strip_comment(line: str) -> str:
    inside = None
    for i, c in enumerate(line):
        if c in "\"'":
            if inside is None:
                inside = c
            elif inside == c:
                inside = None
        elif c == "#" and inside is None and (i == 0 or line[i - 1] in " \t"):
            return line[:i]
    return line


def read_simple_yaml(path: Path) -> dict:
    """Nested maps by indentation, `key: value  # comment`. Lists are ignored (not used here)."""
    root: dict = {}
    stack = [(-1, root)]
    for raw in Path(path).read_text(encoding="utf-8-sig").splitlines():
        line = _strip_comment(raw).rstrip()
        if not line.strip():
            continue
        indent = len(line) - len(line.lstrip(" "))
        body = line.strip()
        if body.startswith("-") or ":" not in body:
            continue
        key, _, value = body.partition(":")
        while len(stack) > 1 and indent <= stack[-1][0]:
            stack.pop()
        parent = stack[-1][1]
        if value.strip() == "":
            new: dict = {}
            parent[key.strip()] = new
            stack.append((indent, new))
        else:
            parent[key.strip()] = _scalar(value)
    return root


def template_profile() -> Path:
    for t in TEMPLATES:
        if t.is_file():
            return t
    raise InputError("No profile template in the repo (profiles/_template.yaml or perfiles/_plantilla.yaml).")


ENV_DATA = ("METHODIC_DATA", "INGENIERO_DATOS")              # the second one is the legacy name
SIBLING_DATA = ("methodic-data", "ingeniero-agentico-datos")  # the second one is the legacy name


def sibling_data() -> Path:
    """../methodic-data, or the legacy ../ingeniero-agentico-datos when the new one does not exist."""
    new, old = (REPO_ROOT.parent / n for n in SIBLING_DATA)
    return old if not new.is_dir() and old.is_dir() else new


def find_profile(name: str | None) -> tuple[Path, str]:
    """Order: direct path · $METHODIC_DATA/(profiles|perfiles) · $INGENIERO_DATOS (legacy) ·
    ../methodic-data/(profiles|perfiles), or the legacy ../ingeniero-agentico-datos/ when the new one
    does not exist · repo template. Returns (path, how), how = arg_path, env:<VAR>, sibling:<folder>
    or template."""
    if name:
        direct = Path(name)
        if direct.is_file():
            return direct, "arg_path"
        fname = name if name.endswith((".yaml", ".yml")) else name + ".yaml"
        candidates = []
        for var in ENV_DATA:
            if os.environ.get(var):
                for sub in ("profiles", "perfiles"):
                    candidates.append((Path(os.environ[var]) / sub / fname, f"env:{var}"))
        sibling = sibling_data()
        for sub in ("profiles", "perfiles"):
            candidates.append((sibling / sub / fname, f"sibling:{sibling.name}"))
        for path, how in candidates:
            if path.is_file():
                return path, how
        raise InputError(tr(f"No encontré el perfil «{name}» (busqué en: ", f"Profile «{name}» not found (looked in: ")
                         + "; ".join(str(c[0]) for c in candidates) + ").")
    return template_profile(), "template"


def profile_origin(how: str) -> str:
    kind, _, detail = how.partition(":")
    return {"arg_path": tr("ruta indicada en --perfil", "path given in --profile"),
            "env": tr(f"variable {detail}", f"{detail} variable"),
            "sibling": tr(f"carpeta hermana {detail}", f"sibling folder {detail}"),
            "template": tr("plantilla (no hay perfil de empresa)", "template (no company profile)")}[kind]


def is_template(path: Path) -> bool:
    return any(t.is_file() and path.resolve() == t.resolve() for t in TEMPLATES)


def _num(txt: str) -> float:
    t = str(txt).strip().replace("$", "").replace(" ", "").replace("−", "-")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:
        t = t.replace(",", ".")
    return float(t)


YEAR_NAMES = ("year", "anio", "año", "ano", "periodo", "period", "t")
FLOW_NAMES = ("flow", "net_flow", "project_flow", "cash_flow", "net_cash_flow", "flujo", "flujo_neto",
              "flujo_proyecto", "flujo_caja", "fcn", "flujo_neto_caja")


def read_project_flow(path: Path) -> tuple[dict[int, float], dict]:
    """Yearly pure project flow. Accepts:
    - the JSON output of economic-evaluation (new English keys: cash_flow.years + net_cash_flow,
      or the old Spanish salida.json: flujo.anios + flujo_de_caja); its tax regime is read too;
    - CSV (separator , ; or tab) with columns year/anio/period and flow/flujo/net_flow;
    - JSON: list of {"year": n, "flow": x}, or {"years": [...], "flows": [...]}, or {"flow": <list>}.
    Numbers without thousands separator. Returns (flow by year, regime {tax, depreciation, source})."""
    if not path.is_file():
        raise InputError(tr(f"No existe el archivo del flujo del proyecto: {path}", f"The project flow file does not exist: {path}"))
    regime = {"tax": None, "depreciation": None, "source": "csv" if path.suffix.lower() != ".json" else "json",
              "assumptions": []}
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict) and isinstance(data.get("cash_flow"), dict):
            par = data.get("parameters") or {}
            regime = {"tax": (par.get("tax") or {}).get("value"),
                      "depreciation": _code((par.get("depreciation") or {}).get("value")),
                      "source": "economic-evaluation",
                      "assumptions": [a for a in data.get("assumptions") or [] if isinstance(a, dict)]}
            data = data["cash_flow"]
        elif isinstance(data, dict) and isinstance(data.get("flujo"), dict):
            par = data.get("parametros") or {}
            regime = {"tax": (par.get("impuesto") or {}).get("valor"),
                      "depreciation": _code((par.get("depreciacion") or {}).get("valor")),
                      "source": "evaluacion-economica",
                      "assumptions": [_from_old_assumption(a) for a in data.get("supuestos") or []
                                      if isinstance(a, dict) and a.get("marca") == "SUPUESTO"]}
            data = data["flujo"]
        years_key = next((k for k in ("years", "anios") if isinstance(data, dict) and k in data), None)
        if years_key:
            series = next((data[k] for k in ("net_cash_flow", "flows", "cash_flow", "flow", "flujos", "flujo_de_caja",
                                             "flujo", "flujo_proyecto") if isinstance(data.get(k), list)), None)
            if series is None:
                raise InputError(tr(f"{path.name}: trae «{years_key}» pero no la serie del flujo (net_cash_flow · flujo_de_caja · flujo).",
                                    f"{path.name}: it has «{years_key}» but not the flow series (net_cash_flow · flujo_de_caja · flow)."))
            return {int(a): float(f) for a, f in zip(data[years_key], series)}, regime
        if isinstance(data, dict):
            for k in ("flow", "project_flow", "rows", "flujo", "flujo_proyecto", "filas"):
                if k in data:
                    data = data[k]
                    break
        res = {}
        for row in data:
            fa = {k.lower(): v for k, v in row.items()}
            a = next(fa[k] for k in YEAR_NAMES if k in fa)
            f = next(fa[k] for k in FLOW_NAMES if k in fa)
            res[int(a)] = float(f)
        return res, regime
    text = path.read_text(encoding="utf-8-sig")
    sep = max((";", ",", "\t"), key=lambda s: text.splitlines()[0].count(s))
    reader = csv.DictReader(text.splitlines(), delimiter=sep)
    cols = {c.strip().lower(): c for c in (reader.fieldnames or [])}
    ca = next((cols[k] for k in YEAR_NAMES if k in cols), None)
    cf = next((cols[k] for k in FLOW_NAMES if k in cols), None)
    if not ca or not cf:
        raise InputError(tr(f"El flujo {path.name} necesita columnas año y flujo; trae {list(cols)}",
                            f"The flow {path.name} needs year and flow columns; it has {list(cols)}"))
    res = {}
    for row in reader:
        if not (row.get(ca) or "").strip():
            continue
        res[int(_num(row[ca]))] = _num(row[cf])
    return res, regime


def _from_old_assumption(a: dict) -> dict:
    """Old Spanish row {parametro, valor, unidad, origen, marca} → {parameter, value, origin, text}."""
    origin = str(a.get("origen") or "")
    low = origin.lower()
    code = ("profile" if low.startswith("perfil") else "template" if low.startswith("plantilla")
            else "default" if low.startswith("por defecto") else "input")
    unit = f" {a['unidad']}" if a.get("unidad") else ""
    return {"parameter": a.get("parametro", ""), "value": a.get("valor"), "origin": code,
            "text": f"SUPUESTO · {a.get('parametro', '')} = {a.get('valor')}{unit} · {origin}"}


# ----------------------------------------------------------------------------------------------
# Assumption registry (traceability: where each number came from)
# ----------------------------------------------------------------------------------------------

_MISSING = object()


class Registry:
    def __init__(self) -> None:
        self.rows: list[dict] = []
        self.warnings: list[str] = []

    def add(self, key, value, unit, origin, assumption, note="", code="input"):
        """code: origin in the common format (input · profile · template · default · script · research:<id>)."""
        self.rows.append(dict(parameter=loc_path(key), key=key, value=value, unit=unit, origin=origin,
                              assumption=bool(assumption), note=note or "", code=code))

    def read(self, d: dict, key: str, path: str, unit: str = "", default=_MISSING, default_note=""):
        full = f"{path}.{key}"
        origin = tr("caso: ", "case: ") + loc_path(full)
        if key in d and d[key] is not None:
            v = d[key]
            if isinstance(v, dict) and "value" in v:
                rid = v.get("research") or v.get("research_id")
                self.add(full, v["value"], unit, origin, bool(v.get("assumption")), v.get("source", ""),
                         f"research:{rid}" if rid else "input")
                return v["value"]
            self.add(full, v, unit, origin, False)
            return v
        if default is _MISSING:
            raise InputError(tr(f"Falta «{loc_path(full)}» en el caso.", f"«{full}» is missing in the case."))
        self.add(full, default, unit, tr("valor por defecto del script", "script default"), True, default_note, "default")
        return default


# ----------------------------------------------------------------------------------------------
# Financial arithmetic
# ----------------------------------------------------------------------------------------------

def rnd(x: float, dec: int) -> float:
    q = Decimal(1).scaleb(-dec)
    return float(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP))


def annuity_factor(i: float, n: int) -> float:
    """a(n, i) = (1 - (1+i)^-n) / i : present value of n unit payments in arrears."""
    return n if i == 0 else (1 - (1 + i) ** -n) / i


def french_installment(amount: float, i: float, n: int) -> float:
    return amount / annuity_factor(i, n)


def french_table(amount, i, n, dec):
    c = rnd(french_installment(amount, i, n), dec)
    balance, rows = amount, []
    for m in range(1, n + 1):
        interest = rnd(balance * i, dec)
        principal = balance if m == n else rnd(c - interest, dec)
        rows.append(dict(month=m, opening_balance=balance, installment=rnd(principal + interest, dec), interest=interest,
                         principal=principal, closing_balance=rnd(balance - principal, dec)))
        balance = rows[-1]["closing_balance"]
    return rows, c


def german_table(amount, i, n, dec):
    a = rnd(amount / n, dec)
    balance, rows = amount, []
    for m in range(1, n + 1):
        interest = rnd(balance * i, dec)
        principal = balance if m == n else a
        rows.append(dict(month=m, opening_balance=balance, installment=rnd(principal + interest, dec), interest=interest,
                         principal=principal, closing_balance=rnd(balance - principal, dec)))
        balance = rows[-1]["closing_balance"]
    return rows, rows[0]["installment"]


def monthly_from_annual(r: float) -> float:
    return (1 + r) ** (1 / 12) - 1


def pv(flows: list[float], r_annual: float) -> float:
    """Present value of a monthly flow (index = month) at an effective annual rate."""
    return sum(f / (1 + r_annual) ** (m / 12) for m, f in enumerate(flows) if f)


def irr_roots(flows: list[float], lo=-0.03, hi=0.15, step=0.0005) -> list[float]:
    """Every monthly IRR between lo and hi (−31 % to +435 % annual), by scan + bisection.
    More than one root = the flow changes sign several times and the IRR is not unique."""
    def npv(r):
        return sum(f / (1 + r) ** m for m, f in enumerate(flows) if f)
    roots, prev = [], None
    for k in range(int(round((hi - lo) / step)) + 1):
        r = lo + k * step
        v = npv(r)
        if prev is not None and (prev[1] < 0) != (v < 0):
            a, b, fa = prev[0], r, prev[1]
            for _ in range(100):
                mid = (a + b) / 2
                fm = npv(mid)
                if (fa < 0) == (fm < 0):
                    a, fa = mid, fm
                else:
                    b = mid
            roots.append((a + b) / 2)
        prev = (r, v)
    return roots


def monthly_irr(flows: list[float]) -> float | None:
    """Monthly IRR of a flow with a single sign change (CAE, implicit rate). None if there is none."""
    roots = irr_roots(flows)
    return roots[0] if roots else None


# ----------------------------------------------------------------------------------------------
# Model: asset flows per alternative
# ----------------------------------------------------------------------------------------------

def year_of_month(m: int) -> int:
    return max(1, math.ceil(m / 12))


def depreciation(base, method, life, accel_life, start_year):
    if base <= 0:
        return {}
    if method == "instant":
        return {start_year: base}
    v = accel_life if method == "accelerated" else life
    if not v or v <= 0:
        raise InputError(tr(f"Depreciación {loc_value(method)} sin vida útil válida.",
                            f"{method} depreciation without a valid useful life."))
    return {start_year + k: base / v for k in range(int(v))}


def build_alternative(alt: dict, P: dict, t: float) -> dict:
    """Monthly asset flow for one alternative (owner, after tax).
    Returns the components per month and the amortization table when there is one."""
    H = P["months"]
    dec = P["dec"]
    fvat = P["vat_factor"]
    payments = [0.0] * (H + 1)
    residual = [0.0] * (H + 1)
    fin_ded: dict[int, float] = {}     # interest, financial costs, loan insurance
    asset_ded: dict[int, float] = {}   # depreciation, leasing installments, rent
    inflows = [0.0] * (H + 1)          # deposit returned, etc.

    def add(d, year, x):
        d[year] = d.get(year, 0.0) + x

    kind = alt["type"]
    price = alt["price"] * fvat
    table, installment, extra = None, None, {}
    dep_base, purchase_year, purchase_month = 0.0, 1, 0

    if kind == "cash":
        payments[0] -= price
        dep_base, purchase_year = price, 1
    elif kind == "loan":
        down = alt["down_payment"]
        amount = price - down
        n, i = alt["term_months"], alt["monthly_rate"]
        loan_tax = rnd(amount * min(alt["loan_tax_rate"] * n, alt["loan_tax_cap"]), dec) \
            if alt["loan_tax_rate"] else 0.0
        fees = alt["upfront_fees"] + loan_tax
        if alt["fees_financed"]:
            amount += fees
            payments[0] -= down
        else:
            payments[0] -= down + fees
        add(fin_ded, 1, fees)
        table, installment = (french_table if alt["system"] == "french" else german_table)(amount, i, n, dec)
        for f in table:
            m = f["month"]
            if m > H:
                break
            payments[m] -= f["installment"] + alt["monthly_insurance"]
            add(fin_ded, year_of_month(m), f["interest"] + alt["monthly_insurance"])
        if n > H:
            balance = table[H - 1]["closing_balance"]
            payments[H] -= balance
            extra["balance_paid_at_horizon"] = balance
        dep_base, purchase_year = price, 1
        # CAE: what really reaches the borrower (amount − fees) against everything paid
        cae_flow = [amount - fees] + [-(f["installment"] + alt["monthly_insurance"]) for f in table]
        tm = monthly_irr(cae_flow)
        extra.update(amount_financed=amount, upfront_fees=fees, loan_tax=loan_tax,
                     cae=None if tm is None else tm * 12,
                     pre_tax_effective_annual_rate=None if tm is None else (1 + tm) ** 12 - 1)
    elif kind == "leasing":
        n = alt["term_months"]
        op_month = alt["option_month"] or n + 1
        if op_month > H:
            raise InputError(tr(f"{alt['name']}: la opción de compra (mes {op_month}) cae después del horizonte ({H} meses). "
                                "Alarga el horizonte o acorta el leasing.",
                                f"{alt['name']}: the purchase option (month {op_month}) falls after the horizon ({H} months). "
                                "Extend the horizon or shorten the leasing."))
        i = alt["monthly_rate"]
        if alt["given_installment"] is not None:
            c = alt["given_installment"] * fvat
        else:
            # P = initial + C·a(n,i) + option·v^op_month ; initial and option may be "one_installment" (k·C)
            k_init = 1.0 if alt["initial_installment"] == "one_installment" else 0.0
            k_op = 1.0 if alt["purchase_option"] == "one_installment" else 0.0
            fixed_init = 0.0 if k_init else float(alt["initial_installment"] or 0) * fvat
            fixed_op = 0.0 if k_op else float(alt["purchase_option"] or 0) * fvat
            v_op = (1 + i) ** -op_month
            c = (price - fixed_init - fixed_op * v_op) / (k_init + annuity_factor(i, n) + k_op * v_op)
            c = rnd(c, dec)
        initial = c if alt["initial_installment"] == "one_installment" else float(alt["initial_installment"] or 0) * fvat
        option = c if alt["purchase_option"] == "one_installment" else float(alt["purchase_option"] or 0) * fvat
        if alt["given_installment"] is not None:
            impl_flow = [price - initial] + [-c] * n + [0.0] * (op_month - n)
            impl_flow[op_month] -= option
            tm = monthly_irr(impl_flow)
            if tm is None:
                raise InputError(tr(f"{alt['name']}: no pude despejar la tasa implícita del leasing.",
                                    f"{alt['name']}: could not solve the implicit leasing rate."))
            i = tm
        payments[0] -= initial
        add(asset_ded, 1, initial)
        balance, table = price - initial, []
        for m in range(1, n + 1):
            interest = rnd(balance * i, dec)
            principal = rnd(c - interest, dec)
            table.append(dict(month=m, opening_balance=balance, installment=c, interest=interest, principal=principal,
                              closing_balance=rnd(balance - principal, dec)))
            balance = table[-1]["closing_balance"]
            payments[m] -= c
            add(asset_ded, year_of_month(m), c)
        wait_months = op_month - n
        op_interest = rnd(balance * ((1 + i) ** wait_months - 1), dec)
        table.append(dict(month=op_month, opening_balance=balance, installment=option, interest=option - balance,
                          principal=balance, closing_balance=0.0, is_option=True,
                          rounding_gap=rnd(option - balance - op_interest, dec)))
        payments[op_month] -= option
        installment = c
        dep_base, purchase_year, purchase_month = option, year_of_month(op_month), op_month
        cae_flow = [price - initial] + [-c] * n + [0.0] * wait_months
        cae_flow[op_month] -= option
        tm = monthly_irr(cae_flow)
        extra.update(amount_financed=price - initial, initial_installment=initial, purchase_option=option,
                     option_month=op_month, implicit_monthly_rate=i,
                     cae=None if tm is None else tm * 12,
                     pre_tax_effective_annual_rate=None if tm is None else (1 + tm) ** 12 - 1)
    elif kind == "rental":
        n = alt["term_months"]
        rent = alt["monthly_rent"] * fvat
        deposit = alt["deposit"]
        payments[0] -= deposit
        for m in range(1, min(n, H) + 1):
            payments[m] -= rent
            add(asset_ded, year_of_month(m), rent)
        inflows[min(n, H)] += deposit
        installment = rent
        if n != H:
            extra["warning"] = tr(f"el arriendo dura {n} meses y el horizonte {H}: no cubre el mismo periodo",
                                  f"the rental lasts {n} months and the horizon {H}: it does not cover the same period")
    else:
        raise InputError(tr(f"Tipo de alternativa desconocido: {kind} (usa contado, credito, leasing, arriendo).",
                            f"Unknown alternative type: {kind} (use {', '.join(TYPES)})."))

    # Owner's depreciation and residual value (cash, loan and leasing with the option exercised)
    dep = depreciation(dep_base, P["dep_method"], P["life"], P["accel_life"], purchase_year) if dep_base else {}
    years_h = P["horizon"]
    for a, d in dep.items():
        if a <= years_h:
            add(asset_ded, a, d)
    if kind != "rental":
        acc_dep = sum(d for a, d in dep.items() if a <= years_h)
        book = dep_base - acc_dep
        mv = P["market_value"]
        sale_tax = t * (mv - book)
        residual[H] += mv - sale_tax
        extra.update(book_value_at_horizon=book, sale_tax=sale_tax,
                     depreciation_base=dep_base, depreciation_start_year=purchase_year)

    fin_shield = [0.0] * (H + 1)
    asset_shield = [0.0] * (H + 1)
    for a in range(1, years_h + 1):
        fin_shield[12 * a] += t * fin_ded.get(a, 0.0)
        asset_shield[12 * a] += t * asset_ded.get(a, 0.0)
    F = [payments[m] + inflows[m] + fin_shield[m] + asset_shield[m] + residual[m] for m in range(H + 1)]
    return dict(name=alt["name"], type=kind, F=F, payments=payments, inflows=inflows, fin_shield=fin_shield,
                asset_shield=asset_shield, residual=residual, fin_ded=fin_ded, asset_ded=asset_ded,
                table=table, installment=installment, extra=extra, price=price, notes=alt.get("notes", []),
                system=alt.get("system"))


# ----------------------------------------------------------------------------------------------
# Parameters
# ----------------------------------------------------------------------------------------------

def prepare(case: dict, case_path: Path, profile_name, rate_arg, tax_arg, project_arg, reg: Registry):
    case = normalize_case(case)
    p_path, how = find_profile(profile_name or case.get("profile"))
    profile = normalize_profile(read_simple_yaml(p_path))
    lang = set_language(profile)
    tpl = is_template(p_path)
    co = profile.get("company", {}) or {}
    fin = profile.get("finance", {}) or {}
    reg.add("profile", p_path.name, "", profile_origin(how), tpl,
            tr("la plantilla no describe a ninguna empresa: sus valores son supuestos",
               "the template describes no company: its values are assumptions") if tpl else "",
            "template" if tpl else "profile")
    pcode = "template" if tpl else "profile"
    from_profile_txt = tr(f"perfil {p_path.name}", f"profile {p_path.name}")

    def from_profile(key, section, arg, flag_es, flag_en, unit):
        v = section.get(key)
        if arg is not None:
            reg.add(key, arg, unit, tr(f"argumento {flag_es}", f"argument {flag_en}"), True,
                    tr("reemplaza el valor del perfil", "replaces the profile value") if v is not None
                    else tr("el perfil lo trae vacío", "the profile has it empty"))
            return arg
        if v is None:
            raise InputError(tr(f"El perfil {p_path.name} no trae «{loc_path(key)}». Pásalo con {flag_es} "
                                "(queda marcado SUPUESTO en la salida).",
                                f"The profile {p_path.name} has no «{key}». Pass it with {flag_en} "
                                "(it is marked ASSUMPTION in the output)."))
        reg.add(key, v, unit, from_profile_txt, tpl, code=pcode)
        return v

    rate = float(from_profile("discount_rate", fin, rate_arg, "--tasa", "--rate", tr("anual efectiva", "effective annual")))
    t = float(from_profile("income_tax", co, tax_arg, "--impuesto", "--tax", tr("fracción", "fraction")))
    currency = co.get("currency") or "CLP"
    dec = int(co.get("currency_decimals") if co.get("currency_decimals") is not None else 0)
    vat = float(co.get("vat") if co.get("vat") is not None else 0.0)
    country = str(co.get("country") or "").upper()
    reg.add("currency", currency, "", from_profile_txt, tpl, code=pcode if co.get("currency") else "default")

    def case_or_profile(d, key, path, unit, v_profile, default):
        if d.get(key) is not None or v_profile is None:
            return reg.read(d, key, path, unit, default)
        reg.add(f"{path}.{key}", v_profile, unit, from_profile_txt, tpl, code=pcode)
        return v_profile

    horizon = int(case_or_profile(case, "horizon_years", "case", tr("años", "years"), fin.get("horizon_years"), 5))
    vat_rec = bool(reg.read(case, "vat_recoverable", "case", tr("sí/no", "yes/no"), True,
                            tr("se supone que la empresa recupera el IVA como crédito fiscal",
                               "the company is assumed to recover the VAT as a tax credit")))
    if not vat_rec:
        reg.add("vat", vat, tr("fracción", "fraction"), from_profile_txt, tpl,
                tr("IVA no recuperable: se suma a precios, cuotas de leasing y rentas",
                   "VAT not recoverable: it is added to prices, leasing installments and rents"), pcode)
    a = case.get("asset") or {}
    price = float(reg.read(a, "net_price", "asset", currency))
    method = case_or_profile(a, "depreciation", "asset", "", fin.get("depreciation"), "straight_line")
    if method not in ("straight_line", "accelerated", "instant"):
        raise InputError(tr(f"Depreciación «{method}» no reconocida (lineal · acelerada · instantanea).",
                            f"Depreciation «{method}» not recognized (straight_line · accelerated · instant)."))
    life = reg.read(a, "useful_life_years", "asset", tr("años", "years")) if method == "straight_line" else a.get("useful_life_years")
    accel_life = reg.read(a, "accelerated_life_years", "asset", tr("años", "years")) if method == "accelerated" else None
    mv = float(reg.read(a, "market_value_at_horizon", "asset", currency))
    P = dict(months=12 * horizon, horizon=horizon, dec=dec, vat_factor=1.0 if vat_rec else 1 + vat,
             dep_method=method, life=life, accel_life=accel_life, market_value=mv, price=price,
             currency=currency, rate=rate, t=t, profile=p_path.name, description=a.get("description", tr("activo", "asset")),
             language=lang, country=country)

    alts = []
    names = set()
    for k, x in enumerate(case.get("alternatives") or []):
        name = x.get("name") or tr(f"alternativa {k + 1}", f"alternative {k + 1}")
        if name in names:
            raise InputError(tr(f"Hay dos alternativas llamadas «{name}».", f"There are two alternatives called «{name}»."))
        names.add(name)
        path = f"alternatives[{name}]"
        kind = x.get("type")
        if kind not in TYPES:
            raise InputError(tr(f"{loc_path(path)}: tipo «{kind}» no reconocido (usa contado, credito, leasing, arriendo).",
                                f"{path}: type «{kind}» not recognized (use {', '.join(TYPES)})."))
        alt = dict(name=name, type=kind, notes=list(x.get("non_comparable_notes") or []))
        # Own price only if the alternative changes it (e.g. a bonus for financing with the brand)
        alt["price"] = float(reg.read(x, "net_price", path, currency)) if "net_price" in x else price
        if kind == "loan":
            alt["system"] = reg.read(x, "system", path, "", "french", tr("cuota fija, la más común", "fixed installment, the most common"))
            if alt["system"] not in ("french", "german"):
                raise InputError(tr(f"{loc_path(path)}: sistema «{alt['system']}» (usa francesa o alemana).",
                                    f"{path}: system «{alt['system']}» (use french or german)."))
            if "down_payment_fraction" in x:
                alt["down_payment"] = rnd(alt["price"] * P["vat_factor"] *
                                          float(reg.read(x, "down_payment_fraction", path,
                                                         tr("fracción del precio", "fraction of the price"))), dec)
            else:
                alt["down_payment"] = float(reg.read(x, "down_payment", path, currency, 0.0, tr("sin pie", "no down payment")))
            alt["term_months"] = int(reg.read(x, "term_months", path, tr("meses", "months")))
            if "monthly_rate" in x:
                alt["monthly_rate"] = float(reg.read(x, "monthly_rate", path, tr("mensual", "monthly")))
            else:
                alt["monthly_rate"] = monthly_from_annual(float(reg.read(x, "annual_rate", path, tr("anual efectiva", "effective annual"))))
            alt["upfront_fees"] = float(reg.read(x, "upfront_fees", path, currency, 0.0, tr("sin gastos", "no fees")))
            alt["fees_financed"] = bool(reg.read(x, "fees_financed", path, tr("sí/no", "yes/no"), False,
                                                 tr("los gastos se pagan al contado", "the fees are paid in cash")))
            alt["monthly_insurance"] = float(reg.read(x, "monthly_insurance", path, currency, 0.0, tr("sin seguro", "no insurance")))
            lt = x.get("loan_tax") or {}
            alt["loan_tax_rate"] = float(reg.read(lt, "monthly_rate", path + ".loan_tax", tr("por mes de plazo", "per month of term"),
                                                  0.0, tr("sin impuesto al crédito", "no loan tax"))) if lt else 0.0
            alt["loan_tax_cap"] = float(reg.read(lt, "cap", path + ".loan_tax", tr("fracción del monto", "fraction of the amount"),
                                                 1.0, tr("sin tope", "no cap"))) if lt else 1.0
        elif kind == "leasing":
            alt["term_months"] = int(reg.read(x, "term_months", path, tr("meses", "months")))
            alt["given_installment"] = None
            if x.get("monthly_installment") is not None:
                alt["given_installment"] = float(reg.read(x, "monthly_installment", path, currency))
                alt["monthly_rate"] = None
            elif "monthly_rate" in x:
                alt["monthly_rate"] = float(reg.read(x, "monthly_rate", path, tr("mensual", "monthly")))
            else:
                alt["monthly_rate"] = monthly_from_annual(float(reg.read(x, "annual_rate", path, tr("anual efectiva", "effective annual"))))
            alt["initial_installment"] = reg.read(x, "initial_installment", path,
                                                  tr(f"{currency} o una_cuota", f"{currency} or one_installment"), "one_installment",
                                                  tr("lo usual en leasing: una cuota al firmar", "usual in leasing: one installment at signing"))
            alt["purchase_option"] = reg.read(x, "purchase_option", path,
                                              tr(f"{currency} o una_cuota", f"{currency} or one_installment"), "one_installment",
                                              tr("lo usual en leasing: opción igual a una cuota",
                                                 "usual in leasing: option equal to one installment"))
            alt["option_month"] = int(reg.read(x, "option_month", path, tr("mes", "month"), alt["term_months"] + 1,
                                               tr("un mes después de la última cuota", "one month after the last installment")))
        elif kind == "rental":
            alt["monthly_rent"] = float(reg.read(x, "monthly_rent", path, currency))
            alt["term_months"] = int(reg.read(x, "term_months", path, tr("meses", "months"), 12 * horizon,
                                              tr("todo el horizonte", "the whole horizon")))
            alt["deposit"] = float(reg.read(x, "deposit", path, currency, 0.0, tr("sin garantía", "no deposit")))
        alts.append(alt)
    if len(alts) < 2:
        raise InputError(tr("El caso necesita al menos dos alternativas para comparar.",
                            "The case needs at least two alternatives to compare."))
    if country and country != "CL" and any(x["type"] == "leasing" for x in alts):
        reg.warnings.append(tr(f"País {country}: el leasing se modela como gasto del arrendatario (regla chilena). Si en {country} "
                               "el arrendatario activa el bien, modélalo como crédito.",
                               f"Country {country}: leasing is modeled as a lessee expense (Chilean rule). If in {country} the "
                               "lessee capitalizes the asset, model it as a loan."))

    project_flow = None
    regime_check = None
    flow_src = project_arg or case.get("project_flow")
    if flow_src:
        rp = Path(flow_src)
        if not rp.is_absolute():
            rp = (Path.cwd() / rp) if project_arg else (case_path.parent / rp)
        yearly, regime = read_project_flow(rp)
        reg.add("project_flow", rp.name, tr("anual", "annual"),
                tr("argumento --proyecto", "argument --project") if project_arg else tr("caso: flujo_proyecto", "case: project_flow"),
                False, tr(f"{len(yearly)} periodos (años {min(yearly)} a {max(yearly)})",
                          f"{len(yearly)} periods (years {min(yearly)} to {max(yearly)})"))
        if max(yearly) > horizon:
            reg.warnings.append(tr(f"El flujo del proyecto llega al año {max(yearly)} y el horizonte es {horizon}: "
                                   "los años siguientes no se consideran.",
                                   f"The project flow reaches year {max(yearly)} and the horizon is {horizon}: "
                                   "later years are ignored."))
        if max(yearly) < horizon:
            reg.warnings.append(tr(f"El flujo del proyecto termina en el año {max(yearly)} y el horizonte es {horizon}: "
                                   "revisa que sean el mismo periodo.",
                                   f"The project flow ends in year {max(yearly)} and the horizon is {horizon}: "
                                   "check that they cover the same period."))
        f0 = yearly.get(0, 0.0)
        if abs(f0 + price * P["vat_factor"]) > 0.25 * price:
            reg.warnings.append(tr(f"El año 0 del proyecto ({f0:,.0f}) no se parece a −precio ({-price:,.0f}): "
                                   "revisa que el proyecto compre este mismo activo.",
                                   f"Year 0 of the project ({f0:,.0f}) does not look like −price ({-price:,.0f}): "
                                   "check that the project buys this same asset."))
        # Tax regime of the project flow against the one of the financing (audit finding of 15-10-2026)
        pt, pd = regime["tax"], regime["depreciation"]
        if pt is None and pd is None:
            reg.add("project_flow_regime", tr("no declarado", "not stated"), "", tr(f"archivo {rp.name}", f"file {rp.name}"), True,
                    tr("el archivo no dice con qué impuesto y depreciación se calculó: se supone el mismo régimen del financiamiento",
                       "the file does not say which tax and depreciation it used: the same regime as the financing is assumed"),
                    "script")
        else:
            tax_ok = pt is None or abs(float(pt) - t) <= 1e-9
            dep_ok = pd is None or pd == method
            regime_check = dict(project_tax=pt, project_depreciation=pd, financing_tax=t, financing_depreciation=method,
                                source=regime["source"], matches=tax_ok and dep_ok)
            if not (tax_ok and dep_ok):
                ptxt = f"{float(pt):.2%}" if pt is not None else "?"
                reg.warnings.append(tr(
                    f"Régimen tributario distinto: el flujo del proyecto ({rp.name}) se calculó con impuesto {ptxt} y "
                    f"depreciación {loc_value(pd) or '?'}, y el financiamiento con impuesto {t:.2%} y depreciación "
                    f"{loc_value(method)}. El VAN del inversionista mezcla los dos: decide por el costo en VP, o corre "
                    "economic-evaluation con el mismo régimen y encadena esa salida.",
                    f"Different tax regime: the project flow ({rp.name}) was computed with tax {ptxt} and {pd or '?'} "
                    f"depreciation, and the financing with tax {t:.2%} and {method} depreciation. The investor NPV mixes "
                    "both: decide by the PV cost, or rerun economic-evaluation with the same regime and chain that output."))
        project_flow = [0.0] * (P["months"] + 1)
        for yr, f in yearly.items():
            if 0 <= yr <= horizon:
                project_flow[12 * yr] += f
    else:
        reg.warnings.append(tr("Sin flujo del proyecto: se compara el costo de financiar; el VAN del inversionista no se calcula.",
                               "No project flow: only the cost of financing is compared; the investor NPV is not computed."))
    P["regime_check"] = regime_check
    P["project_assumptions"] = regime.get("assumptions", []) if flow_src else []
    P["project_file"] = rp.name if flow_src else ""
    return case, P, alts, project_flow


# ----------------------------------------------------------------------------------------------
# Full calculation
# ----------------------------------------------------------------------------------------------

def evaluate(P, alts, project_flow, t=None, rate=None, with_irr=True):
    t = P["t"] if t is None else t
    rate = P["rate"] if rate is None else rate
    ref_alt = dict(name=tr("contado de referencia", "reference cash purchase"), type="cash", price=P["price"], notes=[])
    ref = build_alternative(ref_alt, P, t)
    pv_ref = pv(ref["F"], rate)
    project_npv = pv(project_flow, rate) if project_flow else None
    res = []
    for a in alts:
        r = build_alternative(a, P, t)
        delta = [r["F"][m] - ref["F"][m] for m in range(P["months"] + 1)]
        r["delta"] = delta
        r["pv_cost"] = -pv(r["F"], rate)
        r["advantage_vs_cash"] = pv(r["F"], rate) - pv_ref
        if project_flow:
            r["investor"] = [project_flow[m] + delta[m] for m in range(P["months"] + 1)]
            r["investor_npv"] = pv(r["investor"], rate)
        roots = irr_roots(delta) if with_irr and any(abs(d) > 0.5 for d in delta) else []
        r["multiple_irr"] = len(roots) > 1
        tm = min(roots, key=lambda x: abs(x - monthly_from_annual(rate))) if roots else None
        r["after_tax_effective_rate"] = None if tm is None else (1 + tm) ** 12 - 1
        if r["multiple_irr"]:
            rates = ", ".join(f"{(1 + x) ** 12 - 1:.2%}" for x in roots)
            r["notes"] = list(r["notes"]) + [tr(
                f"el flujo frente al contado cambia de signo varias veces: la tasa efectiva no es única ({rates}); "
                "decide por el costo en VP",
                f"the flow against cash changes sign several times: the effective rate is not unique ({rates}); "
                "decide by the PV cost")]
        elif with_irr and tm is None and any(abs(d) > 0.5 for d in delta):
            r["notes"] = list(r["notes"]) + [tr(
                "sin tasa efectiva: el flujo frente al contado no cambia de signo (domina o es dominada a cualquier tasa); "
                "decide por el costo en VP",
                "no effective rate: the flow against cash does not change sign (it dominates or is dominated at any rate); "
                "decide by the PV cost")]
        res.append(r)
    order = sorted(res, key=lambda r: r["pv_cost"])
    return dict(ref=ref, res=res, order=order, project_npv=project_npv, rate=rate, t=t)


def indifference_rate(P, alts, project_flow, a_name, b_name):
    """Discount rate (annual) at which two alternatives cost the same in present value."""
    def diff(r):
        ev = evaluate(P, alts, project_flow, rate=r, with_irr=False)
        c = {x["name"]: x["pv_cost"] for x in ev["res"]}
        return c[a_name] - c[b_name]
    roots = []
    prev = (0.0, diff(0.0))
    for k in range(1, 121):
        r = k * 0.005
        v = diff(r)
        if (prev[1] < 0) != (v < 0):
            a, b, fa = prev[0], r, prev[1]
            for _ in range(40):
                mid = (a + b) / 2
                fm = diff(mid)
                if (fa < 0) == (fm < 0):
                    a, fa = mid, fm
                else:
                    b = mid
            roots.append((a + b) / 2)
        prev = (r, v)
    return roots


def sensitivity(P, alts, project_flow, extra_taxes):
    rates = sorted({round(max(0.0, P["rate"] + d), 4) for d in (-0.06, -0.03, 0.0, 0.03, 0.06)})
    taxes = []
    for x in [P["t"]] + list(extra_taxes or []):
        if round(float(x), 4) not in taxes:
            taxes.append(round(float(x), 4))
    rows = []
    for ti in taxes:
        for r in rates:
            ev = evaluate(P, alts, project_flow, t=ti, rate=r, with_irr=False)
            rows.append(dict(tax=ti, rate=r, costs={x["name"]: x["pv_cost"] for x in ev["res"]},
                             winner=ev["order"][0]["name"]))
    return rows


# ----------------------------------------------------------------------------------------------
# Outputs: result JSON and Excel
# ----------------------------------------------------------------------------------------------

def _assumption_rows(reg, only_marked=False):
    """Full traceability table (Excel sheet and JSON `traceability`)."""
    out = []
    for f in reg.rows:
        if only_marked and not f["assumption"]:
            continue
        out.append(dict(parameter=f["parameter"], key=f["key"], value=loc_value(f["value"]), unit=f["unit"],
                        origin=f["origin"], mark=mark(f["assumption"]), note=f["note"]))
    return out


def contract_assumptions(reg):
    """The common format of every method: [{parameter, value, origin, text}], only the marked rows."""
    out = []
    for f in reg.rows:
        if not f["assumption"]:
            continue
        v = loc_value(f["value"])
        body = f"{f['parameter']} = {v}" + (f" {f['unit']}" if f["unit"] else "") + f" · {f['origin']}" + \
            (f" · {f['note']}" if f["note"] else "")
        out.append({"parameter": f["parameter"], "value": v, "origin": f["code"], "text": f"{mark(True)} · {body}"})
    return out


def summary_json(P, ev, indif, sens, reg):
    g, s = ev["order"][0], ev["order"][1]
    out = dict(
        method="financial-evaluation", language=P["language"],
        currency=P["currency"], discount_rate=P["rate"], income_tax=P["t"], horizon_years=P["horizon"],
        project_npv=None if ev["project_npv"] is None else round(ev["project_npv"], 2),
        winner=g["name"], runner_up=s["name"], pv_difference=round(s["pv_cost"] - g["pv_cost"], 2),
        indifference_rate_with_runner_up=[round(x, 6) for x in indif],
        alternatives={},
        assumptions=contract_assumptions(reg),
        warnings=reg.warnings,
        project_flow_regime=P.get("regime_check"),
        project_flow_assumptions=P.get("project_assumptions") or [],
        traceability=_assumption_rows(reg),
        sensitivity=[dict(tax=f["tax"], rate=f["rate"], winner=f["winner"]) for f in sens],
    )
    for r in ev["res"]:
        e = r["extra"]
        out["alternatives"][r["name"]] = dict(
            type=r["type"],
            installment=r["installment"],
            pv_cost=round(r["pv_cost"], 2),
            advantage_vs_cash=round(r["advantage_vs_cash"], 2),
            over_winner=round(r["pv_cost"] - g["pv_cost"], 2),
            investor_npv=None if "investor_npv" not in r else round(r["investor_npv"], 2),
            after_tax_effective_rate=None if r["after_tax_effective_rate"] is None
            else round(r["after_tax_effective_rate"], 6),
            cae=None if e.get("cae") is None else round(e["cae"], 6),
            total_paid=round(-sum(r["payments"]), 2),
            total_interest=None if not r["table"] else round(sum(f["interest"] for f in r["table"]), 2),
            total_tax_savings=round(sum(r["fin_shield"]) + sum(r["asset_shield"]), 2),
            net_residual_value=round(sum(r["residual"]), 2),
            notes=list(r["notes"]) + ([e["warning"]] if e.get("warning") else []),
        )
    return out


def sheet_names():
    return [tr("SUPUESTOS", "ASSUMPTIONS"), tr("AMORTIZACION", "AMORTIZATION"),
            tr("FLUJO-INVERSIONISTA", "INVESTOR-CASH-FLOW"), tr("COMPARACION", "COMPARISON")]


def _type_label(r):
    return loc_value(r["type"]) + (f" ({loc_value(r['system'])})" if r["type"] == "loan" else "")


def write_excel(path: Path, P, ev, indif, sens, reg, title):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fmt_m = "#,##0" if P["dec"] == 0 else "#,##0.00"
    fmt_p = "0.00%"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF")
    sup_fill = PatternFill("solid", fgColor="FCE4B6")
    win_fill = PatternFill("solid", fgColor="D9EAD3")
    names = sheet_names()

    wb = Workbook()

    def header(ws, row, texts):
        for j, tx in enumerate(texts, 1):
            c = ws.cell(row=row, column=j, value=tx)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")

    def widths(ws, ws_widths):
        for j, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w

    # ASSUMPTIONS ----------------------------------------------------------------------------
    ws = wb.active
    ws.title = names[0]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = tr("Método de cálculo, no asesoría financiera. Montos en " + P["currency"] +
                  ", netos de IVA salvo que el caso diga lo contrario. Fila naranja = SUPUESTO.",
                  "Calculation method, not financial advice. Amounts in " + P["currency"] +
                  ", net of VAT unless the case says otherwise. Orange row = ASSUMPTION.")
    header(ws, 4, [tr("Parámetro", "Parameter"), tr("Valor", "Value"), tr("Unidad", "Unit"),
                   tr("De dónde vino", "Where it came from"), tr("Marca", "Mark"), tr("Fuente o nota", "Source or note")])
    f = 5
    for x in _assumption_rows(reg):
        vals = [x["parameter"], x["value"], x["unit"], x["origin"], x["mark"], x["note"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v if not isinstance(v, (list, dict)) else json.dumps(v))
            if x["mark"] == mark(True):
                c.fill = sup_fill
        f += 1
    rules = [
        (tr("Tributación", "Taxation"),
         tr("el ahorro tributario de cada año (intereses, depreciación, cuotas de leasing, renta) "
            "se recibe al cierre de ese año; se supone que la empresa tiene utilidades para usarlo",
            "each year's tax saving (interest, depreciation, leasing installments, rent) is received at the "
            "close of that year; the company is assumed to have profits to use it")),
        (tr("Depreciación", "Depreciation"),
         tr("año completo desde el año de compra; en leasing parte cuando se ejerce la opción",
            "full year from the purchase year; in leasing it starts when the option is exercised")),
        (tr("Valor residual", "Residual value"),
         tr("se vende al valor de mercado al horizonte y se paga impuesto sobre la diferencia con el valor libro",
            "sold at market value at the horizon, paying tax on the difference with the book value")),
        ("Leasing", tr("las cuotas y la cuota inicial son gasto tributario del arrendatario (regla chilena); "
                       "otro país puede tratarlo como activo",
                       "the installments and the initial installment are a tax expense of the lessee (Chilean rule); "
                       "another country may treat it as an asset")),
        (tr("Descuento", "Discounting"),
         tr("flujos mensuales descontados con la tasa anual efectiva: v = (1 + r)^(−mes/12)",
            "monthly flows discounted at the effective annual rate: v = (1 + r)^(−month/12)")),
        (tr("Flujo del inversionista", "Investor cash flow"),
         tr("flujo del proyecto + (flujo del activo con la alternativa − flujo del "
            "activo comprado al contado); así no se cuenta dos veces la compra",
            "project flow + (asset flow with the alternative − asset flow bought in cash); "
            "so the purchase is not counted twice")),
    ]
    f += 1
    ws.cell(row=f, column=1, value=tr("Reglas del modelo", "Model rules")).font = bold
    f += 1
    for a, b in rules:
        ws.cell(row=f, column=1, value=a)
        ws.cell(row=f, column=2, value=b)
        f += 1
    if reg.warnings:
        f += 1
        ws.cell(row=f, column=1, value=tr("Avisos", "Warnings")).font = bold
        f += 1
        for a in reg.warnings:
            ws.cell(row=f, column=1, value=a)
            f += 1
    if P.get("project_assumptions"):
        f += 1
        ws.cell(row=f, column=1, value=tr(f"Supuestos del flujo del proyecto (vienen de {P['project_file']})",
                                          f"Project flow assumptions (from {P['project_file']})")).font = bold
        f += 1
        for a in P["project_assumptions"]:
            ws.cell(row=f, column=1, value=a.get("text") or f"{a.get('parameter')} = {a.get('value')}")
            f += 1
    widths(ws, [42, 22, 18, 36, 12, 70])

    # AMORTIZATION ---------------------------------------------------------------------------
    ws = wb.create_sheet(names[1])
    f = 1
    for r in ev["res"]:
        if not r["table"]:
            continue
        e = r["extra"]
        ws.cell(row=f, column=1, value=f"{r['name']} · {_type_label(r)}").font = Font(bold=True, size=12)
        f += 1
        info = [(tr("Monto financiado", "Amount financed"), e.get("amount_financed"), fmt_m)]
        if r["type"] == "loan":
            info += [(tr("Gastos iniciales (incluye impuesto al crédito)", "Upfront fees (include the loan tax)"), e.get("upfront_fees"), fmt_m),
                     (tr("Impuesto al crédito", "Loan tax"), e.get("loan_tax"), fmt_m),
                     (tr("CAE (TIR mensual × 12, con gastos)", "CAE (monthly IRR × 12, with fees)"), e.get("cae"), fmt_p),
                     (tr("Costo efectivo anual antes de impuesto", "Pre-tax effective annual cost"),
                      e.get("pre_tax_effective_annual_rate"), fmt_p)]
        else:
            info += [(tr("Cuota inicial", "Initial installment"), e.get("initial_installment"), fmt_m),
                     (tr("Opción de compra", "Purchase option"), e.get("purchase_option"), fmt_m),
                     (tr("Tasa implícita mensual", "Implicit monthly rate"), e.get("implicit_monthly_rate"), "0.0000%"),
                     (tr("Tasa implícita anual (TIR mensual × 12)", "Implicit annual rate (monthly IRR × 12)"), e.get("cae"), fmt_p)]
        for label, v, fm in info:
            ws.cell(row=f, column=1, value=label)
            c = ws.cell(row=f, column=2, value=v)
            c.number_format = fm
            f += 1
        header(ws, f, [tr("Mes", "Month"), tr("Saldo inicial", "Opening balance"), tr("Cuota", "Installment"),
                       tr("Interés", "Interest"), tr("Amortización", "Principal"), tr("Saldo final", "Closing balance")])
        f += 1
        for x in r["table"]:
            vals = [x["month"], x["opening_balance"], x["installment"], x["interest"], x["principal"], x["closing_balance"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=f, column=j, value=v)
                if j > 1:
                    c.number_format = fmt_m
            if x.get("is_option"):
                ws.cell(row=f, column=7, value=tr("opción de compra", "purchase option"))
            f += 1
        tot = ["Total", None, sum(x["installment"] for x in r["table"]), sum(x["interest"] for x in r["table"]),
               sum(x["principal"] for x in r["table"]), None]
        for j, v in enumerate(tot, 1):
            c = ws.cell(row=f, column=j, value=v)
            c.font = bold
            if j > 1 and v is not None:
                c.number_format = fmt_m
        f += 1
        ties = abs(sum(x["principal"] for x in r["table"]) - e["amount_financed"]) < 0.5
        ws.cell(row=f, column=1, value=tr("Suma de amortizaciones = monto financiado", "Sum of principal = amount financed"))
        ws.cell(row=f, column=2, value="OK" if ties else tr("NO CUADRA", "DOES NOT TIE"))
        f += 3
    widths(ws, [44, 18, 16, 16, 16, 16, 18])

    # INVESTOR CASH FLOW ---------------------------------------------------------------------
    ws = wb.create_sheet(names[2])
    H, years = P["months"], P["horizon"]
    proj = ev["res"][0].get("investor") is not None

    def by_year(series):
        out = [0.0] * (years + 1)
        for m, v in enumerate(series):
            out[0 if m == 0 else year_of_month(m)] += v
        return out

    ws.cell(row=1, column=1, value=tr("Flujo del inversionista por alternativa (montos nominales por año)",
                                      "Investor cash flow by alternative (nominal amounts per year)")).font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=tr("Inversionista = proyecto + precio al contado que no se paga + pagos de la "
                                      "alternativa + su ahorro tributario + su valor residual − (ahorro tributario y "
                                      "residual del contado de referencia). El VAN se calcula mes a mes (tabla al final).",
                                      "Investor = project + cash price not paid + payments of the alternative + its tax "
                                      "saving + its residual value − (tax saving and residual of the reference cash "
                                      "purchase). NPV is computed month by month (table at the end)."))
    f = 4
    ref = ev["ref"]
    ref_sav = by_year([ref["fin_shield"][m] + ref["asset_shield"][m] for m in range(H + 1)])
    ref_res = by_year(ref["residual"])
    for r in ev["res"]:
        ws.cell(row=f, column=1, value=r["name"]).font = Font(bold=True, size=11)
        f += 1
        cols = [tr("Año", "Year"), tr("Flujo del proyecto", "Project flow"),
                tr("+ Precio al contado que no se paga", "+ Cash price not paid"),
                tr("− Pagos de la alternativa", "− Payments of the alternative"),
                tr("+ Devoluciones (garantía)", "+ Refunds (deposit)"),
                tr("+ Ahorro tributario intereses y gastos financieros", "+ Tax saving on interest and financial costs"),
                tr("+ Ahorro tributario depreciación / cuotas / renta", "+ Tax saving on depreciation / installments / rent"),
                tr("+ Valor residual neto", "+ Net residual value"),
                tr("− Ahorro tributario del contado", "− Tax saving of the cash purchase"),
                tr("− Residual del contado", "− Residual of the cash purchase"),
                tr("= Flujo del inversionista", "= Investor cash flow")]
        header(ws, f, cols)
        f += 1
        series = [
            by_year(ev["project_monthly"]) if proj else [None] * (years + 1),
            by_year([ref["price"] if m == 0 else 0.0 for m in range(H + 1)]),
            by_year(r["payments"]), by_year(r["inflows"]), by_year(r["fin_shield"]),
            by_year(r["asset_shield"]), by_year(r["residual"]),
            [-x for x in ref_sav], [-x for x in ref_res],
        ]
        for a in range(years + 1):
            vals = [s[a] for s in series]
            total = sum(v for v in vals if v is not None)
            row = [a] + vals + [total if proj else None]
            for j, v in enumerate(row, 1):
                c = ws.cell(row=f, column=j, value=v)
                if j > 1:
                    c.number_format = fmt_m
            f += 1
        if proj:
            ws.cell(row=f, column=1, value=tr("VAN del inversionista", "Investor NPV")).font = bold
            c = ws.cell(row=f, column=11, value=r["investor_npv"])
            c.number_format, c.font = fmt_m, bold
            f += 1
        f += 2
    if proj:
        ws.cell(row=f, column=1, value=tr("Tabla mensual (con esta se calcula el VAN)",
                                          "Monthly table (NPV is computed with it)")).font = Font(bold=True, size=11)
        f += 1
        cab = [tr("Mes", "Month"), tr("Factor de descuento", "Discount factor"), tr("Flujo del proyecto", "Project flow")] + \
              [tr("Inversionista · ", "Investor · ") + r["name"] for r in ev["res"]]
        header(ws, f, cab)
        f += 1
        for m in range(H + 1):
            fac = (1 + ev["rate"]) ** (-m / 12)
            vals = [m, fac, ev["project_monthly"][m]] + [r["investor"][m] for r in ev["res"]]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=f, column=j, value=v)
                c.number_format = "0.000000" if j == 2 else (fmt_m if j > 2 else "0")
            f += 1
        vals = [tr("VAN", "NPV"), None, ev["project_npv"]] + [r["investor_npv"] for r in ev["res"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v)
            c.font = bold
            if j > 2:
                c.number_format = fmt_m
    widths(ws, [24, 16] + [18] * 9)

    # COMPARISON -------------------------------------------------------------------------------
    ws = wb.create_sheet(names[3])
    g, s = ev["order"][0], ev["order"][1]
    ws["A1"] = tr("Comparación de alternativas de financiamiento", "Comparison of financing alternatives")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = tr(f"Tasa de descuento {ev['rate']:.2%} anual · impuesto {ev['t']:.2%} · horizonte {years} años · {P['description']}",
                  f"Discount rate {ev['rate']:.2%} annual · tax {ev['t']:.2%} · horizon {years} years · {P['description']}")
    cab = [tr("Posición", "Rank"), tr("Alternativa", "Alternative"), tr("Tipo", "Type"), tr("Cuota", "Installment"),
           tr("Total pagado (nominal)", "Total paid (nominal)"), tr("Ahorro tributario total", "Total tax saving"),
           tr("Valor residual neto", "Net residual value"),
           tr("Costo en valor presente después de impuesto", "After-tax present value cost"),
           tr("Ventaja frente al contado (VP)", "Advantage over cash (PV)"),
           tr("Costo efectivo anual después de impuesto", "After-tax effective annual cost"),
           tr("CAE o tasa implícita anual", "CAE or implicit annual rate"), tr("VAN del inversionista", "Investor NPV"),
           tr("Cuánto más cuesta que la ganadora (VP)", "Extra cost over the winner (PV)")]
    header(ws, 4, cab)
    f = 5
    for k, r in enumerate(ev["order"], 1):
        e = r["extra"]
        vals = [k, r["name"], _type_label(r),
                r["installment"], -sum(r["payments"]), sum(r["fin_shield"]) + sum(r["asset_shield"]), sum(r["residual"]),
                r["pv_cost"], r["advantage_vs_cash"], r["after_tax_effective_rate"], e.get("cae"),
                r.get("investor_npv"), r["pv_cost"] - g["pv_cost"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v)
            if j in (4, 5, 6, 7, 8, 9, 12, 13):
                c.number_format = fmt_m
            if j in (10, 11):
                c.number_format = fmt_p
            if k == 1:
                c.fill = win_fill
        f += 1
    f += 1
    ws.cell(row=f, column=1, value=tr("Ganadora", "Winner")).font = bold
    ws.cell(row=f, column=2, value=tr(f"{g['name']}: cuesta {s['pv_cost'] - g['pv_cost']:,.0f} {P['currency']} "
                                      f"menos en valor presente que {s['name']} (segunda).",
                                      f"{g['name']}: costs {s['pv_cost'] - g['pv_cost']:,.0f} {P['currency']} "
                                      f"less in present value than {s['name']} (runner-up)."))
    f += 1
    ws.cell(row=f, column=1, value=tr("Tasa que cambia la decisión", "Rate that flips the decision")).font = bold
    ws.cell(row=f, column=2, value=("; ".join(f"{x:.2%}" for x in indif) +
                                    tr(f" anual: a esa tasa de descuento {g['name']} y {s['name']} empatan",
                                       f" annual: at that discount rate {g['name']} and {s['name']} tie"))
            if indif else tr(f"ninguna entre 0 % y 60 %: {g['name']} gana a {s['name']} en todo ese rango",
                             f"none between 0 % and 60 %: {g['name']} beats {s['name']} over that whole range"))
    f += 1
    ws.cell(row=f, column=1, value=tr("Cómo leer el costo efectivo", "How to read the effective cost")).font = bold
    ws.cell(row=f, column=2, value=tr("Costo efectivo anual después de impuesto = TIR de (flujo de la alternativa − "
                                      "flujo al contado). Si es menor que la tasa de descuento, financiar así conviene "
                                      "más que pagar al contado.",
                                      "After-tax effective annual cost = IRR of (flow of the alternative − cash flow). "
                                      "If it is below the discount rate, financing this way beats paying cash."))
    notes = [(r["name"], n) for r in ev["res"] for n in r["notes"]] + \
            [(r["name"], r["extra"]["warning"]) for r in ev["res"] if r["extra"].get("warning")]
    if notes:
        f += 2
        ws.cell(row=f, column=1, value=tr("No comparable (se declara, no se promedia)", "Not comparable (declared, not averaged)")).font = bold
        f += 1
        for a, n in notes:
            ws.cell(row=f, column=1, value=a)
            ws.cell(row=f, column=2, value=n)
            f += 1
    f += 2
    ws.cell(row=f, column=1, value=tr("Sensibilidad: costo en VP después de impuesto", "Sensitivity: after-tax PV cost")).font = \
        Font(bold=True, size=11)
    f += 1
    alt_names = [r["name"] for r in ev["res"]]
    header(ws, f, [tr("Impuesto", "Tax"), tr("Tasa de descuento", "Discount rate")] + alt_names + [tr("Ganadora", "Winner")])
    f += 1
    for x in sens:
        vals = [x["tax"], x["rate"]] + [x["costs"][n] for n in alt_names] + [x["winner"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=f, column=j, value=v)
            c.number_format = fmt_p if j <= 2 else fmt_m
            if x["tax"] == round(ev["t"], 4) and x["rate"] == round(ev["rate"], 4):
                c.font = bold
        f += 1
    widths(ws, [22, 30, 18, 14, 18, 18, 18, 22, 20, 20, 18, 20, 20])
    wb.save(path)


# ----------------------------------------------------------------------------------------------
# Orchestration
# ----------------------------------------------------------------------------------------------

def output_names(base: str) -> tuple[str, str]:
    return (tr(f"{base}-resultado.json", f"{base}-result.json"),
            tr(f"{base}-evaluacion-financiera.xlsx", f"{base}-financial-evaluation.xlsx"))


def run(case_path: Path, profile=None, rate=None, tax=None, project=None, out=None, quiet=False, case_data=None):
    raw = case_data if case_data is not None else json.loads(case_path.read_text(encoding="utf-8-sig"))
    reg = Registry()
    case, P, alts, project_flow = prepare(raw, case_path, profile, rate, tax, project, reg)
    ev = evaluate(P, alts, project_flow)
    ev["project_monthly"] = project_flow
    g, s = ev["order"][0], ev["order"][1]
    indif = indifference_rate(P, alts, project_flow, g["name"], s["name"])
    sens = sensitivity(P, alts, project_flow, (case.get("sensitivity") or {}).get("tax_rates"))
    result = summary_json(P, ev, indif, sens, reg)
    folder = Path(out) if out else case_path.parent / tr("salidas", "outputs")
    folder.mkdir(parents=True, exist_ok=True)
    base = case.get("case") or case_path.stem
    json_name, xlsx_name = output_names(base)
    (folder / json_name).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    title = tr("Evaluación financiera · ", "Financial evaluation · ") + str(case.get("case", case_path.stem))
    write_excel(folder / xlsx_name, P, ev, indif, sens, reg, title)
    if not quiet:
        print_summary(P, ev, indif, reg, folder / xlsx_name)
    return result, ev, P, alts, reg


def print_summary(P, ev, indif, reg, xlsx_path):
    cur = P["currency"]
    print(tr(f"Tasa de descuento {ev['rate']:.2%} · impuesto {ev['t']:.2%} · horizonte {P['horizon']} años",
             f"Discount rate {ev['rate']:.2%} · tax {ev['t']:.2%} · horizon {P['horizon']} years"))
    sup = [f["parameter"] for f in reg.rows if f["assumption"]]
    if sup:
        print(mark(True) + ": " + ", ".join(sup))
    for a in reg.warnings:
        print(tr("AVISO: ", "WARNING: ") + a)
    if ev["project_npv"] is not None:
        print(tr(f"VAN del proyecto puro: {ev['project_npv']:,.0f} {cur}", f"Pure project NPV: {ev['project_npv']:,.0f} {cur}"))
    print(tr(f"{'Alternativa':34} {'Costo VP d/i':>16} {'Costo ef. d/i':>13} {'CAE/impl.':>10} {'VAN inversionista':>18}",
             f"{'Alternative':34} {'PV cost a/t':>16} {'Eff. cost a/t':>13} {'CAE/impl.':>10} {'Investor NPV':>18}"))
    for r in ev["order"]:
        te = r["after_tax_effective_rate"]
        cae = r["extra"].get("cae")
        npv = r.get("investor_npv")
        print(f"{r['name'][:34]:34} {r['pv_cost']:>16,.0f} {('—' if te is None else f'{te:.2%}'):>13} "
              f"{('—' if cae is None else f'{cae:.2%}'):>10} {('—' if npv is None else f'{npv:,.0f}'):>18}")
    g, s = ev["order"][0], ev["order"][1]
    print(tr(f"GANA {g['name']}: {s['pv_cost'] - g['pv_cost']:,.0f} {cur} menos en VP que {s['name']}.",
             f"WINNER {g['name']}: {s['pv_cost'] - g['pv_cost']:,.0f} {cur} less in PV than {s['name']}."))
    print(tr("Tasa de indiferencia con la segunda: ", "Indifference rate with the runner-up: ")
          + (", ".join(f"{x:.2%}" for x in indif) if indif else tr("ninguna en 0-60 %", "none in 0-60 %")))
    print(tr("Salida: ", "Output: ") + str(xlsx_path))


# ----------------------------------------------------------------------------------------------
# --test: the sample case, against expected.json and a second path
# ----------------------------------------------------------------------------------------------

def _to_spanish(obj):
    """English case → Spanish keys and codes (only to test the alias map)."""
    if isinstance(obj, list):
        return [_to_spanish(x) for x in obj]
    if not isinstance(obj, dict):
        return obj
    out = {}
    for k, v in obj.items():
        nk = {vv: kk for kk, vv in CASE_KEYS.items()}.get(k, k)
        if k in VALUE_KEYS and isinstance(v, str):
            v = ES_VALUE.get(v, v)
        elif k in VALUE_KEYS and isinstance(v, dict):
            v = _to_spanish(v)
            v["valor"] = ES_VALUE.get(v.get("valor"), v.get("valor"))
        elif isinstance(v, (dict, list)) and k != "non_comparable_notes":
            v = _to_spanish(v)
        out[nk] = v
    return out


def run_test() -> int:
    global _LANG
    tdir = SKILL_ROOT / "tests"
    case = tdir / "sample-case.json"
    case_data = json.loads(case.read_text(encoding="utf-8"))
    expected = json.loads((tdir / "expected.json").read_text(encoding="utf-8"))
    profile = str(tdir / "sample-profile.yaml")
    failures: list[str] = []

    def check(cond, msg):
        print(("  ok   " if cond else "  FAIL ") + msg)
        if not cond:
            failures.append(msg)

    with tempfile.TemporaryDirectory() as tmp:
        out, ev, P, alts, reg = run(case, profile=profile, out=tmp, quiet=True)
        print("1. Against expected.json (tolerance 1 peso; rates 1e-6)")
        for name, exp in expected["alternatives"].items():
            got = out["alternatives"].get(name)
            if got is None:
                check(False, f"alternative {name} is missing")
                continue
            for k, v in exp.items():
                g = got.get(k)
                if isinstance(v, float) and abs(v) < 5:
                    check(g is not None and abs(g - v) <= 1e-6, f"{name} · {k}: {g} vs {v}")
                elif isinstance(v, (int, float)):
                    check(g is not None and abs(g - v) <= 1.0, f"{name} · {k}: {g} vs {v}")
                else:
                    check(g == v, f"{name} · {k}: {g} vs {v}")
        check(out["winner"] == expected["winner"], f"winner {out['winner']} vs {expected['winner']}")
        check(abs(out["pv_difference"] - expected["pv_difference"]) <= 1.0,
              f"difference {out['pv_difference']} vs {expected['pv_difference']}")
        check(out["project_npv"] is not None and abs(out["project_npv"] - expected["project_npv"]) <= 1.0,
              f"project NPV {out['project_npv']} vs {expected['project_npv']}")

        print("2. Second path: annuity formula against the summed table, to the peso")
        for r in ev["res"]:
            if not r["table"]:
                continue
            e, tab = r["extra"], r["table"]
            sum_pr = sum(x["principal"] for x in tab)
            check(abs(sum_pr - e["amount_financed"]) < 0.5,
                  f"{r['name']}: sum of principal {sum_pr:,.0f} = amount {e['amount_financed']:,.0f}")
            check(abs(sum(x["installment"] for x in tab) - sum(x["interest"] for x in tab) - sum_pr) < 0.5,
                  f"{r['name']}: installments − interest = principal")
            if r["type"] == "loan" and r["system"] == "french":
                alt = next(a for a in alts if a["name"] == r["name"])
                i, n, L = alt["monthly_rate"], alt["term_months"], e["amount_financed"]
                c_formula = L * i * (1 + i) ** n / ((1 + i) ** n - 1)
                check(abs(round(c_formula) - tab[0]["installment"]) <= 0,
                      f"{r['name']}: formula installment {c_formula:,.2f} → {round(c_formula):,} = table {tab[0]['installment']:,.0f}")
                check(abs(tab[-1]["installment"] - tab[0]["installment"]) <= n,
                      f"{r['name']}: the last installment absorbs the rounding ({tab[-1]['installment'] - tab[0]['installment']:+,.0f})")
                pv_inst = sum(x["installment"] / (1 + i) ** x["month"] for x in tab)
                check(abs(pv_inst - L) <= 0.5 * n,
                      f"{r['name']}: PV of the installments at the rate = amount ({pv_inst - L:+.2f}; "
                      f"tolerance half a peso per rounded installment)")
            if r["type"] == "loan" and r["system"] == "german":
                alt = next(a for a in alts if a["name"] == r["name"])
                i, n, L = alt["monthly_rate"], alt["term_months"], e["amount_financed"]
                int_formula = i * L * (n + 1) / 2
                check(abs(int_formula - sum(x["interest"] for x in tab)) <= n,
                      f"{r['name']}: interest by formula i·L·(n+1)/2 = {int_formula:,.0f} ≈ table")
            if r["type"] == "leasing":
                check(abs(tab[-1].get("rounding_gap", 0)) <= len(tab),
                      f"{r['name']}: the option closes the balance (gap {tab[-1].get('rounding_gap', 0):+,.0f})")
        print("3. Model identities")
        for r in ev["res"]:
            npv_id = ev["project_npv"] + r["advantage_vs_cash"]
            check(abs(npv_id - r["investor_npv"]) <= 1.0,
                  f"{r['name']}: investor NPV = project NPV + advantage over cash")
        cash = next((r for r in ev["res"] if r["type"] == "cash"), None)
        if cash:
            check(abs(cash["investor_npv"] - ev["project_npv"]) <= 1.0, "cash: investor NPV = pure project NPV")
        from openpyxl import load_workbook
        wb = load_workbook(Path(tmp) / output_names(case_data["case"])[1])
        check(wb.sheetnames == ["SUPUESTOS", "AMORTIZACION", "FLUJO-INVERSIONISTA", "COMPARACION"],
              f"Excel sheets (es): {wb.sheetnames}")
        check(any(x["key"] == "project_flow_regime" and x["mark"] == "SUPUESTO" for x in out["traceability"]),
              "a CSV flow without tax regime: the regime is marked SUPUESTO")
        valid = ("input", "profile", "template", "default", "script")
        check(all(set(a) == {"parameter", "value", "origin", "text"} and (a["origin"] in valid or a["origin"].startswith("research:"))
                  and a["text"].startswith("SUPUESTO · ") for a in out["assumptions"]),
              "JSON assumptions are [{parameter, value, origin, text}] with a known origin")

        print("4. Profile without a rate: the script requires it and marks it as an assumption")
        tpl = template_profile()
        try:
            run(case, profile=str(tpl), out=tmp, quiet=True)
            check(False, "with the template and no --rate it had to stop")
        except InputError as ex:
            check("--tasa" in str(ex) or "--rate" in str(ex), "without --rate it stops and asks for the argument")
        o2, *_ = run(case, profile=str(tpl), rate=0.12, out=tmp, quiet=True)
        check(any(s["key"] == "discount_rate" and s["mark"] in ("SUPUESTO", "ASSUMPTION") and "--" in s["origin"]
                  for s in o2["traceability"]) and any(a["origin"] == "template" for a in o2["assumptions"]),
              "with --rate the rate is marked as an assumption; template values have origin template")

        print("5. Project flow from economic-evaluation (new English output and old Spanish salida.json)")
        flow = {0: -24000000, 1: 6800000, 2: 6800000, 3: 6800000, 4: 6800000, 5: 15221429}
        new_json = {"method": "economic-evaluation",
                    "assumptions": [{"parameter": "Occupancy", "value": 0.55, "origin": "input",
                                     "text": "ASSUMPTION · Occupancy = 0.55"}],
                    "parameters": {"tax": {"value": 0.25}, "depreciation": {"value": "straight_line"}},
                    "cash_flow": {"years": list(flow), "net_cash_flow": list(flow.values())}}
        old_json = {"parametros": {"impuesto": {"valor": 0.27}, "depreciacion": {"valor": "lineal"}},
                    "supuestos": [{"parametro": "Ocupación", "valor": 0.55, "unidad": "", "origen": "entrada · ingresos[0]",
                                   "marca": "SUPUESTO"},
                                  {"parametro": "Precio", "valor": 45000, "unidad": "", "origen": "entrada", "marca": "dato"}],
                    "flujo": {"anios": list(flow), "flujo_de_caja": list(flow.values())}}
        (Path(tmp) / "eco-new.json").write_text(json.dumps(new_json), encoding="utf-8")
        (Path(tmp) / "eco-old.json").write_text(json.dumps(old_json), encoding="utf-8")
        o3, *_ = run(case, profile=profile, project=str(Path(tmp) / "eco-new.json"), out=tmp, quiet=True)
        check(abs(o3["project_npv"] - out["project_npv"]) <= 1e-6, "new English output: same project NPV as the CSV")
        check(o3["project_flow_regime"]["matches"] is False and any("25.00%" in w for w in o3["warnings"]),
              "tax 25 % in the project and 27 % in the financing: warning in the JSON")
        o4, *_ = run(case, profile=profile, project=str(Path(tmp) / "eco-old.json"), out=tmp, quiet=True)
        check(abs(o4["project_npv"] - out["project_npv"]) <= 1e-6 and o4["project_flow_regime"]["matches"] is True
              and not o4["warnings"], "old Spanish salida.json with the same regime: read, no warning")
        check(o3["project_flow_assumptions"][0]["parameter"] == "Occupancy"
              and [a["parameter"] for a in o4["project_flow_assumptions"]] == ["Ocupación"]
              and o4["project_flow_assumptions"][0]["origin"] == "input",
              "project flow assumptions carried over (new `assumptions` and old `supuestos`)")

        print("6. Spanish keys and language en")
        o5, *_ = run(case, profile=profile, out=tmp, quiet=True, case_data=_to_spanish(case_data))
        same = all(o5["alternatives"][n][k] == out["alternatives"][n][k] for n in out["alternatives"]
                   for k in ("pv_cost", "investor_npv", "installment", "cae", "after_tax_effective_rate"))
        check(same and o5["winner"] == out["winner"], "a case with Spanish keys gives the same figures")
        en_profile = Path(tmp) / "sample-profile-en.yaml"
        en_profile.write_text("language: en\n" + (tdir / "sample-profile.yaml").read_text(encoding="utf-8"), encoding="utf-8")
        o6, *_ = run(case, profile=str(en_profile), out=tmp, quiet=True)
        wb = load_workbook(Path(tmp) / output_names(case_data["case"])[1])
        check(wb.sheetnames == ["ASSUMPTIONS", "AMORTIZATION", "INVESTOR-CASH-FLOW", "COMPARISON"],
              f"Excel sheets (en): {wb.sheetnames}")
        check(all(x["text"].startswith("ASSUMPTION · ") for x in o6["assumptions"]) and o6["pv_difference"] == out["pv_difference"],
              "en: marker ASSUMPTION, same figures")

    print("7. Profile lookup: $METHODIC_DATA · legacy $INGENIERO_DATOS · ../methodic-data · legacy sibling")
    saved_env = {v: os.environ.pop(v) for v in ENV_DATA if v in os.environ}
    saved_root = REPO_ROOT
    with tempfile.TemporaryDirectory() as tmp:
        try:
            globals()["REPO_ROOT"] = Path(tmp) / "lookup" / "repo"
            for folder, var, how in (("ingeniero-agentico-datos", None, "sibling:ingeniero-agentico-datos"),
                                     ("methodic-data", None, "sibling:methodic-data"),
                                     ("env-old", "INGENIERO_DATOS", "env:INGENIERO_DATOS"),
                                     ("env-new", "METHODIC_DATA", "env:METHODIC_DATA")):
                d = Path(tmp) / "lookup" / folder / "profiles"
                d.mkdir(parents=True)
                (d / "co.yaml").write_text("language: es\n", encoding="utf-8")
                if var:
                    os.environ[var] = str(Path(tmp) / "lookup" / folder)
                got = find_profile("co")
                check(got == (d / "co.yaml", how), f"profile «co» found through {how} ({profile_origin(got[1])})")
        finally:
            globals()["REPO_ROOT"] = saved_root
            for v in ENV_DATA:
                os.environ.pop(v, None)
            os.environ.update(saved_env)
    _LANG = "es"
    print("OK" if not failures else f"FAIL ({len(failures)})")
    return 0 if not failures else 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Financial evaluation: cash, loan, leasing or rental.")
    ap.add_argument("--case", dest="case", help="case JSON (see templates/case.json); Spanish keys are accepted")
    ap.add_argument("--project", dest="project", help="pure project flow: economic-evaluation output JSON, or CSV/JSON with year and flow")
    ap.add_argument("--profile", dest="profile", help="profile name (looked up in METHODIC_DATA and ../methodic-data; the legacy names are still read) or path")
    ap.add_argument("--rate", dest="rate", type=float, help="annual discount rate (required if the profile does not have it)")
    ap.add_argument("--tax", dest="tax", type=float, help="income tax rate (required if the profile does not have it)")
    ap.add_argument("--out", dest="out", help="output folder (default <case folder>/salidas, or /outputs with language en)")
    ap.add_argument("--test", dest="test", action="store_true", help="run the sample case and check that it ties")
    # Spanish flags: hidden aliases (same dest).
    ap.add_argument("--caso", dest="case", help=argparse.SUPPRESS)
    ap.add_argument("--proyecto", dest="project", help=argparse.SUPPRESS)
    ap.add_argument("--perfil", dest="profile", help=argparse.SUPPRESS)
    ap.add_argument("--tasa", dest="rate", type=float, help=argparse.SUPPRESS)
    ap.add_argument("--impuesto", dest="tax", type=float, help=argparse.SUPPRESS)
    ap.add_argument("--salida", dest="out", help=argparse.SUPPRESS)
    ap.add_argument("--probar", dest="test", action="store_true", help=argparse.SUPPRESS)
    a = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if a.test:
        return run_test()
    if not a.case:
        ap.error("missing --case (or use --test)")
    try:
        run(Path(a.case), a.profile, a.rate, a.tax, a.project, a.out)
    except InputError as ex:
        print(f"ERROR: {ex}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
