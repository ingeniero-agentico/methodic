# -*- coding: utf-8 -*-
"""
map_context · stage 0 (map the context) of ingeniero-agentico.

    python map_context.py <folder> [<folder> ...] --out <context/>
                          [--exclude PATTERN ...] [--threshold 0.20] [--max-words 3000]
                          [--max-mb 25] [--core path1,path2] [--graphify]
    python map_context.py --out <context/> --mark-read
    python map_context.py --test

Walks the folders of a task and leaves in <out>:

- graph.json        nodes (files) and links between them, in graphify's shape
                    (nodes: [{id, label, type, ...}], links: [{source, target, relation}]).
- fingerprints.json path -> size, date and short sha1 of each file (what decides the mode).
- changes.json      the changes not yet read, accumulated run by run until --mark-read.
- MAP.md            what is there, where and from when, in <= 3,000 words, with the MODE on top.

No LLM: the one-line summary comes from the file's structure (frontmatter, title, docstring,
keys, sheets, metadata). Same input, same output.

The MODE is decided against the previous fingerprints and the changes not yet read:
- no fingerprints (or no map, or other roots) -> first-run;
- accumulated changes (new + changed + deleted) <= threshold -> reread, with the list;
- more -> first-run.

Contexts written before v1.1 (Spanish names: contexto/, MAPA.md, grafo.json, huellas.json,
cambios.json, Spanish keys) are read as they are and keep their file names; old flags
(--salida, --nucleo, --marcar-leido, --probar...) still work as hidden aliases.

Standard library only; openpyxl (.xlsx sheets) and pypdf (.pdf title) if installed.
"""
import argparse
import ast
import datetime as dt
import fnmatch
import hashlib
import json
import os
import posixpath
import re
import shutil
import struct
import sys
import tempfile
import time
import urllib.parse
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve()
SKILL = HERE.parents[1]
VERSION = 2          # 1 = Spanish keys (before v1.1); still read

# File names of a context folder. A context keeps the layout it was created with.
LAYOUTS = {
    "en": {"map": "MAP.md", "graph": "graph.json", "fingerprints": "fingerprints.json",
           "changes": "changes.json"},
    "es": {"map": "MAPA.md", "graph": "grafo.json", "fingerprints": "huellas.json",
           "changes": "cambios.json"},
}
MODES = ("first-run", "reread")
OLD_MODES = {"primera-vez": "first-run", "relectura": "reread"}

EXCLUDE_DIRS = {
    ".git", "node_modules", "__pycache__", ".venv", "venv", ".mypy_cache", ".pytest_cache",
    ".ruff_cache", ".obsidian", ".trash", "graphify-out", ".idea", ".vscode",
}
EXCLUDE_FILES = {"thumbs.db", "desktop.ini", ".ds_store"}
EXCLUDE_PREFIXES = ("~$", ".~lock")

KIND_BY_EXT = {
    ".md": "markdown", ".markdown": "markdown", ".txt": "text",
    ".json": "json", ".yaml": "yaml", ".yml": "yaml", ".toml": "config", ".ini": "config",
    ".cfg": "config", ".py": "python", ".ps1": "script", ".sh": "script", ".bat": "script",
    ".cmd": "script", ".js": "script", ".ts": "script", ".sql": "script",
    ".html": "web", ".htm": "web", ".css": "web", ".xml": "data",
    ".csv": "table", ".tsv": "table", ".xlsx": "excel", ".xlsm": "excel", ".xls": "excel",
    ".docx": "word", ".doc": "word", ".pptx": "presentation", ".pdf": "pdf",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".gif": "image", ".webp": "image",
    ".svg": "image", ".bmp": "image", ".zip": "archive", ".base": "obsidian",
    ".canvas": "obsidian", ".dxf": "cad", ".dwg": "cad", ".step": "cad", ".stp": "cad",
}
TEXT_EXT = {
    ".md", ".markdown", ".txt", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".py", ".ps1",
    ".sh", ".bat", ".cmd", ".js", ".ts", ".sql", ".html", ".htm", ".css", ".xml", ".csv", ".tsv",
    ".base", ".canvas", ".svg",
}
# Files searched for links (the rest are only summarized).
LINK_EXT = TEXT_EXT - {".csv", ".tsv", ".svg", ".css"}
TEXT_LIMIT = 2_000_000            # bytes read from a text file
SUMMARY_WORDS = 30                # cap of the summary kept in the graph
MD_WORDS = 60                     # a .md also carries its sections
DOC = {"markdown", "text", "word", "pdf", "excel", "presentation"}
KEY_NAMES = {"leeme", "readme", "skill", "claude", "indice", "index", "mapa", "map", "arquitectura",
             "architecture", "guia", "guide", "resultado", "result", "agents", "prd"}

CITABLE_EXT = ("md|markdown|txt|json|ya?ml|toml|py|ps1|sh|bat|js|ts|sql|html?|css|csv|tsv|"
               "xlsx|xlsm|xls|docx|pptx|pdf|png|jpe?g|gif|webp|svg|zip|base|canvas|dxf|dwg|step|stp")
RE_URL = re.compile(r"(?:https?|ftp|mailto):\S+", re.I)
RE_WIKI = re.compile(r"!?\[\[([^\]\|#\n]+)(?:#[^\]\|\n]*)?(?:\|[^\]\n]*)?\]\]")
RE_MDLINK = re.compile(r"!?\[[^\]\n]*\]\(\s*<?([^)\s>]+)>?(?:\s+\"[^\"]*\")?\s*\)")
RE_BACKTICK = re.compile(r"`([^`\n]{1,240})`")
RE_PATH = re.compile(r"(?<![\w/\\.:])((?:[A-Za-z]:[\\/])?[\w\-.~/\\]*[\w\-]\.(?:" + CITABLE_EXT
                     + r"))(?![\w])", re.I)


class InputError(Exception):
    """An error that stops the mapping. Fix the call and run again."""


# ------------------------------------------------------------------ utilities
def _long(p):
    """Absolute path with the \\\\?\\ prefix on Windows (paths over 260 characters "do not exist")."""
    p = os.path.abspath(p)
    if os.name == "nt" and not p.startswith("\\\\?\\"):
        if p.startswith("\\\\"):
            return "\\\\?\\UNC\\" + p[2:]
        return "\\\\?\\" + p
    return p


def _short(p):
    if p.startswith("\\\\?\\UNC\\"):
        return "\\\\" + p[8:]
    if p.startswith("\\\\?\\"):
        return p[4:]
    return p


def _posix(p):
    return p.replace("\\", "/")


def _words(text):
    return len(re.findall(r"\S+", text))


def _trim(text, n):
    words = text.split()
    if len(words) <= n:
        return " ".join(words)
    return " ".join(words[:n]).rstrip(".,;:") + "…"


def _date(ns):
    return dt.datetime.fromtimestamp(ns / 1e9).strftime("%d-%m-%Y")


def _iso(ns):
    return dt.datetime.fromtimestamp(ns / 1e9).isoformat(timespec="seconds")


def _size(n):
    if n < 1024:
        return f"{n} B"
    if n < 1024 * 1024:
        return f"{round(n / 1024)} KB"
    return f"{n / 1024 / 1024:.1f} MB"


def _pct(x):
    return f"{x * 100:.1f}%"


def _plural(n, word):
    return f"{n} {word}{'s' if n != 1 else ''}"


def _read_text(full, limit=TEXT_LIMIT):
    with open(full, "rb") as fh:
        data = fh.read(limit)
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _sha1(full):
    h = hashlib.sha1()
    with open(full, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()[:12]


def _clean_md(s):
    s = re.sub(r"!?\[\[([^\]\|]+)\|([^\]]+)\]\]", r"\2", s)
    s = re.sub(r"!?\[\[([^\]]+)\]\]", r"\1", s)
    s = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"(\*\*|__|~~)(.+?)\1", r"\2", s)
    s = re.sub(r"(?<![\w*])([*_])(?=\S)(.+?)(?<=\S)\1(?![\w*])", r"\2", s)
    s = re.sub(r"<[^>]+>", "", s)
    return re.sub(r"\s+", " ", s).strip()


def _first_sentence(s):
    m = re.search(r"(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚÑ¿¡«(\d])", s)
    return s[:m.start()] if m else s


def normalize_mode(mode):
    """first-run | reread; the old Spanish words (primera-vez, relectura) are accepted."""
    return OLD_MODES.get(mode, mode)


# ------------------------------------------------------------------ frontmatter
def frontmatter(text):
    """Returns (dict, body). Minimal YAML: scalars, [a, b] lists and dash lists."""
    if not text.startswith("---"):
        return {}, text
    m = re.match(r"---[ \t]*\r?\n(.*?)\r?\n---[ \t]*(?:\r?\n|$)", text, re.S)
    if not m:
        return {}, text
    data, key = {}, None
    for line in m.group(1).splitlines():
        mm = re.match(r"^([A-Za-z_][\w\-]*):\s*(.*)$", line)
        if mm:
            key = mm.group(1)
            data[key] = mm.group(2).strip()
            continue
        if key is None or not line.strip():
            continue
        s = line.strip()
        if s.startswith("- "):
            prev = data.get(key)
            lst = prev if isinstance(prev, list) else []
            lst.append(s[2:].strip().strip("\"'"))
            data[key] = lst
        elif isinstance(data.get(key), str) and line[:1] in " \t":
            data[key] = (data[key] + " " + s).strip()
    for k, v in list(data.items()):
        if isinstance(v, str):
            v = re.sub(r"^[>|][+-]?\s*", "", v).strip()
            if v.startswith("[") and v.endswith("]"):
                v = [x.strip().strip("\"'") for x in v[1:-1].split(",") if x.strip()]
            elif len(v) >= 2 and v[0] == v[-1] == '"':
                try:
                    v = json.loads(v)
                except ValueError:
                    v = v[1:-1]
            elif len(v) >= 2 and v[0] == v[-1] == "'":
                v = v[1:-1].replace("''", "'")
            data[k] = v
    return data, text[m.end():]


# ------------------------------------------------------------------ summaries
RE_FIELD = re.compile(r"^\|?\s*\*\*([^*|]{1,30}?)\s*:?\s*\*\*\s*:?\s*\|?\s*(.*?)\s*\|?\s*$")
KEY_FIELDS = {"status", "purpose", "replaces", "client", "client / site",
              "estado", "para qué", "para que", "reemplaza a", "cliente", "cliente / obra"}
REPLACES_FIELDS = {"replaces", "reemplaza a", "reemplaza"}
REPLACES_FM = ("replaces", "reemplaza_a", "reemplaza")


def summary_md(text):
    """Title · status fields · first sentence (or description) · ## sections."""
    fm, body = frontmatter(text)
    title, sentence, in_code = None, None, False
    fields, sections = [], []
    for line in body.splitlines():
        s = line.strip()
        if s.startswith("```") or s.startswith("~~~"):
            in_code = not in_code
            continue
        if in_code or not s:
            continue
        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            clean = _clean_md(s.lstrip("#").strip())
            if title is None:
                title = clean
            elif level == 2:
                sections.append(_trim(re.sub(r"^[\d.]+\s*[·.\-)]?\s*", "", clean), 5))
            continue
        m = RE_FIELD.match(s)
        if m:                                     # metadata "**Key:** value" or "| **Key** | value |"
            key = m.group(1).strip().rstrip(":").lower()
            if key in KEY_FIELDS and m.group(2) and len(fields) < 3:
                fields.append(f"{key}: {_trim(_clean_md(m.group(2)), 8)}")
            continue
        if sentence is not None:
            continue
        if s.startswith(("|", "<!--", "---", "***", "![", "<")) or re.match(r"^[-=]{3,}$", s):
            continue
        s = re.sub(r"^(>\s*)+", "", s)
        s = re.sub(r"^([-*+]|\d+[.)])\s+", "", s)
        s = _clean_md(s)
        if len(s) >= 3 and not re.match(r"^[\wÁÉÍÓÚÑáéíóúñ /]{1,25}:\s", s):
            sentence = _trim(_first_sentence(s), 25)
    title = title or (fm.get("title") or fm.get("titulo") or fm.get("name") or None)
    desc = fm.get("description") or fm.get("descripcion") or fm.get("resumen") or fm.get("summary")
    body_sum = _trim(desc, 25) if isinstance(desc, str) and desc else sentence
    parts = [p for p in [title] + fields + [body_sum] if p]
    if len(sections) >= 2:
        parts.append("sections: " + "; ".join(sections[:6]) + (" …" if len(sections) > 6 else ""))
    return _trim(" · ".join(parts), MD_WORDS), fm


def replaces_values(text, fm):
    """Raw values of «Replaces:» / «Reemplaza a:» (body field or frontmatter key)."""
    values = []
    for k in REPLACES_FM:
        v = fm.get(k)
        if isinstance(v, str) and v.strip():
            values.append(v)
        elif isinstance(v, list):
            values += [str(x) for x in v if str(x).strip()]
    body = frontmatter(text)[1]
    in_code = False
    for line in body.splitlines():
        s = line.strip()
        if s.startswith("```") or s.startswith("~~~"):
            in_code = not in_code
            continue
        if in_code:
            continue
        m = RE_FIELD.match(s)
        if m and m.group(1).strip().rstrip(":").lower() in REPLACES_FIELDS and m.group(2).strip():
            values.append(m.group(2))
    return values


def summary_json(text):
    try:
        obj = json.loads(text)
    except (ValueError, RecursionError):
        return "JSON that could not be read"
    if isinstance(obj, dict):
        keys = list(obj.keys())
        extra = f" (+{len(keys) - 8})" if len(keys) > 8 else ""
        return "keys: " + ", ".join(map(str, keys[:8])) + extra
    if isinstance(obj, list):
        s = f"list of {len(obj)} items"
        if obj and isinstance(obj[0], dict):
            s += "; keys: " + ", ".join(map(str, list(obj[0].keys())[:6]))
        return s
    return f"{type(obj).__name__} value"


def summary_yaml(text):
    keys = re.findall(r"^([A-Za-z_][\w\-]*)\s*:", text, re.M)
    if not keys:
        return ""
    extra = f" (+{len(keys) - 8})" if len(keys) > 8 else ""
    return "keys: " + ", ".join(keys[:8]) + extra


def summary_py(text):
    try:
        doc = ast.get_docstring(ast.parse(text))
    except (SyntaxError, ValueError, RecursionError):
        doc = None
    if doc:
        for line in doc.strip().splitlines():
            if line.strip():
                return line.strip()
    return summary_comment(text)


def summary_comment(text):
    for line in text.splitlines()[:40]:
        s = line.strip()
        if s.startswith("#!") or "coding" in s[:20]:
            continue
        m = re.match(r"^(?:#|//|REM\s|::|<#|--)\s*(.+)$", s, re.I)
        if m and len(m.group(1).strip("#=-* ")) >= 3:
            return m.group(1).strip("#=-* ")
    return ""


def summary_text(text):
    for line in text.splitlines():
        if line.strip():
            return _trim(line.strip(), SUMMARY_WORDS)
    return ""


def summary_csv(text):
    lines = [x for x in text.splitlines() if x.strip()]
    if not lines:
        return ""
    return f"{_plural(len(lines) - 1, 'row')}; columns: " + _trim(lines[0], 12)


def summary_html(text):
    m = re.search(r"<title[^>]*>(.*?)</title>", text, re.I | re.S)
    return _clean_md(m.group(1)) if m else ""


def summary_xlsx(full):
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        return "", "openpyxl is not installed: name only"
    try:
        wb = openpyxl.load_workbook(full, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return f"Excel that could not be opened ({type(e).__name__})", None
    try:
        sheets = []
        a1 = None
        for i, ws in enumerate(wb.worksheets):
            dim = f" ({ws.max_row}×{ws.max_column})" if ws.max_row and ws.max_column else ""
            sheets.append(f"{ws.title}{dim}")
            if i == 0:
                try:
                    row = next(ws.iter_rows(min_row=1, max_row=1, max_col=3, values_only=True), ())
                    a1 = next((str(v) for v in row if isinstance(v, str) and v.strip()), None)
                except Exception:  # noqa: BLE001
                    a1 = None
        s = "sheets: " + ", ".join(sheets[:8]) + (f" (+{len(sheets) - 8})" if len(sheets) > 8 else "")
        if a1:
            s += f'; first row "{_trim(a1, 8)}"'
        return s, None
    finally:
        wb.close()


def summary_docx(full):
    try:
        with zipfile.ZipFile(full) as z:
            title = None
            if "docProps/core.xml" in z.namelist():
                core = z.read("docProps/core.xml").decode("utf-8", "replace")
                m = re.search(r"<dc:title>([^<]+)</dc:title>", core)
                title = m.group(1).strip() if m else None
            para = None
            if "word/document.xml" in z.namelist():
                with z.open("word/document.xml") as fh:
                    xml = fh.read(2_000_000).decode("utf-8", "replace")
                for p in re.finditer(r"<w:p[ >].*?</w:p>", xml, re.S):
                    t = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", p.group(0))).strip()
                    if len(t) >= 3:
                        para = t
                        break
        parts = list(dict.fromkeys(x for x in (title, para) if x))
        return _trim(" · ".join(parts), SUMMARY_WORDS)
    except (zipfile.BadZipFile, OSError, KeyError):
        return ""


def summary_pdf(full):
    try:
        import logging  # noqa: PLC0415
        import pypdf  # noqa: PLC0415
        logging.getLogger("pypdf").setLevel(logging.ERROR)
    except ImportError:
        return "", "pypdf is not installed: name only"
    try:
        r = pypdf.PdfReader(full)
        n = len(r.pages)
        title = None
        try:
            meta = r.metadata
            title = (meta.title or "").strip() if meta else None
        except Exception:  # noqa: BLE001
            title = None
        if not title and n:
            try:
                text = r.pages[0].extract_text() or ""
                for line in text.splitlines():
                    if len(line.strip()) >= 4:
                        title = 'first line: "' + _trim(line.strip(), 14) + '"'
                        break
            except Exception:  # noqa: BLE001
                pass
        pages = _plural(n, "page")
        return (f"{title} · {pages}" if title else pages), None
    except Exception as e:  # noqa: BLE001
        return f"PDF that could not be read ({type(e).__name__})", None


def summary_image(full, ext):
    try:
        with open(full, "rb") as fh:
            head = fh.read(65536)
        w = h = None
        if ext == ".png" and head[:8] == b"\x89PNG\r\n\x1a\n":
            w, h = struct.unpack(">II", head[16:24])
        elif ext == ".gif" and head[:4] == b"GIF8":
            w, h = struct.unpack("<HH", head[6:10])
        elif ext in (".jpg", ".jpeg") and head[:2] == b"\xff\xd8":
            i = 2
            while i + 9 < len(head):
                if head[i] != 0xFF:
                    i += 1
                    continue
                mark = head[i + 1]
                length = struct.unpack(">H", head[i + 2:i + 4])[0]
                if mark in (0xC0, 0xC1, 0xC2):
                    h, w = struct.unpack(">HH", head[i + 5:i + 9])
                    break
                i += 2 + length
        return f"image {w}×{h}" if w and h else "image"
    except (OSError, struct.error):
        return "image"


def summary_zip(full):
    try:
        with zipfile.ZipFile(full) as z:
            return f"archive with {_plural(len(z.namelist()), 'file')}"
    except (zipfile.BadZipFile, OSError):
        return "archive"


def summarize(full, ext, text):
    """(summary, frontmatter, warning). Never calls a model."""
    fm, warning = {}, None
    if ext in (".md", ".markdown"):
        res, fm = summary_md(text)
    elif ext in (".json", ".canvas"):
        res = summary_json(text)
    elif ext in (".yaml", ".yml", ".toml", ".base"):
        res = summary_yaml(text)
    elif ext == ".py":
        res = summary_py(text)
    elif ext in (".ps1", ".sh", ".bat", ".cmd", ".js", ".ts", ".sql", ".ini", ".cfg"):
        res = summary_comment(text)
    elif ext in (".html", ".htm"):
        res = summary_html(text)
    elif ext in (".csv", ".tsv"):
        res = summary_csv(text)
    elif ext == ".txt":
        res = summary_text(text)
    elif ext in (".xlsx", ".xlsm"):
        res, warning = summary_xlsx(full)
    elif ext == ".docx":
        res = summary_docx(full)
    elif ext == ".pdf":
        res, warning = summary_pdf(full)
    elif ext in (".png", ".jpg", ".jpeg", ".gif"):
        res = summary_image(full, ext)
    elif ext == ".zip":
        res = summary_zip(full)
    else:
        res = ""
    cap = MD_WORDS if ext in (".md", ".markdown") else SUMMARY_WORDS
    return _trim(res or "", cap), fm, warning


# ------------------------------------------------------------------ walk
def _matches(name, rel, patterns):
    return any(fnmatch.fnmatch(name, p) or fnmatch.fnmatch(rel, p) for p in patterns)


def _is_link(full):
    try:
        if os.path.islink(full):
            return True
        isj = getattr(os.path, "isjunction", None)
        return bool(isj and isj(full))
    except OSError:
        return False


def walk(root, patterns, out_long, max_bytes, excl):
    base = _long(root)
    for dirpath, dirnames, filenames in os.walk(base):
        rel_dir = _posix(dirpath[len(base):].lstrip("\\/"))
        keep = []
        for d in sorted(dirnames):
            full = os.path.join(dirpath, d)
            rel = posixpath.join(rel_dir, d) if rel_dir else d
            if os.path.normcase(full) == os.path.normcase(out_long):
                excl["output"].append(rel)
            elif d in EXCLUDE_DIRS:
                excl["dirs"].append(rel)
            elif _matches(d, rel, patterns):
                excl["pattern"].append(rel + "/")
            elif _is_link(full):
                excl["links"].append(rel)
            else:
                keep.append(d)
        dirnames[:] = keep
        for f in sorted(filenames):
            full = os.path.join(dirpath, f)
            rel = posixpath.join(rel_dir, f) if rel_dir else f
            if f.lower() in EXCLUDE_FILES or f.startswith(EXCLUDE_PREFIXES):
                excl["system"].append(rel)
                continue
            if _matches(f, rel, patterns):
                excl["pattern"].append(rel)
                continue
            try:
                st = os.stat(full)
            except OSError:
                excl["unreadable"].append(rel)
                continue
            ext = os.path.splitext(f)[1].lower()
            if ext not in TEXT_EXT and st.st_size > max_bytes:
                excl["binaries"].append(rel)
                continue
            yield rel, full, st, ext


def _prefixes(roots):
    if len(roots) == 1:
        return [""]
    seen, out = Counter(), []
    for r in roots:
        b = os.path.basename(r.rstrip("\\/")) or "root"
        seen[b] += 1
        out.append(f"{b}~{seen[b]}/" if seen[b] > 1 else f"{b}/")
    return out


# ------------------------------------------------------------------ links
class Resolver:
    def __init__(self, ids, roots, prefixes):
        self.ids = set(ids)
        self.lower = {i.lower(): i for i in ids}
        self.by_name = defaultdict(list)
        self.by_stem = defaultdict(list)
        for i in ids:
            name = posixpath.basename(i).lower()
            self.by_name[name].append(i)
            self.by_stem[os.path.splitext(name)[0]].append(i)
        self.roots = [(_posix(r).rstrip("/").lower(), p) for r, p in zip(roots, prefixes)]

    def _exact(self, c):
        if c in self.ids:
            return c
        return self.lower.get(c.lower())

    def resolve(self, cand, origin, prefix, wiki=False):
        c = urllib.parse.unquote(cand.strip().strip("\"'<>()[]")).replace("\\", "/")
        c = c.split("#")[0].split("?")[0].strip()
        if not c or c.endswith("/"):
            return None
        if re.match(r"^[A-Za-z]:/", c) or c.startswith("/"):
            cl = c.lower()
            for root, pref in self.roots:
                if cl.startswith(root + "/"):
                    return self._exact(posixpath.normpath(pref + c[len(root) + 1:]))
            return None
        c = re.sub(r"^(\./)+", "", c)
        base = posixpath.dirname(origin)
        tries = [posixpath.normpath(posixpath.join(base, c)),
                 posixpath.normpath(posixpath.join(prefix, c)) if prefix else posixpath.normpath(c)]
        for t in tries:
            for v in (t, t + ".md"):
                hit = self._exact(v)
                if hit:
                    return hit
        name = posixpath.basename(c).lower()
        if "/" not in c:
            lst = self.by_name.get(name) or (self.by_stem.get(name) if wiki else None) or []
            if not lst and wiki:
                lst = self.by_name.get(name + ".md", [])
            return lst[0] if len(lst) == 1 else None
        # partial path cited from another root: unique suffix
        cand_ids = self.by_name.get(name) or (self.by_name.get(name + ".md", []) if wiki else [])
        suffix = "/" + c.lower().lstrip("/")
        hits = [i for i in cand_ids if ("/" + i.lower()).endswith(suffix)
                or ("/" + i.lower()).endswith(suffix + ".md")]
        return hits[0] if len(hits) == 1 else None


def extract_links(text, fm, origin, prefix, res):
    """List of (target, relation, confidence). Target = file id or "skill:<name>"."""
    out = []
    skills = []
    for k in ("skills", "skill"):
        v = fm.get(k)
        if isinstance(v, str) and v:
            skills += [x.strip() for x in v.split(",") if x.strip()]
        elif isinstance(v, list):
            skills += [str(x) for x in v if str(x).strip()]
    for s in skills:
        target = res.resolve(f"skills/{s}/SKILL.md", origin, prefix)
        out.append((target or f"skill:{s}", "uses_skill", "EXTRACTED"))
    no_url = RE_URL.sub(" ", text)
    for m in RE_WIKI.finditer(no_url):
        d = res.resolve(m.group(1), origin, prefix, wiki=True)
        if d:
            out.append((d, "wikilink", "EXTRACTED"))
    for m in RE_MDLINK.finditer(no_url):
        d = res.resolve(m.group(1), origin, prefix)
        if d:
            out.append((d, "link", "EXTRACTED"))
    cands = [m.group(1) for m in RE_BACKTICK.finditer(no_url)]
    cands += [m.group(1) for m in RE_PATH.finditer(no_url)]
    for c in cands:
        if "." not in c:
            continue
        d = res.resolve(c, origin, prefix)
        if d:
            out.append((d, "mention", "INFERRED"))
    return out


def resolve_replaced(values, origin, prefix, res):
    """Files that «Replaces:» / «Reemplaza a:» names. Links, backticks, paths, or a bare name."""
    hits = []
    for v in values:
        found = []
        for m in RE_WIKI.finditer(v):
            found.append(res.resolve(m.group(1), origin, prefix, wiki=True))
        for m in RE_MDLINK.finditer(v):
            found.append(res.resolve(m.group(1), origin, prefix))
        for m in RE_BACKTICK.finditer(v):
            found.append(res.resolve(m.group(1), origin, prefix, wiki=True))
        for m in RE_PATH.finditer(v):
            found.append(res.resolve(m.group(1), origin, prefix))
        found = [f for f in found if f]
        if not found:                       # a bare name: «LIST-001 (the raw transcription)»
            first = re.split(r"[\s,;(]", _clean_md(v).strip(), maxsplit=1)[0].strip("`'\".:")
            if first:
                f = res.resolve(first, origin, prefix, wiki=True)
                if f:
                    found.append(f)
        for f in found:
            if f != origin and f not in hits:
                hits.append(f)
    return hits


# ------------------------------------------------------------------ fingerprints, changes and mode
def read_json(path):
    try:
        with open(_long(path), encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        return None


def layout_of(out):
    """'en' (context/, MAP.md, fingerprints.json...) or 'es' (a context written before v1.1)."""
    out_long = _long(out)
    if os.path.exists(os.path.join(out_long, LAYOUTS["en"]["fingerprints"])):
        return "en"
    for key in ("fingerprints", "map"):
        if os.path.exists(os.path.join(out_long, LAYOUTS["es"][key])):
            return "es"
    return "en"


def context_files(out):
    lay = layout_of(out)
    return lay, {k: os.path.join(out, v) for k, v in LAYOUTS[lay].items()}


FP_KEYS = {"generado": "generated", "raices": "roots", "archivos": "files"}
FP_FILE_KEYS = {"tam": "size", "fecha": "date", "resumen": "summary"}
CH_KEYS = {"leido_el": "read_at", "primera_vez_pendiente": "first_run_pending",
           "primera_vez_desde": "first_run_since", "corridas": "runs"}
CH_RUN_KEYS = {"fecha": "date", "nuevos": "new", "cambiados": "changed", "borrados": "deleted"}


def _rename(d, keys):
    return {keys.get(k, k): v for k, v in d.items()}


def normalize_fingerprints(fp):
    """fingerprints.json in English keys; a huellas.json (Spanish keys) is accepted."""
    if not isinstance(fp, dict):
        return None
    fp = _rename(fp, FP_KEYS)
    if isinstance(fp.get("files"), dict):
        fp["files"] = {i: _rename(x, FP_FILE_KEYS) if isinstance(x, dict) else x
                       for i, x in fp["files"].items()}
    return fp


def normalize_changes(ch):
    """changes.json in English keys; a cambios.json (Spanish keys) is accepted."""
    if not isinstance(ch, dict):
        return None
    ch = _rename(ch, CH_KEYS)
    if isinstance(ch.get("runs"), list):
        ch["runs"] = [_rename(r, CH_RUN_KEYS) if isinstance(r, dict) else r for r in ch["runs"]]
    return ch


def delta(prev, current):
    """Changes of this run against the previous one: (new, changed, deleted)."""
    new = sorted(i for i in current if i not in prev)
    deleted = sorted(i for i in prev if i not in current)
    changed = sorted(i for i in current if i in prev and (
        current[i]["sha1"] != prev[i].get("sha1") or current[i]["size"] != prev[i].get("size")))
    return new, changed, deleted


def net(runs):
    """Accumulated state of each file across the unread runs.

    new -> changed stays new; new -> deleted disappears (it was never in the base);
    deleted -> new is changed."""
    state = {}
    for r in runs:
        for i in r.get("new", []):
            state[i] = "changed" if state.get(i) == "deleted" else "new"
        for i in r.get("changed", []):
            state[i] = "new" if state.get(i) == "new" else "changed"
        for i in r.get("deleted", []):
            if state.get(i) == "new":
                state.pop(i)
            else:
                state[i] = "deleted"
    return sorted(state.items())


def empty_changes(base=None):
    return {"version": VERSION, "read_at": None, "base": base, "first_run_pending": False,
            "first_run_since": None, "runs": []}


def decide_mode(prev_fp, changes, current, roots, has_map, threshold, now):
    """Decides the mode against what has accumulated UNREAD (changes.json).

    Returns (mode, reason, run_delta, accumulated, ratio, updated changes)."""
    no_delta = ([], [], [])
    if not prev_fp or "files" not in prev_fp:
        c = empty_changes()
        c.update(first_run_pending=True, first_run_since=now)
        return "first-run", "no previous fingerprints", no_delta, [], 1.0, c
    prev = prev_fp["files"]
    if [r.lower() for r in prev_fp.get("roots", [])] != [r.lower() for r in roots]:
        c = empty_changes()
        c.update(first_run_pending=True, first_run_since=now)
        return "first-run", "the mapped folders changed", no_delta, [], 1.0, c
    c = changes if isinstance(changes, dict) and "runs" in changes else empty_changes(len(prev))
    c["version"] = VERSION
    if c.get("base") is None:
        c["base"] = len(prev)
    d = delta(prev, current)
    if any(d):
        c["runs"].append({"date": now, "new": d[0], "changed": d[1], "deleted": d[2]})
    accumulated = net(c["runs"])
    ratio = len(accumulated) / max(c["base"], 1)
    if not has_map:
        c.update(first_run_pending=True, first_run_since=c.get("first_run_since") or now)
        return "first-run", "fingerprints exist but the map is missing", d, accumulated, ratio, c
    if c.get("first_run_pending"):
        reason = (f"the first run of {c.get('first_run_since') or '?'} has not been marked read "
                  "(`--mark-read` when consolidating)")
        return "first-run", reason, d, accumulated, ratio, c
    since = f" since {c['read_at']}" if c.get("read_at") else ""
    if ratio <= threshold:
        reason = (f"{len(accumulated)} of {c['base']} files changed{since} without being read "
                  f"({_pct(ratio)}, threshold {_pct(threshold)})"
                  + ("" if accumulated else "; this map is enough"))
        return "reread", reason, d, accumulated, ratio, c
    c.update(first_run_pending=True, first_run_since=now)
    reason = (f"{len(accumulated)} of {c['base']} files changed{since} without being read "
              f"({_pct(ratio)}, above the threshold {_pct(threshold)})")
    return "first-run", reason, d, accumulated, ratio, c


def mark_read(out):
    """Empties what has accumulated: orchestrate calls it when consolidating. The base becomes the
    last run."""
    out = _short(os.path.abspath(out))
    _, paths = context_files(out)
    fp = normalize_fingerprints(read_json(paths["fingerprints"]))
    if not isinstance(fp, dict) or not isinstance(fp.get("files"), dict):
        raise InputError(f"No fingerprints in {out}: map the folders first")
    c = empty_changes(len(fp["files"]))
    c["read_at"] = dt.datetime.now().strftime("%d-%m-%Y %H:%M")
    with open(_long(paths["changes"]), "w", encoding="utf-8", newline="\n") as fh:
        json.dump(c, fh, ensure_ascii=False, indent=1)
    return c


# ------------------------------------------------------------------ core
CORE_NAMES = ["leeme", "readme", "trabajo", "job", "claude", "indice", "index"]
CORE_OUT = {"image", "archive", "cad"}
CORE_MAX = 3


def choose_core(files, inbound, pinned=(), replaced=None):
    """Up to 3 files that are always read. Deterministic rule:
    1) those pinned with --core; 2) key names (LEEME/README, TRABAJO/JOB, CLAUDE, INDICE/INDEX),
    shallowest first; 3) the most linked (no images, archives or CAD); tie: the most recent.
    A file that another declares replaced («Replaces:» / «Reemplaza a:») never enters by rule."""
    replaced = replaced or {}

    def stem(i):
        return os.path.splitext(posixpath.basename(i))[0].lower()

    chosen = [(f, "pinned by hand") for f in pinned]
    used = {i for i, _ in chosen}
    by_name = sorted((i for i in files if stem(i) in CORE_NAMES and i not in used and i not in replaced),
                     key=lambda i: (i.count("/"), CORE_NAMES.index(stem(i)), -inbound[i],
                                    -files[i]["mtime_ns"], i))
    for i in by_name:
        if len(chosen) >= CORE_MAX:
            break
        chosen.append((i, "key name"))
        used.add(i)
    rest = sorted((i for i in files if i not in used and i not in replaced
                   and files[i]["kind"] not in CORE_OUT),
                  key=lambda i: (-inbound[i], -files[i]["mtime_ns"], i))
    for i in rest:
        if len(chosen) >= CORE_MAX:
            break
        n = inbound[i]
        chosen.append((i, f"{_plural(n, 'inbound link')}" if n else "the most recent"))
    return chosen[:CORE_MAX]


def _check_pinned(pinned, files):
    lower = {i.lower(): i for i in files}
    out = []
    for f in pinned:
        c = re.sub(r"^(\./)+", "", _posix(f.strip()))
        if not c:
            continue
        hit = lower.get(c.lower())
        if not hit:
            raise InputError(f"--core: \"{f}\" is not in the map (write it as the map's relative path)")
        if hit not in out:
            out.append(hit)
    if len(out) > CORE_MAX:
        raise InputError(f"--core: at most {CORE_MAX} files")
    return out


# ------------------------------------------------------------------ MAP.md
LEVELS = [
    dict(res=60, k=None, size=True, cap=25, depth=2),
    dict(res=30, k=None, size=True, cap=25, depth=2),
    dict(res=14, k=None, size=False, cap=15, depth=2),
    dict(res=12, k=10, size=False, cap=10, depth=2),
    dict(res=10, k=6, size=False, cap=10, depth=2),
    dict(res=10, k=4, size=False, cap=8, depth=1),
    dict(res=8, k=2, size=False, cap=8, depth=1),
    dict(res=8, k=1, size=False, cap=6, depth=1),
    dict(res=8, k=0, size=False, cap=6, depth=1),
    dict(res=6, k=0, size=False, cap=4, depth=0),
]


def _folder(i):
    d = posixpath.dirname(i)
    return d + "/" if d else "./"


def _kinds(lst, n=4):
    c = Counter(x["kind"] for x in lst)
    return ", ".join(f"{v} {k}" for k, v in c.most_common(n)) + (" …" if len(c) > n else "")


def render_map(ctx, lv):
    nodes = ctx["nodes"]
    L = []
    a = L.append
    a("# Context map")
    a("")
    a(f"**MODE: {ctx['mode']}** · {ctx['reason']}")
    a("")
    a(f"Generated on {ctx['generated']} by `map_context.py` (no LLM). Roots: "
      + " · ".join(f"`{r}`" for r in ctx["roots"]))
    a(f"{_plural(len(nodes), 'file')} · {_size(ctx['bytes'])} · from {ctx['since']} to {ctx['until']} · "
      f"{_plural(ctx['n_links'], 'link')} between files")
    a("By kind: " + _kinds(nodes.values(), 12))
    if ctx["excluded"]:
        a("Excluded: " + ctx["excluded"])
    a("")
    a("## Core · always read")
    a("")
    a("Reread = these files + those under «To reread» + this map.")
    for i, why in ctx["core"]:
        r = _trim(nodes[i]["summary"], min(lv["res"], 30)) if i in nodes else ""
        a(f"- `{i}` ({why})" + (f" · {r}" if r else ""))
    replaced = ctx["replaced"]
    for i in sorted(replaced)[:lv["cap"]]:
        a(f"- left out: `{i}` (replaced by `{replaced[i]}`)")
    a("")
    if ctx["reread"] or ctx["deleted"]:
        a("## To reread")
        a("")
        if ctx["reread_since"]:
            a(f"Changes accumulated without being read since {ctx['reread_since']}.")
        lim = lv["cap"] * 2
        for i, state in ctx["reread"][:lim]:
            r = _trim(nodes[i]["summary"], lv["res"]) if i in nodes else ""
            a(f"- `{i}` ({state})" + (f" · {r}" if r else ""))
        for i in ctx["deleted"][:lim]:
            a(f"- `{i}` (deleted)")
        rest = max(len(ctx["reread"]) - lim, 0) + max(len(ctx["deleted"]) - lim, 0)
        if rest:
            a(f"- +{rest} more: the full list is in `{ctx['graph_name']}` (graph.changes)")
        a("")

    # Folders, grouped at the depth of the level
    a("## Folders")
    a("")
    groups = defaultdict(list)
    for n in nodes.values():
        parts = posixpath.dirname(n["id"]).split("/") if posixpath.dirname(n["id"]) else []
        key = "/".join(parts[:lv["depth"]]) + "/" if parts and lv["depth"] else "./"
        if lv["depth"] == 0:
            key = "./"
        groups[key].append(n)
    for c in sorted(groups):
        g = groups[c]
        last = max(x["mtime_ns"] for x in g)
        a(f"- `{c}` · {_plural(len(g), 'file')} ({_kinds(g)}) · last {_date(last)}")
    a("")

    score = ctx["score"]
    if lv["k"] != 0:
        a("## Files")
        a("")
        by_folder = defaultdict(list)
        for n in nodes.values():
            by_folder[_folder(n["id"])].append(n)
        for c in sorted(by_folder):
            g = by_folder[c]
            a(f"### `{c}`")
            if lv["k"] is None or len(g) <= lv["k"]:
                chosen, rest = g, []
            else:
                order = sorted(g, key=lambda x: (-score[x["id"]], x["id"]))
                chosen, rest = order[:lv["k"]], order[lv["k"]:]
            for n in sorted(chosen, key=lambda x: x["id"].lower()):
                parts = [f"`{posixpath.basename(n['id'])}`"]
                if lv["size"]:
                    parts.append(_size(n["size"]))
                parts.append(_date(n["mtime_ns"]))
                r = _trim(n["summary"], lv["res"])
                if r:
                    parts.append(r)
                a("- " + " · ".join(parts))
            if rest:
                a(f"- +{len(rest)} more ({_kinds(rest)})")
            a("")
    else:
        a("## Key files")
        a("")
        top = sorted(nodes.values(), key=lambda x: (-score[x["id"]], x["id"]))[:lv["cap"] * 3]
        for n in top:
            r = _trim(n["summary"], lv["res"])
            a(f"- `{n['id']}` · {_date(n['mtime_ns'])}" + (f" · {r}" if r else ""))
        a(f"- (the other {max(len(nodes) - len(top), 0)} are in `{ctx['graph_name']}`)")
        a("")

    a("## Most recent")
    a("")
    for n in sorted(nodes.values(), key=lambda x: (-x["mtime_ns"], x["id"]))[:lv["cap"] // 2 or 1]:
        hour = dt.datetime.fromtimestamp(n["mtime_ns"] / 1e9).strftime("%H:%M")
        a(f"- {_date(n['mtime_ns'])} {hour} · `{n['id']}`")
    a("")

    a("## Most linked")
    a("")
    inbound = ctx["inbound"]
    top = sorted((i for i in inbound if inbound[i]), key=lambda i: (-inbound[i], i))[:lv["cap"] // 2 or 1]
    if top:
        for i in top:
            a(f"- `{i}` ← {inbound[i]}")
    else:
        a("- none: the files do not cite each other")
    a("")

    orph = ctx["orphans"]
    a(f"## Orphans ({len(orph)})")
    a("")
    a("Nobody links them and they link nobody.")
    for i in orph[:lv["cap"]]:
        a(f"- `{i}`")
    if len(orph) > lv["cap"]:
        a(f"- +{len(orph) - lv['cap']} more")
    a("")

    no_sum = ctx["no_summary"]
    a(f"## No summary ({len(no_sum)})")
    a("")
    if no_sum:
        a("The script could not summarize them; the first ones weigh most in the graph.")
        for i in no_sum[:min(lv["cap"], 10)]:
            a(f"- `{i}`")
        if len(no_sum) > min(lv["cap"], 10):
            a(f"- +{len(no_sum) - min(lv['cap'], 10)} more")
    else:
        a("Every file has a summary.")
    a("")
    for note in ctx["notes"]:
        a(note)
        a("")
    return "\n".join(L).rstrip() + "\n"


def build_map(ctx, max_words):
    text = ""
    for lv in LEVELS:
        text = render_map(ctx, lv)
        if _words(text) <= max_words:
            return text, lv
    # Last resort: cut by lines until it fits.
    lines, count, out = text.splitlines(), 0, []
    for ln in lines:
        p = _words(ln)
        if count + p > max_words - 20:
            break
        out.append(ln)
        count += p
    out.append("")
    out.append(f"(Map truncated to fit the word cap; the full detail is in `{ctx['graph_name']}`.)")
    return "\n".join(out) + "\n", LEVELS[-1]


# ------------------------------------------------------------------ main
def map_context(folders, out, exclude=(), threshold=0.20, max_words=3000, max_mb=25.0,
                graphify=False, core=()):
    t0 = time.perf_counter()
    now = dt.datetime.now().strftime("%d-%m-%Y %H:%M")
    roots = []
    for c in folders:
        p = _short(os.path.abspath(c))
        if not os.path.isdir(_long(p)):
            raise InputError(f"Folder not found: {c}")
        roots.append(p)
    prefixes = _prefixes(roots)
    out = _short(os.path.abspath(out))
    out_long = _long(out)
    layout, paths = context_files(out)
    excl = defaultdict(list)
    files = {}
    for root, pref in zip(roots, prefixes):
        for rel, full, st, ext in walk(root, list(exclude), out_long, max_mb * 1024 * 1024, excl):
            i = pref + rel
            files[i] = {"id": i, "full": full, "pref": pref, "ext": ext, "size": st.st_size,
                        "mtime_ns": st.st_mtime_ns,
                        "kind": KIND_BY_EXT.get(ext, ext.lstrip(".") or "no-extension")}

    prev_fp = normalize_fingerprints(read_json(paths["fingerprints"]))
    prev_files = (prev_fp or {}).get("files", {}) if isinstance(prev_fp, dict) else {}
    if not isinstance(prev_files, dict):
        prev_files = {}
    same_version = isinstance(prev_fp, dict) and prev_fp.get("version") == VERSION
    warnings = set()

    # Fingerprint: the sha1 is reused if size and date did not change (the file is not read again).
    for i, x in files.items():
        p = prev_files.get(i)
        if p and p.get("size") == x["size"] and p.get("mtime_ns") == x["mtime_ns"] and p.get("sha1"):
            x["sha1"] = p["sha1"]
            x["reuse"] = same_version and "summary" in p
            x["prev_summary"] = p.get("summary", "")
        else:
            try:
                x["sha1"] = _sha1(x["full"])
            except OSError:
                x["sha1"] = "unreadable"
            x["reuse"] = False

    pinned = _check_pinned(core, files)
    has_map = any(os.path.exists(os.path.join(out_long, LAYOUTS[k]["map"])) for k in LAYOUTS)
    mode, reason, run_delta, accumulated, ratio, changes = decide_mode(
        prev_fp if isinstance(prev_fp, dict) else None,
        normalize_changes(read_json(paths["changes"])), files, roots, has_map, threshold, now)
    new = [i for i, e in accumulated if e == "new"]
    changed = [i for i, e in accumulated if e == "changed"]
    deleted = [i for i, e in accumulated if e == "deleted"]

    res = Resolver(list(files), roots, prefixes)
    links = {}
    replaced = {}                       # replaced file -> the file that replaces it
    for i, x in files.items():
        text, fm = "", {}
        if x["ext"] in TEXT_EXT:
            try:
                text = _read_text(x["full"])
            except OSError:
                text = ""
        if x["reuse"] and x["ext"] not in TEXT_EXT:
            summary, warning = x["prev_summary"], None
        else:
            try:
                summary, fm, warning = summarize(x["full"], x["ext"], text)
            except Exception as e:  # noqa: BLE001  (an odd file does not stop the map)
                summary, fm, warning = "", {}, f"{type(e).__name__} while summarizing"
        if x["ext"] in (".md", ".markdown") and not fm:
            fm = frontmatter(text)[0]
        if warning:
            warnings.add(warning)
        x["summary"] = summary
        if x["ext"] in (".md", ".markdown") and text:
            for old in resolve_replaced(replaces_values(text, fm), i, x["pref"], res):
                replaced.setdefault(old, i)
        if x["ext"] in LINK_EXT and text:
            for target, rel, conf in extract_links(text, fm, i, x["pref"], res):
                if target == i:
                    continue
                key = (i, target)
                order = {"uses_skill": 0, "wikilink": 1, "link": 2, "mention": 3}
                if key not in links or order[rel] < order[links[key][0]]:
                    links[key] = (rel, conf)

    # Degrees and relevance score
    inbound, outbound = Counter(), Counter()
    for (s, t) in links:
        outbound[s] += 1
        inbound[t] += 1
    mt_max = max((x["mtime_ns"] for x in files.values()), default=0)
    score = {}
    for i, x in files.items():
        stem = os.path.splitext(posixpath.basename(i))[0].lower()
        p = inbound[i] * 3 + outbound[i]
        p += 2 if x["kind"] in DOC else 0
        p += 3 if stem in KEY_NAMES or stem.startswith(tuple(KEY_NAMES)) else 0
        p += 1 if mt_max - x["mtime_ns"] < 30 * 86400 * 10**9 else 0
        score[i] = p
    orphans = sorted(i for i in files if not inbound[i] and not outbound[i])
    no_summary = sorted((i for i, x in files.items() if not x["summary"]),
                        key=lambda i: (-score[i], i))

    reread = [(i, e) for i, e in accumulated if e != "deleted"]
    if mode == "first-run":
        reread_map, deleted_map = [], []
    else:
        reread_map, deleted_map = reread, deleted
    core_sel = choose_core(files, inbound, pinned, replaced)

    # Readable exclusions
    parts = []
    if excl["output"]:
        parts.append("the output folder")
    if excl["dirs"]:
        names = sorted({posixpath.basename(x) for x in excl["dirs"]})
        parts.append(f"{_plural(len(excl['dirs']), 'system folder')} ({', '.join(names[:5])})")
    if excl["links"]:
        parts.append(f"{len(excl['links'])} links or junctions (not followed)")
    if excl["binaries"]:
        parts.append(f"{len(excl['binaries'])} binaries over {max_mb:g} MB")
    if excl["pattern"]:
        parts.append(f"{len(excl['pattern'])} by --exclude")
    if excl["unreadable"]:
        parts.append(f"{len(excl['unreadable'])} unreadable")

    notes = []
    if warnings:
        notes.append("Warnings: " + "; ".join(sorted(warnings)) + ".")
    graphify_path = shutil.which("graphify") if graphify else None
    if graphify:
        if graphify_path:
            notes.append(
                "## Enrich with graphify\n\n"
                "graphify is installed. This map does not call it (graphify extracts concepts with a "
                "model, and that costs). If the task justifies it, in Claude Code: `/graphify <folder>`; "
                f"it leaves `graphify-out/graph.json`, with the same shape as `{LAYOUTS[layout]['graph']}` "
                "(nodes, links), and `graphify merge-graphs` joins them.")
        else:
            notes.append("graphify is not installed: the map stays structural only.")

    nodes = {}
    for i, x in files.items():
        nodes[i] = {"id": i, "label": posixpath.basename(i), "type": "file", "kind": x["kind"],
                    "file_type": x["kind"], "source_file": i, "folder": _folder(i),
                    "size": x["size"], "mtime_ns": x["mtime_ns"], "date": _iso(x["mtime_ns"]),
                    "summary": x["summary"], "in_links": inbound[i], "out_links": outbound[i]}
        if i in replaced:
            nodes[i]["replaced_by"] = replaced[i]
    generated = now
    ctx = {
        "core": core_sel, "replaced": replaced, "graph_name": LAYOUTS[layout]["graph"],
        "reread_since": changes["runs"][0]["date"] if changes.get("runs") else None,
        "mode": mode, "reason": reason, "generated": generated, "roots": roots, "nodes": nodes,
        "bytes": sum(x["size"] for x in files.values()),
        "since": _date(min((x["mtime_ns"] for x in files.values()), default=0)),
        "until": _date(mt_max), "n_links": len(links), "excluded": " · ".join(parts),
        "reread": reread_map, "deleted": deleted_map, "score": score, "inbound": inbound,
        "orphans": orphans, "no_summary": no_summary, "notes": notes,
    }
    map_text, level = build_map(ctx, max_words)
    words = _words(map_text)

    # graph.json
    external = sorted({t for (_, t) in links if t not in nodes})
    graph = {
        "directed": True, "multigraph": False,
        "graph": {"tool": "map_context", "version": VERSION, "generated": generated,
                  "roots": roots, "mode": mode, "reason": reason,
                  "core": [i for i, _ in core_sel],
                  "replaced": dict(sorted(replaced.items())),
                  "changes": {"new": new, "changed": changed, "deleted": deleted},
                  "changes_this_run": {"new": run_delta[0], "changed": run_delta[1],
                                       "deleted": run_delta[2]}},
        "nodes": [{k: v for k, v in n.items() if k != "mtime_ns"} for n in
                  sorted(nodes.values(), key=lambda n: n["id"])]
        + [{"id": e, "label": e.split(":", 1)[1], "type": "skill", "file_type": "skill",
            "external": True} for e in external],
        "links": [{"source": s, "target": t, "relation": r, "confidence": c}
                  for (s, t), (r, c) in sorted(links.items())],
    }
    fingerprints = {
        "version": VERSION, "generated": generated, "roots": roots,
        "files": {i: {"size": x["size"], "mtime_ns": x["mtime_ns"], "date": _iso(x["mtime_ns"]),
                      "sha1": x["sha1"], "summary": x["summary"]}
                  for i, x in sorted(files.items())},
    }
    os.makedirs(out_long, exist_ok=True)
    for key, content in (("graph", graph), ("fingerprints", fingerprints), ("changes", changes)):
        with open(_long(paths[key]), "w", encoding="utf-8", newline="\n") as fh:
            json.dump(content, fh, ensure_ascii=False, indent=1)
    with open(_long(paths["map"]), "w", encoding="utf-8", newline="\n") as fh:
        fh.write(map_text)

    return {
        "mode": mode, "reason": reason, "new": new, "changed": changed,
        "deleted": deleted, "n_changes": len(accumulated), "core": core_sel, "replaced": replaced,
        "pending_runs": len(changes.get("runs", [])), "layout": layout,
        "ratio": ratio, "words": words, "files": len(files), "links": len(links),
        "link_list": sorted((s, t, r) for (s, t), (r, _) in links.items()),
        "orphans": orphans, "no_summary": no_summary, "excluded": {k: list(v) for k, v in excl.items()},
        "level": LEVELS.index(level), "seconds": time.perf_counter() - t0,
        "map": paths["map"], "summaries": {i: x["summary"] for i, x in files.items()},
    }


def final_line(r):
    return f"MODE: {r['mode']} · changed {r['n_changes']} · map {r['words']} words"


def report(r):
    print(f"Files: {r['files']} · links: {r['links']} · orphans: {len(r['orphans'])} · "
          f"no summary: {len(r['no_summary'])}")
    print(f"Reason for the mode: {r['reason']}")
    for i, why in r["core"]:
        print(f"  core: {i} ({why})")
    for i in sorted(r["replaced"]):
        print(f"  out of the core: {i} (replaced by {r['replaced'][i]})")
    if r["mode"] == "reread":
        for i in r["new"]:
            print(f"  reread (new): {i}")
        for i in r["changed"]:
            print(f"  reread (changed): {i}")
        for i in r["deleted"]:
            print(f"  deleted: {i}")
    legacy = " (context written before v1.1: Spanish file names kept)" if r["layout"] == "es" else ""
    print(f"Map: {r['map']} (compression level {r['level']}){legacy}")
    print(f"Time: {r['seconds']:.2f} s")
    print(final_line(r))


# ------------------------------------------------------------------ test
def _write_legacy_context(out, fp, roots, base):
    """Turns an English context into one written before v1.1 (Spanish names and keys)."""
    huellas = {"version": 1, "generado": fp["generated"], "raices": roots,
               "archivos": {i: {"tam": x["size"], "mtime_ns": x["mtime_ns"], "fecha": x["date"],
                                "sha1": x["sha1"], "resumen": x["summary"]}
                            for i, x in fp["files"].items()}}
    cambios = {"version": 1, "leido_el": None, "base": base, "primera_vez_pendiente": False,
               "primera_vez_desde": None, "corridas": []}
    for name in os.listdir(out):
        os.remove(os.path.join(out, name))
    for name, content in (("huellas.json", huellas), ("cambios.json", cambios), ("grafo.json", {})):
        with open(os.path.join(out, name), "w", encoding="utf-8", newline="\n") as fh:
            json.dump(content, fh, ensure_ascii=False, indent=1)
    with open(os.path.join(out, "MAPA.md"), "w", encoding="utf-8", newline="\n") as fh:
        fh.write("# Mapa de contexto\n\n**MODO: relectura** · mapa anterior a la v1.1\n")


def run_test():
    case = json.loads((SKILL / "tests" / "sample-case.json").read_text(encoding="utf-8"))
    exp = json.loads((SKILL / "tests" / "expected.json").read_text(encoding="utf-8"))
    failures = []

    def check(cond, msg):
        print(("  ok    " if cond else "  FAIL  ") + msg)
        if not cond:
            failures.append(msg)

    def write_tree(root, tree):
        for rel, content in tree.items():
            path = os.path.join(root, *rel.split("/"))
            os.makedirs(os.path.dirname(path), exist_ok=True)
            with open(path, "w", encoding="utf-8", newline="\n") as fh:
                fh.write(content)

    tmp = tempfile.mkdtemp(prefix="map-context-")
    try:
        root = os.path.join(tmp, "task")
        write_tree(root, case["files"])
        has_xlsx = False
        try:
            import openpyxl  # noqa: PLC0415
            wb = openpyxl.Workbook()
            wb.active.title = "SUMMARY"
            wb.active["A1"] = "Fictitious purchase"
            wb.create_sheet("DETAIL")
            wb.save(os.path.join(root, "data", "sheet.xlsx"))
            has_xlsx = True
        except ImportError:
            pass
        out = os.path.join(root, "context")      # inside the root: it must exclude itself
        n_exp = exp["first"]["files"] + (1 if has_xlsx else 0)

        print("1 · first run")
        r = map_context([root], out)
        check(r["mode"] == "first-run", f"mode first-run (got {r['mode']})")
        check(r["files"] == n_exp, f"{n_exp} files (got {r['files']})")
        for s, t, rel in exp["first"]["links"]:
            check((s, t, rel) in r["link_list"], f"link {s} → {t} ({rel})")
        check(r["links"] == len(exp["first"]["links"]),
              f"{len(exp['first']['links'])} links in total (got {r['links']})")
        for i, frag in exp["first"]["summaries"].items():
            check(frag in r["summaries"].get(i, ""), f"summary of {i} contains \"{frag}\"")
        if has_xlsx:
            check("SUMMARY" in r["summaries"].get("data/sheet.xlsx", ""), "the xlsx summary lists its sheets")
        for h in exp["first"]["orphans"]:
            check(h in r["orphans"], f"orphan {h}")
        exc = [posixpath.basename(x) for x in r["excluded"].get("dirs", [])]
        for d in exp["first"]["excluded_dirs"]:
            check(d in exc, f"excluded folder {d}")
        for f in ("MAP.md", "graph.json", "fingerprints.json", "changes.json"):
            check(os.path.exists(os.path.join(out, f)), f"wrote {f}")
        g = read_json(os.path.join(out, "graph.json"))
        check(bool(g) and {"id", "label", "type"} <= set(g["nodes"][0])
              and {"source", "target", "relation"} <= set(g["links"][0]), "graph.json in nodes/links shape")
        map_md = Path(out, "MAP.md").read_text(encoding="utf-8")
        check(map_md.splitlines()[2].startswith("**MODE: first-run**"), "MAP.md opens with the MODE")
        check(r["words"] <= 3000, f"map ≤ 3,000 words ({r['words']})")
        core = [i for i, _ in r["core"]]
        check(core == exp["core"]["rule"], f"core by rule {exp['core']['rule']} (got {core})")
        check("## Core · always read" in map_md and map_md.index("## Core") < map_md.index("## Folders"),
              "MAP.md carries «Core · always read» at the top")
        print("  " + final_line(r))

        def edit(rel, text="\nChange.\n"):
            with open(os.path.join(root, *rel.split("/")), "a", encoding="utf-8") as fh:
                fh.write(text)

        def pending():
            ch = normalize_changes(read_json(os.path.join(out, "changes.json"))) or {}
            return [i for i, _ in net(ch.get("runs", []))]

        print("1b · another run without marking read")
        r = map_context([root], out)
        check(r["mode"] == "first-run" and "marked read" in r["reason"],
              f"stays first-run until it is marked read (got {r['mode']})")
        print("  " + final_line(r))

        print("2 · --mark-read and no changes")
        c = mark_read(out)
        check(c["base"] == n_exp and not c["runs"] and not c["first_run_pending"],
              f"mark read leaves base {n_exp} and nothing pending")
        r = map_context([root], out)
        check(r["mode"] == "reread" and r["n_changes"] == 0,
              f"reread with 0 changed (got {r['mode']}, {r['n_changes']})")
        check(r["excluded"].get("output") == ["context"] and r["files"] == n_exp,
              "the output folder (inside the root) is not mapped")
        print("  " + final_line(r))

        print("3 · one file changes")
        mod = exp["one_change"]["modify"]
        edit(mod, "\nA fourth quote arrived.\n")
        r = map_context([root], out)
        check(r["mode"] == "reread", f"mode reread (got {r['mode']})")
        check(r["changed"] == [mod] and not r["new"] and not r["deleted"],
              f"reread exactly {mod} (got {r['new'] + r['changed'] + r['deleted']})")
        map_md = Path(out, "MAP.md").read_text(encoding="utf-8")
        check(f"- `{mod}` (changed)" in map_md, "MAP.md lists it under «To reread»")
        print("  " + final_line(r))

        print("3b · another run without changes: what was not read stays")
        r = map_context([root], out)
        check(r["mode"] == "reread" and r["changed"] == [mod],
              f"{mod} stays pending (got {r['changed']})")
        check(r["pending_runs"] == 1, "an empty run is not recorded in changes.json")
        print("  " + final_line(r))

        print("3c · another one changes: it accumulates")
        other = exp["accumulated"]["modify"]
        edit(other)
        r = map_context([root], out, threshold=exp["accumulated"]["threshold"])
        check(r["mode"] == "reread" and r["changed"] == sorted([mod, other]),
              f"accumulates {sorted([mod, other])} (got {r['changed']})")
        check(pending() == sorted([mod, other]) and r["pending_runs"] == 2,
              "changes.json keeps the two unread runs")
        print("  " + final_line(r))

        print("3d · --mark-read empties what accumulated")
        mark_read(out)
        check(pending() == [], "changes.json is empty")
        r = map_context([root], out)
        check(r["mode"] == "reread" and r["n_changes"] == 0, f"reread with 0 (got {r['n_changes']})")
        print("  " + final_line(r))

        print("4 · many change")
        for rel in exp["many_changes"]["modify"]:
            edit(rel)
        for rel in exp["many_changes"]["delete"]:
            os.remove(os.path.join(root, *rel.split("/")))
        r = map_context([root], out)
        check(r["mode"] == "first-run", f"above the threshold it goes back to first-run (got {r['mode']}, "
              f"{_pct(r['ratio'])})")
        print("  " + final_line(r))
        r = map_context([root], out)
        check(r["mode"] == "first-run", "and stays first-run until it is marked read")

        print("5 · --core pins files by hand")
        r = map_context([root], out, core=exp["core"]["pin"])
        core = [(i, m) for i, m in r["core"]]
        check([i for i, _ in core] == exp["core"]["with_pin"] and core[0][1] == "pinned by hand",
              f"core {exp['core']['with_pin']} (got {[i for i, _ in core]})")
        try:
            map_context([root], out, core=["does/not-exist.md"])
            check(False, "--core with a path that does not exist aborts")
        except InputError:
            check(True, "--core with a path that does not exist aborts")

        print("6 · large tree: the map compresses")
        large = os.path.join(tmp, "large")
        nc, na = exp["large"]["folders"], exp["large"]["files_per_folder"]
        for c in range(nc):
            for k in range(na):
                path = os.path.join(large, f"area-{c:02d}", f"sub-{k % 3}", f"note-{k:03d}.md")
                os.makedirs(os.path.dirname(path), exist_ok=True)
                with open(path, "w", encoding="utf-8") as fh:
                    fh.write(f"# Note {k} of area {c}\n\nA fictitious note with enough text to "
                             f"fill the one-line summary of file number {k}. See [[note-{(k + 1) % na:03d}]].\n")
        r = map_context([large], os.path.join(tmp, "context-large"))
        check(r["files"] == nc * na, f"{nc * na} files (got {r['files']})")
        check(r["words"] <= 3000, f"map ≤ 3,000 words with {nc * na} files ({r['words']}, "
              f"level {r['level']})")
        print("  " + final_line(r))

        print("7 · a replaced file stays out of the core («Replaces:» and «Reemplaza a:»)")
        rep = exp["replaced"]
        for variant, tree in rep["trees"].items():
            vroot = os.path.join(tmp, f"replaced-{variant}")
            write_tree(vroot, tree)
            r = map_context([vroot], os.path.join(tmp, f"context-replaced-{variant}"))
            core = [i for i, _ in r["core"]]
            check(r["replaced"] == rep["replaced"][variant],
                  f"{variant}: detects {rep['replaced'][variant]} (got {r['replaced']})")
            check(core == rep["core"][variant], f"{variant}: core {rep['core'][variant]} (got {core})")
            old = next(iter(rep["replaced"][variant]))
            r = map_context([vroot], os.path.join(tmp, f"context-replaced-{variant}"), core=[old])
            check(old in [i for i, _ in r["core"]], f"{variant}: --core can still pin {old}")

        print("8 · a context written before v1.1 (Spanish names and keys) is read as it is")
        lroot = os.path.join(tmp, "legacy")
        write_tree(lroot, case["files"])
        lout = os.path.join(tmp, "legacy-out", "contexto")
        map_context([lroot], lout)
        mark_read(lout)
        fp = read_json(os.path.join(lout, "fingerprints.json"))
        _write_legacy_context(lout, fp, fp["roots"], len(fp["files"]))
        r = map_context([lroot], lout)
        check(r["layout"] == "es" and r["mode"] == "reread" and r["n_changes"] == 0,
              f"huellas.json + cambios.json: reread with 0 changes (got {r['layout']}, {r['mode']}, "
              f"{r['n_changes']})")
        names = sorted(os.listdir(lout))
        check(names == ["MAPA.md", "cambios.json", "grafo.json", "huellas.json"],
              f"keeps its Spanish file names, no English duplicates (got {names})")
        check(Path(lout, "MAPA.md").read_text(encoding="utf-8").splitlines()[2].startswith("**MODE: reread**"),
              "MAPA.md is rewritten with the current map")
        with open(os.path.join(lroot, *exp["one_change"]["modify"].split("/")), "a", encoding="utf-8") as fh:
            fh.write("\nChange.\n")
        r = map_context([lroot], lout)
        check(r["mode"] == "reread" and r["changed"] == [exp["one_change"]["modify"]],
              f"a change after the legacy context is detected (got {r['changed']})")
        c = mark_read(lout)
        check(c["base"] == exp["first"]["files"] and not os.path.exists(os.path.join(lout, "changes.json")),
              "--mark-read writes cambios.json in a legacy context")
        check(normalize_mode("relectura") == "reread" and normalize_mode("primera-vez") == "first-run",
              "the old mode words are accepted when reading")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    if failures:
        print(f"FAIL · {len(failures)} checks do not match")
        return 1
    print("OK")
    return 0


def _alias(ap, old, dest, **kw):
    """A pre-v1.1 Spanish flag that still works and does not show in --help."""
    ap.add_argument(old, dest=dest, help=argparse.SUPPRESS, default=argparse.SUPPRESS, **kw)


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description="Map folders: MAP.md + graph.json + fingerprints.json, no LLM.")
    ap.add_argument("folders", nargs="*", help="the task's folders")
    ap.add_argument("--out", help="context/ folder where the files are written")
    ap.add_argument("--exclude", action="append", default=[], metavar="PATTERN",
                    help="glob pattern on the name or the relative path (repeatable)")
    ap.add_argument("--threshold", type=float, default=0.20,
                    help="fraction of changes up to which the mode is reread")
    ap.add_argument("--max-words", type=int, default=3000)
    ap.add_argument("--max-mb", type=float, default=25.0, help="binaries larger than this are excluded")
    ap.add_argument("--graphify", action="store_true", help="adds the note on how to enrich with /graphify")
    ap.add_argument("--core", default="", metavar="PATH1,PATH2",
                    help="up to 3 files (map relative path) that are always read; the rule fills the rest")
    ap.add_argument("--mark-read", action="store_true",
                    help="empties the accumulated changes of <out>/changes.json (orchestrate calls it "
                         "when consolidating)")
    ap.add_argument("--test", action="store_true", help="runs the test on the sample case")
    _alias(ap, "--salida", "out")
    _alias(ap, "--excluir", "exclude", action="append")
    _alias(ap, "--umbral", "threshold", type=float)
    _alias(ap, "--max-palabras", "max_words", type=int)
    _alias(ap, "--nucleo", "core")
    _alias(ap, "--marcar-leido", "mark_read", action="store_true")
    _alias(ap, "--probar", "test", action="store_true")
    a = ap.parse_args(argv)
    if a.test:
        return run_test()
    if a.mark_read:
        if not a.out or a.folders:
            ap.error("--mark-read goes alone, with --out")
        try:
            c = mark_read(a.out)
        except InputError as e:
            print(f"ERROR: {e}", file=sys.stderr)
            return 2
        print(f"Marked read on {c['read_at']}: base {c['base']} files, accumulated changes at 0.")
        return 0
    if not a.folders or not a.out:
        ap.error("the folders or --out are missing")
    try:
        r = map_context(a.folders, a.out, exclude=a.exclude, threshold=a.threshold,
                        max_words=a.max_words, max_mb=a.max_mb, graphify=a.graphify,
                        core=[x for x in a.core.split(",") if x.strip()])
    except InputError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 2
    report(r)
    return 0


if __name__ == "__main__":
    sys.exit(main())
