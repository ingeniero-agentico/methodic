# -*- coding: utf-8 -*-
"""
costo-y-precio · costo real por ítem, precio por recargo y punto de equilibrio.

    python costo_precio.py <caso.json> [--perfil <ruta|nombre>] [--salida costo.xlsx] [--json res.json]
    python costo_precio.py --probar

Qué calcula (todo neto, sin impuesto al valor agregado):

  (a) Costo directo por ítem desde líneas de órdenes de compra o cotizaciones adjudicadas.
      Si el ítem se compró a varios proveedores, el costo es el PROMEDIO PONDERADO por
      cantidad. Si se compró en la unidad del proveedor (tira, rollo, caja) y sobra, el
      sobrante se carga al costo de la unidad vendida: costo total ÷ unidades vendidas.
  (b) Gastos generales como fracción del costo directo (costeo por absorción simplificado).
  (c) Precio por RECARGO sobre el costo (0,30 = costo × 1,30), por ítem o global, y al lado
      el MARGEN REAL sobre la venta: (precio − costo) ÷ precio. Nunca se mezclan los nombres.
  (d) Margen de contribución y punto de equilibrio de un servicio o producto; si se declaran
      unidades producidas y vendidas, la utilidad por costeo variable y por absorción.

Solo biblioteca estándar + openpyxl.
"""
import argparse
import decimal
import json
import math
import os
import sys
import tempfile

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover
    sys.exit("Falta openpyxl: pip install openpyxl")

AQUI = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(AQUI)
REPO = os.path.dirname(os.path.dirname(SKILL))
PLANTILLA_PERFIL = os.path.join(REPO, "perfiles", "_plantilla.yaml")

SUPUESTO = "SUPUESTO"


class ErrorCaso(Exception):
    """Un dato que falta o no cuadra: se pregunta, no se inventa."""


# --------------------------------------------------------------------------- perfil (YAML)

def _escalar(txt):
    t = txt.strip()
    if t == "" or t in ("~", "null", "Null", "NULL"):
        return None
    if t in ("[]",):
        return []
    if t in ("{}",):
        return {}
    if t in ("true", "True"):
        return True
    if t in ("false", "False"):
        return False
    if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
        return t[1:-1]
    try:
        return int(t)
    except ValueError:
        pass
    try:
        return float(t)
    except ValueError:
        return t


def _sin_comentario(linea):
    """Quita el comentario (# precedido de espacio) respetando comillas."""
    dentro = None
    for i, ch in enumerate(linea):
        if ch in "\"'":
            dentro = None if dentro == ch else (ch if dentro is None else dentro)
        elif ch == "#" and dentro is None and (i == 0 or linea[i - 1] in " \t"):
            return linea[:i].rstrip()
    return linea.rstrip()


def leer_yaml(texto):
    """
    Lector mínimo del subconjunto de YAML que usan los perfiles: mapas anidados por
    sangría, escalares, listas con «- » (de escalares o de mapas) y [] / {} vacíos.
    No hace falta PyYAML.
    """
    lineas = []
    for cruda in texto.splitlines():
        s = _sin_comentario(cruda)
        if s.strip():
            lineas.append((len(s) - len(s.lstrip(" ")), s.strip()))

    def bloque(i, sangria):
        if i < len(lineas) and lineas[i][1].startswith("- "):
            lista = []
            while i < len(lineas) and lineas[i][0] == sangria and lineas[i][1].startswith("- "):
                resto = lineas[i][1][2:]
                if ":" in resto and not resto.startswith(("'", '"')):
                    # mapa dentro de la lista: la primera clave va en la misma línea
                    sub = sangria + 2
                    lineas[i] = (sub, resto)
                    valor, i = bloque(i, sub)
                    lista.append(valor)
                else:
                    lista.append(_escalar(resto))
                    i += 1
            return lista, i
        mapa = {}
        while i < len(lineas) and lineas[i][0] == sangria:
            clave, _, valor = lineas[i][1].partition(":")
            clave = clave.strip().strip("\"'")
            i += 1
            if valor.strip() == "":
                if i < len(lineas) and lineas[i][0] > sangria:
                    mapa[clave], i = bloque(i, lineas[i][0])
                elif i < len(lineas) and lineas[i][0] == sangria and lineas[i][1].startswith("- "):
                    mapa[clave], i = bloque(i, sangria)
                else:
                    mapa[clave] = None
            else:
                mapa[clave] = _escalar(valor)
        return mapa, i

    if not lineas:
        return {}
    datos, _ = bloque(0, lineas[0][0])
    return datos


def _dirs_perfiles():
    """Dónde buscar perfiles, en orden: INGENIERO_DATOS y la carpeta hermana del repo."""
    dirs = []
    env = os.environ.get("INGENIERO_DATOS")
    if env:
        dirs.append(os.path.join(env, "perfiles"))
    dirs.append(os.path.join(os.path.dirname(REPO), "ingeniero-agentico-datos", "perfiles"))
    return dirs


def cargar_perfil(arg=None):
    """
    Devuelve (perfil, origen). Orden: --perfil (ruta o nombre) · INGENIERO_DATOS/perfiles ·
    ../ingeniero-agentico-datos/perfiles · plantilla del repo (y entonces todo lo que salga
    de ella queda marcado SUPUESTO).
    """
    candidatos = []
    if arg:
        if os.path.isfile(arg):
            candidatos.append(arg)
        else:
            nombre = arg if arg.endswith((".yaml", ".yml")) else arg + ".yaml"
            candidatos += [os.path.join(d, nombre) for d in _dirs_perfiles()]
        for c in candidatos:
            if os.path.isfile(c):
                with open(c, encoding="utf-8-sig") as fh:
                    return leer_yaml(fh.read()), os.path.basename(c)
        raise ErrorCaso(f"No encontré el perfil «{arg}» (busqué en: {', '.join(candidatos)})")
    with open(PLANTILLA_PERFIL, encoding="utf-8-sig") as fh:
        return leer_yaml(fh.read()), "plantilla"


def _g(d, *ruta):
    for k in ruta:
        if not isinstance(d, dict):
            return None
        d = d.get(k)
    return d


# --------------------------------------------------------------------------- cálculo

def redondear(x, dec):
    """
    Redondeo comercial (mitad hacia arriba), igual que ROUND de una planilla; round() de Python
    no lo es (redondea al par y arrastra el ruido binario: 18.932,355 saldría 18.932,35).
    """
    q = decimal.Decimal(1).scaleb(-dec)
    return float(decimal.Decimal(repr(round(x, 9))).quantize(q, rounding=decimal.ROUND_HALF_UP))


def _num(x, campo, donde, minimo=None, estricto=False):
    if x is None or x == "":
        raise ErrorCaso(f"{donde}: falta «{campo}»")
    try:
        v = float(x)
    except (TypeError, ValueError):
        raise ErrorCaso(f"{donde}: «{campo}» no es número ({x!r})")
    if minimo is not None and (v <= minimo if estricto else v < minimo):
        raise ErrorCaso(f"{donde}: «{campo}» = {v} fuera de rango")
    return v


def calcular(caso, perfil, origen_perfil="perfil"):
    avisos, supuestos = [], []
    de_plantilla = origen_perfil == "plantilla"
    etiqueta_perfil = f"perfil {origen_perfil}" if not de_plantilla else f"plantilla del repo · {SUPUESTO}"

    moneda = caso.get("moneda") or _g(perfil, "empresa", "moneda")
    if not moneda:
        raise ErrorCaso("No hay moneda: ni el caso ni el perfil la declaran")
    dec = caso.get("decimales_moneda", _g(perfil, "empresa", "decimales_moneda"))
    dec = int(dec) if dec is not None else (0 if moneda == "CLP" else 2)
    supuestos.append({"parametro": "moneda", "valor": f"{moneda} · {dec} decimales",
                      "origen": "caso" if caso.get("moneda") else etiqueta_perfil})

    # Recargo global: caso > perfil. Si no hay, cada ítem tiene que traer el suyo.
    if caso.get("recargo_global") is not None:
        recargo_global, origen_rg = _num(caso["recargo_global"], "recargo_global", "caso"), "caso"
    elif _g(perfil, "precios", "recargo_por_defecto") is not None:
        recargo_global = float(_g(perfil, "precios", "recargo_por_defecto"))
        origen_rg = etiqueta_perfil
    else:
        recargo_global, origen_rg = None, None
    if recargo_global is not None:
        supuestos.append({"parametro": "recargo global (sobre costo)", "valor": recargo_global,
                          "origen": origen_rg, "formato": "pct"})

    # Gastos generales: caso > perfil > 0 marcado SUPUESTO.
    if caso.get("gastos_generales") is not None:
        gg, origen_gg = _num(caso["gastos_generales"], "gastos_generales", "caso", 0), "caso"
    elif _g(perfil, "precios", "gastos_generales") is not None:
        gg, origen_gg = float(_g(perfil, "precios", "gastos_generales")), etiqueta_perfil
    else:
        gg, origen_gg = 0.0, f"{SUPUESTO}: ni el caso ni el perfil los declaran; se usa 0 %"
        avisos.append("Gastos generales no declarados: se calculó con 0 % (SUPUESTO). Pregúntalos.")
    supuestos.append({"parametro": "gastos generales (fracción del costo directo)", "valor": gg,
                      "origen": origen_gg, "formato": "pct"})

    # Tipos de cambio declarados: nunca se inventan.
    tcs = {}
    for mon, d in (caso.get("tipos_de_cambio") or {}).items():
        if mon.startswith("_"):
            continue
        if not isinstance(d, dict):
            d = {"valor": d}
        v = _num(d.get("valor"), "valor", f"tipo de cambio {mon}", 0, estricto=True)
        fuente = d.get("fuente") or ""
        fecha = d.get("fecha") or ""
        tcs[mon] = v
        origen = f"caso · {fuente or 'fuente no declarada'}" + (f" · {fecha}" if fecha else "")
        if not fuente or not fecha:
            avisos.append(f"Tipo de cambio {mon}: falta {'fuente' if not fuente else 'fecha'}; decláralo.")
        supuestos.append({"parametro": f"tipo de cambio {mon}/{moneda}", "valor": redondear(v, 2),
                          "origen": origen, "formato": "tc"})

    compras = caso.get("compras") or []
    if not compras:
        raise ErrorCaso("El caso no trae «compras» (líneas de OC o cotizaciones adjudicadas)")

    # Ítems: si el caso no los declara, se crea uno por cada «item» de las compras.
    items_decl = caso.get("items")
    if items_decl:
        items = {}
        for k, it in enumerate(items_decl, start=1):
            if not it.get("id"):
                raise ErrorCaso(f"items[{k}]: falta «id»")
            if it["id"] in items:
                raise ErrorCaso(f"items: «{it['id']}» repetido")
            items[it["id"]] = dict(it)
    else:
        items = {}
        for l in compras:
            iid = l.get("item")
            if iid and iid not in items:
                items[iid] = {"id": iid, "descripcion": l.get("descripcion") or iid,
                              "unidad_venta": l.get("unidad_venta") or l.get("unidad") or "un"}

    lineas = []
    for k, l in enumerate(compras, start=1):
        donde = f"compras[{k}] ({l.get('doc', '?')} línea {l.get('linea', '?')})"
        iid = l.get("item")
        if not iid:
            raise ErrorCaso(f"{donde}: falta «item» (a qué ítem de venta se carga)")
        if iid not in items:
            raise ErrorCaso(f"{donde}: el ítem «{iid}» no está en «items»")
        if not l.get("doc"):
            raise ErrorCaso(f"{donde}: falta «doc» (OC o cotización de origen)")
        cant = _num(l.get("cantidad"), "cantidad", donde, 0, estricto=True)
        precio = _num(l.get("precio_unit"), "precio_unit", donde, 0)
        factor = _num(l.get("factor_unidad", 1), "factor_unidad", donde, 0, estricto=True)
        mon = l.get("moneda") or moneda
        if mon == moneda:
            tc = 1.0
        elif mon in tcs:
            tc = tcs[mon]
        else:
            raise ErrorCaso(f"{donde}: viene en {mon} y el caso no declara tipo de cambio {mon}")
        total_orig = cant * precio
        total_local = total_orig * tc
        lineas.append({
            "doc": str(l["doc"]), "linea": l.get("linea", ""), "proveedor": l.get("proveedor", ""),
            "item": iid, "descripcion": l.get("descripcion") or items[iid].get("descripcion", iid),
            "cantidad": cant, "unidad": l.get("unidad", "un"), "factor_unidad": factor,
            "cantidad_uv": cant * factor, "precio_unit": precio, "moneda": mon, "tc": tc,
            "precio_local": precio * tc, "costo_uv_local": precio * tc / factor,
            "total_orig": total_orig, "total_local": total_local,
        })

    margen_alerta = caso.get("margen_minimo_alerta")
    filas = []
    for iid, it in items.items():
        ls = [x for x in lineas if x["item"] == iid]
        donde = f"ítem «{iid}»"
        if not ls:
            raise ErrorCaso(f"{donde}: no tiene ninguna línea de compra; sin costo no hay precio")
        total = sum(x["total_local"] for x in ls)
        comprado = sum(x["cantidad_uv"] for x in ls)
        if it.get("cantidad_venta") not in (None, ""):
            venta_q = _num(it["cantidad_venta"], "cantidad_venta", donde, 0, estricto=True)
            origen_q = "caso"
        else:
            venta_q, origen_q = comprado, "= lo comprado (sin sobrante)"
        if venta_q > comprado + 1e-9:
            raise ErrorCaso(f"{donde}: se venden {venta_q:g} y se compraron {comprado:g} "
                            f"{it.get('unidad_venta', '')}: la compra no cubre la venta")
        sobrante = comprado - venta_q
        for x in ls:
            x["peso"] = x["total_local"] / total if total else 0.0
        pond = total / comprado
        cdu = total / venta_q
        ctu = cdu * (1 + gg)
        manual = it.get("precio_manual")
        if manual not in (None, ""):
            precio = _num(manual, "precio_manual", donde, 0, estricto=True)
            recargo = precio / ctu - 1
            origen_r = "precio manual (el recargo se recalcula)"
        else:
            if it.get("recargo") not in (None, ""):
                recargo, origen_r = _num(it["recargo"], "recargo", donde, -1, estricto=True), "ítem"
            elif recargo_global is not None:
                recargo, origen_r = recargo_global, f"global ({origen_rg})"
            else:
                raise ErrorCaso(f"{donde}: no tiene recargo y no hay recargo global en el caso ni en el perfil")
            precio = redondear(ctu * (1 + recargo), dec)
        margen = (precio - ctu) / precio if precio else None
        unit_uv = [x["costo_uv_local"] for x in ls]
        provs = sorted({x["proveedor"] for x in ls if x["proveedor"]})
        disp = (max(unit_uv) - min(unit_uv)) / min(unit_uv) if min(unit_uv) else 0.0
        if len(provs) > 1 and disp > 0.15:
            avisos.append(f"{donde}: entre proveedores el costo por unidad difiere {disp:.1%}; "
                          "confirma que sea el mismo producto")
        if sobrante > 1e-9:
            avisos.append(f"{donde}: sobran {sobrante:g} {it.get('unidad_venta', '')} "
                          f"(se compró en la unidad del proveedor); su costo va en la unidad vendida")
        if margen_alerta is not None and margen is not None and margen < float(margen_alerta):
            avisos.append(f"{donde}: margen real {margen:.1%} bajo el mínimo de alerta {float(margen_alerta):.1%}")
        filas.append({
            "id": iid, "descripcion": it.get("descripcion", iid), "unidad_venta": it.get("unidad_venta", "un"),
            "cantidad_venta": venta_q, "origen_cantidad": origen_q, "cantidad_comprada_uv": comprado,
            "sobrante_uv": sobrante, "costo_sobrante": sobrante * pond, "costo_directo_total": total, "costo_ponderado_uv": pond,
            "costo_directo_unit": cdu, "gastos_generales": gg, "costo_total_unit": ctu,
            "recargo": recargo, "origen_recargo": origen_r, "precio_unit": precio, "margen_real": margen,
            "venta_total": precio * venta_q, "costo_total": ctu * venta_q,
            "utilidad": precio * venta_q - ctu * venta_q, "proveedores": provs,
            "n_proveedores": len(provs), "dispersion_costo": disp,
            "documentos": sorted({f"{x['doc']}:{x['linea']}" for x in ls}),
        })

    cd = sum(f["costo_directo_total"] for f in filas)
    ct = sum(f["costo_total"] for f in filas)
    vt = sum(f["venta_total"] for f in filas)
    totales = {"costo_directo": cd, "gastos_generales": ct - cd, "costo_total": ct,
               "venta_total": vt, "utilidad": vt - ct,
               "margen_global": (vt - ct) / vt if vt else None,
               "recargo_efectivo": vt / ct - 1 if ct else None}

    equilibrio = [punto_equilibrio(e) for e in (caso.get("equilibrio") or [])]

    if de_plantilla:
        avisos.insert(0, "Se usó la PLANTILLA de perfil: moneda, recargo y gastos son SUPUESTOS. "
                         "Pasa --perfil con el de la empresa.")
    return {"caso": caso.get("caso", ""), "perfil": origen_perfil, "moneda": moneda, "decimales": dec,
            "items": filas, "lineas": lineas, "totales": totales, "equilibrio": equilibrio,
            "gg_supuesto": SUPUESTO in str(origen_gg), "supuestos": supuestos, "avisos": avisos}


def punto_equilibrio(e):
    """
    Margen de contribución y punto de equilibrio. Con unidades producidas y vendidas,
    además la utilidad por costeo variable y por absorción (sin inventario inicial).
    """
    nom = e.get("nombre", "servicio")
    donde = f"equilibrio «{nom}»"
    p = _num(e.get("precio_unit"), "precio_unit", donde, 0, estricto=True)
    cv = _num(e.get("costo_variable_unit"), "costo_variable_unit", donde, 0)
    cfp = _num(e.get("costos_fijos_produccion", 0), "costos_fijos_produccion", donde, 0)
    cfo = _num(e.get("costos_fijos_operacion", 0), "costos_fijos_operacion", donde, 0)
    mc = p - cv
    if mc <= 0:
        raise ErrorCaso(f"{donde}: el precio no cubre el costo variable (margen de contribución {mc:g}); no hay equilibrio")
    cf = cfp + cfo
    pe = cf / mc
    r = {"nombre": nom, "unidad": e.get("unidad", "un"), "precio_unit": p, "costo_variable_unit": cv,
         "costos_fijos_produccion": cfp, "costos_fijos_operacion": cfo, "costos_fijos": cf,
         "margen_contribucion_unit": mc, "razon_contribucion": mc / p,
         "pe_unidades": pe, "pe_unidades_enteras": math.ceil(pe - 1e-9), "pe_ventas": pe * p}
    v = e.get("volumen_esperado")
    if v not in (None, ""):
        v = _num(v, "volumen_esperado", donde, 0, estricto=True)
        r.update({"volumen_esperado": v, "utilidad_esperada": mc * v - cf,
                  "margen_seguridad_unidades": v - pe, "margen_seguridad": (v - pe) / v,
                  "costo_unit_absorcion": cv + cfp / v})
    qp, qv = e.get("unidades_producidas"), e.get("unidades_vendidas")
    if qp not in (None, "") and qv not in (None, ""):
        qp = _num(qp, "unidades_producidas", donde, 0, estricto=True)
        qv = _num(qv, "unidades_vendidas", donde, 0)
        if qv > qp:
            raise ErrorCaso(f"{donde}: se venden {qv:g} y se producen {qp:g}; sin inventario inicial no cuadra")
        cu_abs = cv + cfp / qp
        u_var = mc * qv - cf
        u_abs = (p - cu_abs) * qv - cfo
        inv = qp - qv
        r.update({"unidades_producidas": qp, "unidades_vendidas": qv, "costo_unit_absorcion": cu_abs,
                  "utilidad_variable": u_var, "utilidad_absorcion": u_abs,
                  "diferencia_utilidad": u_abs - u_var, "fijo_en_inventario": cfp / qp * inv,
                  "inventario_final_variable": cv * inv, "inventario_final_absorcion": cu_abs * inv})
    return r


# --------------------------------------------------------------------------- Excel

AZUL = "1F3A5F"
F_TIT = Font(bold=True, size=13, color="FFFFFF")
F_HEAD = Font(bold=True, size=10, color="FFFFFF")
F_BOLD = Font(bold=True)
F_NOTA = Font(italic=True, size=9)
FILL_TIT = PatternFill("solid", fgColor=AZUL)
FILL_SUP = PatternFill("solid", fgColor="FCE4D6")    # naranjo = SUPUESTO
FILL_PARTIDO = PatternFill("solid", fgColor="DCE6F1")  # azul claro = comprado a varios proveedores
PCT = "0.0%"


def _fmt_dinero(dec):
    return "#,##0" if dec == 0 else "#,##0." + "0" * dec


def _cabecera(ws, fila, cols):
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=fila, column=i, value=h)
        c.font, c.fill = F_HEAD, FILL_TIT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 42


def _titulo(ws, texto, ncols, nota):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=texto)
    c.font, c.fill = F_TIT, FILL_TIT
    ws.row_dimensions[1].height = 24
    ws.cell(row=2, column=1, value=nota).font = F_NOTA


def escribir_excel(res, ruta):
    dec = res["decimales"]
    DIN = _fmt_dinero(dec)
    DIN2 = "#,##0.00"
    mon = res["moneda"]
    wb = openpyxl.Workbook()

    # ---- COSTO Y PRECIO
    ws = wb.active
    ws.title = "COSTO Y PRECIO"
    cols = ["ÍTEM", "DESCRIPCIÓN", "UN. VENTA", "CANT. VENDIDA", "CANT. COMPRADA\n(un. venta)",
            "SOBRANTE", "PROV.", "COSTO DIRECTO\nTOTAL", "COSTO PONDERADO\n$/un comprada",
            "COSTO DIRECTO\n$/un vendida", "GASTOS\nGENERALES %", "COSTO TOTAL\n$/un",
            "RECARGO %\n(sobre costo)", "ORIGEN DEL\nRECARGO", "PRECIO UNIT.\nVENTA",
            "MARGEN REAL\n(sobre venta)", "VENTA TOTAL", "UTILIDAD", "DE DÓNDE VIENE\n(doc:línea)",
            "COSTO DEL\nSOBRANTE"]
    _titulo(ws, f"COSTO Y PRECIO · {res['caso']}", len(cols),
            f"Interno: lleva costos. Montos netos en {mon}. RECARGO = sobre el costo (0,30 → costo × 1,30); "
            "MARGEN REAL = (precio − costo) ÷ precio. Fondo azul: comprado a más de un proveedor (costo ponderado). "
            "Naranjo: SUPUESTO.")
    _cabecera(ws, 4, cols)
    r = 5
    for f in res["items"]:
        vals = [f["id"], f["descripcion"], f["unidad_venta"], f["cantidad_venta"], f["cantidad_comprada_uv"],
                f["sobrante_uv"], f["n_proveedores"], redondear(f["costo_directo_total"], 2),
                redondear(f["costo_ponderado_uv"], 2), redondear(f["costo_directo_unit"], 2), f["gastos_generales"],
                redondear(f["costo_total_unit"], 2), f["recargo"], f["origen_recargo"], f["precio_unit"],
                f["margen_real"], redondear(f["venta_total"], 2), redondear(f["utilidad"], 2), " · ".join(f["documentos"]),
                redondear(f["costo_sobrante"], 2)]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)
        for c in (8, 17, 18, 20):
            ws.cell(row=r, column=c).number_format = DIN
        for c in (9, 10, 12):
            ws.cell(row=r, column=c).number_format = DIN2
        ws.cell(row=r, column=15).number_format = DIN
        for c in (11, 13, 16):
            ws.cell(row=r, column=c).number_format = PCT
        if f["n_proveedores"] > 1:
            for c in (1, 2, 9):
                ws.cell(row=r, column=c).fill = FILL_PARTIDO
        if res["gg_supuesto"]:
            ws.cell(row=r, column=11).fill = FILL_SUP
        ws.cell(row=r, column=15).font = F_BOLD
        r += 1
    t = res["totales"]
    r += 1
    for etiqueta, col, valor, fmt in (("COSTO DIRECTO", 8, t["costo_directo"], DIN),
                                      ("COSTO TOTAL (con gastos generales)", 12, t["costo_total"], DIN),
                                      ("VENTA TOTAL", 17, t["venta_total"], DIN),
                                      ("UTILIDAD", 18, t["utilidad"], DIN)):
        ws.cell(row=r, column=2, value=etiqueta).font = F_BOLD
        c = ws.cell(row=r, column=col, value=redondear(valor, 2))
        c.number_format, c.font = fmt, F_BOLD
        r += 1
    ws.cell(row=r, column=2, value="RECARGO EFECTIVO (venta ÷ costo total − 1)").font = F_BOLD
    c = ws.cell(row=r, column=13, value=t["recargo_efectivo"])
    c.number_format, c.font = PCT, F_BOLD
    r += 1
    ws.cell(row=r, column=2, value="MARGEN REAL GLOBAL (utilidad ÷ venta)").font = F_BOLD
    c = ws.cell(row=r, column=16, value=t["margen_global"])
    c.number_format, c.font = PCT, F_BOLD
    for col, w in zip("ABCDEFGHIJKLMNOPQRST",
                      [12, 40, 9, 10, 12, 9, 6, 14, 15, 14, 10, 13, 10, 22, 13, 12, 14, 13, 28, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"

    # ---- TRAZABILIDAD
    wt = wb.create_sheet("TRAZABILIDAD")
    cols = ["DOC", "LÍNEA", "PROVEEDOR", "ÍTEM", "DESCRIPCIÓN EN EL DOC", "CANT.", "UN. PROVEEDOR",
            "FACTOR\n(un. venta por un.)", "CANT.\n(un. venta)", "PRECIO UNIT.\nORIGINAL", "MONEDA",
            "TIPO DE\nCAMBIO", f"PRECIO UNIT.\n{mon}", "TOTAL\nORIGINAL", f"TOTAL\n{mon}", "PESO EN\nEL ÍTEM"]
    _titulo(wt, "TRAZABILIDAD · de qué documento y línea viene cada costo", len(cols),
            "Una fila por línea de OC o cotización adjudicada. Moneda extranjera con 2 decimales y el tipo de "
            "cambio declarado en el caso (hoja SUPUESTOS).")
    _cabecera(wt, 4, cols)
    r = 5
    for x in res["lineas"]:
        ext = x["moneda"] != mon
        vals = [x["doc"], x["linea"], x["proveedor"], x["item"], x["descripcion"], x["cantidad"], x["unidad"],
                x["factor_unidad"], x["cantidad_uv"], x["precio_unit"], x["moneda"], x["tc"],
                redondear(x["precio_local"], 2), redondear(x["total_orig"], 2), redondear(x["total_local"], 2), x.get("peso")]
        for i, v in enumerate(vals, start=1):
            wt.cell(row=r, column=i, value=v)
        wt.cell(row=r, column=10).number_format = DIN2
        wt.cell(row=r, column=12).number_format = DIN2
        wt.cell(row=r, column=13).number_format = DIN2
        wt.cell(row=r, column=14).number_format = DIN2 if ext else DIN
        wt.cell(row=r, column=15).number_format = DIN2
        wt.cell(row=r, column=16).number_format = PCT
        r += 1
    for col, w in zip("ABCDEFGHIJKLMNOP", [12, 7, 22, 12, 40, 8, 12, 11, 10, 13, 8, 10, 13, 13, 14, 9]):
        wt.column_dimensions[col].width = w
    wt.freeze_panes = "A5"

    # ---- EQUILIBRIO
    if res["equilibrio"]:
        we = wb.create_sheet("EQUILIBRIO")
        _titulo(we, "MARGEN DE CONTRIBUCIÓN Y PUNTO DE EQUILIBRIO", 2 + len(res["equilibrio"]),
                "Costeo variable: el fijo es gasto del período. Costeo por absorción: el fijo de producción se "
                "reparte en las unidades producidas y lo no vendido queda en inventario.")
        _cabecera(we, 4, ["CONCEPTO", "UNIDAD"] + [e["nombre"] for e in res["equilibrio"]])
        filas = [("Precio unitario", "$/un", "precio_unit", DIN),
                 ("Costo variable unitario", "$/un", "costo_variable_unit", DIN),
                 ("Margen de contribución unitario", "$/un", "margen_contribucion_unit", DIN),
                 ("Razón de contribución", "%", "razon_contribucion", PCT),
                 ("Costos fijos de producción (período)", "$", "costos_fijos_produccion", DIN),
                 ("Costos fijos de operación (período)", "$", "costos_fijos_operacion", DIN),
                 ("Punto de equilibrio", "un", "pe_unidades", "#,##0.00"),
                 ("Punto de equilibrio (unidades enteras)", "un", "pe_unidades_enteras", "#,##0"),
                 ("Ventas de equilibrio", "$", "pe_ventas", DIN),
                 ("Volumen esperado", "un", "volumen_esperado", "#,##0.##"),
                 ("Utilidad al volumen esperado", "$", "utilidad_esperada", DIN),
                 ("Margen de seguridad", "%", "margen_seguridad", PCT),
                 ("Costo unitario por absorción", "$/un", "costo_unit_absorcion", "#,##0.00"),
                 ("Unidades producidas", "un", "unidades_producidas", "#,##0.##"),
                 ("Unidades vendidas", "un", "unidades_vendidas", "#,##0.##"),
                 ("Utilidad · costeo variable", "$", "utilidad_variable", DIN),
                 ("Utilidad · costeo por absorción", "$", "utilidad_absorcion", DIN),
                 ("Diferencia (absorción − variable)", "$", "diferencia_utilidad", DIN),
                 ("Fijo de producción que quedó en inventario", "$", "fijo_en_inventario", DIN)]
        r = 5
        for nombre, un, clave, fmt in filas:
            if not any(clave in e for e in res["equilibrio"]):
                continue
            we.cell(row=r, column=1, value=nombre)
            we.cell(row=r, column=2, value=un)
            for j, e in enumerate(res["equilibrio"], start=3):
                if clave in e:
                    v = e[clave]
                    c = we.cell(row=r, column=j, value=round(v, 4) if isinstance(v, float) else v)
                    c.number_format = fmt
            r += 1
        we.column_dimensions["A"].width = 44
        we.column_dimensions["B"].width = 8
        for j in range(len(res["equilibrio"])):
            we.column_dimensions[chr(ord("C") + j)].width = 22

    # ---- SUPUESTOS
    wsu = wb.create_sheet("SUPUESTOS")
    _titulo(wsu, "SUPUESTOS Y PARÁMETROS · de dónde salió cada uno", 3,
            f"Perfil usado: {res['perfil']}. Lo marcado SUPUESTO no vino del usuario: confírmalo antes de cotizar.")
    _cabecera(wsu, 4, ["PARÁMETRO", "VALOR", "ORIGEN"])
    r = 5
    for s in res["supuestos"]:
        wsu.cell(row=r, column=1, value=s["parametro"])
        c = wsu.cell(row=r, column=2, value=s["valor"])
        if s.get("formato") == "pct":
            c.number_format = PCT
        elif s.get("formato") == "tc":
            c.number_format = "#,##0.00"
        o = wsu.cell(row=r, column=3, value=s["origen"])
        if SUPUESTO in str(s["origen"]):
            for cc in (1, 2, 3):
                wsu.cell(row=r, column=cc).fill = FILL_SUP
        r += 1
    if res["avisos"]:
        r += 1
        wsu.cell(row=r, column=1, value="AVISOS").font = F_BOLD
        for a in res["avisos"]:
            r += 1
            wsu.cell(row=r, column=1, value=a)
    wsu.column_dimensions["A"].width = 48
    wsu.column_dimensions["B"].width = 18
    wsu.column_dimensions["C"].width = 60

    wb.save(ruta)


# --------------------------------------------------------------------------- salida y prueba

def resumen_json(res):
    """Lo comparable: redondeado como se muestra."""
    dec = res["decimales"]
    out = {"moneda": res["moneda"], "items": {}, "totales": {}, "equilibrio": {}}
    for f in res["items"]:
        out["items"][f["id"]] = {
            "cantidad_venta": f["cantidad_venta"], "cantidad_comprada_uv": f["cantidad_comprada_uv"],
            "sobrante_uv": round(f["sobrante_uv"], 6), "costo_sobrante": redondear(f["costo_sobrante"], 2),
            "n_proveedores": f["n_proveedores"],
            "costo_directo_total": redondear(f["costo_directo_total"], 2),
            "costo_ponderado_uv": redondear(f["costo_ponderado_uv"], 2),
            "costo_directo_unit": redondear(f["costo_directo_unit"], 2),
            "costo_total_unit": redondear(f["costo_total_unit"], 2),
            "recargo": round(f["recargo"], 6), "origen_recargo": f["origen_recargo"],
            "precio_unit": redondear(f["precio_unit"], dec), "margen_real": round(f["margen_real"], 6),
            "venta_total": redondear(f["venta_total"], 2), "utilidad": redondear(f["utilidad"], 2)}
    out["totales"] = {k: (round(v, 6) if k in ("margen_global", "recargo_efectivo") else redondear(v, 2))
                      for k, v in res["totales"].items() if v is not None}
    for e in res["equilibrio"]:
        out["equilibrio"][e["nombre"]] = {k: (round(v, 6) if isinstance(v, float) else v)
                                          for k, v in e.items() if k not in ("nombre", "unidad")}
    return out


def comparar(obtenido, esperado, ruta="", fallas=None):
    fallas = [] if fallas is None else fallas
    if isinstance(esperado, dict):
        for k, v in esperado.items():
            if k.startswith("_"):
                continue
            if not isinstance(obtenido, dict) or k not in obtenido:
                fallas.append(f"{ruta}/{k}: falta en lo obtenido")
            else:
                comparar(obtenido[k], v, f"{ruta}/{k}", fallas)
    elif isinstance(esperado, (int, float)) and not isinstance(esperado, bool):
        tol = 1e-4 if abs(esperado) < 10 else 0.01
        if not isinstance(obtenido, (int, float)) or abs(obtenido - esperado) > tol:
            fallas.append(f"{ruta}: esperado {esperado}, obtenido {obtenido}")
    elif obtenido != esperado:
        fallas.append(f"{ruta}: esperado {esperado!r}, obtenido {obtenido!r}")
    return fallas


def imprimir(res):
    dec = res["decimales"]
    print(f"\nCaso: {res['caso']} · perfil: {res['perfil']} · moneda {res['moneda']}")
    w = min(max([4] + [len(str(f["id"])) for f in res["items"]]), 34)
    print(f"{'ÍTEM':<{w}} {'CANT':>8} {'SOBRA':>6} {'COSTO/u':>12} {'RECARGO':>8} {'PRECIO':>12} {'MARGEN':>7}")
    for f in res["items"]:
        print(f"{str(f['id'])[:w]:<{w}} {f['cantidad_venta']:>8g} {f['sobrante_uv']:>6g} {f['costo_total_unit']:>12,.2f} "
              f"{f['recargo']:>8.1%} {f['precio_unit']:>12,.{dec}f} {f['margen_real']:>7.1%}")
    t = res["totales"]
    print(f"Costo total {t['costo_total']:,.2f} · venta {t['venta_total']:,.{dec}f} · "
          f"recargo efectivo {t['recargo_efectivo']:.1%} · margen real {t['margen_global']:.1%}")
    for e in res["equilibrio"]:
        print(f"Equilibrio {e['nombre']}: MC {e['margen_contribucion_unit']:,.2f}/{e['unidad']} · "
              f"PE {e['pe_unidades']:,.2f} {e['unidad']} ({e['pe_unidades_enteras']} enteras)")
    for a in res["avisos"]:
        print(f"  [!] {a}")


def probar():
    dir_p = os.path.join(SKILL, "pruebas")
    with open(os.path.join(dir_p, "caso-ficticio.json"), encoding="utf-8") as fh:
        caso = json.load(fh)
    with open(os.path.join(dir_p, "esperado.json"), encoding="utf-8") as fh:
        esperado = json.load(fh)
    perfil, origen = cargar_perfil(os.path.join(dir_p, "perfil-ficticio.yaml"))
    res = calcular(caso, perfil, origen)
    fallas = comparar(resumen_json(res), esperado)

    # La plantilla del repo tiene que poder leerse.
    plantilla, _ = cargar_perfil(None)
    if _g(plantilla, "precios", "recargo_por_defecto") != 0.30:
        fallas.append("plantilla de perfil: no se leyó precios.recargo_por_defecto")

    # Los errores que no se pueden tragar en silencio.
    for nombre, mod in (("sin tipo de cambio", lambda c: c.pop("tipos_de_cambio")),
                        ("compra no cubre la venta", lambda c: c["items"][1].update(cantidad_venta=43)),
                        ("línea sin documento", lambda c: c["compras"][0].pop("doc"))):
        c2 = json.loads(json.dumps(caso))
        mod(c2)
        try:
            calcular(c2, perfil, origen)
            fallas.append(f"no abortó: {nombre}")
        except ErrorCaso:
            pass

    # El Excel se escribe y trae lo que dice.
    with tempfile.TemporaryDirectory() as tmp:
        ruta = os.path.join(tmp, "prueba.xlsx")
        escribir_excel(res, ruta)
        wb = openpyxl.load_workbook(ruta)
        hojas = set(wb.sheetnames)
        for h in ("COSTO Y PRECIO", "TRAZABILIDAD", "EQUILIBRIO", "SUPUESTOS"):
            if h not in hojas:
                fallas.append(f"Excel: falta la hoja {h}")
        ws = wb["COSTO Y PRECIO"]
        precios = {ws.cell(r, 1).value: ws.cell(r, 15).value for r in range(5, 5 + len(res["items"]))}
        for iid, e in esperado["items"].items():
            if precios.get(iid) != e["precio_unit"]:
                fallas.append(f"Excel: precio de {iid} = {precios.get(iid)}, esperado {e['precio_unit']}")
        wb.close()

    imprimir(res)
    if fallas:
        print("\nFALLA")
        for f in fallas:
            print("  -", f)
        return 1
    print(f"\nOK · {len(esperado['items'])} ítems, totales, {len(esperado['equilibrio'])} equilibrios, "
          "3 errores provocados y el Excel cuadran con esperado.json")
    return 0


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Costo real por ítem, precio por recargo y punto de equilibrio")
    ap.add_argument("caso", nargs="?", help="caso.json (ver plantillas/caso.json)")
    ap.add_argument("--perfil", help="ruta al perfil YAML o nombre de la empresa (busca en INGENIERO_DATOS "
                                     "y en ../ingeniero-agentico-datos/perfiles)")
    ap.add_argument("--salida", help="Excel de salida (.xlsx)")
    ap.add_argument("--json", help="escribe el resumen comparable en este JSON")
    ap.add_argument("--probar", action="store_true", help="corre la prueba del caso ficticio")
    a = ap.parse_args(argv)
    if a.probar:
        return probar()
    if not a.caso:
        ap.error("falta el caso (o usa --probar)")
    try:
        with open(a.caso, encoding="utf-8-sig") as fh:
            caso = json.load(fh)
        perfil, origen = cargar_perfil(a.perfil or caso.get("perfil"))
        res = calcular(caso, perfil, origen)
    except ErrorCaso as e:
        print(f"ERROR: {e}")
        return 2
    imprimir(res)
    salida = a.salida or os.path.splitext(a.caso)[0] + "_costo-y-precio.xlsx"
    escribir_excel(res, salida)
    print(f"\nExcel: {salida}")
    if a.json:
        with open(a.json, "w", encoding="utf-8") as fh:
            json.dump(resumen_json(res), fh, ensure_ascii=False, indent=2)
        print(f"JSON:  {a.json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
