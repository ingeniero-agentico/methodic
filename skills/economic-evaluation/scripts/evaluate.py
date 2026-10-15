#!/usr/bin/env python3
"""economic-evaluation · is the project worth doing? (pure project, no financing)

Builds the pure project cash flow (Chilean-style layout: non-cash charges added back, book value,
market salvage value, working capital) and computes NPV, IRR, payback, NPV index and a
one-variable sensitivity (each variable ±x % and the critical value that makes NPV zero).
Writes an Excel workbook and a JSON file. Sheet names, labels, the assumption marker and the
console text follow the profile `language` (es by default: SUPUESTOS · FLUJO · INDICADORES ·
SENSIBILIDAD; en: ASSUMPTIONS · CASH FLOW · INDICATORS · SENSITIVITY).

Usage:
    python evaluate.py input.json [--profile <name|path>] [--rate 0.12] [--tax 0.27]
                       [--delta 0.10] [--sweep "revenues[0].occupancy=0.3,0.5,0.7"]
                       [--out folder]
    python evaluate.py --test

Spanish flags still work as hidden aliases (--perfil --tasa --impuesto --barrer --salida
--probar), and Spanish input and profile keys are accepted through an alias map.

Standard library + openpyxl (openpyxl only for the Excel file).
"""
from __future__ import annotations

import argparse
import copy
import json
import math
import os
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
SKILL = HERE.parent
REPO = SKILL.parent.parent

# ---------------------------------------------------------------------------
# Output language (profile `language`, default es)
# ---------------------------------------------------------------------------

_LANG = "es"


def tr(es: str, en: str) -> str:
    """Text in the current output language."""
    return en if _LANG == "en" else es


def mark(assumption: bool) -> str:
    if assumption:
        return tr("SUPUESTO", "ASSUMPTION")
    return tr("dato", "data")


class InputError(Exception):
    pass


# ---------------------------------------------------------------------------
# Spanish → English alias maps (input, profile, values)
# ---------------------------------------------------------------------------

INPUT_KEYS = {
    "proyecto": "project", "empresa": "company", "nota": "note", "notas": "notes", "fecha": "date",
    "horizonte_anios": "horizon_years", "perdidas_tributarias": "tax_losses",
    "metodo_desecho": "salvage_method", "ingresos": "revenues", "costos_variables": "variable_costs",
    "costos_fijos": "fixed_costs", "inversiones": "investments", "capital_trabajo": "working_capital",
    "nombre": "name", "unidad": "unit", "precio": "price", "cantidad": "quantity",
    "capacidad": "capacity", "ocupacion": "occupancy", "crecimiento_precio": "price_growth",
    "crecimiento_cantidad": "quantity_growth", "fuente": "source", "supuesto": "assumption",
    "por_unidad_de": "per_unit_of", "monto_por_unidad": "amount_per_unit", "crecimiento": "growth",
    "monto_anual": "annual_amount", "anio": "year", "monto": "amount", "vida_util": "useful_life",
    "vida_util_acelerada": "accelerated_life", "valor_desecho": "salvage_value",
    "recupera": "recovered", "valor": "value", "depreciable": "depreciable", "investigacion": "research",
}
# Keys whose string value is a code with a Spanish alias.
VALUE_KEYS = {"tax_losses", "salvage_method", "depreciation"}
VALUES = {"arrastra": "carry_forward", "compensa": "offset", "comercial": "market",
          "contable": "book", "lineal": "straight_line", "acelerada": "accelerated"}

PROFILE_SECTIONS = {"empresa": "company", "compras": "purchasing", "precios": "pricing",
                    "finanzas": "finance", "idioma": "language"}
PROFILE_KEYS = {"nombre": "name", "rut": "tax_id", "pais": "country", "moneda": "currency",
                "decimales_moneda": "currency_decimals", "iva": "vat", "impuesto_renta": "income_tax",
                "proveedores": "suppliers", "plazo_entrega_dias": "lead_time_days",
                "desempate": "tie_break", "recargo_por_defecto": "default_markup",
                "gastos_generales": "overhead", "tasa_descuento": "discount_rate",
                "horizonte_anios": "horizon_years", "depreciacion": "depreciation"}
PROFILE_VALUES = {"mayor_disponibilidad": "highest_availability", "lineal": "straight_line",
                  "acelerada": "accelerated"}

# English → Spanish, to write labels in Spanish (language es) and for the test.
ES_KEY = {v: k for k, v in INPUT_KEYS.items()}
ES_KEY.update({"rate": "tasa", "tax": "impuesto"})
ES_VALUE = {v: k for k, v in VALUES.items()}
SWEEP_KEYS = dict(INPUT_KEYS, tasa="rate", impuesto="tax")


def _value_code(v):
    if isinstance(v, str):
        return VALUES.get(v, v)
    return v


def normalize_input(obj, parent: str | None = None):
    """Spanish keys (and code values) → English. Keys starting with _ are kept as they are."""
    if isinstance(obj, list):
        return [normalize_input(x, parent) for x in obj]
    if not isinstance(obj, dict):
        return obj
    out: dict = {}
    # Spanish keys first, so an English key given as well wins.
    for k in sorted(obj, key=lambda k: k not in INPUT_KEYS):
        v = obj[k]
        nk = k if k.startswith("_") else INPUT_KEYS.get(k, k)
        if nk == "assumption" and isinstance(v, list):
            v = [INPUT_KEYS.get(x, x) if isinstance(x, str) else x for x in v]
        elif nk in VALUE_KEYS:
            if isinstance(v, dict):
                v = normalize_input(v, nk)
                v["value"] = _value_code(v.get("value"))
            else:
                v = _value_code(v)
        elif isinstance(v, (dict, list)):
            v = normalize_input(v, nk)
        out[nk] = v
    return out


def normalize_profile(raw: dict) -> dict:
    out: dict = {}
    for k in sorted(raw, key=lambda k: k not in PROFILE_SECTIONS):
        v = raw[k]
        nk = PROFILE_SECTIONS.get(k, k)
        if isinstance(v, dict):
            dest = out.setdefault(nk, {})
            for ik in sorted(v, key=lambda k: k not in PROFILE_KEYS):
                iv = v[ik]
                nik = PROFILE_KEYS.get(ik, ik)
                if nik in ("tie_break", "depreciation") and isinstance(iv, str):
                    iv = PROFILE_VALUES.get(iv, iv)
                dest[nik] = iv
        else:
            out[nk] = v
    return out


def es_path(path: str) -> str:
    """'revenues[0].occupancy' → 'ingresos[0].ocupacion' (only names outside brackets)."""
    parts = []
    for seg in path.split("."):
        base, br, rest = seg.partition("[")
        parts.append(ES_KEY.get(base, base) + (br + rest if br else ""))
    return ".".join(parts)


def loc_path(path: str) -> str:
    return es_path(path) if _LANG == "es" else path


def loc_value(v):
    if _LANG == "es" and isinstance(v, str):
        return ES_VALUE.get(v, v)
    return v


def canonical_key(key: str) -> str:
    """A sweep key typed in Spanish or English → the English key."""
    return re.sub(r"[A-Za-z_]+", lambda mt: SWEEP_KEYS.get(mt.group(0), mt.group(0)), key.strip())


# ---------------------------------------------------------------------------
# Profile (simple YAML: nested maps and scalars; lists are ignored)
# ---------------------------------------------------------------------------


def _scalar(txt: str):
    t = txt.strip()
    if t in ("", "~", "null", "Null", "NULL"):
        return None
    if t in ("true", "True"):
        return True
    if t in ("false", "False"):
        return False
    if t == "[]":
        return []
    if t == "{}":
        return {}
    if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
        return t[1:-1]
    try:
        return int(t)
    except ValueError:
        pass
    try:
        return float(t)
    except ValueError:
        return t


def _strip_comment(line: str) -> str:
    inside = None
    for i, c in enumerate(line):
        if c in "\"'":
            if inside is None:
                inside = c
            elif inside == c:
                inside = None
        elif c == "#" and inside is None and (i == 0 or line[i - 1] in " \t"):
            return line[:i]
    return line


def read_simple_yaml(path: Path) -> dict:
    root: dict = {}
    stack = [(-1, root)]
    for raw in Path(path).read_text(encoding="utf-8-sig").splitlines():
        line = _strip_comment(raw).rstrip()
        if not line.strip():
            continue
        indent = len(line) - len(line.lstrip(" "))
        body = line.strip()
        if body.startswith("-") or ":" not in body:
            continue  # lists: this method does not use them
        key, _, value = body.partition(":")
        while len(stack) > 1 and indent <= stack[-1][0]:
            stack.pop()
        parent = stack[-1][1]
        if value.strip() == "":
            new: dict = {}
            parent[key.strip()] = new
            stack.append((indent, new))
        else:
            parent[key.strip()] = _scalar(value)
    return root


def _data_folders():
    folders = []
    if os.environ.get("INGENIERO_DATOS"):
        folders.append(Path(os.environ["INGENIERO_DATOS"]))
    folders.append(REPO.parent / "ingeniero-agentico-datos")
    return folders


def find_profile(arg: str | None):
    """Order: --profile (path or name) · $INGENIERO_DATOS · ../ingeniero-agentico-datos
    (profiles/ first, then perfiles/) · repo template (profiles/_template.yaml, then
    perfiles/_plantilla.yaml). Returns (path, how, name)."""
    folders = _data_folders()
    if arg:
        p = Path(arg)
        if p.is_file():
            return p, "arg_path", p.name
        name = arg if arg.endswith((".yaml", ".yml")) else arg + ".yaml"
        for d in folders:
            for sub in ("profiles", "perfiles"):
                q = d / sub / name
                if q.is_file():
                    return q, "name", q.name
        raise InputError(tr(f"No encontré el perfil «{arg}» (ni como ruta ni en <datos>/profiles/ o perfiles/).",
                            f"Profile «{arg}» not found (neither as a path nor in <data>/profiles/ or perfiles/)."))
    for d in folders:
        found: dict = {}
        for sub in ("perfiles", "profiles"):  # profiles/ wins on equal names
            dp = d / sub
            if dp.is_dir():
                for x in dp.glob("*.yaml"):
                    if not x.name.startswith("_"):
                        found[x.stem] = x
        if len(found) == 1:
            x = next(iter(found.values()))
            return x, "name", x.name
        if len(found) > 1:
            names = ", ".join(sorted(found))
            raise InputError(tr(f"Hay varios perfiles ({names}): elige uno con --profile <nombre>.",
                                f"There are several profiles ({names}): choose one with --profile <name>."))
    for template in (REPO / "profiles" / "_template.yaml", REPO / "perfiles" / "_plantilla.yaml"):
        if template.is_file():
            return template, "template", template.name
    return None, "none", ""


def profile_label(how: str, name: str) -> str:
    if how == "arg_path":
        return tr(f"perfil {name} (--perfil)", f"profile {name} (--profile)")
    if how == "name":
        return tr(f"perfil {name}", f"profile {name}")
    if how == "template":
        return tr("plantilla (sin perfil de empresa)", "template (no company profile)")
    return tr("sin perfil", "no profile")


def set_language(profile: dict):
    global _LANG
    lang = (profile.get("language") or "es")
    lang = str(lang).strip().lower()
    if lang not in ("es", "en"):
        raise InputError(f"language «{lang}»: use es or en.")
    _LANG = lang
    return lang


# ---------------------------------------------------------------------------
# Parameters: profile + arguments + input, each with its origin
# ---------------------------------------------------------------------------


def _setting(inp: dict, key: str):
    """A top-level setting given as a plain value or as {value, assumption, source}.
    Returns (value, given, assumption flag or None, source)."""
    v = inp.get(key)
    if isinstance(v, dict):
        return v.get("value"), v.get("value") is not None, bool(v.get("assumption", False)), v.get("source") or ""
    return v, v is not None and v != "", None, ""


def resolve_parameters(profile: dict, prof: str, inp: dict, rate=None, tax=None, pcode="profile"):
    """pcode: origin code of the profile values (profile, or template when it is the repo template)."""
    co = profile.get("company") or {}
    fin = profile.get("finance") or {}
    p = {}

    def put(key, value, origin, assumption, code):
        p[key] = {"value": value, "origin": origin, "assumption": assumption, "origin_code": code}

    # discount rate
    if rate is not None:
        extra = tr(" (reemplaza la del perfil)", " (replaces the profile's)") if fin.get("discount_rate") is not None \
            else tr(" (el perfil no la trae)", " (the profile does not have it)")
        put("rate", float(rate), tr("argumento --tasa", "argument --rate") + extra, True, "input")
    elif fin.get("discount_rate") is not None:
        put("rate", float(fin["discount_rate"]), f"{prof} · " + tr("finanzas.tasa_descuento", "finance.discount_rate"), False, pcode)
    else:
        raise InputError(tr("El perfil no trae finanzas.tasa_descuento: pasa --tasa (la salida la marca SUPUESTO).",
                            "The profile has no finance.discount_rate: pass --rate (the output marks it ASSUMPTION)."))

    # income tax
    if tax is not None:
        extra = tr(" (reemplaza la del perfil)", " (replaces the profile's)") if co.get("income_tax") is not None \
            else tr(" (el perfil no la trae)", " (the profile does not have it)")
        put("tax", float(tax), tr("argumento --impuesto", "argument --tax") + extra, True, "input")
    elif co.get("income_tax") is not None:
        put("tax", float(co["income_tax"]), f"{prof} · " + tr("empresa.impuesto_renta", "company.income_tax"), False, pcode)
    else:
        raise InputError(tr("El perfil no trae empresa.impuesto_renta: pasa --impuesto (la salida la marca SUPUESTO).",
                            "The profile has no company.income_tax: pass --tax (the output marks it ASSUMPTION)."))

    # horizon
    hv, hgiven, hsup, hsrc = _setting(inp, "horizon_years")
    if hgiven:
        put("horizon", int(hv), tr("entrada · horizonte_anios", "input · horizon_years") + (f" · {hsrc}" if hsrc else ""),
            bool(hsup), "input")
    elif fin.get("horizon_years") is not None:
        put("horizon", int(fin["horizon_years"]), f"{prof} · " + tr("finanzas.horizonte_anios", "finance.horizon_years"), False, pcode)
    else:
        raise InputError(tr("Falta el horizonte: ni la entrada ni el perfil traen horizonte_anios.",
                            "Missing horizon: neither the input nor the profile has horizon_years."))

    dep = fin.get("depreciation") or "straight_line"
    if dep not in ("straight_line", "accelerated"):
        raise InputError(tr(f"Depreciación «{dep}» no reconocida (lineal · acelerada).",
                            f"Depreciation «{dep}» not recognized (straight_line · accelerated)."))
    put("depreciation", dep, f"{prof} · " + tr("finanzas.depreciacion", "finance.depreciation") if fin.get("depreciation")
        else tr("por defecto (lineal)", "default (straight_line)"), not fin.get("depreciation"),
        pcode if fin.get("depreciation") else "default")
    put("currency", co.get("currency") or "CLP", f"{prof} · " + tr("empresa.moneda", "company.currency"), not co.get("currency"),
        pcode if co.get("currency") else "default")
    dec = co.get("currency_decimals")
    put("decimals", int(dec) if dec is not None else 0, f"{prof} · " + tr("empresa.decimales_moneda", "company.currency_decimals"),
        dec is None, "default" if dec is None else pcode)

    # tax losses: a modeling choice that says whether the company has other profits. Given as a
    # plain value it carries no source, so it is an ASSUMPTION (audit finding of 15-10-2026).
    lv, lgiven, lsup, lsrc = _setting(inp, "tax_losses")
    losses = lv if lgiven else "carry_forward"
    if losses not in ("carry_forward", "offset"):
        raise InputError(tr("perdidas_tributarias debe ser «arrastra» o «compensa».",
                            "tax_losses must be «carry_forward» or «offset»."))
    if not lgiven:
        put("tax_losses", losses, tr("por defecto (arrastra)", "default (carry_forward)"), True, "default")
    elif lsup is None:
        put("tax_losses", losses, tr("entrada · perdidas_tributarias (sin fuente: supone si la empresa tiene otras utilidades)",
                                     "input · tax_losses (no source: it assumes whether the company has other profits)"), True, "input")
    else:
        put("tax_losses", losses, tr("entrada · perdidas_tributarias", "input · tax_losses") + (f" · {lsrc}" if lsrc else ""), lsup, "input")

    sv, sgiven, ssup, ssrc = _setting(inp, "salvage_method")
    method = sv if sgiven else "market"
    if method not in ("market", "book"):
        raise InputError(tr("metodo_desecho debe ser «comercial» o «contable».",
                            "salvage_method must be «market» or «book»."))
    put("salvage_method", method, (tr("entrada · metodo_desecho", "input · salvage_method") + (f" · {ssrc}" if ssrc else ""))
        if sgiven else tr("por defecto (comercial)", "default (market)"), bool(ssup),
        "input" if sgiven else "default")
    return p


def build_model(inp: dict, params: dict) -> dict:
    """Normalized numeric model (what the sensitivity changes)."""
    m = {
        "N": params["horizon"]["value"],
        "rate": params["rate"]["value"],
        "tax": params["tax"]["value"],
        "depreciation": params["depreciation"]["value"],
        "tax_losses": params["tax_losses"]["value"],
        "salvage_method": params["salvage_method"]["value"],
        "revenues": copy.deepcopy(inp.get("revenues") or []),
        "variable_costs": copy.deepcopy(inp.get("variable_costs") or []),
        "fixed_costs": copy.deepcopy(inp.get("fixed_costs") or []),
        "investments": copy.deepcopy(inp.get("investments") or []),
        "working_capital": copy.deepcopy(inp.get("working_capital") or {}),
    }
    names = [r.get("name") for r in m["revenues"]]
    if len(set(names)) != len(names):
        raise InputError(tr("Dos ingresos tienen el mismo nombre.", "Two revenues have the same name."))
    for rev in m["revenues"]:
        if "price" not in rev:
            raise InputError(tr(f"Ingreso «{rev.get('name')}» sin precio.", f"Revenue «{rev.get('name')}» has no price."))
        if "quantity" not in rev and not ("capacity" in rev and "occupancy" in rev):
            raise InputError(tr(f"Ingreso «{rev.get('name')}»: falta cantidad, o capacidad + ocupacion.",
                                f"Revenue «{rev.get('name')}»: missing quantity, or capacity + occupancy."))
    for vc in m["variable_costs"]:
        if vc.get("per_unit_of") not in names:
            raise InputError(tr(f"Costo variable «{vc.get('name')}»: por_unidad_de debe ser el nombre de un ingreso.",
                                f"Variable cost «{vc.get('name')}»: per_unit_of must be the name of a revenue."))
    for inv in m["investments"]:
        if int(inv.get("year", 0)) >= m["N"]:
            raise InputError(tr(f"Inversión «{inv.get('name')}» en el año {inv.get('year')}: debe ser anterior al horizonte.",
                                f"Investment «{inv.get('name')}» in year {inv.get('year')}: it must come before the horizon."))
        if inv.get("depreciable", True) and not inv.get("useful_life"):
            raise InputError(tr(f"Inversión «{inv.get('name')}» depreciable sin vida_util.",
                                f"Depreciable investment «{inv.get('name')}» without useful_life."))
    return m


# ---------------------------------------------------------------------------
# Pure project cash flow
# ---------------------------------------------------------------------------


def _growth(item, field, t):
    g = float(item.get(field, 0) or 0)
    return (1 + g) ** (t - 1) if t >= 1 else 0.0


def quantity_of(rev):
    if "quantity" in rev:
        return float(rev["quantity"])
    return float(rev["capacity"]) * float(rev["occupancy"])


def cash_flow(m: dict) -> dict:
    N = m["N"]
    T = m["tax"]
    years = list(range(N + 1))
    zero = lambda: [0.0] * (N + 1)
    rows = []  # (code, label, values, origin)

    # revenues
    tot_rev = zero()
    qty = {}
    for k, rev in enumerate(m["revenues"]):
        q = quantity_of(rev)
        v = zero()
        cq = zero()
        for t in years[1:]:
            cq[t] = q * _growth(rev, "quantity_growth", t)
            v[t] = float(rev["price"]) * _growth(rev, "price_growth", t) * cq[t]
            tot_rev[t] += v[t]
        qty[rev["name"]] = cq
        form = tr("capacidad × ocupación", "capacity × occupancy") if "quantity" not in rev else tr("cantidad", "quantity")
        rows.append(("revenue", tr("Ingreso · ", "Revenue · ") + rev["name"], v,
                     tr(f"entrada · ingresos[{k}]: precio × {form}", f"input · revenues[{k}]: price × {form}")))
    # variable costs
    tot_vc = zero()
    for k, vc in enumerate(m["variable_costs"]):
        v = zero()
        for t in years[1:]:
            v[t] = -float(vc["amount_per_unit"]) * _growth(vc, "growth", t) * qty[vc["per_unit_of"]][t]
            tot_vc[t] += v[t]
        rows.append(("variable_cost", tr("Costo variable · ", "Variable cost · ") + vc["name"], v,
                     tr(f"entrada · costos_variables[{k}]: monto_por_unidad × cantidad de «{vc['per_unit_of']}»",
                        f"input · variable_costs[{k}]: amount_per_unit × quantity of «{vc['per_unit_of']}»")))
    # fixed costs
    tot_fc = zero()
    for k, fc in enumerate(m["fixed_costs"]):
        v = zero()
        for t in years[1:]:
            v[t] = -float(fc["annual_amount"]) * _growth(fc, "growth", t)
            tot_fc[t] += v[t]
        rows.append(("fixed_cost", tr("Costo fijo · ", "Fixed cost · ") + fc["name"], v,
                     tr(f"entrada · costos_fijos[{k}]: monto_anual", f"input · fixed_costs[{k}]: annual_amount")))
    # depreciation, sale and book value
    tot_dep = zero()
    sale = zero()
    book = zero()
    investment = zero()
    assets = []
    for k, inv in enumerate(m["investments"]):
        a = int(inv.get("year", 0))
        amount = float(inv["amount"])
        investment[a] -= amount
        acc_dep = 0.0
        if inv.get("depreciable", True):
            life = inv["useful_life"]
            if m["depreciation"] == "accelerated" and inv.get("accelerated_life"):
                life = inv["accelerated_life"]
            quota = amount / float(life)
            for t in range(a + 1, min(a + int(math.ceil(life)), N) + 1):
                d = min(quota, amount - acc_dep)
                tot_dep[t] -= d
                acc_dep += d
        bv = amount - acc_dep
        sv = float(inv.get("salvage_value", 0) or 0) if m["salvage_method"] == "market" else bv
        sale[N] += sv
        book[N] -= bv
        assets.append({"name": inv["name"], "amount": amount, "year": a, "accumulated_depreciation": acc_dep,
                       "final_book_value": bv, "salvage_value": sv})
    ebt = [tot_rev[t] + tot_vc[t] + tot_fc[t] + tot_dep[t] + sale[t] + book[t] for t in years]
    taxes = zero()
    loss = 0.0
    for t in years[1:]:
        if m["tax_losses"] == "offset":
            taxes[t] = -T * ebt[t]
        else:
            base = ebt[t] + loss
            if base > 0:
                taxes[t] = -T * base
                loss = 0.0
            else:
                loss = base
    net = [ebt[t] + taxes[t] for t in years]
    dep_back = [-x for x in tot_dep]
    book_back = [-x for x in book]
    wc = m["working_capital"] or {}
    inv_wc = zero()
    rec_wc = zero()
    if wc.get("amount"):
        a = int(wc.get("year", 0))
        inv_wc[a] -= float(wc["amount"])
        if wc.get("recovered", True):
            rec_wc[N] += float(wc["amount"])
    flow = [net[t] + dep_back[t] + book_back[t] + investment[t] + inv_wc[t] + rec_wc[t] for t in years]

    losses_txt = loc_value(m["tax_losses"])
    method_txt = loc_value(m["salvage_method"])
    body = list(rows)
    body += [
        ("total", tr("= Ingresos totales", "= Total revenue"), tot_rev, tr("suma de los ingresos", "sum of the revenues")),
        ("total", tr("− Costos variables totales", "− Total variable costs"), tot_vc,
         tr("suma de los costos variables", "sum of the variable costs")),
        ("total", tr("− Costos fijos totales", "− Total fixed costs"), tot_fc, tr("suma de los costos fijos", "sum of the fixed costs")),
        ("dep", tr("− Depreciación", "− Depreciation"), tot_dep,
         tr("inversiones: monto ÷ vida útil (perfil · depreciación)", "investments: amount ÷ useful life (profile · depreciation)")),
        ("sale", tr("+ Venta de activos (valor de desecho)", "+ Sale of assets (salvage value)"), sale,
         tr(f"inversiones: valor_desecho (método {method_txt})", f"investments: salvage_value (method {method_txt})")),
        ("book", tr("− Valor libro de los activos vendidos", "− Book value of the assets sold"), book,
         tr("monto − depreciación acumulada", "amount − accumulated depreciation")),
        ("ebt", tr("= Utilidad antes de impuesto", "= Profit before tax"), ebt, tr("suma de lo anterior", "sum of the above")),
        ("tax", tr("− Impuesto", "− Tax"), taxes,
         tr(f"tasa {T:.2%} × utilidad (pérdidas: {losses_txt})", f"rate {T:.2%} × profit (losses: {losses_txt})")),
        ("net", tr("= Utilidad neta", "= Net profit"), net, tr("utilidad antes de impuesto − impuesto", "profit before tax − tax")),
        ("adjustment", tr("+ Depreciación (no es salida de caja)", "+ Depreciation (not a cash outflow)"), dep_back,
         tr("se devuelve lo restado arriba", "adds back what was subtracted above")),
        ("adjustment", tr("+ Valor libro (no es salida de caja)", "+ Book value (not a cash outflow)"), book_back,
         tr("se devuelve lo restado arriba", "adds back what was subtracted above")),
        ("inv", tr("− Inversión en activos", "− Investment in assets"), investment,
         tr("entrada · inversiones (monto, año)", "input · investments (amount, year)")),
        ("inv", tr("− Inversión en capital de trabajo", "− Investment in working capital"), inv_wc,
         tr("entrada · capital_trabajo.monto", "input · working_capital.amount")),
        ("inv", tr("+ Recuperación del capital de trabajo", "+ Recovery of working capital"), rec_wc,
         tr("entrada · capital_trabajo (se recupera al final)", "input · working_capital (recovered at the end)")),
        ("flow", tr("= FLUJO DE CAJA DEL PROYECTO PURO", "= PURE PROJECT CASH FLOW"), flow,
         tr("utilidad neta + ajustes − inversiones + recuperaciones", "net profit + adjustments − investments + recoveries")),
    ]
    return {"years": years, "rows": body, "flow": flow, "assets": assets,
            "total_investment": [investment[t] + inv_wc[t] for t in years], "final_loss": loss}


# ---------------------------------------------------------------------------
# Indicators
# ---------------------------------------------------------------------------


def npv(flow, r):
    return sum(f / (1 + r) ** t for t, f in enumerate(flow))


def sign_changes(flow):
    s = [1 if f > 0 else -1 for f in flow if abs(f) > 1e-9]
    return sum(1 for a, b in zip(s, s[1:]) if a != b)


def irr_roots(flow):
    grid = [-0.95 + 0.005 * i for i in range(int(1.95 / 0.005) + 1)] + [1.0 + 0.05 * i for i in range(1, 181)]
    roots = []
    prev_r, prev_v = None, None
    for r in grid:
        v = npv(flow, r)
        if v == 0:
            roots.append(r)
        elif prev_v is not None and prev_v * v < 0:
            a, b, va = prev_r, r, prev_v
            for _ in range(200):
                c = (a + b) / 2
                vc = npv(flow, c)
                if va * vc <= 0:
                    b = c
                else:
                    a, va = c, vc
            roots.append((a + b) / 2)
        prev_r, prev_v = r, v
    return roots


def payback(flow, r=None):
    acc = 0.0
    for t, f in enumerate(flow):
        fd = f if r is None else f / (1 + r) ** t
        new = acc + fd
        if t > 0 and acc < 0 <= new:
            return (t - 1) + (-acc / fd)
        acc = new
    return None


def indicators(cf: dict, m: dict) -> dict:
    flow = cf["flow"]
    r = m["rate"]
    v = npv(flow, r)
    roots = irr_roots(flow)
    cs = sign_changes(flow)
    warnings = []
    irr = None
    if cs > 1:
        warnings.append(tr(f"El flujo cambia {cs} veces de signo: puede haber varias TIR (o ninguna). Decide con el VAN.",
                           f"The flow changes sign {cs} times: there may be several IRRs (or none). Decide with the NPV."))
    if len(roots) == 1:
        irr = roots[0]
    elif len(roots) > 1:
        warnings.append(tr("Hay varias TIR: ", "There are several IRRs: ") + ", ".join(f"{x:.4%}" for x in roots)
                        + tr(". La TIR no sirve para decidir aquí.", ". The IRR cannot be used to decide here."))
    else:
        warnings.append(tr("No hay TIR en el rango −95 % a 1.000 %.", "No IRR in the range −95 % to 1,000 %."))
    inv_pv = -sum(x / (1 + r) ** t for t, x in enumerate(cf["total_investment"]))
    index = v / inv_pv if inv_pv > 0 else None
    pb = payback(flow)
    dpb = payback(flow, r)
    if pb is None:
        warnings.append(tr("La inversión no se recupera dentro del horizonte (ni sin descontar).",
                           "The investment is not recovered within the horizon (not even undiscounted)."))
    if cf["final_loss"] < 0:
        warnings.append(tr("Queda pérdida tributaria sin usar al final del horizonte (no se valoriza).",
                           "An unused tax loss remains at the end of the horizon (it is not valued)."))
    return {"npv": v, "irr": irr, "irr_roots": roots, "sign_changes": cs, "payback": pb, "discounted_payback": dpb,
            "npv_index": index, "investment_pv": inv_pv, "rate": r, "warnings": warnings,
            "decision": tr("CONVIENE (VAN ≥ 0)", "ACCEPT (NPV ≥ 0)") if v >= 0 else tr("NO CONVIENE (VAN < 0)", "REJECT (NPV < 0)")}


# ---------------------------------------------------------------------------
# Sensitivity
# ---------------------------------------------------------------------------


def sensitive_variables(m: dict):
    """(key, label, route). route = (section, index, field) or (field,)."""
    vs = []
    for k, rev in enumerate(m["revenues"]):
        vs.append((f"revenues[{k}].price", f"{rev['name']} · " + tr("precio", "price"), ("revenues", k, "price")))
        if "quantity" in rev:
            vs.append((f"revenues[{k}].quantity", f"{rev['name']} · " + tr("cantidad", "quantity"), ("revenues", k, "quantity")))
        else:
            vs.append((f"revenues[{k}].occupancy", f"{rev['name']} · " + tr("ocupación", "occupancy"), ("revenues", k, "occupancy")))
    for k, vc in enumerate(m["variable_costs"]):
        vs.append((f"variable_costs[{k}].amount_per_unit", f"{vc['name']} · " + tr("por unidad", "per unit"),
                   ("variable_costs", k, "amount_per_unit")))
    for k, fc in enumerate(m["fixed_costs"]):
        vs.append((f"fixed_costs[{k}].annual_amount", f"{fc['name']} · " + tr("anual", "annual"), ("fixed_costs", k, "annual_amount")))
    for k, inv in enumerate(m["investments"]):
        vs.append((f"investments[{k}].amount", f"{inv['name']} · " + tr("inversión", "investment"), ("investments", k, "amount")))
        if m["salvage_method"] == "market":
            vs.append((f"investments[{k}].salvage_value", f"{inv['name']} · " + tr("valor de desecho", "salvage value"),
                       ("investments", k, "salvage_value")))
    if (m["working_capital"] or {}).get("amount"):
        vs.append(("working_capital.amount", tr("Capital de trabajo", "Working capital"), ("working_capital", None, "amount")))
    vs.append(("rate", tr("Tasa de descuento", "Discount rate"), ("rate",)))
    vs.append(("tax", tr("Tasa de impuesto", "Tax rate"), ("tax",)))
    return [v for v in vs if (_get(m, v[2]) or 0) != 0]


def _get(m, route):
    if len(route) == 1:
        return m[route[0]]
    sec, k, field = route
    obj = m[sec] if k is None else m[sec][k]
    return obj.get(field)


def _with(m, route, value):
    m2 = copy.deepcopy(m)
    if len(route) == 1:
        m2[route[0]] = value
    else:
        sec, k, field = route
        (m2[sec] if k is None else m2[sec][k])[field] = value
    return m2


def npv_with(m, route, value):
    m2 = _with(m, route, value)
    return npv(cash_flow(m2)["flow"], m2["rate"])


def critical_value(m, route):
    """Value of the variable that makes NPV zero, taking the crossing closest to the base value."""
    base = float(_get(m, route))
    factors = [i * 0.05 for i in range(0, 61)] + [3.5, 4, 5, 6, 8, 10, 15, 20, 30, 50, 100]
    points = [(f, npv_with(m, route, base * f)) for f in factors]
    spans = [(a, b) for a, b in zip(points, points[1:]) if a[1] == 0 or a[1] * b[1] < 0]
    if not spans:
        return None
    (fa, va), (fb, vb) = min(spans, key=lambda sp: min(abs(sp[0][0] - 1), abs(sp[1][0] - 1)))
    if va == 0:
        return base * fa
    for _ in range(200):
        fc = (fa + fb) / 2
        vc = npv_with(m, route, base * fc)
        if va * vc <= 0:
            fb = fc
        else:
            fa, va = fc, vc
    return base * (fa + fb) / 2


def sensitivity(m: dict, delta: float):
    rows = []
    base_npv = npv(cash_flow(m)["flow"], m["rate"])
    for key, label, route in sensitive_variables(m):
        base = float(_get(m, route))
        v_minus = npv_with(m, route, base * (1 - delta))
        v_plus = npv_with(m, route, base * (1 + delta))
        crit = critical_value(m, route)
        note = ""
        if crit is None:
            side = tr("positivo", "positive") if base_npv >= 0 else tr("negativo", "negative")
            note = tr(f"sin valor crítico: el VAN sigue {side} desde 0 hasta 100 veces el valor base",
                      f"no critical value: NPV stays {side} from 0 to 100 times the base value")
        elif route[-1] == "occupancy" and crit > 1:
            note = tr("fuera de rango: haría falta más de 100 % de ocupación", "out of range: it would take more than 100 % occupancy")
        elif crit < 0:
            note = tr("valor crítico negativo", "negative critical value")
        rows.append({"key": key, "variable": label, "base": base, "npv_minus": v_minus, "npv_plus": v_plus,
                     "range": abs(v_plus - v_minus), "critical": crit,
                     "change_to_critical": (crit / base - 1) if crit is not None and base else None, "note": note})
    rows.sort(key=lambda f: -f["range"])
    return {"delta": delta, "base_npv": base_npv, "rows": rows}


def sweep(m: dict, spec: str):
    key, _, values = spec.partition("=")
    key = canonical_key(key)
    routes = {c: r for c, _, r in sensitive_variables(m)}
    routes.update({"rate": ("rate",), "tax": ("tax",)})
    if key not in routes:
        raise InputError(tr(f"--barrer: no conozco la variable «{loc_path(key)}». Usa una clave de la hoja SENSIBILIDAD.",
                            f"--sweep: unknown variable «{key}». Use a key from the SENSITIVITY sheet."))
    route = routes[key]
    out = []
    for txt in values.split(","):
        x = float(txt)
        m2 = _with(m, route, x)
        cf = cash_flow(m2)
        ind = indicators(cf, m2)
        out.append({"value": x, "npv": ind["npv"], "irr": ind["irr"]})
    return {"key": key, "points": out}


# ---------------------------------------------------------------------------
# Assumptions (traceability of every input)
# ---------------------------------------------------------------------------


def _is_assumption(item, field):
    s = item.get("assumption", False)
    if isinstance(s, list):
        return field in s
    return bool(s)


def assumptions_table(inp, params):
    t = []
    for key, label, unit in (("rate", tr("Tasa de descuento", "Discount rate"), tr("anual", "annual")),
                             ("tax", tr("Tasa de impuesto a la renta", "Income tax rate"), tr("fracción", "fraction")),
                             ("horizon", tr("Horizonte de evaluación", "Evaluation horizon"), tr("años", "years")),
                             ("depreciation", tr("Depreciación", "Depreciation"), ""),
                             ("tax_losses", tr("Pérdidas tributarias", "Tax losses"), ""),
                             ("salvage_method", tr("Valor de desecho", "Salvage value"), tr("método", "method")),
                             ("currency", tr("Moneda", "Currency"), ""),
                             ("decimals", tr("Decimales de la moneda", "Currency decimals"), "")):
        p = params[key]
        t.append({"parameter": label, "value": loc_value(p["value"]), "unit": unit, "origin": p["origin"],
                  "mark": mark(p["assumption"]), "origin_code": p["origin_code"]})
    # The input note states assumptions (net or gross of VAT, real or nominal money, what the
    # client pays apart) that change the flow: it goes to the table, marked, not hidden.
    notes = []
    if inp.get("note"):
        notes.append((tr("entrada · nota", "input · note"), inp["note"]))
    for k, n in enumerate(inp.get("notes") or []):
        notes.append((tr(f"entrada · notas[{k}]", f"input · notes[{k}]"), n))
    for origin, text in notes:
        t.append({"parameter": tr("Nota de la entrada (supuestos del caso)", "Input note (case assumptions)"),
                  "value": text, "unit": "", "origin": origin, "mark": mark(True), "origin_code": "input"})
    fields = {
        "revenues": ("price", "quantity", "capacity", "occupancy", "price_growth", "quantity_growth"),
        "variable_costs": ("amount_per_unit", "growth"),
        "fixed_costs": ("annual_amount", "growth"),
        "investments": ("amount", "year", "useful_life", "accelerated_life", "salvage_value"),
    }
    for sec, fs in fields.items():
        for k, item in enumerate(inp.get(sec) or []):
            for c in fs:
                if c in item:
                    t.append({"parameter": f"{item.get('name')} · {loc_path(c)}", "value": item[c],
                              "unit": item.get("unit", "") if c in ("quantity", "capacity", "price") else "",
                              "origin": tr("entrada · ", "input · ") + loc_path(f"{sec}[{k}].{c}")
                              + (f" · {item['source']}" if item.get("source") else ""),
                              "mark": mark(_is_assumption(item, c)), "origin_code": _code_of(item)})
    wc = inp.get("working_capital") or {}
    if wc.get("amount") is not None:
        t.append({"parameter": tr("Capital de trabajo · monto", "Working capital · amount"), "value": wc["amount"], "unit": "",
                  "origin": tr("entrada · capital_trabajo.monto", "input · working_capital.amount")
                  + (f" · {wc['source']}" if wc.get("source") else ""),
                  "mark": mark(_is_assumption(wc, "amount")), "origin_code": _code_of(wc)})
    return t


def _code_of(item):
    """input, or research:<id> when the item names the research datum it came from."""
    rid = item.get("research") or item.get("research_id")
    return f"research:{rid}" if rid else "input"


def contract_assumptions(table):
    """The common format of every method: [{parameter, value, origin, text}], only the marked rows."""
    out = []
    for r in table:
        if r["mark"] != mark(True):
            continue
        body = f"{r['parameter']}: {r['value']}" if r["parameter"].startswith(tr("Nota", "Input note"))             else f"{r['parameter']} = {r['value']}" + (f" {r['unit']}" if r["unit"] else "") + f" · {r['origin']}"
        out.append({"parameter": r["parameter"], "value": r["value"], "origin": r["origin_code"],
                    "text": f"{mark(True)} · {body}"})
    return out


# ---------------------------------------------------------------------------
# Full evaluation
# ---------------------------------------------------------------------------


def evaluate(inp, profile, how="none", prof_name="", rate=None, tax=None, delta=0.10, sweeps=None):
    """inp and profile may come with Spanish or English keys."""
    profile = normalize_profile(profile or {})
    lang = set_language(profile)
    inp = normalize_input(inp)
    prof = profile_label(how, prof_name)
    params = resolve_parameters(profile, prof, inp, rate, tax, "template" if how == "template" else "profile")
    m = build_model(inp, params)
    cf = cash_flow(m)
    ind = indicators(cf, m)
    sens = sensitivity(m, delta)
    sw = [sweep(m, s) for s in (sweeps or [])]
    return {"project": inp.get("project", ""), "company": inp.get("company") or (profile.get("company") or {}).get("name", ""),
            "profile": prof, "language": lang, "note": inp.get("note", ""), "parameters": params,
            "assumptions": assumptions_table(inp, params), "cash_flow": cf, "indicators": ind,
            "sensitivity": sens, "sweeps": sw, "_model": m}


# ---------------------------------------------------------------------------
# Output: Excel and JSON
# ---------------------------------------------------------------------------


def sheet_names():
    return [tr("SUPUESTOS", "ASSUMPTIONS"), tr("FLUJO", "CASH FLOW"), tr("INDICADORES", "INDICATORS"),
            tr("SENSIBILIDAD", "SENSITIVITY")]


def write_excel(res: dict, path: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    dec = res["parameters"]["decimals"]["value"]
    fmt_m = "#,##0" if dec == 0 else "#,##0." + "0" * dec
    bold = Font(bold=True)
    grey = PatternFill("solid", fgColor="E7E6E6")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    names = sheet_names()
    wb = Workbook()

    # ASSUMPTIONS
    ws = wb.active
    ws.title = names[0]
    ws.append([tr("Evaluación económica · ", "Economic evaluation · ") + f"{res['project']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([tr(f"Empresa: {res['company']} · {res['profile']} · proyecto puro (sin financiamiento)",
                  f"Company: {res['company']} · {res['profile']} · pure project (no financing)")])
    ws.append([])
    ws.append([tr("Parámetro", "Parameter"), tr("Valor", "Value"), tr("Unidad", "Unit"), tr("Origen", "Origin"), tr("Marca", "Mark")])
    for c in ws[4]:
        c.font = bold
        c.fill = grey
    for s in res["assumptions"]:
        ws.append([s["parameter"], s["value"], s["unit"], s["origin"], s["mark"]])
        if s["mark"] == mark(True):
            for c in ws[ws.max_row]:
                c.fill = yellow
    for col, w in zip("ABCDE", (38, 16, 10, 90, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    # CASH FLOW
    wf = wb.create_sheet(names[1])
    years = res["cash_flow"]["years"]
    wf.append([tr("Concepto", "Item")] + [tr(f"Año {t}", f"Year {t}") for t in years] + [tr("Origen", "Origin")])
    for c in wf[1]:
        c.font = bold
        c.fill = grey
    for code, label, values, origin in res["cash_flow"]["rows"]:
        wf.append([label] + [round(v, 6) for v in values] + [origin])
        row = wf[wf.max_row]
        for c in row[1:len(years) + 1]:
            c.number_format = fmt_m
        if code in ("ebt", "net", "flow", "total"):
            for c in row:
                c.font = bold
        if code == "flow":
            for c in row:
                c.fill = grey
    wf.column_dimensions["A"].width = 46
    for i in range(len(years)):
        wf.column_dimensions[chr(66 + i)].width = 15
    wf.column_dimensions[chr(66 + len(years))].width = 70
    wf.freeze_panes = "B2"

    # INDICATORS
    wi = wb.create_sheet(names[2])
    ind = res["indicators"]
    wi.append([tr("Indicador", "Indicator"), tr("Valor", "Value"), tr("Cómo se calcula", "How it is calculated")])
    for c in wi[1]:
        c.font = bold
        c.fill = grey
    not_rec = tr("no se recupera", "not recovered")
    rows = [
        (tr("Decisión", "Decision"), ind["decision"], tr("VAN ≥ 0 conviene; VAN < 0 no conviene", "NPV ≥ 0 accept; NPV < 0 reject"), None),
        (tr("Tasa de descuento", "Discount rate"), ind["rate"],
         res["parameters"]["rate"]["origin"] + (" · " + mark(True) if res["parameters"]["rate"]["assumption"] else ""), "0.00%"),
        (tr("VAN", "NPV"), ind["npv"], tr("Σ flujo_t ÷ (1 + tasa)^t, t = 0 … N (hoja FLUJO, última fila)",
                                          "Σ flow_t ÷ (1 + rate)^t, t = 0 … N (CASH FLOW sheet, last row)"), fmt_m),
        (tr("TIR", "IRR"), ind["irr"] if ind["irr"] is not None else tr("no definida", "undefined"),
         tr("tasa que deja el VAN en cero (bisección)", "rate that makes NPV zero (bisection)"), "0.00%"),
        (tr("Cambios de signo del flujo", "Sign changes of the flow"), ind["sign_changes"],
         tr("más de 1: puede haber varias TIR", "more than 1: there may be several IRRs"), None),
        (tr("Período de recuperación (años)", "Payback period (years)"), ind["payback"] if ind["payback"] is not None else not_rec,
         tr("flujo acumulado sin descontar; interpolación dentro del año", "undiscounted cumulative flow; interpolated within the year"),
         "0.00"),
        (tr("Período de recuperación descontado (años)", "Discounted payback period (years)"),
         ind["discounted_payback"] if ind["discounted_payback"] is not None else not_rec,
         tr("flujo acumulado descontado a la tasa", "cumulative flow discounted at the rate"), "0.00"),
        (tr("Inversión en valor presente", "Investment in present value"), ind["investment_pv"],
         tr("activos + capital de trabajo, descontados", "assets + working capital, discounted"), fmt_m),
        (tr("IVAN", "NPV index"), ind["npv_index"] if ind["npv_index"] is not None else tr("no aplica", "not applicable"),
         tr("VAN ÷ inversión en valor presente", "NPV ÷ investment in present value"), "0.000"),
    ]
    for name, value, how, f in rows:
        wi.append([name, value, how])
        if f and isinstance(value, (int, float)):
            wi.cell(row=wi.max_row, column=2).number_format = f
    wi.append([])
    for a in ind["warnings"]:
        wi.append([tr("AVISO", "WARNING"), a])
        wi.cell(row=wi.max_row, column=1).font = Font(bold=True, color="C00000")
    wi.column_dimensions["A"].width = 42
    wi.column_dimensions["B"].width = 22
    wi.column_dimensions["C"].width = 70

    # SENSITIVITY
    wsn = wb.create_sheet(names[3])
    sens = res["sensitivity"]
    d = sens["delta"]
    wsn.append([tr(f"Una variable a la vez, ±{d:.0%}; las demás en su valor base. Orden: de la que más mueve el VAN a la que menos.",
                   f"One variable at a time, ±{d:.0%}; the rest at their base value. Order: from the one that moves NPV most to the least.")])
    wsn.append([tr(f"VAN base: {sens['base_npv']:,.0f}", f"Base NPV: {sens['base_npv']:,.0f}")])
    wsn.append([])
    wsn.append([tr("Clave", "Key"), "Variable", tr("Valor base", "Base value"), tr(f"VAN con −{d:.0%}", f"NPV at −{d:.0%}"),
                tr(f"VAN con +{d:.0%}", f"NPV at +{d:.0%}"), tr("Rango del VAN", "NPV range"),
                tr("Valor crítico (VAN = 0)", "Critical value (NPV = 0)"), tr("Cambio hasta el crítico", "Change to the critical value"),
                tr("Nota", "Note")])
    for c in wsn[4]:
        c.font = bold
        c.fill = grey
    for f in sens["rows"]:
        wsn.append([loc_path(f["key"]), f["variable"], f["base"], f["npv_minus"], f["npv_plus"], f["range"],
                    f["critical"] if f["critical"] is not None else tr("sin cruce", "no crossing"), f["change_to_critical"], f["note"]])
        r = wsn.max_row
        for col in (4, 5, 6):
            wsn.cell(row=r, column=col).number_format = fmt_m
        is_frac = f["key"].endswith("occupancy") or f["key"] in ("rate", "tax")
        for col in (3, 7):
            wsn.cell(row=r, column=col).number_format = "0.00%" if is_frac else fmt_m
        wsn.cell(row=r, column=8).number_format = "0.0%"
    for b in res["sweeps"]:
        wsn.append([])
        wsn.append([tr("Barrido de ", "Sweep of ") + loc_path(b["key"]), "", tr("Valor", "Value"), tr("VAN", "NPV"), tr("TIR", "IRR")])
        for c in wsn[wsn.max_row]:
            c.font = bold
        for p in b["points"]:
            wsn.append(["", "", p["value"], p["npv"], p["irr"] if p["irr"] is not None else tr("no definida", "undefined")])
            wsn.cell(row=wsn.max_row, column=4).number_format = fmt_m
            wsn.cell(row=wsn.max_row, column=5).number_format = "0.00%"
    for col, w in zip("ABCDEFGHI", (34, 34, 16, 16, 16, 16, 20, 18, 50)):
        wsn.column_dimensions[col].width = w
    for row in wsn.iter_rows(min_row=1, max_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=False)
    wb.save(path)


def to_json(res: dict) -> dict:
    """English keys. financial-evaluation reads cash_flow.years + cash_flow.net_cash_flow and the
    tax regime from parameters.tax and parameters.depreciation."""
    cf = res["cash_flow"]
    return {
        "method": "economic-evaluation",
        "project": res["project"], "company": res["company"], "profile": res["profile"], "language": res["language"],
        "note": res["note"], "parameters": res["parameters"],
        "assumptions": contract_assumptions(res["assumptions"]),
        "traceability": [{k: v for k, v in r.items() if k != "origin_code"} for r in res["assumptions"]],
        "cash_flow": {"years": cf["years"],
                      "rows": [{"code": c, "item": e, "values": v, "origin": o} for c, e, v, o in cf["rows"]],
                      "net_cash_flow": cf["flow"], "assets": cf["assets"]},
        "indicators": res["indicators"], "sensitivity": res["sensitivity"], "sweeps": res["sweeps"],
    }


def text_summary(res: dict) -> str:
    ind = res["indicators"]
    dec = res["parameters"]["decimals"]["value"]
    cur = res["parameters"]["currency"]["value"]
    if _LANG == "en":
        fm = lambda x: f"{x:,.{dec}f}"
        pc = lambda x: "undefined" if x is None else f"{x * 100:.2f} %"
        yr = lambda x: "not recovered" if x is None else f"{x:.2f} years"
        idx = "not applicable" if ind["npv_index"] is None else f"{ind['npv_index']:.3f}"
        chg = lambda x: f" ({x * 100:+.1f} %)"
    else:
        fm = lambda x: f"{x:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
        pc = lambda x: "no definida" if x is None else f"{x * 100:.2f} %".replace(".", ",")
        yr = lambda x: "no se recupera" if x is None else f"{x:.2f} años".replace(".", ",")
        idx = "no aplica" if ind["npv_index"] is None else f"{ind['npv_index']:.3f}".replace(".", ",")
        chg = lambda x: f" ({x * 100:+.1f} %)".replace(".", ",")
    lines = [tr("Proyecto: ", "Project: ") + f"{res['project']} · {res['profile']}",
             tr("Tasa: ", "Rate: ") + pc(ind["rate"]) + (f" ({mark(True)})" if res["parameters"]["rate"]["assumption"] else ""),
             tr("Flujo", "Flow") + f" ({cur}): " + " | ".join(fm(x) for x in res["cash_flow"]["flow"]),
             tr("VAN: ", "NPV: ") + f"{fm(ind['npv'])} {cur} → {ind['decision']}",
             tr("TIR: ", "IRR: ") + pc(ind["irr"]) + tr(" · PRI: ", " · Payback: ") + yr(ind["payback"])
             + tr(" · PRI descontado: ", " · Discounted payback: ") + yr(ind["discounted_payback"])
             + tr(" · IVAN: ", " · NPV index: ") + idx]
    for a in ind["warnings"]:
        lines.append(tr("AVISO: ", "WARNING: ") + a)
    lines.append(tr("Sensibilidad (de la que más pesa a la que menos):", "Sensitivity (from the heaviest to the lightest):"))
    for f in res["sensitivity"]["rows"][:6]:
        frac = f["key"].endswith("occupancy") or f["key"] in ("rate", "tax")
        crit = tr("sin cruce", "no crossing") if f["critical"] is None else (pc(f["critical"]) if frac else fm(f["critical"]))
        change = "" if f["change_to_critical"] is None else chg(f["change_to_critical"])
        lines.append(f"  {f['variable']}: " + tr("crítico ", "critical ") + f"{crit}{change}" + (f" · {f['note']}" if f["note"] else ""))
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Test
# ---------------------------------------------------------------------------


def _to_spanish(obj, parent=None):
    """English input → Spanish keys and codes (only to test the alias map)."""
    if isinstance(obj, list):
        return [_to_spanish(x, parent) for x in obj]
    if not isinstance(obj, dict):
        return obj
    out = {}
    for k, v in obj.items():
        nk = ES_KEY.get(k, k) if k not in ("rate", "tax") else k
        if k == "assumption" and isinstance(v, list):
            v = [ES_KEY.get(x, x) for x in v]
        elif k in VALUE_KEYS and isinstance(v, str):
            v = ES_VALUE.get(v, v)
        elif isinstance(v, (dict, list)):
            v = _to_spanish(v, k)
        out[nk] = v
    return out


def run_test() -> int:
    global _LANG
    tdir = SKILL / "tests"
    case = json.loads((tdir / "sample-case.json").read_text(encoding="utf-8"))
    profile = read_simple_yaml(tdir / "sample-profile.yaml")
    expected = json.loads((tdir / "expected.json").read_text(encoding="utf-8"))
    res = evaluate(case, profile, "name", "sample-profile.yaml")
    ind = res["indicators"]
    failures = []

    def near(name, got, exp, tol):
        if got is None or abs(got - exp) > tol:
            failures.append(f"{name}: got {got}, expected {exp} (±{tol})")

    tol_m = expected["currency_tolerance"]
    for t, (a, b) in enumerate(zip(res["cash_flow"]["flow"], expected["cash_flow"])):
        near(f"flow year {t}", a, b, tol_m)
    near("NPV", ind["npv"], expected["npv"], tol_m)
    near("IRR", ind["irr"], expected["irr"], expected["rate_tolerance"])
    near("payback", ind["payback"], expected["payback"], 1e-4)
    near("discounted payback", ind["discounted_payback"], expected["discounted_payback"], 1e-4)
    near("NPV index", ind["npv_index"], expected["npv_index"], 1e-4)
    crit = {f["key"]: f["critical"] for f in res["sensitivity"]["rows"]}
    for key, val in expected["critical_values"].items():
        near(f"critical {key}", crit.get(key), val, expected["critical_tolerance"].get(key, 1e-6))
    # the critical value makes NPV zero (every variable)
    m = res["_model"]
    routes = {c: r for c, _, r in sensitive_variables(m)}
    for f in res["sensitivity"]["rows"]:
        if f["critical"] is not None:
            v0 = npv_with(m, routes[f["key"]], f["critical"])
            if abs(v0) > 1.0:
                failures.append(f"NPV at the critical value of {f['key']} = {v0:.4f}, not zero")
    # the critical value of the rate is the IRR
    near("critical rate = IRR", crit.get("rate"), ind["irr"], 1e-7)
    # flow with three sign changes: IRR 10 %, 20 % and 30 %
    roots = irr_roots([-1000, 3600, -4310, 1716])
    if len(roots) != 3 or any(abs(a - b) > 1e-7 for a, b in zip(roots, (0.10, 0.20, 0.30))):
        failures.append(f"several IRRs: got {roots}, expected 10 %, 20 % and 30 %")
    if sign_changes([-1000, 3600, -4310, 1716]) != 3:
        failures.append("sign changes miscounted")
    # profile without a rate: --rate is required and marked as an assumption
    no_rate = copy.deepcopy(profile)
    no_rate["finance"]["discount_rate"] = None
    try:
        evaluate(case, no_rate, "name", "no-rate")
        failures.append("with a null rate in the profile it did not require --rate")
    except InputError:
        pass
    r2 = evaluate(case, no_rate, "name", "no-rate", rate=0.12)
    if not r2["parameters"]["rate"]["assumption"]:
        failures.append("the rate given by argument was not marked as an assumption")
    if abs(r2["indicators"]["npv"] - ind["npv"]) > 1e-6:
        failures.append("--rate 0.12 does not give the same NPV as the profile with 0.12")
    # audit fix: the input note and the tax-loss choice reach the assumptions, marked
    sup = res["assumptions"]
    if not any(s["value"] == case["note"] and s["mark"] == "SUPUESTO" for s in sup):
        failures.append("the input note is not in the assumptions marked SUPUESTO")
    ja = to_json(res)["assumptions"]
    if not ja or any(set(a) != {"parameter", "value", "origin", "text"} for a in ja) or             any(a["origin"] not in ("input", "profile", "template", "default", "script") and
                not a["origin"].startswith("research:") for a in ja):
        failures.append("JSON assumptions are not [{parameter, value, origin, text}] with a known origin")
    if not any(a["origin"] == "default" and a["text"].startswith("SUPUESTO · ") for a in to_json(r2)["assumptions"]) and             not any(a["origin"] == "input" for a in ja):
        failures.append("assumption origins not set")
    if not res["parameters"]["tax_losses"]["assumption"]:
        failures.append("tax_losses given without a source is not marked as an assumption")
    # Spanish keys (old private cases) give exactly the same result
    es_case = _to_spanish(case)
    es_profile = {"empresa": {"moneda": "CLP", "decimales_moneda": 0, "impuesto_renta": 0.27},
                  "finanzas": {"tasa_descuento": 0.12, "horizonte_anios": 5, "depreciacion": "lineal"}}
    if normalize_input(es_case) != normalize_input(case):
        failures.append("the Spanish alias map does not give back the English input")
    r3 = evaluate(es_case, es_profile, "name", "es", sweeps=["ingresos[0].ocupacion=0.5,0.55"])
    if r3["indicators"]["npv"] != ind["npv"] or [f["critical"] for f in r3["sensitivity"]["rows"]] != \
            [f["critical"] for f in res["sensitivity"]["rows"]]:
        failures.append("Spanish keys give a different NPV or different critical values")
    if abs(r3["sweeps"][0]["points"][1]["npv"] - ind["npv"]) > 1e-6 or r3["sweeps"][0]["key"] != "revenues[0].occupancy":
        failures.append("a Spanish sweep key is not read as revenues[0].occupancy")
    # Excel, es and en
    try:
        import tempfile
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            res = evaluate(case, profile, "name", "sample-profile.yaml")
            write_excel(res, Path(tmp) / "test-es.xlsx")
            sheets = load_workbook(Path(tmp) / "test-es.xlsx").sheetnames
            if sheets != ["SUPUESTOS", "FLUJO", "INDICADORES", "SENSIBILIDAD"]:
                failures.append(f"Excel sheets (es): {sheets}")
            en_profile = dict(copy.deepcopy(profile), language="en")
            ren = evaluate(case, en_profile, "name", "sample-profile.yaml")
            write_excel(ren, Path(tmp) / "test-en.xlsx")
            wb = load_workbook(Path(tmp) / "test-en.xlsx")
            if wb.sheetnames != ["ASSUMPTIONS", "CASH FLOW", "INDICATORS", "SENSITIVITY"]:
                failures.append(f"Excel sheets (en): {wb.sheetnames}")
            marks = {c.value for c in wb["ASSUMPTIONS"]["E"][4:]}
            if "ASSUMPTION" not in marks or "SUPUESTO" in marks:
                failures.append(f"en marker: {marks}")
            if ren["indicators"]["npv"] != ind["npv"] or not ren["indicators"]["decision"].startswith("ACCEPT"):
                failures.append("language en changes the NPV or the decision text")
            print(text_summary(ren))
            print()
    except ImportError:
        failures.append("openpyxl is missing for the Excel file")
    _LANG = "es"
    res = evaluate(case, profile, "name", "sample-profile.yaml")
    print(text_summary(res))
    print()
    if failures:
        for f in failures:
            print("FAIL:", f)
        print("FAIL")
        return 1
    print("OK")
    return 0


# ---------------------------------------------------------------------------


def main(argv=None):
    ap = argparse.ArgumentParser(description="Economic evaluation of the pure project: cash flow, NPV, IRR, sensitivity.")
    ap.add_argument("input", nargs="?", help="case JSON (see templates/input.json); Spanish keys are accepted")
    ap.add_argument("--profile", dest="profile", help="profile name (in <data>/profiles/ or perfiles/) or path to a .yaml")
    ap.add_argument("--rate", dest="rate", type=float, help="annual discount rate (fraction); marked as an assumption")
    ap.add_argument("--tax", dest="tax", type=float, help="income tax rate (fraction); marked as an assumption")
    ap.add_argument("--delta", type=float, default=0.10, help="sensitivity change (0.10 = ±10 %%)")
    ap.add_argument("--sweep", dest="sweep", action="append",
                    help='sweep one variable: "revenues[0].occupancy=0.3,0.5,0.7" (Spanish keys accepted)')
    ap.add_argument("--out", dest="out", help="output folder (default: the input's folder)")
    ap.add_argument("--test", dest="test", action="store_true", help="run the test with the sample case")
    # Spanish flags: hidden aliases (same dest).
    ap.add_argument("--perfil", dest="profile", help=argparse.SUPPRESS)
    ap.add_argument("--tasa", dest="rate", type=float, help=argparse.SUPPRESS)
    ap.add_argument("--impuesto", dest="tax", type=float, help=argparse.SUPPRESS)
    ap.add_argument("--barrer", dest="sweep", action="append", help=argparse.SUPPRESS)
    ap.add_argument("--salida", dest="out", help=argparse.SUPPRESS)
    ap.add_argument("--probar", dest="test", action="store_true", help=argparse.SUPPRESS)
    a = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if a.test:
        return run_test()
    if not a.input:
        ap.error("missing the input (or use --test)")
    try:
        in_path = Path(a.input)
        inp = json.loads(in_path.read_text(encoding="utf-8-sig"))
        p_path, how, p_name = find_profile(a.profile)
        profile = read_simple_yaml(p_path) if p_path else {}
        res = evaluate(inp, profile, how, p_name, a.rate, a.tax, a.delta, a.sweep)
    except InputError as e:
        print("ERROR:", e)
        return 2
    folder = Path(a.out) if a.out else in_path.parent
    folder.mkdir(parents=True, exist_ok=True)
    stem = in_path.stem
    if "input" in stem:
        base = stem.replace("input", "output")
    elif "entrada" in stem:
        base = stem.replace("entrada", "salida")
    else:
        base = stem + tr("-salida", "-output")
    (folder / f"{base}.json").write_text(json.dumps(to_json(res), ensure_ascii=False, indent=2), encoding="utf-8")
    write_excel(res, folder / f"{base}.xlsx")
    print(text_summary(res))
    print(tr(f"\nEscribí {base}.xlsx y {base}.json en la carpeta de salida.", f"\nWrote {base}.xlsx and {base}.json in the output folder."))
    return 0


if __name__ == "__main__":
    sys.exit(main())
