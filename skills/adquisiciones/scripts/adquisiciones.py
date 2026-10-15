# -*- coding: utf-8 -*-
"""
adquisiciones · comparar cotizaciones y asignar la compra en cascada.

    python adquisiciones.py <entrada.json> --salida <ADQUISICIONES.xlsx>
                            [--perfil <perfil.yaml>] [--empresa <nombre>] [--json <resultado.json>]
    python adquisiciones.py --probar

Qué calcula (el criterio lo pone quien llena la entrada; el script solo calcula):

1. Valida cada cotización contra su subtotal impreso. Si las líneas no suman lo
   impreso, aborta: falta una línea, sobra una o un total está mal copiado.
2. Deriva el unitario efectivo: total_línea ÷ cantidad. El unitario impreso es
   referencia, nunca dato (hay proveedores que imprimen lista y descuentan en el total).
3. Convierte a la moneda del perfil con el tipo de cambio que declara la cotización.
4. Reparte en cascada: al de menor unitario se le compra lo que tiene, el saldo pasa
   al siguiente y lo que nadie cubre queda POR COTIZAR. Empate de precio: lo decide
   el perfil (`compras.desempate`).
5. Compra en la unidad del proveedor: si vende tiras de 6 m, se compran tiras enteras
   y lo que sobra se declara y se paga (queda en el costo).

Busca el perfil en este orden: --perfil · variable INGENIERO_DATOS · carpeta hermana
../ingeniero-agentico-datos/perfiles/ · perfiles/_plantilla.yaml del repo.

Solo biblioteca estándar + openpyxl.
"""
import json
import math
import os
import shutil
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

EPS = 1e-9
AQUI = Path(__file__).resolve()
SKILL = AQUI.parents[1]
REPO = AQUI.parents[3] if len(AQUI.parents) > 3 else SKILL


class ErrorEntrada(Exception):
    """Error que impide calcular. Se corrige la entrada y se vuelve a correr."""


# ---------------------------------------------------------------- YAML mínimo
def _escalar(s):
    s = s.strip()
    if s in ("", "~", "null", "Null", "NULL"):
        return None
    if s in ("true", "True"):
        return True
    if s in ("false", "False"):
        return False
    if s == "[]":
        return []
    if s == "{}":
        return {}
    if len(s) >= 2 and s[0] == s[-1] and s[0] in "\"'":
        return s[1:-1]
    if s.startswith("[") and s.endswith("]"):
        return [_escalar(x) for x in s[1:-1].split(",") if x.strip()]
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return s


def _sin_comentario(linea):
    dentro, q = False, ""
    for i, ch in enumerate(linea):
        if ch in "\"'":
            if not dentro:
                dentro, q = True, ch
            elif ch == q:
                dentro = False
        elif ch == "#" and not dentro and (i == 0 or linea[i - 1] in " \t"):
            return linea[:i].rstrip()
    return linea.rstrip()


def leer_yaml(texto):
    """Subconjunto de YAML que usan los perfiles: mapas anidados, listas y escalares."""
    filas = []
    for cruda in texto.splitlines():
        lin = _sin_comentario(cruda)
        if lin.strip():
            filas.append((len(lin) - len(lin.lstrip(" ")), lin.strip()))

    def bloque(i, ind):
        if i >= len(filas):
            return None, i
        if filas[i][1].startswith("- ") or filas[i][1] == "-":
            lista = []
            while i < len(filas) and filas[i][0] == ind and filas[i][1].startswith("-"):
                resto = filas[i][1][1:].strip()
                if not resto:
                    val, i = bloque(i + 1, filas[i + 1][0]) if i + 1 < len(filas) else (None, i + 1)
                    lista.append(val)
                    continue
                if ":" in resto and not resto.startswith(("\"", "'")):
                    # elemento mapa: «- clave: valor» y sus continuaciones más adentro
                    sub_ind = ind + 2
                    filas[i] = (sub_ind, resto)
                    val, i = bloque(i, sub_ind)
                    lista.append(val)
                else:
                    lista.append(_escalar(resto))
                    i += 1
            return lista, i
        mapa = {}
        while i < len(filas) and filas[i][0] == ind and not filas[i][1].startswith("- "):
            clave, _, valor = filas[i][1].partition(":")
            clave = clave.strip().strip("\"'")
            if valor.strip():
                mapa[clave] = _escalar(valor)
                i += 1
            elif i + 1 < len(filas) and filas[i + 1][0] > ind:
                mapa[clave], i = bloque(i + 1, filas[i + 1][0])
            elif i + 1 < len(filas) and filas[i + 1][0] == ind and filas[i + 1][1].startswith("- "):
                mapa[clave], i = bloque(i + 1, ind)
            else:
                mapa[clave] = None
                i += 1
        return mapa, i

    datos, _ = bloque(0, filas[0][0]) if filas else ({}, 0)
    return datos or {}


# ---------------------------------------------------------------- perfil
DEFECTO = {"moneda": "CLP", "decimales_moneda": 0, "iva": 0.19, "desempate": "mayor_disponibilidad"}
DESEMPATES = ("mayor_disponibilidad", "menor_plazo")


def _elegir(carpeta, empresa):
    yamls = sorted(p for p in Path(carpeta).glob("*.yaml") if not p.name.startswith("_"))
    if empresa:
        for p in yamls:
            if p.stem.lower() == empresa.lower():
                return p
        raise ErrorEntrada(f"No hay perfil «{empresa}» en {Path(carpeta).name}/. "
                           f"Hay: {', '.join(p.stem for p in yamls) or 'ninguno'}")
    if len(yamls) == 1:
        return yamls[0]
    if len(yamls) > 1:
        raise ErrorEntrada("Hay varios perfiles; elige uno con --empresa: "
                           + ", ".join(p.stem for p in yamls))
    return None


def buscar_perfil(ruta=None, empresa=None):
    """Devuelve (ruta, es_plantilla). Orden: --perfil · INGENIERO_DATOS · hermana · plantilla."""
    if ruta:
        p = Path(ruta)
        if not p.is_file():
            raise ErrorEntrada(f"--perfil no existe: {ruta}")
        return p, p.name.startswith("_plantilla")
    env = os.environ.get("INGENIERO_DATOS")
    if env:
        p = Path(env)
        if p.is_file():
            return p, False
        for c in (p / "perfiles", p):
            if c.is_dir():
                hallado = _elegir(c, empresa)
                if hallado:
                    return hallado, False
    hermana = REPO.parent / "ingeniero-agentico-datos" / "perfiles"
    if hermana.is_dir():
        hallado = _elegir(hermana, empresa)
        if hallado:
            return hallado, False
    plantilla = REPO / "perfiles" / "_plantilla.yaml"
    return (plantilla if plantilla.is_file() else None), True


def cargar_perfil(ruta=None, empresa=None):
    p, es_plantilla = buscar_perfil(ruta, empresa)
    crudo = leer_yaml(p.read_text(encoding="utf-8-sig")) if p else {}
    emp = crudo.get("empresa") or {}
    com = crudo.get("compras") or {}
    perfil = {"origen": p.name if p else "(sin archivo: valores por defecto)",
              "es_plantilla": es_plantilla or p is None, "supuestos": [],
              "proveedores": com.get("proveedores") or []}
    for campo, fuente in (("moneda", emp), ("decimales_moneda", emp), ("iva", emp),
                          ("desempate", com)):
        v = fuente.get(campo)
        if v in (None, ""):
            if campo in ("moneda", "iva") and not perfil["es_plantilla"]:
                raise ErrorEntrada(f"El perfil {perfil['origen']} no trae `{campo}`. "
                                   f"Se pregunta a la empresa; no se inventa.")
            v = DEFECTO[campo]
            perfil["supuestos"].append(f"SUPUESTO · {campo} = {v} (el perfil no lo trae)")
        perfil[campo] = v
    if perfil["es_plantilla"]:
        perfil["supuestos"].insert(0, "SUPUESTO · se usó la plantilla del perfil, no el de la "
                                      "empresa: moneda, IVA y desempate son los de ejemplo")
    perfil["moneda"] = str(perfil["moneda"]).upper()
    perfil["decimales_moneda"] = int(perfil["decimales_moneda"])
    perfil["iva"] = float(perfil["iva"])
    if perfil["desempate"] not in DESEMPATES:
        raise ErrorEntrada(f"desempate «{perfil['desempate']}» no existe. Usa: {', '.join(DESEMPATES)}")
    return perfil


# ---------------------------------------------------------------- utilidades
def num(v, campo, donde):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        raise ErrorEntrada(f"{donde}: `{campo}` no es un número")
    if isinstance(v, (int, float)):
        return float(v)
    raise ErrorEntrada(f"{donde}: `{campo}` debe ser número (vino «{v}»). "
                       f"Escríbelo sin separador de miles: 1234567.89")


def fmt(x, dec=0):
    """Número a la chilena para notas de texto: 1.234.567 · 1.234,56."""
    if x is None:
        return "—"
    s = f"{x:,.{dec}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def g(x):
    """Cantidad sin ceros de más."""
    return fmt(x, 0) if abs(x - round(x)) < 1e-6 else fmt(x, 2).rstrip("0").rstrip(",")


# ---------------------------------------------------------------- preparar
def preparar(entrada, perfil):
    """Valida la entrada y la deja como una lista de ofertas comparables."""
    errores, avisos, supuestos = [], [], []
    req = entrada.get("requerimiento") or []
    cots = entrada.get("cotizaciones") or []
    if not req:
        errores.append("La entrada no trae `requerimiento`.")
    if not cots:
        errores.append("La entrada no trae `cotizaciones`.")
    ns = [r.get("n") for r in req]
    for n in sorted({n for n in ns if ns.count(n) > 1}, key=str):
        errores.append(f"N° de ítem repetido en el requerimiento: {n}")
    for r in req:
        if num(r.get("cantidad"), "cantidad", f"ítem {r.get('n')}") in (None, 0.0):
            errores.append(f"Ítem {r.get('n')}: sin cantidad requerida")
    validos = set(ns)

    ids = [c.get("id") for c in cots]
    for i in sorted({i for i in ids if ids.count(i) > 1}, key=str):
        errores.append(f"Cotización con id repetido: {i}")

    base = perfil["moneda"]
    ofertas, cuadratura, impresos, sin_item = [], [], [], []
    sin_disp = 0
    for c in cots:
        cid, prov = c.get("id"), c.get("proveedor")
        donde = f"Cotización {cid}"
        if not cid or not prov:
            errores.append(f"{donde}: falta `id` o `proveedor`")
            continue
        moneda = (c.get("moneda") or "").upper()
        if not moneda:
            moneda = base
            supuestos.append(f"SUPUESTO · {cid}: no declara moneda; se asume {base}")
        tc = 1.0
        if moneda != base:
            tc = num(c.get("tipo_cambio"), "tipo_cambio", donde)
            if not tc:
                errores.append(f"{donde}: cotiza en {moneda} y no declara `tipo_cambio`. El cambio "
                               f"lo pone quien paga (el que cobra su banco), no una web.")
                tc = None
            elif not c.get("fecha_cambio"):
                avisos.append(f"{donde}: tipo de cambio sin `fecha_cambio`; no se podrá reconstruir.")
        if not c.get("identificador"):
            avisos.append(f"{donde} ({prov}): sin identificador tributario. Sin él no se emite la orden.")
        dec_cot = 2 if moneda != base else perfil["decimales_moneda"]
        suma, vistos, desc_vistas = 0.0, set(), defaultdict(list)
        for k, l in enumerate(c.get("lineas") or [], 1):
            lid = str(l.get("linea", k))
            dl = f"{donde} línea {lid}"
            if lid in vistos:
                errores.append(f"{dl}: número de línea repetido en la misma cotización")
            vistos.add(lid)
            cant = num(l.get("cantidad"), "cantidad", dl)
            tot = num(l.get("total_linea"), "total_linea", dl)
            imp = num(l.get("unitario_impreso"), "unitario_impreso", dl)
            if not cant or cant <= 0:
                errores.append(f"{dl}: sin cantidad. Una línea sin cantidad no se compara ni se compra.")
                continue
            sup_tot = False
            if tot is None:
                if imp is None:
                    errores.append(f"{dl}: sin total_linea ni unitario_impreso")
                    continue
                tot, sup_tot = imp * cant, True
                supuestos.append(f"SUPUESTO · {dl}: no imprime total; se usó unitario impreso × cantidad")
            suma += tot
            desc_vistas[(l.get("descripcion") or "").strip().upper()].append(lid)
            uef_orig = tot / cant
            if imp and abs(imp - uef_orig) > max(abs(imp) * 0.005, 0.01):
                impresos.append({"cotizacion": cid, "proveedor": prov, "linea": lid,
                                 "impreso": imp, "efectivo": uef_orig, "moneda": moneda,
                                 "dif": uef_orig / imp - 1})
            item = l.get("item")
            if item in (None, ""):
                sin_item.append({"cotizacion": cid, "proveedor": prov, "linea": lid,
                                 "descripcion": l.get("descripcion") or "", "total": tot, "moneda": moneda})
                continue
            if item not in validos:
                errores.append(f"{dl} apunta al ítem {item}, que no existe en el requerimiento")
                continue
            uv = l.get("unidad_venta") or None
            contenido = 1.0
            if uv:
                contenido = num(uv.get("contenido", uv.get("largo")), "unidad_venta.contenido", dl)
                if not contenido or contenido <= 0:
                    errores.append(f"{dl}: `unidad_venta` sin `contenido` (cuánto del ítem trae "
                                   f"cada unidad que vende el proveedor)")
                    continue
                if not uv.get("textual"):
                    avisos.append(f"{dl}: `unidad_venta` sin `textual`; es la prueba de dónde salió "
                                  f"el {g(contenido)}.")
            disp = num(l.get("disponible"), "disponible", dl)
            disp_sup = False
            if disp is None:
                disp, disp_sup = cant, True
                sin_disp += 1
            q_req = round(cant * contenido, 6)
            ofertas.append({
                "item": item, "proveedor": prov, "vendedor": (c.get("vendedor") or "").strip(),
                "cotizacion": cid, "linea": lid, "fecha": c.get("fecha") or "",
                "plazo_dias": c.get("plazo_dias"), "identificador": c.get("identificador") or "",
                "moneda": moneda, "tipo_cambio": tc, "cantidad": cant, "contenido": contenido,
                "uv": uv, "q_req": q_req, "disp_unid": disp, "disp_req": round(disp * contenido, 6),
                "disp_supuesta": disp_sup, "total_orig": tot, "total_sup": sup_tot,
                "u_orig": tot / q_req, "u": (tot * tc / q_req) if tc else None,
                "descarte": (l.get("descarte") or "").strip(), "nota": l.get("nota") or "",
                "descripcion": l.get("descripcion") or "",
            })
        for d, lids in desc_vistas.items():
            if d and len(lids) > 1:
                avisos.append(f"{donde}: «{d[:40]}» aparece en las líneas {', '.join(lids)}. No es error "
                              f"si venía en dos partidas: van las dos.")
        sub = num(c.get("subtotal_impreso"), "subtotal_impreso", donde)
        tol = max(2 * 10 ** -dec_cot, 0.05) if dec_cot else 2.0
        estado = "SIN ANCLA"
        if sub is None:
            avisos.append(f"{donde}: sin `subtotal_impreso`; la extracción no se pudo verificar completa.")
        elif abs(suma - sub) > tol:
            errores.append(f"{donde}: NO CUADRA. Las líneas suman {fmt(suma, dec_cot)} y el documento "
                           f"imprime {fmt(sub, dec_cot)} (diferencia {fmt(suma - sub, dec_cot)}). "
                           f"Falta una línea, sobra una o un total está mal copiado.")
            estado = "NO CUADRA"
        else:
            estado = "CUADRA"
        cuadratura.append({"cotizacion": cid, "proveedor": prov, "moneda": moneda,
                           "tipo_cambio": tc if moneda != base else None,
                           "lineas": len(c.get("lineas") or []), "suma": round(suma, 2),
                           "subtotal_impreso": sub, "estado": estado,
                           "vence": c.get("vence") or "", "condiciones": c.get("condiciones") or "",
                           "plazo_dias": c.get("plazo_dias"), "archivo": c.get("archivo") or "",
                           "identificador": c.get("identificador") or "",
                           "vendedor": (c.get("vendedor") or "").strip(), "fecha": c.get("fecha") or ""})
    if sin_disp:
        supuestos.append(f"SUPUESTO · {sin_disp} línea(s) sin `disponible`: se asume que el proveedor "
                         f"entrega todo lo que cotizó")

    # Contraoferta: mismo ítem + mismo proveedor + mismo vendedor en DOS cotizaciones.
    # La más nueva reemplaza a la vieja; no compiten entre sí.
    reemplazos = []
    grupos = defaultdict(set)
    for o in ofertas:
        grupos[(o["item"], o["proveedor"], o["vendedor"])].add((o["fecha"], o["cotizacion"]))
    fuera = set()
    for (item, prov, vend), cs in grupos.items():
        if len(cs) < 2:
            continue
        fechas = [f for f, _ in cs]
        if any(not f for f in fechas) or len(set(fechas)) != len(fechas):
            errores.append(f"Ítem {item} · {prov}{' · ' + vend if vend else ''}: aparece en "
                           f"{len(cs)} cotizaciones sin fecha que las ordene. Si es contraoferta, "
                           f"pon `fecha`; si son vendedores distintos, pon `vendedor`.")
            continue
        nueva = max(cs)[1]
        for f, cid in cs:
            if cid != nueva:
                fuera.add((item, cid))
        antes = [o["u"] for o in ofertas if o["item"] == item and (item, o["cotizacion"]) in fuera
                 and o["proveedor"] == prov and o["u"] is not None]
        ahora = [o["u"] for o in ofertas if o["item"] == item and o["cotizacion"] == nueva and o["u"] is not None]
        reemplazos.append({"item": item, "proveedor": prov, "vendedor": vend, "vale": nueva,
                           "antes": min(antes) if antes else None, "ahora": min(ahora) if ahora else None})
    for o in ofertas:
        o["reemplazada"] = (o["item"], o["cotizacion"]) in fuera

    usados = {o["proveedor"] for o in ofertas}
    for c in cots:
        if c.get("proveedor") and c["proveedor"] not in usados:
            avisos.append(f"{c['proveedor']} ({c.get('id')}): ninguna línea quedó asignada a un ítem. "
                          f"¿Falta emparejar su cotización?")
    con_oferta = {o["item"] for o in ofertas if not o["descarte"] and not o["reemplazada"]}
    for r in req:
        if r.get("n") not in con_oferta:
            avisos.append(f"Ítem {r.get('n')} «{r.get('item')}»: nadie lo cubre")
    if errores:
        raise ErrorEntrada("\n".join(errores))
    return ofertas, cuadratura, impresos, reemplazos, avisos, supuestos, sin_item


# ---------------------------------------------------------------- reparto
def plazo_de(o, perfil):
    if o.get("plazo_dias") not in (None, ""):
        return float(o["plazo_dias"])
    for p in perfil.get("proveedores") or []:
        if isinstance(p, dict) and o["identificador"] and str(p.get("rut")) == o["identificador"]:
            if p.get("plazo_entrega_dias") not in (None, ""):
                return float(p["plazo_entrega_dias"])
    return math.inf


def etiqueta_uv(o):
    uv = o["uv"]
    if not uv:
        return ""
    nombre = uv.get("nombre") or "unidad"
    return f"{nombre} de {g(o['contenido'])} {uv.get('unidad', '')}".strip()


def cascada(requerimiento, ofertas, perfil):
    vivas = [o for o in ofertas if not o["descarte"] and not o["reemplazada"]]
    por_item = defaultdict(list)
    for o in vivas:
        por_item[o["item"]].append(o)
    criterio = perfil["desempate"]

    def clave(o):
        if criterio == "menor_plazo":
            return (o["u"], plazo_de(o, perfil), -o["disp_req"])
        return (o["u"], -o["disp_req"])

    plan = []
    for r in requerimiento:
        n, falta = r["n"], float(r["cantidad"])
        cands = sorted(por_item.get(n, []), key=clave)   # sort estable: a igual clave, orden de entrada
        tramos, orden = [], 0
        for o in cands:
            if falta <= EPS:
                break
            if o["disp_req"] <= EPS:
                continue
            toma, sobra, unidades = min(falta, o["disp_req"]), 0.0, None
            notas = []
            if o["uv"]:
                # Se compra en la unidad del proveedor: tiras, barras, cajas enteras.
                unidades = math.ceil(round(toma / o["contenido"], 6) - EPS)
                tope = math.floor(round(o["disp_unid"], 6) + EPS)
                if unidades > tope:
                    unidades = tope
                if unidades <= 0:
                    continue
                toma_ent = round(unidades * o["contenido"], 6)
                sobra = max(0.0, round(toma_ent - toma, 6))
                toma = toma_ent
                if sobra > EPS:
                    notas.append(f"sobran {g(sobra)} {r.get('unidad', '')}".rstrip()
                                 + " por comprar unidades enteras (se pagan)")
            else:
                unidades = toma / o["contenido"]
            empatados = [x for x in cands if x is not o and abs(x["u"] - o["u"]) < 1e-9]
            if empatados:
                otros = ", ".join(f"{x['proveedor']} ({x['cotizacion']})" for x in empatados)
                notas.append(f"empate de precio con {otros}; desempate: {criterio.replace('_', ' ')}")
            if o["disp_req"] + EPS < float(r["cantidad"]):
                notas.append(f"disponible {g(o['disp_req'])} de {g(float(r['cantidad']))}")
            if o["disp_supuesta"]:
                notas.append("disponible SUPUESTO = lo cotizado")
            if o["nota"]:
                notas.append(o["nota"])
            orden += 1
            subtotal = round(toma * o["u"], 2)
            traza = (f"{o['cotizacion']} línea {o['linea']}: {fmt(o['total_orig'], 2 if o['moneda'] != perfil['moneda'] else perfil['decimales_moneda'])} {o['moneda']}"
                     f" ÷ {g(o['q_req'])}"
                     + (f" ({g(o['cantidad'])} × {g(o['contenido'])})" if o["uv"] else "")
                     + (f" × {fmt(o['tipo_cambio'], 2)}" if o["moneda"] != perfil["moneda"] else "")
                     + f" = {fmt(o['u'], 2)} por {r.get('unidad') or 'unidad'}")
            tramos.append({
                "orden": orden, "proveedor": o["proveedor"], "vendedor": o["vendedor"],
                "cotizacion": o["cotizacion"], "linea": o["linea"],
                "unidades": round(unidades, 6), "unidad_venta": etiqueta_uv(o),
                "cantidad": round(toma, 6), "sobra": sobra,
                "unitario": round(o["u"], 2), "subtotal": subtotal,
                "moneda_origen": o["moneda"], "unitario_origen": round(o["u_orig"] * (o["contenido"] if o["uv"] else 1), 4),
                "subtotal_origen": round(toma * o["u_orig"], 2) if o["moneda"] != perfil["moneda"] else None,
                "traza": traza, "nota": " · ".join(notas),
            })
            falta = max(0.0, round(falta - toma, 6))
        motivo = ""
        if falta > EPS:
            descartadas = [o for o in ofertas if o["item"] == n and o["descarte"]]
            if cands:
                motivo = "nadie cubre el saldo"
            elif descartadas:
                motivo = f"solo hubo ofertas descartadas ({len(descartadas)})"
            else:
                motivo = "nadie lo cotizó"
            tramos.append({"orden": orden + 1, "proveedor": "POR COTIZAR", "cotizacion": "",
                           "linea": "", "cantidad": falta, "unitario": None, "subtotal": None,
                           "traza": "", "nota": motivo})
        plan.append({"req": r, "tramos": tramos, "faltante": falta, "motivo": motivo,
                     "candidatos": cands})
    return plan


def resumir(plan, perfil):
    neto = 0.0
    por_prov = {}
    for p in plan:
        for t in p["tramos"]:
            if t["subtotal"] is None:
                continue
            neto = round(neto + t["subtotal"], 2)
            k = t["proveedor"] + (f" · {t['vendedor']}" if t.get("vendedor") else "")
            e = por_prov.setdefault(k, {"proveedor": k, "cotizaciones": [], "lineas": 0, "neto": 0.0})
            e["lineas"] += 1
            e["neto"] = round(e["neto"] + t["subtotal"], 2)
            if t["cotizacion"] not in e["cotizaciones"]:
                e["cotizaciones"].append(t["cotizacion"])
    iva = round(neto * perfil["iva"], 2)
    completos = sum(1 for p in plan if p["faltante"] <= EPS)
    return {"neto": neto, "iva": iva, "total": round(neto + iva, 2), "items": len(plan),
            "completos": completos, "incompletos": len(plan) - completos}, \
        sorted(por_prov.values(), key=lambda e: -e["neto"])


def calcular(entrada, perfil):
    ofertas, cuadratura, impresos, reemplazos, avisos, supuestos, sin_item = preparar(entrada, perfil)
    plan = cascada(entrada["requerimiento"], ofertas, perfil)
    totales, por_prov = resumir(plan, perfil)
    return {
        "titulo": entrada.get("titulo", ""),
        "perfil": {k: perfil[k] for k in ("origen", "moneda", "decimales_moneda", "iva", "desempate")},
        "totales": totales,
        "por_proveedor": por_prov,
        "items": [{"n": p["req"]["n"], "item": p["req"].get("item", ""),
                   "cantidad": float(p["req"]["cantidad"]), "unidad": p["req"].get("unidad", ""),
                   "faltante": p["faltante"], "tramos": p["tramos"]} for p in plan],
        "faltantes": [{"n": p["req"]["n"], "item": p["req"].get("item", ""),
                       "cantidad": p["faltante"], "motivo": p["motivo"]}
                      for p in plan if p["faltante"] > EPS],
        "descartes": [{"item": o["item"], "proveedor": o["proveedor"], "cotizacion": o["cotizacion"],
                       "linea": o["linea"], "motivo": o["descarte"]} for o in ofertas if o["descarte"]],
        "contraofertas": reemplazos,
        "lineas_sin_item": sin_item,
        "impreso_vs_efectivo": impresos,
        "cuadratura": cuadratura,
        "supuestos": perfil["supuestos"] + supuestos,
        "avisos": avisos,
        "_ofertas": ofertas, "_plan": plan,
    }


# ---------------------------------------------------------------- Excel
def escribir_excel(res, destino):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    dec = res["perfil"]["decimales_moneda"]
    DIN = "#,##0" + ("." + "0" * dec if dec else "")
    DIN2 = "#,##0.00"
    AZUL, GRIS, VERDE, ROJO, AMAR = "1F3A5F", "E7E6E6", "E2F0D9", "C00000", "FFF2CC"
    thin = Side(style="thin", color="BFBFBF")
    caja = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_tit, f_cab, f_neg = Font(bold=True, size=13, color="FFFFFF"), Font(bold=True, color="FFFFFF"), Font(bold=True)
    relleno = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731

    wb = openpyxl.Workbook()

    def hoja(nombre, titulo, ancho):
        ws = wb.create_sheet(nombre)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
        c = ws.cell(row=1, column=1, value=titulo)
        c.font, c.fill = f_tit, relleno(AZUL)
        return ws

    def cabecera(ws, fila, cols, anchos=None):
        for j, t in enumerate(cols, 1):
            c = ws.cell(row=fila, column=j, value=t)
            c.font, c.fill, c.border = f_cab, relleno(AZUL), caja
            c.alignment = Alignment(wrap_text=True, vertical="center")
        if anchos:
            for j, w in enumerate(anchos, 1):
                ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(row=fila + 1, column=1)

    mon = res["perfil"]["moneda"]
    t = res["totales"]

    # RESUMEN
    ws = hoja("RESUMEN", f"ADQUISICIONES · {res['titulo']}", 6)
    ws.cell(row=2, column=1, value=f"Perfil: {res['perfil']['origen']} · moneda {mon} · IVA "
                                   f"{res['perfil']['iva']:.0%} · desempate {res['perfil']['desempate']}")
    filas = [("TOTAL NETO", t["neto"]), (f"IVA {res['perfil']['iva']:.0%}", t["iva"]),
             ("TOTAL CON IVA", t["total"])]
    for i, (k, v) in enumerate(filas):
        ws.cell(row=4 + i, column=1, value=k).font = f_neg
        c = ws.cell(row=4 + i, column=2, value=v)
        c.number_format = DIN
    ws.cell(row=7, column=1, value="Ítems cubiertos al 100 %")
    ws.cell(row=7, column=2, value=f"{t['completos']} de {t['items']}")
    r = 9
    cabecera(ws, r, ["PROVEEDOR", "COTIZACIONES", "LÍNEAS", f"NETO {mon}", f"IVA", "% DE LA COMPRA"],
             [34, 22, 9, 16, 14, 14])
    for e in res["por_proveedor"]:
        r += 1
        ws.cell(row=r, column=1, value=e["proveedor"])
        ws.cell(row=r, column=2, value=", ".join(e["cotizaciones"]))
        ws.cell(row=r, column=3, value=e["lineas"])
        ws.cell(row=r, column=4, value=e["neto"]).number_format = DIN
        ws.cell(row=r, column=5, value=round(e["neto"] * res["perfil"]["iva"], 2)).number_format = DIN
        ws.cell(row=r, column=6, value=e["neto"] / t["neto"] if t["neto"] else None).number_format = "0.0%"
    r += 2
    ws.cell(row=r, column=1, value="LO QUE NADIE CUBRE (POR COTIZAR)").font = Font(bold=True, color=ROJO)
    if not res["faltantes"]:
        ws.cell(row=r + 1, column=1, value="Nada: todo el requerimiento quedó asignado.")
        r += 1
    for f in res["faltantes"]:
        r += 1
        ws.cell(row=r, column=1, value=f"Ítem {f['n']} · {f['item']}")
        ws.cell(row=r, column=2, value=f"falta {g(f['cantidad'])}")
        ws.cell(row=r, column=3, value=f["motivo"])
    r += 2
    ws.cell(row=r, column=1, value="SUPUESTOS (no vinieron del usuario)").font = f_neg
    for s in res["supuestos"] or ["Ninguno."]:
        r += 1
        ws.cell(row=r, column=1, value=s)

    # ASIGNACION
    ws = hoja("ASIGNACION", "ASIGNACIÓN EN CASCADA · al más barato lo que tenga, el saldo al siguiente", 14)
    cabecera(ws, 3, ["N°", "ÍTEM", "CANT. REQ.", "UN.", "ORD.", "PROVEEDOR", "COTIZACIÓN", "LÍNEA",
                     "UNIDADES DEL PROVEEDOR", "CANT. ASIGNADA", "SOBRA", f"UNITARIO {mon}",
                     f"SUBTOTAL {mon}", "DE DÓNDE SALE (traza) · NOTA"],
             [5, 38, 10, 6, 5, 20, 12, 7, 22, 12, 8, 13, 15, 80])
    r = 3
    for it in res["items"]:
        for i, tr in enumerate(it["tramos"]):
            r += 1
            if i == 0:
                ws.cell(row=r, column=1, value=it["n"])
                ws.cell(row=r, column=2, value=it["item"])
                ws.cell(row=r, column=3, value=it["cantidad"])
                ws.cell(row=r, column=4, value=it["unidad"])
            ws.cell(row=r, column=5, value=f"{tr['orden']}°")
            ws.cell(row=r, column=6, value=tr["proveedor"])
            ws.cell(row=r, column=7, value=tr["cotizacion"])
            ws.cell(row=r, column=8, value=tr["linea"])
            if tr.get("unidad_venta"):
                ws.cell(row=r, column=9, value=f"{g(tr['unidades'])} × {tr['unidad_venta']}")
            ws.cell(row=r, column=10, value=tr["cantidad"])
            if tr.get("sobra"):
                ws.cell(row=r, column=11, value=tr["sobra"])
            if tr["unitario"] is not None:
                ws.cell(row=r, column=12, value=tr["unitario"]).number_format = DIN2
            if tr["subtotal"] is None:
                c = ws.cell(row=r, column=13, value="FALTA")
                c.font = Font(bold=True, color=ROJO)
                for j in range(1, 15):
                    ws.cell(row=r, column=j).fill = relleno("FBE5E5")
            else:
                ws.cell(row=r, column=13, value=tr["subtotal"]).number_format = DIN
            ws.cell(row=r, column=14, value=" · ".join(x for x in (tr["traza"], tr["nota"]) if x))
            if len(it["tramos"]) > 1 and tr["subtotal"] is not None:
                for j in range(5, 14):
                    ws.cell(row=r, column=j).fill = relleno(AMAR)
    r += 1
    ws.cell(row=r, column=2, value="TOTAL NETO").font = f_neg
    c = ws.cell(row=r, column=13, value=t["neto"])
    c.font, c.number_format = f_neg, DIN
    ws.cell(row=r + 1, column=2, value="Amarillo = ítem partido entre proveedores. Rojo = lo que falta.")

    # COMPARATIVA
    provs = []
    for o in res["_ofertas"]:
        k = (o["proveedor"], o["cotizacion"])
        if k not in provs:
            provs.append(k)
    ancho = 3 + len(provs) + 3
    ws = hoja("COMPARATIVA", f"UNITARIO EFECTIVO EN {mon} (total ÷ cantidad; nunca el impreso)", ancho)
    cabecera(ws, 3, ["N°", "ÍTEM", "CANT."] + [f"{p}\n{c}" for p, c in provs]
             + ["MEJOR", "PEOR", "DIF. %"], [5, 38, 8] + [16] * len(provs) + [13, 13, 9])
    ws.row_dimensions[3].height = 32
    r = 3
    gano = {(it["n"], tr["cotizacion"]) for it in res["items"] for tr in it["tramos"] if tr["subtotal"] is not None}
    for it in res["items"]:
        r += 1
        ws.cell(row=r, column=1, value=it["n"])
        ws.cell(row=r, column=2, value=it["item"])
        ws.cell(row=r, column=3, value=it["cantidad"])
        vals = []
        for j, (p, cid) in enumerate(provs, 4):
            mias = [o for o in res["_ofertas"] if o["item"] == it["n"] and o["cotizacion"] == cid]
            vivas = [o for o in mias if not o["descarte"] and not o["reemplazada"]]
            c = ws.cell(row=r, column=j)
            if vivas:
                u = min(o["u"] for o in vivas)
                vals.append(u)
                c.value, c.number_format = round(u, 2), DIN2
                if (it["n"], cid) in gano:
                    c.fill = relleno(VERDE)
            elif mias:
                c.value, c.fill = ("DESCARTADA" if any(o["descarte"] for o in mias) else "REEMPLAZADA"), relleno(GRIS)
            else:
                c.value = "—"
            c.border = caja
        if vals:
            mejor, peor = min(vals), max(vals)
            ws.cell(row=r, column=ancho - 2, value=round(mejor, 2)).number_format = DIN2
            ws.cell(row=r, column=ancho - 1, value=round(peor, 2)).number_format = DIN2
            ws.cell(row=r, column=ancho, value=(peor / mejor - 1) if len(vals) > 1 and mejor else None
                    ).number_format = "0.0%"
        else:
            ws.cell(row=r, column=ancho - 2, value="NADIE").font = Font(bold=True, color=ROJO)
    ws.cell(row=r + 2, column=2, value="Verde = recibió asignación. Gris = línea fuera del reparto "
                                       "(motivo en VERIFICACION).")

    # PROVEEDORES
    ws = hoja("PROVEEDORES", "COTIZACIONES · cuadratura contra el subtotal impreso", 13)
    cabecera(ws, 3, ["COTIZACIÓN", "PROVEEDOR", "VENDEDOR", "IDENTIFICADOR", "FECHA", "MONEDA",
                     "TIPO CAMBIO", "LÍNEAS", "SUMA LÍNEAS", "SUBTOTAL IMPRESO", "ESTADO",
                     f"ASIGNADO NETO {mon}", "VENCE · CONDICIONES · PLAZO · ARCHIVO"],
             [12, 22, 14, 15, 11, 8, 11, 7, 15, 15, 11, 15, 60])
    asign = defaultdict(float)
    for it in res["items"]:
        for tr in it["tramos"]:
            if tr["subtotal"] is not None:
                asign[tr["cotizacion"]] += tr["subtotal"]
    r = 3
    for q in res["cuadratura"]:
        r += 1
        vals = [q["cotizacion"], q["proveedor"], q["vendedor"], q["identificador"] or "FALTA",
                q["fecha"], q["moneda"], q["tipo_cambio"], q["lineas"], q["suma"],
                q["subtotal_impreso"], q["estado"], round(asign.get(q["cotizacion"], 0.0), 2),
                " · ".join(str(x) for x in (q["vence"], q["condiciones"],
                                            f"{q['plazo_dias']} días" if q["plazo_dias"] not in (None, "") else "",
                                            q["archivo"]) if x)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            if j in (9, 10):
                c.number_format = DIN2
            if j == 12:
                c.number_format = DIN
        est = ws.cell(row=r, column=11)
        est.fill = relleno(VERDE if q["estado"] == "CUADRA" else ("FBE5E5" if q["estado"] == "NO CUADRA" else AMAR))

    # VERIFICACION
    ws = hoja("VERIFICACION", "VERIFICACIÓN · leer antes de emitir órdenes de compra", 6)
    for j, w in enumerate([5, 22, 14, 10, 14, 90], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    r = 3

    def bloque(titulo, filas, cols):
        nonlocal r
        ws.cell(row=r, column=1, value=titulo).font = f_neg
        r += 1
        if not filas:
            ws.cell(row=r, column=2, value="Nada que reportar.")
            r += 2
            return
        for j, h in enumerate(cols, 1):
            c = ws.cell(row=r, column=j, value=h)
            c.font, c.fill = f_cab, relleno(AZUL)
        for f in filas:
            r += 1
            for j, v in enumerate(f, 1):
                ws.cell(row=r, column=j, value=v)
        r += 2

    pedidos = sum(it["cantidad"] for it in res["items"])
    cubiertos = sum(tr["cantidad"] for it in res["items"] for tr in it["tramos"] if tr["subtotal"] is not None)
    sobras = sum(tr.get("sobra") or 0 for it in res["items"] for tr in it["tramos"])
    faltan = sum(f["cantidad"] for f in res["faltantes"])
    bloque("1 · CUADRATURA DE CANTIDADES (pedido = asignado − sobra + falta)",
           [["", "pedido", g(pedidos), "", "", ""], ["", "asignado", g(cubiertos), "", "", ""],
            ["", "sobra por unidad entera", g(sobras), "", "", ""], ["", "falta", g(faltan), "", "", ""],
            ["", "cuadra", "SÍ" if abs(pedidos - (cubiertos - sobras + faltan)) < 1e-4 else "NO", "", "", ""]],
           ["", "CONCEPTO", "CANTIDAD", "", "", ""])
    bloque("2 · LÍNEAS DESCARTADAS POR ESPECIFICACIÓN (fuera del reparto)",
           [[d["item"], d["proveedor"], d["cotizacion"], d["linea"], "", d["motivo"]] for d in res["descartes"]],
           ["ÍTEM", "PROVEEDOR", "COTIZACIÓN", "LÍNEA", "", "MOTIVO"])
    bloque("3 · UNITARIO IMPRESO ≠ EFECTIVO (descuento aplicado en el total; se usó el efectivo)",
           [["", x["proveedor"], x["cotizacion"], x["linea"], f"{x['dif'] * 100:+.1f} %".replace(".", ","),
             f"impreso {fmt(x['impreso'], 2)} · efectivo {fmt(x['efectivo'], 2)} {x['moneda']}"]
            for x in res["impreso_vs_efectivo"]],
           ["", "PROVEEDOR", "COTIZACIÓN", "LÍNEA", "DIF.", "DETALLE"])
    bloque("4 · CONTRAOFERTAS (la cotización más nueva reemplaza a la vieja)",
           [[x["item"], x["proveedor"], x["vale"], "", "",
             f"antes {fmt(x['antes'], 2)} → ahora {fmt(x['ahora'], 2)}"] for x in res["contraofertas"]],
           ["ÍTEM", "PROVEEDOR", "VALE", "", "", "DETALLE"])
    bloque("5 · LÍNEAS COTIZADAS SIN ÍTEM (flete, embalaje, lo no pedido): no entran al reparto; decidir aparte",
           [["", x["proveedor"], x["cotizacion"], x["linea"], fmt(x["total"], 2), f"{x['descripcion']} ({x['moneda']})"]
            for x in res["lineas_sin_item"]],
           ["", "PROVEEDOR", "COTIZACIÓN", "LÍNEA", "TOTAL", "DESCRIPCIÓN"])
    bloque("6 · SUPUESTOS", [["", "", "", "", "", s] for s in res["supuestos"]], ["", "", "", "", "", "SUPUESTO"])
    bloque("7 · AVISOS", [["", "", "", "", "", a] for a in res["avisos"]], ["", "", "", "", "", "AVISO"])

    del wb["Sheet"]
    wb.save(destino)


# ---------------------------------------------------------------- salida JSON
def a_json(res):
    return {k: v for k, v in res.items() if not k.startswith("_")}


def imprimir(res):
    mon, dec = res["perfil"]["moneda"], res["perfil"]["decimales_moneda"]
    t = res["totales"]
    print(f"Perfil       : {res['perfil']['origen']} ({mon}, IVA {res['perfil']['iva']:.0%}, "
          f"desempate {res['perfil']['desempate']})")
    print(f"TOTAL NETO   : {fmt(t['neto'], 2)} {mon}   IVA {fmt(t['iva'], 2)}   TOTAL {fmt(t['total'], 2)}")
    print(f"Ítems 100 %  : {t['completos']} de {t['items']}   incompletos: {t['incompletos']}")
    print(f"Órdenes      : {len(res['por_proveedor'])}")
    for e in res["por_proveedor"]:
        print(f"   {e['proveedor'][:28]:<28} {e['lineas']:>3} líneas  {fmt(e['neto'], dec):>16}")
    for f in res["faltantes"]:
        print(f"POR COTIZAR  : ítem {f['n']} «{f['item'][:40]}» falta {g(f['cantidad'])} ({f['motivo']})")
    for s in res["supuestos"]:
        print(f"   {s}")
    if res["avisos"]:
        print(f"Avisos ({len(res['avisos'])}), detalle en VERIFICACION:")
        for a in res["avisos"][:10]:
            print(f"   [!] {a}")


# ---------------------------------------------------------------- prueba
def probar():
    caso = SKILL / "pruebas" / "caso-ficticio.json"
    esperado = json.loads((SKILL / "pruebas" / "esperado.json").read_text(encoding="utf-8"))
    plantilla = REPO / "perfiles" / "_plantilla.yaml"
    perfil = cargar_perfil(str(plantilla) if plantilla.is_file() else None)
    if not plantilla.is_file():
        perfil.update({"origen": "(defecto)", **DEFECTO})
    entrada = json.loads(caso.read_text(encoding="utf-8"))
    res = calcular(entrada, perfil)
    fallas = []

    def igual(a, b, que):
        if isinstance(b, (int, float)) and not isinstance(b, bool):
            if a is None or abs(float(a) - float(b)) > 0.01:
                fallas.append(f"{que}: salió {a}, se esperaba {b}")
        elif a != b:
            fallas.append(f"{que}: salió {a!r}, se esperaba {b!r}")

    for k, v in esperado["totales"].items():
        igual(res["totales"].get(k), v, f"totales.{k}")
    got_prov = {e["proveedor"]: e for e in res["por_proveedor"]}
    igual(sorted(got_prov), sorted(e["proveedor"] for e in esperado["por_proveedor"]), "proveedores con orden")
    for e in esperado["por_proveedor"]:
        if e["proveedor"] in got_prov:
            igual(got_prov[e["proveedor"]]["neto"], e["neto"], f"neto de {e['proveedor']}")
            igual(got_prov[e["proveedor"]]["lineas"], e["lineas"], f"líneas de {e['proveedor']}")
    got_items = {it["n"]: it for it in res["items"]}
    for it in esperado["items"]:
        tr_got = [t for t in got_items[it["n"]]["tramos"]]
        igual(len(tr_got), len(it["tramos"]), f"ítem {it['n']} n° de tramos")
        for a, b in zip(tr_got, it["tramos"]):
            for k, v in b.items():
                igual(a.get(k), v, f"ítem {it['n']} tramo {b.get('orden')} {k}")
    igual([f["n"] for f in res["faltantes"]], [f["n"] for f in esperado["faltantes"]], "faltantes")
    igual(len(res["impreso_vs_efectivo"]), esperado["impreso_distinto_de_efectivo"], "líneas con impreso ≠ efectivo")
    igual(len(res["descartes"]), esperado["descartes"], "descartes")

    # Las trampas: la validación tiene que atajar estos errores.
    import copy
    malo = copy.deepcopy(entrada)
    malo["cotizaciones"][0]["lineas"][1]["total_linea"] += 1000
    try:
        calcular(malo, perfil)
        fallas.append("un total mal copiado no hizo abortar la cuadratura contra el subtotal impreso")
    except ErrorEntrada as e:
        if "NO CUADRA" not in str(e):
            fallas.append(f"abortó por otra causa: {e}")
    malo = copy.deepcopy(entrada)
    malo["cotizaciones"][0]["lineas"][0]["item"] = 99
    try:
        calcular(malo, perfil)
        fallas.append("una línea que apunta a un ítem inexistente no abortó")
    except ErrorEntrada:
        pass
    malo = copy.deepcopy(entrada)
    malo["cotizaciones"][0]["moneda"] = "USD"
    try:
        calcular(malo, perfil)
        fallas.append("una cotización en otra moneda sin tipo de cambio no abortó")
    except ErrorEntrada:
        pass

    tmp = Path(tempfile.mkdtemp(prefix="adquisiciones-"))
    try:
        escribir_excel(res, tmp / "prueba.xlsx")
        import openpyxl
        hojas = openpyxl.load_workbook(tmp / "prueba.xlsx").sheetnames
        for h in ("RESUMEN", "ASIGNACION", "COMPARATIVA", "PROVEEDORES", "VERIFICACION"):
            if h not in hojas:
                fallas.append(f"falta la hoja {h} en el Excel")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    imprimir(res)
    print()
    if fallas:
        print("FALLA")
        for f in fallas:
            print(f"   [X] {f}")
        return 1
    print("OK · el caso ficticio cuadra con pruebas/esperado.json y las tres trampas abortan")
    return 0


# ---------------------------------------------------------------- main
def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    import argparse
    ap = argparse.ArgumentParser(description="Comparar cotizaciones y asignar la compra en cascada.")
    ap.add_argument("entrada", nargs="?")
    ap.add_argument("--salida", help="libro Excel de salida")
    ap.add_argument("--json", help="resultado en JSON (por defecto, junto al Excel)")
    ap.add_argument("--perfil")
    ap.add_argument("--empresa", help="nombre del perfil cuando hay varios")
    ap.add_argument("--probar", action="store_true")
    a = ap.parse_args(argv)
    try:
        if a.probar:
            return probar()
        if not a.entrada or not a.salida:
            ap.print_help()
            return 1
        perfil = cargar_perfil(a.perfil, a.empresa)
        entrada = json.loads(Path(a.entrada).read_text(encoding="utf-8-sig"))
        res = calcular(entrada, perfil)
    except ErrorEntrada as e:
        print("ERRORES · no se genera nada hasta corregirlos:")
        for linea in str(e).splitlines():
            print(f"   [X] {linea}")
        return 2
    escribir_excel(res, a.salida)
    destino_json = a.json or str(Path(a.salida).with_suffix(".json"))
    Path(destino_json).write_text(json.dumps(a_json(res), ensure_ascii=False, indent=1), encoding="utf-8")
    imprimir(res)
    print(f"\nLibro: {a.salida}\nJSON : {destino_json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
