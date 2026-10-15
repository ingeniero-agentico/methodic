#!/usr/bin/env python3
"""evaluacion-economica · ¿conviene el proyecto? (proyecto puro, sin financiamiento)

Arma el flujo de caja del proyecto puro a la chilena y calcula VAN, TIR, período de
recuperación, IVAN y la sensibilidad (variable por variable ±x % y el valor crítico que deja
el VAN en cero). Escribe un Excel (SUPUESTOS · FLUJO · INDICADORES · SENSIBILIDAD) y un JSON.

Uso:
    python evaluar.py entrada.json [--perfil <nombre|ruta>] [--tasa 0.12] [--impuesto 0.27]
                      [--delta 0.10] [--barrer "ingresos[0].ocupacion=0.3,0.5,0.7"]
                      [--salida carpeta]
    python evaluar.py --probar

Solo biblioteca estándar + openpyxl (openpyxl solo para el Excel).
"""
from __future__ import annotations

import argparse
import copy
import json
import math
import os
import sys
from pathlib import Path

AQUI = Path(__file__).resolve().parent
SKILL = AQUI.parent
REPO = SKILL.parent.parent

# ---------------------------------------------------------------------------
# Perfil (YAML simple: mapas anidados y escalares; las listas se ignoran)
# ---------------------------------------------------------------------------


def _escalar(txt: str):
    t = txt.strip()
    if t in ("", "~", "null", "Null", "NULL"):
        return None
    if t in ("true", "True"):
        return True
    if t in ("false", "False"):
        return False
    if t == "[]":
        return []
    if t == "{}":
        return {}
    if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
        return t[1:-1]
    try:
        return int(t)
    except ValueError:
        pass
    try:
        return float(t)
    except ValueError:
        return t


def _sin_comentario(linea: str) -> str:
    dentro = None
    for i, c in enumerate(linea):
        if c in "\"'":
            if dentro is None:
                dentro = c
            elif dentro == c:
                dentro = None
        elif c == "#" and dentro is None and (i == 0 or linea[i - 1] in " \t"):
            return linea[:i]
    return linea


def leer_yaml_simple(ruta: Path) -> dict:
    raiz: dict = {}
    pila = [(-1, raiz)]
    for cruda in Path(ruta).read_text(encoding="utf-8-sig").splitlines():
        linea = _sin_comentario(cruda).rstrip()
        if not linea.strip():
            continue
        sangria = len(linea) - len(linea.lstrip(" "))
        cuerpo = linea.strip()
        if cuerpo.startswith("-") or ":" not in cuerpo:
            continue  # listas: este método no las usa
        clave, _, valor = cuerpo.partition(":")
        while len(pila) > 1 and sangria <= pila[-1][0]:
            pila.pop()
        padre = pila[-1][1]
        if valor.strip() == "":
            nuevo: dict = {}
            padre[clave.strip()] = nuevo
            pila.append((sangria, nuevo))
        else:
            padre[clave.strip()] = _escalar(valor)
    return raiz


def buscar_perfil(arg: str | None):
    """Orden: --perfil (ruta o nombre) · $INGENIERO_DATOS · ../ingeniero-agentico-datos · plantilla."""
    carpetas = []
    if os.environ.get("INGENIERO_DATOS"):
        carpetas.append(Path(os.environ["INGENIERO_DATOS"]))
    carpetas.append(REPO.parent / "ingeniero-agentico-datos")
    if arg:
        p = Path(arg)
        if p.is_file():
            return p, f"perfil {p.name} (--perfil)"
        for d in carpetas:
            q = d / "perfiles" / f"{arg}.yaml"
            if q.is_file():
                return q, f"perfil {q.name}"
        raise ErrorEntrada(f"No encontré el perfil «{arg}» (ni como ruta ni en <datos>/perfiles/).")
    for d in carpetas:
        dp = d / "perfiles"
        if not dp.is_dir():
            continue
        perfiles = sorted(x for x in dp.glob("*.yaml") if not x.name.startswith("_"))
        if len(perfiles) == 1:
            return perfiles[0], f"perfil {perfiles[0].name}"
        if len(perfiles) > 1:
            nombres = ", ".join(x.stem for x in perfiles)
            raise ErrorEntrada(f"Hay varios perfiles ({nombres}): elige uno con --perfil <nombre>.")
    plantilla = REPO / "perfiles" / "_plantilla.yaml"
    if plantilla.is_file():
        return plantilla, "plantilla (sin perfil de empresa)"
    return None, "sin perfil"


class ErrorEntrada(Exception):
    pass


# ---------------------------------------------------------------------------
# Parámetros: perfil + argumentos + entrada, cada uno con su origen
# ---------------------------------------------------------------------------


def resolver_parametros(perfil: dict, origen_perfil: str, entrada: dict, tasa=None, impuesto=None):
    emp = perfil.get("empresa") or {}
    fin = perfil.get("finanzas") or {}
    p = {}

    def poner(clave, valor, origen, supuesto):
        p[clave] = {"valor": valor, "origen": origen, "supuesto": supuesto}

    # tasa de descuento
    if tasa is not None:
        extra = " (reemplaza la del perfil)" if fin.get("tasa_descuento") is not None else " (el perfil no la trae)"
        poner("tasa", float(tasa), "argumento --tasa" + extra, True)
    elif fin.get("tasa_descuento") is not None:
        poner("tasa", float(fin["tasa_descuento"]), f"{origen_perfil} · finanzas.tasa_descuento", False)
    else:
        raise ErrorEntrada("El perfil no trae finanzas.tasa_descuento: pasa --tasa (la salida la marca SUPUESTO).")

    # impuesto a la renta
    if impuesto is not None:
        extra = " (reemplaza la del perfil)" if emp.get("impuesto_renta") is not None else " (el perfil no la trae)"
        poner("impuesto", float(impuesto), "argumento --impuesto" + extra, True)
    elif emp.get("impuesto_renta") is not None:
        poner("impuesto", float(emp["impuesto_renta"]), f"{origen_perfil} · empresa.impuesto_renta", False)
    else:
        raise ErrorEntrada("El perfil no trae empresa.impuesto_renta: pasa --impuesto (la salida la marca SUPUESTO).")

    # horizonte
    if entrada.get("horizonte_anios") is not None:
        poner("horizonte", int(entrada["horizonte_anios"]), "entrada · horizonte_anios", False)
    elif fin.get("horizonte_anios") is not None:
        poner("horizonte", int(fin["horizonte_anios"]), f"{origen_perfil} · finanzas.horizonte_anios", False)
    else:
        raise ErrorEntrada("Falta el horizonte: ni la entrada ni el perfil traen horizonte_anios.")

    dep = fin.get("depreciacion") or "lineal"
    poner("depreciacion", dep, f"{origen_perfil} · finanzas.depreciacion" if fin.get("depreciacion") else "por defecto (lineal)", not fin.get("depreciacion"))
    poner("moneda", emp.get("moneda") or "CLP", f"{origen_perfil} · empresa.moneda", not emp.get("moneda"))
    dec = emp.get("decimales_moneda")
    poner("decimales", int(dec) if dec is not None else 0, f"{origen_perfil} · empresa.decimales_moneda", dec is None)
    perd = entrada.get("perdidas_tributarias") or "arrastra"
    if perd not in ("arrastra", "compensa"):
        raise ErrorEntrada("perdidas_tributarias debe ser «arrastra» o «compensa».")
    poner("perdidas", perd, "entrada · perdidas_tributarias" if entrada.get("perdidas_tributarias") else "por defecto (arrastra)", not entrada.get("perdidas_tributarias"))
    metodo = entrada.get("metodo_desecho") or "comercial"
    if metodo not in ("comercial", "contable"):
        raise ErrorEntrada("metodo_desecho debe ser «comercial» o «contable».")
    poner("metodo_desecho", metodo, "entrada · metodo_desecho" if entrada.get("metodo_desecho") else "por defecto (comercial)", False)
    return p


def construir_modelo(entrada: dict, params: dict) -> dict:
    """Modelo numérico normalizado (lo que la sensibilidad modifica)."""
    m = {
        "N": params["horizonte"]["valor"],
        "tasa": params["tasa"]["valor"],
        "impuesto": params["impuesto"]["valor"],
        "depreciacion": params["depreciacion"]["valor"],
        "perdidas": params["perdidas"]["valor"],
        "metodo_desecho": params["metodo_desecho"]["valor"],
        "ingresos": copy.deepcopy(entrada.get("ingresos") or []),
        "costos_variables": copy.deepcopy(entrada.get("costos_variables") or []),
        "costos_fijos": copy.deepcopy(entrada.get("costos_fijos") or []),
        "inversiones": copy.deepcopy(entrada.get("inversiones") or []),
        "capital_trabajo": copy.deepcopy(entrada.get("capital_trabajo") or {}),
    }
    nombres = [i.get("nombre") for i in m["ingresos"]]
    if len(set(nombres)) != len(nombres):
        raise ErrorEntrada("Dos ingresos tienen el mismo nombre.")
    for ing in m["ingresos"]:
        if "precio" not in ing:
            raise ErrorEntrada(f"Ingreso «{ing.get('nombre')}» sin precio.")
        if "cantidad" not in ing and not ("capacidad" in ing and "ocupacion" in ing):
            raise ErrorEntrada(f"Ingreso «{ing.get('nombre')}»: falta cantidad, o capacidad + ocupacion.")
    for cv in m["costos_variables"]:
        if cv.get("por_unidad_de") not in nombres:
            raise ErrorEntrada(f"Costo variable «{cv.get('nombre')}»: por_unidad_de debe ser el nombre de un ingreso.")
    for inv in m["inversiones"]:
        if int(inv.get("anio", 0)) >= m["N"]:
            raise ErrorEntrada(f"Inversión «{inv.get('nombre')}» en el año {inv.get('anio')}: debe ser anterior al horizonte.")
        if inv.get("depreciable", True) and not inv.get("vida_util"):
            raise ErrorEntrada(f"Inversión «{inv.get('nombre')}» depreciable sin vida_util.")
    return m


# ---------------------------------------------------------------------------
# Flujo de caja del proyecto puro
# ---------------------------------------------------------------------------


def _crec(item, campo, t):
    g = float(item.get(campo, 0) or 0)
    return (1 + g) ** (t - 1) if t >= 1 else 0.0


def cantidad_de(ing):
    if "cantidad" in ing:
        return float(ing["cantidad"])
    return float(ing["capacidad"]) * float(ing["ocupacion"])


def flujo_de_caja(m: dict) -> dict:
    N = m["N"]
    T = m["impuesto"]
    anios = list(range(N + 1))
    cero = lambda: [0.0] * (N + 1)
    filas = []  # (clave, etiqueta, valores, origen)

    # ingresos
    tot_ing = cero()
    cant = {}
    for k, ing in enumerate(m["ingresos"]):
        q = cantidad_de(ing)
        v = cero()
        cq = cero()
        for t in anios[1:]:
            cq[t] = q * _crec(ing, "crecimiento_cantidad", t)
            v[t] = float(ing["precio"]) * _crec(ing, "crecimiento_precio", t) * cq[t]
            tot_ing[t] += v[t]
        cant[ing["nombre"]] = cq
        forma = "capacidad × ocupación" if "cantidad" not in ing else "cantidad"
        filas.append(("ingreso", f"Ingreso · {ing['nombre']}", v, f"entrada · ingresos[{k}]: precio × {forma}"))
    # costos variables
    tot_cv = cero()
    for k, cv in enumerate(m["costos_variables"]):
        v = cero()
        for t in anios[1:]:
            v[t] = -float(cv["monto_por_unidad"]) * _crec(cv, "crecimiento", t) * cant[cv["por_unidad_de"]][t]
            tot_cv[t] += v[t]
        filas.append(("costo_variable", f"Costo variable · {cv['nombre']}", v, f"entrada · costos_variables[{k}]: monto_por_unidad × cantidad de «{cv['por_unidad_de']}»"))
    # costos fijos
    tot_cf = cero()
    for k, cf in enumerate(m["costos_fijos"]):
        v = cero()
        for t in anios[1:]:
            v[t] = -float(cf["monto_anual"]) * _crec(cf, "crecimiento", t)
            tot_cf[t] += v[t]
        filas.append(("costo_fijo", f"Costo fijo · {cf['nombre']}", v, f"entrada · costos_fijos[{k}]: monto_anual"))
    # depreciación, venta y valor libro
    tot_dep = cero()
    venta = cero()
    libro = cero()
    inversion = cero()
    detalle_activos = []
    for k, inv in enumerate(m["inversiones"]):
        a = int(inv.get("anio", 0))
        monto = float(inv["monto"])
        inversion[a] -= monto
        dep_acum = 0.0
        if inv.get("depreciable", True):
            vida = inv["vida_util"]
            if m["depreciacion"] == "acelerada" and inv.get("vida_util_acelerada"):
                vida = inv["vida_util_acelerada"]
            cuota = monto / float(vida)
            for t in range(a + 1, min(a + int(math.ceil(vida)), N) + 1):
                d = min(cuota, monto - dep_acum)
                tot_dep[t] -= d
                dep_acum += d
        vl = monto - dep_acum
        vd = float(inv.get("valor_desecho", 0) or 0) if m["metodo_desecho"] == "comercial" else vl
        venta[N] += vd
        libro[N] -= vl
        detalle_activos.append({"nombre": inv["nombre"], "monto": monto, "anio": a, "depreciacion_acumulada": dep_acum,
                                "valor_libro_final": vl, "valor_desecho": vd})
    uai = [tot_ing[t] + tot_cv[t] + tot_cf[t] + tot_dep[t] + venta[t] + libro[t] for t in anios]
    imp = cero()
    perdida = 0.0
    for t in anios[1:]:
        if m["perdidas"] == "compensa":
            imp[t] = -T * uai[t]
        else:
            base = uai[t] + perdida
            if base > 0:
                imp[t] = -T * base
                perdida = 0.0
            else:
                perdida = base
    un = [uai[t] + imp[t] for t in anios]
    ajuste_dep = [-x for x in tot_dep]
    ajuste_libro = [-x for x in libro]
    kt = m["capital_trabajo"] or {}
    inv_kt = cero()
    rec_kt = cero()
    if kt.get("monto"):
        a = int(kt.get("anio", 0))
        inv_kt[a] -= float(kt["monto"])
        if kt.get("recupera", True):
            rec_kt[N] += float(kt["monto"])
    flujo = [un[t] + ajuste_dep[t] + ajuste_libro[t] + inversion[t] + inv_kt[t] + rec_kt[t] for t in anios]

    cuerpo = list(filas)
    cuerpo += [
        ("total", "= Ingresos totales", tot_ing, "suma de los ingresos"),
        ("total", "− Costos variables totales", tot_cv, "suma de los costos variables"),
        ("total", "− Costos fijos totales", tot_cf, "suma de los costos fijos"),
        ("dep", "− Depreciación", tot_dep, "inversiones: monto ÷ vida útil (perfil · depreciación)"),
        ("venta", "+ Venta de activos (valor de desecho)", venta, f"inversiones: valor_desecho (método {m['metodo_desecho']})"),
        ("libro", "− Valor libro de los activos vendidos", libro, "monto − depreciación acumulada"),
        ("uai", "= Utilidad antes de impuesto", uai, "suma de lo anterior"),
        ("imp", "− Impuesto", imp, f"tasa {T:.2%} × utilidad (pérdidas: {m['perdidas']})"),
        ("un", "= Utilidad neta", un, "utilidad antes de impuesto − impuesto"),
        ("ajuste", "+ Depreciación (no es salida de caja)", ajuste_dep, "se devuelve lo restado arriba"),
        ("ajuste", "+ Valor libro (no es salida de caja)", ajuste_libro, "se devuelve lo restado arriba"),
        ("inv", "− Inversión en activos", inversion, "entrada · inversiones (monto, año)"),
        ("inv", "− Inversión en capital de trabajo", inv_kt, "entrada · capital_trabajo.monto"),
        ("inv", "+ Recuperación del capital de trabajo", rec_kt, "entrada · capital_trabajo (se recupera al final)"),
        ("flujo", "= FLUJO DE CAJA DEL PROYECTO PURO", flujo, "utilidad neta + ajustes − inversiones + recuperaciones"),
    ]
    return {"anios": anios, "filas": cuerpo, "flujo": flujo, "activos": detalle_activos,
            "inversion_total": [inversion[t] + inv_kt[t] for t in anios], "perdida_final": perdida}


# ---------------------------------------------------------------------------
# Indicadores
# ---------------------------------------------------------------------------


def van(flujo, r):
    return sum(f / (1 + r) ** t for t, f in enumerate(flujo))


def cambios_de_signo(flujo):
    s = [1 if f > 0 else -1 for f in flujo if abs(f) > 1e-9]
    return sum(1 for a, b in zip(s, s[1:]) if a != b)


def raices_tir(flujo):
    malla = [-0.95 + 0.005 * i for i in range(int(1.95 / 0.005) + 1)] + [1.0 + 0.05 * i for i in range(1, 181)]
    raices = []
    previo_r, previo_v = None, None
    for r in malla:
        v = van(flujo, r)
        if v == 0:
            raices.append(r)
        elif previo_v is not None and previo_v * v < 0:
            a, b, va = previo_r, r, previo_v
            for _ in range(200):
                c = (a + b) / 2
                vc = van(flujo, c)
                if va * vc <= 0:
                    b = c
                else:
                    a, va = c, vc
            raices.append((a + b) / 2)
        previo_r, previo_v = r, v
    return raices


def recuperacion(flujo, r=None):
    acum = 0.0
    for t, f in enumerate(flujo):
        fd = f if r is None else f / (1 + r) ** t
        nuevo = acum + fd
        if t > 0 and acum < 0 <= nuevo:
            return (t - 1) + (-acum / fd)
        acum = nuevo
    return None


def indicadores(fc: dict, m: dict) -> dict:
    flujo = fc["flujo"]
    r = m["tasa"]
    v = van(flujo, r)
    raices = raices_tir(flujo)
    cs = cambios_de_signo(flujo)
    avisos = []
    tir = None
    if cs > 1:
        avisos.append(f"El flujo cambia {cs} veces de signo: puede haber varias TIR (o ninguna). Decide con el VAN.")
    if len(raices) == 1:
        tir = raices[0]
    elif len(raices) > 1:
        avisos.append("Hay varias TIR: " + ", ".join(f"{x:.4%}" for x in raices) + ". La TIR no sirve para decidir aquí.")
    else:
        avisos.append("No hay TIR en el rango −95 % a 1.000 %.")
    inv_vp = -sum(x / (1 + r) ** t for t, x in enumerate(fc["inversion_total"]))
    ivan = v / inv_vp if inv_vp > 0 else None
    pri = recuperacion(flujo)
    prid = recuperacion(flujo, r)
    if pri is None:
        avisos.append("La inversión no se recupera dentro del horizonte (ni sin descontar).")
    if fc["perdida_final"] < 0:
        avisos.append("Queda pérdida tributaria sin usar al final del horizonte (no se valoriza).")
    return {"van": v, "tir": tir, "tir_raices": raices, "cambios_de_signo": cs, "pri": pri, "pri_descontado": prid,
            "ivan": ivan, "inversion_vp": inv_vp, "tasa": r, "avisos": avisos,
            "decision": "CONVIENE (VAN ≥ 0)" if v >= 0 else "NO CONVIENE (VAN < 0)"}


# ---------------------------------------------------------------------------
# Sensibilidad
# ---------------------------------------------------------------------------


def variables_sensibles(m: dict):
    """(clave, etiqueta, ruta). ruta = (sección, índice, campo) o (campo,)."""
    vs = []
    for k, ing in enumerate(m["ingresos"]):
        vs.append((f"ingresos[{k}].precio", f"{ing['nombre']} · precio", ("ingresos", k, "precio")))
        if "cantidad" in ing:
            vs.append((f"ingresos[{k}].cantidad", f"{ing['nombre']} · cantidad", ("ingresos", k, "cantidad")))
        else:
            vs.append((f"ingresos[{k}].ocupacion", f"{ing['nombre']} · ocupación", ("ingresos", k, "ocupacion")))
    for k, cv in enumerate(m["costos_variables"]):
        vs.append((f"costos_variables[{k}].monto_por_unidad", f"{cv['nombre']} · por unidad", ("costos_variables", k, "monto_por_unidad")))
    for k, cf in enumerate(m["costos_fijos"]):
        vs.append((f"costos_fijos[{k}].monto_anual", f"{cf['nombre']} · anual", ("costos_fijos", k, "monto_anual")))
    for k, inv in enumerate(m["inversiones"]):
        vs.append((f"inversiones[{k}].monto", f"{inv['nombre']} · inversión", ("inversiones", k, "monto")))
        if m["metodo_desecho"] == "comercial":
            vs.append((f"inversiones[{k}].valor_desecho", f"{inv['nombre']} · valor de desecho", ("inversiones", k, "valor_desecho")))
    if (m["capital_trabajo"] or {}).get("monto"):
        vs.append(("capital_trabajo.monto", "Capital de trabajo", ("capital_trabajo", None, "monto")))
    vs.append(("tasa", "Tasa de descuento", ("tasa",)))
    vs.append(("impuesto", "Tasa de impuesto", ("impuesto",)))
    return [v for v in vs if (_leer(m, v[2]) or 0) != 0]


def _leer(m, ruta):
    if len(ruta) == 1:
        return m[ruta[0]]
    sec, k, campo = ruta
    obj = m[sec] if k is None else m[sec][k]
    return obj.get(campo)


def _con(m, ruta, valor):
    m2 = copy.deepcopy(m)
    if len(ruta) == 1:
        m2[ruta[0]] = valor
    else:
        sec, k, campo = ruta
        (m2[sec] if k is None else m2[sec][k])[campo] = valor
    return m2


def van_con(m, ruta, valor):
    m2 = _con(m, ruta, valor)
    return van(flujo_de_caja(m2)["flujo"], m2["tasa"])


def valor_critico(m, ruta):
    """Valor de la variable que deja el VAN en cero, buscando el cruce más cercano al valor base."""
    base = float(_leer(m, ruta))
    factores = [i * 0.05 for i in range(0, 61)] + [3.5, 4, 5, 6, 8, 10, 15, 20, 30, 50, 100]
    puntos = [(f, van_con(m, ruta, base * f)) for f in factores]
    tramos = [(a, b) for a, b in zip(puntos, puntos[1:]) if a[1] == 0 or a[1] * b[1] < 0]
    if not tramos:
        return None
    (fa, va), (fb, vb) = min(tramos, key=lambda tr: min(abs(tr[0][0] - 1), abs(tr[1][0] - 1)))
    if va == 0:
        return base * fa
    for _ in range(200):
        fc = (fa + fb) / 2
        vc = van_con(m, ruta, base * fc)
        if va * vc <= 0:
            fb = fc
        else:
            fa, va = fc, vc
    return base * (fa + fb) / 2


def sensibilidad(m: dict, delta: float):
    filas = []
    van_base = van(flujo_de_caja(m)["flujo"], m["tasa"])
    for clave, etiqueta, ruta in variables_sensibles(m):
        base = float(_leer(m, ruta))
        vmen = van_con(m, ruta, base * (1 - delta))
        vmas = van_con(m, ruta, base * (1 + delta))
        crit = valor_critico(m, ruta)
        nota = ""
        if crit is None:
            lado = "positivo" if van_base >= 0 else "negativo"
            nota = f"sin valor crítico: el VAN sigue {lado} desde 0 hasta 100 veces el valor base"
        elif ruta[-1] == "ocupacion" and crit > 1:
            nota = "fuera de rango: haría falta más de 100 % de ocupación"
        elif crit < 0:
            nota = "valor crítico negativo"
        filas.append({"clave": clave, "variable": etiqueta, "base": base, "van_menos": vmen, "van_mas": vmas,
                      "rango": abs(vmas - vmen), "critico": crit,
                      "cambio_critico": (crit / base - 1) if crit is not None and base else None, "nota": nota})
    filas.sort(key=lambda f: -f["rango"])
    return {"delta": delta, "van_base": van_base, "filas": filas}


def barrido(m: dict, spec: str):
    clave, _, valores = spec.partition("=")
    clave = clave.strip()
    mapa = {c: r for c, _, r in variables_sensibles(m)}
    mapa.update({"tasa": ("tasa",), "impuesto": ("impuesto",)})
    if clave not in mapa:
        raise ErrorEntrada(f"--barrer: no conozco la variable «{clave}». Usa una clave de la hoja SENSIBILIDAD.")
    ruta = mapa[clave]
    salida = []
    for txt in valores.split(","):
        x = float(txt)
        m2 = _con(m, ruta, x)
        fc = flujo_de_caja(m2)
        ind = indicadores(fc, m2)
        salida.append({"valor": x, "van": ind["van"], "tir": ind["tir"]})
    return {"clave": clave, "puntos": salida}


# ---------------------------------------------------------------------------
# Supuestos (trazabilidad de cada entrada)
# ---------------------------------------------------------------------------


def _es_supuesto(item, campo):
    s = item.get("supuesto", False)
    if isinstance(s, list):
        return campo in s
    return bool(s)


def tabla_supuestos(entrada, params):
    t = []
    for clave, etiqueta, unidad in (("tasa", "Tasa de descuento", "anual"), ("impuesto", "Tasa de impuesto a la renta", "fracción"),
                                    ("horizonte", "Horizonte de evaluación", "años"), ("depreciacion", "Depreciación", ""),
                                    ("perdidas", "Pérdidas tributarias", ""), ("metodo_desecho", "Valor de desecho", "método"),
                                    ("moneda", "Moneda", ""), ("decimales", "Decimales de la moneda", "")):
        p = params[clave]
        t.append({"parametro": etiqueta, "valor": p["valor"], "unidad": unidad, "origen": p["origen"],
                  "marca": "SUPUESTO" if p["supuesto"] else "dato"})
    campos = {
        "ingresos": ("precio", "cantidad", "capacidad", "ocupacion", "crecimiento_precio", "crecimiento_cantidad"),
        "costos_variables": ("monto_por_unidad", "crecimiento"),
        "costos_fijos": ("monto_anual", "crecimiento"),
        "inversiones": ("monto", "anio", "vida_util", "vida_util_acelerada", "valor_desecho"),
    }
    for sec, cs in campos.items():
        for k, item in enumerate(entrada.get(sec) or []):
            for c in cs:
                if c in item:
                    t.append({"parametro": f"{item.get('nombre')} · {c}", "valor": item[c], "unidad": item.get("unidad", "") if c in ("cantidad", "capacidad", "precio") else "",
                              "origen": f"entrada · {sec}[{k}].{c}" + (f" · {item['fuente']}" if item.get("fuente") else ""),
                              "marca": "SUPUESTO" if _es_supuesto(item, c) else "dato"})
    kt = entrada.get("capital_trabajo") or {}
    if kt.get("monto") is not None:
        t.append({"parametro": "Capital de trabajo · monto", "valor": kt["monto"], "unidad": "",
                  "origen": "entrada · capital_trabajo.monto" + (f" · {kt['fuente']}" if kt.get("fuente") else ""),
                  "marca": "SUPUESTO" if _es_supuesto(kt, "monto") else "dato"})
    return t


# ---------------------------------------------------------------------------
# Evaluación completa
# ---------------------------------------------------------------------------


def evaluar(entrada, perfil, origen_perfil, tasa=None, impuesto=None, delta=0.10, barrer=None):
    params = resolver_parametros(perfil, origen_perfil, entrada, tasa, impuesto)
    m = construir_modelo(entrada, params)
    fc = flujo_de_caja(m)
    ind = indicadores(fc, m)
    sens = sensibilidad(m, delta)
    barr = [barrido(m, b) for b in (barrer or [])]
    return {"proyecto": entrada.get("proyecto", ""), "empresa": entrada.get("empresa") or (perfil.get("empresa") or {}).get("nombre", ""),
            "perfil": origen_perfil, "parametros": params, "supuestos": tabla_supuestos(entrada, params),
            "flujo": fc, "indicadores": ind, "sensibilidad": sens, "barridos": barr}


# ---------------------------------------------------------------------------
# Salida: Excel y JSON
# ---------------------------------------------------------------------------


def escribir_excel(res: dict, ruta: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    dec = res["parametros"]["decimales"]["valor"]
    fmt_m = "#,##0" if dec == 0 else "#,##0." + "0" * dec
    neg = Font(bold=True)
    gris = PatternFill("solid", fgColor="E7E6E6")
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()

    # SUPUESTOS
    ws = wb.active
    ws.title = "SUPUESTOS"
    ws.append([f"Evaluación económica · {res['proyecto']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Empresa: {res['empresa']} · {res['perfil']} · proyecto puro (sin financiamiento)"])
    ws.append([])
    ws.append(["Parámetro", "Valor", "Unidad", "Origen", "Marca"])
    for c in ws[4]:
        c.font = neg
        c.fill = gris
    for s in res["supuestos"]:
        ws.append([s["parametro"], s["valor"], s["unidad"], s["origen"], s["marca"]])
        if s["marca"] == "SUPUESTO":
            for c in ws[ws.max_row]:
                c.fill = amarillo
    for col, w in zip("ABCDE", (38, 16, 10, 90, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    # FLUJO
    wf = wb.create_sheet("FLUJO")
    anios = res["flujo"]["anios"]
    wf.append(["Concepto"] + [f"Año {t}" for t in anios] + ["Origen"])
    for c in wf[1]:
        c.font = neg
        c.fill = gris
    for clave, etiqueta, valores, origen in res["flujo"]["filas"]:
        wf.append([etiqueta] + [round(v, 6) for v in valores] + [origen])
        fila = wf[wf.max_row]
        for c in fila[1:len(anios) + 1]:
            c.number_format = fmt_m
        if clave in ("uai", "un", "flujo", "total"):
            for c in fila:
                c.font = neg
        if clave == "flujo":
            for c in fila:
                c.fill = gris
    wf.column_dimensions["A"].width = 46
    for i in range(len(anios)):
        wf.column_dimensions[chr(66 + i)].width = 15
    wf.column_dimensions[chr(66 + len(anios))].width = 70
    wf.freeze_panes = "B2"

    # INDICADORES
    wi = wb.create_sheet("INDICADORES")
    ind = res["indicadores"]
    wi.append(["Indicador", "Valor", "Cómo se calcula"])
    for c in wi[1]:
        c.font = neg
        c.fill = gris
    filas = [
        ("Decisión", ind["decision"], "VAN ≥ 0 conviene; VAN < 0 no conviene", None),
        ("Tasa de descuento", ind["tasa"], res["parametros"]["tasa"]["origen"] + (" · SUPUESTO" if res["parametros"]["tasa"]["supuesto"] else ""), "0.00%"),
        ("VAN", ind["van"], "Σ flujo_t ÷ (1 + tasa)^t, t = 0 … N (hoja FLUJO, última fila)", fmt_m),
        ("TIR", ind["tir"] if ind["tir"] is not None else "no definida", "tasa que deja el VAN en cero (bisección)", "0.00%"),
        ("Cambios de signo del flujo", ind["cambios_de_signo"], "más de 1: puede haber varias TIR", None),
        ("Período de recuperación (años)", ind["pri"] if ind["pri"] is not None else "no se recupera", "flujo acumulado sin descontar; interpolación dentro del año", "0.00"),
        ("Período de recuperación descontado (años)", ind["pri_descontado"] if ind["pri_descontado"] is not None else "no se recupera", "flujo acumulado descontado a la tasa", "0.00"),
        ("Inversión en valor presente", ind["inversion_vp"], "activos + capital de trabajo, descontados", fmt_m),
        ("IVAN", ind["ivan"] if ind["ivan"] is not None else "no aplica", "VAN ÷ inversión en valor presente", "0.000"),
    ]
    for nombre, valor, como, f in filas:
        wi.append([nombre, valor, como])
        if f and isinstance(valor, (int, float)):
            wi.cell(row=wi.max_row, column=2).number_format = f
    wi.append([])
    for a in ind["avisos"]:
        wi.append(["AVISO", a])
        wi.cell(row=wi.max_row, column=1).font = Font(bold=True, color="C00000")
    wi.column_dimensions["A"].width = 42
    wi.column_dimensions["B"].width = 22
    wi.column_dimensions["C"].width = 70

    # SENSIBILIDAD
    wsn = wb.create_sheet("SENSIBILIDAD")
    sens = res["sensibilidad"]
    d = sens["delta"]
    wsn.append([f"Una variable a la vez, ±{d:.0%}; las demás en su valor base. Orden: de la que más mueve el VAN a la que menos."])
    wsn.append([f"VAN base: {sens['van_base']:,.0f}"])
    wsn.append([])
    wsn.append(["Clave", "Variable", "Valor base", f"VAN con −{d:.0%}", f"VAN con +{d:.0%}", "Rango del VAN",
                "Valor crítico (VAN = 0)", "Cambio hasta el crítico", "Nota"])
    for c in wsn[4]:
        c.font = neg
        c.fill = gris
    for f in sens["filas"]:
        wsn.append([f["clave"], f["variable"], f["base"], f["van_menos"], f["van_mas"], f["rango"],
                    f["critico"] if f["critico"] is not None else "sin cruce", f["cambio_critico"], f["nota"]])
        r = wsn.max_row
        for col in (4, 5, 6):
            wsn.cell(row=r, column=col).number_format = fmt_m
        es_frac = f["clave"].endswith("ocupacion") or f["clave"] in ("tasa", "impuesto")
        for col in (3, 7):
            wsn.cell(row=r, column=col).number_format = "0.00%" if es_frac else fmt_m
        wsn.cell(row=r, column=8).number_format = "0.0%"
    for b in res["barridos"]:
        wsn.append([])
        wsn.append([f"Barrido de {b['clave']}", "", "Valor", "VAN", "TIR"])
        for c in wsn[wsn.max_row]:
            c.font = neg
        for p in b["puntos"]:
            wsn.append(["", "", p["valor"], p["van"], p["tir"] if p["tir"] is not None else "no definida"])
            wsn.cell(row=wsn.max_row, column=4).number_format = fmt_m
            wsn.cell(row=wsn.max_row, column=5).number_format = "0.00%"
    for col, w in zip("ABCDEFGHI", (34, 34, 16, 16, 16, 16, 20, 18, 50)):
        wsn.column_dimensions[col].width = w
    for row in wsn.iter_rows(min_row=1, max_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=False)
    wb.save(ruta)


def a_json(res: dict) -> dict:
    fc = res["flujo"]
    return {
        "proyecto": res["proyecto"], "empresa": res["empresa"], "perfil": res["perfil"],
        "parametros": res["parametros"], "supuestos": res["supuestos"],
        "flujo": {"anios": fc["anios"], "filas": [{"concepto": e, "valores": v, "origen": o} for _, e, v, o in fc["filas"]],
                  "flujo_de_caja": fc["flujo"], "activos": fc["activos"]},
        "indicadores": res["indicadores"], "sensibilidad": res["sensibilidad"], "barridos": res["barridos"],
    }


def resumen_texto(res: dict) -> str:
    ind = res["indicadores"]
    dec = res["parametros"]["decimales"]["valor"]
    mon = res["parametros"]["moneda"]["valor"]
    fm = lambda x: f"{x:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    pc = lambda x: "no definida" if x is None else f"{x * 100:.2f} %".replace(".", ",")
    an = lambda x: "no se recupera" if x is None else f"{x:.2f} años".replace(".", ",")
    lin = [f"Proyecto: {res['proyecto']} · {res['perfil']}",
           f"Tasa: {pc(ind['tasa'])}" + (" (SUPUESTO)" if res["parametros"]["tasa"]["supuesto"] else ""),
           f"Flujo ({mon}): " + " | ".join(fm(x) for x in res["flujo"]["flujo"]),
           f"VAN: {fm(ind['van'])} {mon} → {ind['decision']}",
           f"TIR: {pc(ind['tir'])} · PRI: {an(ind['pri'])} · PRI descontado: {an(ind['pri_descontado'])} · IVAN: "
           + ("no aplica" if ind["ivan"] is None else f"{ind['ivan']:.3f}".replace(".", ","))]
    for a in ind["avisos"]:
        lin.append("AVISO: " + a)
    lin.append("Sensibilidad (de la que más pesa a la que menos):")
    for f in res["sensibilidad"]["filas"][:6]:
        frac = f["clave"].endswith("ocupacion") or f["clave"] in ("tasa", "impuesto")
        crit = "sin cruce" if f["critico"] is None else (pc(f["critico"]) if frac else fm(f["critico"]))
        cambio = "" if f["cambio_critico"] is None else f" ({f['cambio_critico'] * 100:+.1f} %)".replace(".", ",")
        lin.append(f"  {f['variable']}: crítico {crit}{cambio}" + (f" · {f['nota']}" if f["nota"] else ""))
    return "\n".join(lin)


# ---------------------------------------------------------------------------
# Prueba
# ---------------------------------------------------------------------------


def probar() -> int:
    pr = SKILL / "pruebas"
    entrada = json.loads((pr / "caso-ficticio.json").read_text(encoding="utf-8"))
    perfil = leer_yaml_simple(pr / "perfil-ficticio.yaml")
    esperado = json.loads((pr / "esperado.json").read_text(encoding="utf-8"))
    res = evaluar(entrada, perfil, "perfil perfil-ficticio.yaml")
    ind = res["indicadores"]
    fallas = []

    def cerca(nombre, obtenido, esperado_, tol):
        if obtenido is None or abs(obtenido - esperado_) > tol:
            fallas.append(f"{nombre}: obtuve {obtenido}, esperaba {esperado_} (±{tol})")

    tol_m = esperado["tolerancia_moneda"]
    for t, (a, b) in enumerate(zip(res["flujo"]["flujo"], esperado["flujo"])):
        cerca(f"flujo año {t}", a, b, tol_m)
    cerca("VAN", ind["van"], esperado["van"], tol_m)
    cerca("TIR", ind["tir"], esperado["tir"], esperado["tolerancia_tasa"])
    cerca("PRI", ind["pri"], esperado["pri"], 1e-4)
    cerca("PRI descontado", ind["pri_descontado"], esperado["pri_descontado"], 1e-4)
    cerca("IVAN", ind["ivan"], esperado["ivan"], 1e-4)
    crit = {f["clave"]: f["critico"] for f in res["sensibilidad"]["filas"]}
    for clave, val in esperado["criticos"].items():
        cerca(f"crítico {clave}", crit.get(clave), val, esperado["tolerancia_critico"].get(clave, 1e-6))
    # el valor crítico deja el VAN en cero (todas las variables)
    m = construir_modelo(entrada, res["parametros"])
    rutas = {c: r for c, _, r in variables_sensibles(m)}
    for f in res["sensibilidad"]["filas"]:
        if f["critico"] is not None:
            v0 = van_con(m, rutas[f["clave"]], f["critico"])
            if abs(v0) > 1.0:
                fallas.append(f"VAN en el crítico de {f['clave']} = {v0:.4f}, no cero")
    # el crítico de la tasa es la TIR
    cerca("crítico de la tasa = TIR", crit.get("tasa"), ind["tir"], 1e-7)
    # flujo con tres cambios de signo: TIR 10 %, 20 % y 30 %
    raices = raices_tir([-1000, 3600, -4310, 1716])
    if len(raices) != 3 or any(abs(a - b) > 1e-7 for a, b in zip(raices, (0.10, 0.20, 0.30))):
        fallas.append(f"varias TIR: obtuve {raices}, esperaba 10 %, 20 % y 30 %")
    if cambios_de_signo([-1000, 3600, -4310, 1716]) != 3:
        fallas.append("cambios de signo mal contados")
    # perfil sin tasa: tiene que exigir --tasa y marcar SUPUESTO
    sin_tasa = copy.deepcopy(perfil)
    sin_tasa["finanzas"]["tasa_descuento"] = None
    try:
        evaluar(entrada, sin_tasa, "perfil sin tasa")
        fallas.append("con tasa null en el perfil no exigió --tasa")
    except ErrorEntrada:
        pass
    r2 = evaluar(entrada, sin_tasa, "perfil sin tasa", tasa=0.12)
    if not r2["parametros"]["tasa"]["supuesto"]:
        fallas.append("la tasa por argumento no quedó marcada SUPUESTO")
    if abs(r2["indicadores"]["van"] - ind["van"]) > 1e-6:
        fallas.append("--tasa 0.12 no da el mismo VAN que el perfil con 0.12")
    # Excel
    try:
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            escribir_excel(res, Path(tmp) / "prueba.xlsx")
            from openpyxl import load_workbook
            hojas = load_workbook(Path(tmp) / "prueba.xlsx").sheetnames
            if hojas != ["SUPUESTOS", "FLUJO", "INDICADORES", "SENSIBILIDAD"]:
                fallas.append(f"hojas del Excel: {hojas}")
    except ImportError:
        fallas.append("falta openpyxl para el Excel")

    print(resumen_texto(res))
    print()
    if fallas:
        for f in fallas:
            print("FALLA:", f)
        print("FALLA")
        return 1
    print("OK")
    return 0


# ---------------------------------------------------------------------------


def main(argv=None):
    ap = argparse.ArgumentParser(description="Evaluación económica del proyecto puro: flujo de caja, VAN, TIR, sensibilidad.")
    ap.add_argument("entrada", nargs="?", help="JSON del caso (ver plantillas/entrada.json)")
    ap.add_argument("--perfil", help="nombre del perfil (en <datos>/perfiles/) o ruta a un .yaml")
    ap.add_argument("--tasa", type=float, help="tasa de descuento anual (fracción); queda marcada SUPUESTO")
    ap.add_argument("--impuesto", type=float, help="tasa de impuesto a la renta (fracción); queda marcada SUPUESTO")
    ap.add_argument("--delta", type=float, default=0.10, help="variación de la sensibilidad (0.10 = ±10 %%)")
    ap.add_argument("--barrer", action="append", help='barrido de una variable: "ingresos[0].ocupacion=0.3,0.5,0.7"')
    ap.add_argument("--salida", help="carpeta de salida (por defecto, la de la entrada)")
    ap.add_argument("--probar", action="store_true", help="corre la prueba del caso ficticio")
    a = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if a.probar:
        return probar()
    if not a.entrada:
        ap.error("falta la entrada (o usa --probar)")
    try:
        ruta_e = Path(a.entrada)
        entrada = json.loads(ruta_e.read_text(encoding="utf-8-sig"))
        ruta_p, origen = buscar_perfil(a.perfil)
        perfil = leer_yaml_simple(ruta_p) if ruta_p else {}
        res = evaluar(entrada, perfil, origen, a.tasa, a.impuesto, a.delta, a.barrer)
    except ErrorEntrada as e:
        print("ERROR:", e)
        return 2
    carpeta = Path(a.salida) if a.salida else ruta_e.parent
    carpeta.mkdir(parents=True, exist_ok=True)
    base = ruta_e.stem.replace("entrada", "salida") if "entrada" in ruta_e.stem else ruta_e.stem + "-salida"
    (carpeta / f"{base}.json").write_text(json.dumps(a_json(res), ensure_ascii=False, indent=2), encoding="utf-8")
    escribir_excel(res, carpeta / f"{base}.xlsx")
    print(resumen_texto(res))
    print(f"\nEscribí {base}.xlsx y {base}.json en la carpeta de salida.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
