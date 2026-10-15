#!/usr/bin/env python3
"""recorrer.py · la parte mecánica de la auditoría (etapa 4 de ingeniero-agentico).

Vuelve a correr el script de un método con la MISMA entrada, pero dejando la salida en una
carpeta temporal, y compara cifra por cifra lo que salió ahora contra la salida archivada.
No escribe nada fuera de la carpeta temporal y avisa si el script tocó las carpetas del caso.

    python recorrer.py correr --script <metodo.py> --original <salida.json> [--original <salida.xlsx>]
                              [--tol-abs 0.01] [--tol-rel 1e-9] [--ignorar <ruta>] -- <args con {SALIDA}>
    python recorrer.py manifiesto <contexto/ejecuciones/ | una-ejecucion.json>
    python recorrer.py valor <archivo.json> <ruta>            (ej. indicadores.van)
    python recorrer.py cruzar <a.json> <ruta_a> <b.json> <ruta_b> [--tol-abs 1]
    python recorrer.py --probar

{SALIDA} en los argumentos se reemplaza por la carpeta temporal. Cada --original se compara con
el archivo del mismo nombre en esa carpeta (o con otro nombre: --original ruta=nombre_nuevo).
Rutas de la forma  alternativas["Leasing 36 cuotas"].cae  o  filas[0].valor.

Salida: 0 = todo igual · 1 = hay diferencias o cambios en el caso · 2 = el script falló o falta algo.
Solo biblioteca estándar; si openpyxl está instalado, también compara los .xlsx celda por celda.
"""
from __future__ import annotations

import argparse
import fnmatch
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True
REPO = Path(__file__).resolve().parents[3]
MARCA = "{SALIDA}"


# ------------------------------------------------------------------ rutas dentro de un JSON
def partir_ruta(ruta: str) -> list:
    """'a.b[0]["c d"].e' -> ['a', 'b', 0, 'c d', 'e']"""
    pasos, i, n = [], 0, len(ruta)
    actual = ""
    while i < n:
        ch = ruta[i]
        if ch == ".":
            if actual:
                pasos.append(actual)
            actual = ""
            i += 1
        elif ch == "[":
            if actual:
                pasos.append(actual)
            actual = ""
            fin = ruta.index("]", i) if ruta[i + 1] not in "\"'" else ruta.index(ruta[i + 1] + "]", i + 2) + 1
            dentro = ruta[i + 1:fin]
            if dentro[:1] in "\"'":
                pasos.append(dentro[1:-1])
            else:
                pasos.append(int(dentro))
            i = fin + 1
        else:
            actual += ch
            i += 1
    if actual:
        pasos.append(actual)
    return pasos


def leer_valor(datos, ruta: str):
    v = datos
    for p in partir_ruta(ruta):
        if isinstance(p, int):
            if not isinstance(v, list) or p >= len(v):
                raise KeyError(f"{ruta}: no hay índice [{p}]")
            v = v[p]
        else:
            if not isinstance(v, dict) or p not in v:
                raise KeyError(f"{ruta}: no existe la clave «{p}»")
            v = v[p]
    return v


def cargar_json(ruta) -> object:
    return json.loads(Path(ruta).read_text(encoding="utf-8-sig"))


# ------------------------------------------------------------------ comparar
def numeros_iguales(a: float, b: float, tol_abs: float, tol_rel: float) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if math.isnan(a) and math.isnan(b):
        return True
    return abs(a - b) <= tol_abs or abs(a - b) <= tol_rel * max(abs(a), abs(b))


def es_numero(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def comparar_valores(a, b, ruta, difs, cuenta, tol_abs, tol_rel, ignorar, reemplazos):
    if any(fnmatch.fnmatchcase(ruta, patron) for patron in ignorar):
        return
    if isinstance(a, dict) and isinstance(b, dict):
        for k in a:
            sub = f"{ruta}.{k}" if ruta else str(k)
            if k not in b:
                difs.append((sub, a[k], "<no está>"))
            else:
                comparar_valores(a[k], b[k], sub, difs, cuenta, tol_abs, tol_rel, ignorar, reemplazos)
        for k in b:
            if k not in a:
                difs.append((f"{ruta}.{k}" if ruta else str(k), "<no está>", b[k]))
        return
    if isinstance(a, list) and isinstance(b, list):
        if len(a) != len(b):
            difs.append((f"{ruta} (largo)", len(a), len(b)))
        for i, (x, y) in enumerate(zip(a, b)):
            comparar_valores(x, y, f"{ruta}[{i}]", difs, cuenta, tol_abs, tol_rel, ignorar, reemplazos)
        return
    cuenta[0] += 1
    if es_numero(a) and es_numero(b):
        if not numeros_iguales(float(a), float(b), tol_abs, tol_rel):
            difs.append((ruta, a, b))
        return
    if isinstance(a, str) and isinstance(b, str):
        for viejo, nuevo in reemplazos:
            b = b.replace(viejo, nuevo)
        if a != b:
            difs.append((ruta, a, b))
        return
    if a != b:
        difs.append((ruta, a, b))


def comparar_json(orig: Path, nuevo: Path, tol_abs, tol_rel, ignorar, reemplazos):
    difs, cuenta = [], [0]
    comparar_valores(cargar_json(orig), cargar_json(nuevo), "", difs, cuenta, tol_abs, tol_rel, ignorar, reemplazos)
    return difs, cuenta[0]


def comparar_xlsx(orig: Path, nuevo: Path, tol_abs, tol_rel, reemplazos):
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        return None, 0
    wa = openpyxl.load_workbook(orig)
    wb = openpyxl.load_workbook(nuevo)
    difs, cuenta = [], 0
    if wa.sheetnames != wb.sheetnames:
        difs.append(("hojas", wa.sheetnames, wb.sheetnames))
    for nombre in wa.sheetnames:
        if nombre not in wb.sheetnames:
            continue
        ha, hb = wa[nombre], wb[nombre]
        filas = max(ha.max_row, hb.max_row)
        cols = max(ha.max_column, hb.max_column)
        for r in range(1, filas + 1):
            for c in range(1, cols + 1):
                ca, cb = ha.cell(row=r, column=c), hb.cell(row=r, column=c)
                va, vb = ca.value, cb.value
                if va is None and vb is None:
                    continue
                cuenta += 1
                donde = f"{nombre}!{ca.coordinate}"
                if es_numero(va) and es_numero(vb):
                    if not numeros_iguales(float(va), float(vb), tol_abs, tol_rel):
                        difs.append((donde, va, vb))
                else:
                    if isinstance(vb, str):
                        for viejo, nuevo_txt in reemplazos:
                            vb = vb.replace(viejo, nuevo_txt)
                    if va != vb:
                        difs.append((donde, va, vb))
                fa = ca.fill.fgColor.rgb if ca.fill and ca.fill.fill_type else None
                fb = cb.fill.fgColor.rgb if cb.fill and cb.fill.fill_type else None
                if fa != fb:
                    difs.append((donde + " (color)", fa, fb))
    return difs, cuenta


# ------------------------------------------------------------------ vigilar las carpetas del caso
def foto(carpetas) -> dict:
    estado = {}
    for c in carpetas:
        c = Path(c)
        if not c.is_dir():
            continue
        for raiz, dirs, archivos in os.walk(c):
            dirs[:] = [d for d in dirs if d != "__pycache__"]
            for a in archivos:
                p = Path(raiz) / a
                try:
                    st = p.stat()
                except OSError:
                    continue
                estado[str(p)] = (st.st_size, st.st_mtime_ns)
    return estado


def cambios(antes: dict, despues: dict) -> list:
    out = []
    for p in sorted(set(antes) | set(despues)):
        if p not in antes:
            out.append(f"NUEVO      {p}")
        elif p not in despues:
            out.append(f"BORRADO    {p}")
        elif antes[p] != despues[p]:
            out.append(f"MODIFICADO {p}")
    return out


# ------------------------------------------------------------------ correr y comparar
def resolver_script(script: str) -> Path:
    p = Path(script)
    if not p.is_absolute():
        p = (REPO / p) if (REPO / p).exists() else p.resolve()
    return p


def preparar_args(args, tmp: Path) -> list:
    out = []
    for a in args:
        a = a.replace(MARCA, str(tmp))
        if not os.path.isabs(a) and not a.startswith("-") and os.path.exists(a):
            a = os.path.abspath(a)
        out.append(a)
    return out


def recorrer(script, originales, args, tol_abs=0.01, tol_rel=1e-9, ignorar=(), vigilar=(),
             conservar=False, max_difs=25, imprimir=print) -> dict:
    """Corre el script en temporal y compara. Devuelve un informe (dict)."""
    ruta_script = resolver_script(script)
    pares = []
    for o in originales:
        if "=" in o and not Path(o).exists():
            o, nombre = o.rsplit("=", 1)
        else:
            nombre = Path(o).name
        pares.append((Path(o).resolve(), nombre))
    informe = {"script": str(ruta_script), "args": list(args), "archivos": [], "cambios_en_caso": [],
               "codigo": None, "veredicto": None}

    faltan = [str(o) for o, _ in pares if not o.exists()]
    if not ruta_script.exists():
        faltan.append(str(ruta_script))
    if faltan:
        informe["veredicto"] = "FALTA"
        informe["error"] = "no existe: " + " · ".join(faltan)
        imprimir(f"FALTA · no existe: {' · '.join(faltan)}")
        return informe
    if not any(MARCA in a for a in args):
        informe["veredicto"] = "FALTA"
        informe["error"] = f"ningún argumento lleva {MARCA}: la salida iría a la carpeta del caso"
        imprimir(f"FALTA · {informe['error']}. No se corre.")
        return informe

    carpetas = set(vigilar) or {str(o.parent) for o, _ in pares}
    for a in args:
        if not a.startswith("-") and MARCA not in a and os.path.isfile(a):
            carpetas.add(str(Path(a).resolve().parent))
    tmp = Path(tempfile.mkdtemp(prefix="auditar-"))
    try:
        argv = [sys.executable, "-B", str(ruta_script)] + preparar_args(args, tmp)
        antes = foto(carpetas)
        env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")
        proc = subprocess.run(argv, cwd=tmp, capture_output=True, env=env)
        despues = foto(carpetas)
        informe["codigo"] = proc.returncode
        informe["cambios_en_caso"] = cambios(antes, despues)
        salida_txt = proc.stdout.decode("utf-8", "replace")
        error_txt = proc.stderr.decode("utf-8", "replace")
        imprimir(f"RECORRER · {ruta_script.name}")
        imprimir("Comando  : python " + " ".join(f'"{x}"' if " " in x else x for x in argv[2:]))
        imprimir(f"Temporal : {tmp}{'  (se conserva)' if conservar else ''}")
        imprimir(f"Código   : {proc.returncode}")
        if proc.returncode != 0:
            informe["veredicto"] = "FALLA"
            informe["stdout"], informe["stderr"] = salida_txt, error_txt
            imprimir("El script FALLÓ. Salida tal cual:")
            imprimir(salida_txt.rstrip())
            imprimir(error_txt.rstrip())
            return informe
        reemplazos = [(str(tmp), MARCA)]
        hay_dif = False
        for orig, nombre in pares:
            nuevo = tmp / nombre
            reg = {"original": str(orig), "nuevo": nombre, "comparados": 0, "diferencias": []}
            if not nuevo.exists():
                reg["estado"] = "NO GENERADO"
                hay_dif = True
                imprimir(f"[NO GENERADO] {nombre}: el script no dejó este archivo")
                informe["archivos"].append(reg)
                continue
            suf = orig.suffix.lower()
            if suf == ".json":
                difs, n = comparar_json(orig, nuevo, tol_abs, tol_rel, ignorar, reemplazos)
            elif suf == ".xlsx":
                difs, n = comparar_xlsx(orig, nuevo, tol_abs, tol_rel, reemplazos)
                if difs is None:
                    reg["estado"] = "SIN COMPARAR"
                    imprimir(f"[SIN COMPARAR] {nombre}: falta openpyxl")
                    informe["archivos"].append(reg)
                    continue
            else:
                a, b = orig.read_bytes(), nuevo.read_bytes()
                difs, n = ([] if a == b else [("bytes", len(a), len(b))]), 1
            reg["comparados"] = n
            reg["diferencias"] = [{"donde": d, "original": x, "nuevo": y} for d, x, y in difs]
            reg["estado"] = "IGUAL" if not difs else "DISTINTO"
            hay_dif = hay_dif or bool(difs)
            imprimir(f"[{reg['estado']}] {orig.name}: {n} valores comparados, {len(difs)} diferencias")
            for d, x, y in difs[:max_difs]:
                dif = ""
                if es_numero(x) and es_numero(y):
                    dif = f" · dif {y - x:+,.6g}"
                imprimir(f"      {d}: original {x!r} · ahora {y!r}{dif}")
            if len(difs) > max_difs:
                imprimir(f"      … y {len(difs) - max_difs} más")
            informe["archivos"].append(reg)
        if informe["cambios_en_caso"]:
            imprimir("CAMBIOS EN LAS CARPETAS DEL CASO (el script escribió fuera del temporal):")
            for c in informe["cambios_en_caso"]:
                imprimir(f"      {c}")
        else:
            imprimir(f"Vigilancia: sin cambios en {len(carpetas)} carpeta(s) del caso")
        informe["veredicto"] = "DISTINTO" if hay_dif or informe["cambios_en_caso"] else "IGUAL"
        imprimir(f"VEREDICTO MECÁNICO: {informe['veredicto']}")
        return informe
    finally:
        if not conservar:
            shutil.rmtree(tmp, ignore_errors=True)


def codigo_de(informe) -> int:
    return {"IGUAL": 0, "DISTINTO": 1}.get(informe["veredicto"], 2)


# ------------------------------------------------------------------ manifiestos del ejecutor
def leer_manifiestos(ruta: Path) -> list:
    if ruta.is_dir():
        return [(p, cargar_json(p)) for p in sorted(ruta.glob("*.json"))]
    return [(ruta, cargar_json(ruta))]


def correr_manifiestos(ruta: Path, tol_abs, tol_rel, conservar) -> int:
    peor = 0
    for p, m in leer_manifiestos(ruta):
        print(f"\n=== {p.name} · {m.get('id', '')} · {m.get('metodo', '')}")
        inf = recorrer(m["script"], m["salidas"], m["args"], m.get("tol_abs", tol_abs),
                       m.get("tol_rel", tol_rel), m.get("ignorar", ()), conservar=conservar)
        c = codigo_de(inf)
        cd = m.get("cifra_decisiva")
        if c == 0 and cd and cd.get("archivo") and cd.get("ruta"):
            try:
                real = leer_valor(cargar_json(cd["archivo"]), cd["ruta"])
                ok = (numeros_iguales(float(real), float(cd["valor"]), max(tol_abs, 0.5), tol_rel)
                      if es_numero(real) and es_numero(cd.get("valor")) else real == cd.get("valor"))
                print(f"Cifra decisiva declarada {cd.get('nombre', cd['ruta'])} = {cd.get('valor')!r} · "
                      f"en la salida = {real!r} · {'COINCIDE' if ok else 'NO COINCIDE'}")
                if not ok:
                    c = max(c, 1)
            except (KeyError, OSError, ValueError) as e:
                print(f"Cifra decisiva: no se pudo leer ({e})")
                c = max(c, 1)
        peor = max(peor, c)
    return peor


# ------------------------------------------------------------------ prueba
def probar() -> int:
    fallas = []

    def chequear(cond, texto):
        print(("  ok   " if cond else "  MAL  ") + texto)
        if not cond:
            fallas.append(texto)

    silencio = lambda *_a, **_k: None  # noqa: E731
    base = Path(tempfile.mkdtemp(prefix="recorrer-prueba-"))
    try:
        caso = base / "caso"
        caso.mkdir()
        script = base / "metodo_ficticio.py"
        script.write_text(
            "import json, sys, pathlib\n"
            "e = json.loads(pathlib.Path(sys.argv[1]).read_text(encoding='utf-8'))\n"
            "if e.get('romper'): print('ERROR: falta el campo precio', file=sys.stderr); sys.exit(2)\n"
            "van = sum(f / (1 + e['tasa']) ** t for t, f in enumerate(e['flujos']))\n"
            "s = pathlib.Path(sys.argv[2]); s.mkdir(parents=True, exist_ok=True)\n"
            "out = {'empresa': 'Ferretería Andes', 'indicadores': {'van': round(van, 2)},\n"
            "       'alternativas': {'Leasing 36 cuotas': {'cae': 0.2046}},\n"
            "       'supuestos': [{'parametro': 'tasa', 'marca': 'SUPUESTO'}]}\n"
            "(s / 'salida.json').write_text(json.dumps(out, ensure_ascii=False), encoding='utf-8')\n"
            "if e.get('ensuciar'): (pathlib.Path(sys.argv[1]).parent / 'basura.txt').write_text('x')\n",
            encoding="utf-8")
        entrada = caso / "entrada.json"
        entrada.write_text(json.dumps({"tasa": 0.12, "flujos": [-1000, 400, 400, 400]}), encoding="utf-8")
        subprocess.run([sys.executable, "-B", str(script), str(entrada), str(caso)], check=True)
        archivada = caso / "salida.json"

        inf = recorrer(str(script), [str(archivada)], [str(entrada), MARCA], imprimir=silencio)
        chequear(inf["veredicto"] == "IGUAL", "misma entrada, mismo script: IGUAL")
        chequear(inf["archivos"][0]["comparados"] == 5, "compara los 5 valores hoja del JSON")

        d = cargar_json(archivada)
        d["indicadores"]["van"] += 37.0
        archivada.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        inf = recorrer(str(script), [str(archivada)], [str(entrada), MARCA], imprimir=silencio)
        difs = inf["archivos"][0]["diferencias"]
        chequear(inf["veredicto"] == "DISTINTO" and difs and difs[0]["donde"] == "indicadores.van",
                 "una salida archivada alterada en $37: DISTINTO en indicadores.van")
        inf = recorrer(str(script), [str(archivada)], [str(entrada), MARCA], tol_abs=50, imprimir=silencio)
        chequear(inf["veredicto"] == "IGUAL", "con tolerancia de 50 la misma diferencia pasa")

        rota = caso / "entrada-rota.json"
        rota.write_text(json.dumps({"romper": True}), encoding="utf-8")
        inf = recorrer(str(script), [str(archivada)], [str(rota), MARCA], imprimir=silencio)
        chequear(inf["veredicto"] == "FALLA" and "falta el campo precio" in inf.get("stderr", ""),
                 "si el script falla, devuelve FALLA con el error tal cual")

        intacta = archivada.read_text(encoding="utf-8")
        inf = recorrer(str(script), [str(archivada)], [str(entrada), str(caso)], imprimir=silencio)
        chequear(inf["veredicto"] == "FALTA" and archivada.read_text(encoding="utf-8") == intacta,
                 f"sin {MARCA} en los argumentos no corre (no pisa la salida archivada)")

        sucia = caso / "entrada-sucia.json"
        sucia.write_text(json.dumps({"tasa": 0.12, "flujos": [-1000, 400, 400, 400], "ensuciar": True}),
                         encoding="utf-8")
        inf = recorrer(str(script), [str(archivada) + "=salida.json"], [str(sucia), MARCA],
                       tol_abs=50, imprimir=silencio)
        chequear(inf["veredicto"] == "DISTINTO" and any("basura.txt" in c for c in inf["cambios_en_caso"]),
                 "si el script escribe en la carpeta del caso, la vigilancia lo detecta")

        datos = {"alternativas": {"Leasing 36 cuotas": {"cae": 0.2046}}, "filas": [{"v": 1}, {"v": 2}]}
        chequear(leer_valor(datos, 'alternativas["Leasing 36 cuotas"].cae') == 0.2046,
                 "ruta con comillas y espacios")
        chequear(leer_valor(datos, "filas[1].v") == 2, "ruta con índice")

        manif = base / "ejecuciones"
        manif.mkdir()
        (manif / "metodo-ficticio.json").write_text(json.dumps({
            "id": "prueba", "metodo": "ficticio", "script": str(script),
            "args": [str(entrada), MARCA], "salidas": [str(archivada)], "tol_abs": 50,
            "cifra_decisiva": {"nombre": "VAN", "archivo": str(archivada), "ruta": "indicadores.van",
                               "valor": d["indicadores"]["van"]}}), encoding="utf-8")
        viejo = sys.stdout
        sys.stdout = open(os.devnull, "w", encoding="utf-8")
        try:
            c = correr_manifiestos(manif, 0.01, 1e-9, False)
        finally:
            sys.stdout.close()
            sys.stdout = viejo
        chequear(c == 0, "manifiesto del ejecutor: re-corre y la cifra decisiva coincide")
    finally:
        shutil.rmtree(base, ignore_errors=True)

    print("OK" if not fallas else f"FALLA ({len(fallas)})")
    return 0 if not fallas else 1


# ------------------------------------------------------------------ main
def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    argv = list(sys.argv[1:] if argv is None else argv)
    if "--probar" in argv:
        return probar()
    extra = []
    if "--" in argv:
        i = argv.index("--")
        argv, extra = argv[:i], argv[i + 1:]
    ap = argparse.ArgumentParser(description="Re-corre un método en temporal y compara la salida cifra por cifra.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    pc = sub.add_parser("correr")
    pc.add_argument("--script", required=True, help="script del método (ruta absoluta o relativa al repo)")
    pc.add_argument("--original", action="append", required=True, help="salida archivada (repetible)")
    pc.add_argument("--tol-abs", type=float, default=0.01)
    pc.add_argument("--tol-rel", type=float, default=1e-9)
    pc.add_argument("--ignorar", action="append", default=[], help="ruta a ignorar (admite * y ?)")
    pc.add_argument("--vigilar", action="append", default=[], help="carpeta del caso a vigilar")
    pc.add_argument("--informe", help="guarda el informe en JSON (fuera del caso)")
    pc.add_argument("--conservar", action="store_true", help="no borra la carpeta temporal")
    pm = sub.add_parser("manifiesto")
    pm.add_argument("ruta")
    pm.add_argument("--tol-abs", type=float, default=0.01)
    pm.add_argument("--tol-rel", type=float, default=1e-9)
    pm.add_argument("--conservar", action="store_true")
    pv = sub.add_parser("valor")
    pv.add_argument("archivo")
    pv.add_argument("ruta")
    px = sub.add_parser("cruzar")
    px.add_argument("a")
    px.add_argument("ruta_a")
    px.add_argument("b")
    px.add_argument("ruta_b")
    px.add_argument("--tol-abs", type=float, default=0.01)
    px.add_argument("--tol-rel", type=float, default=1e-9)
    a = ap.parse_args(argv)

    if a.cmd == "correr":
        inf = recorrer(a.script, a.original, extra, a.tol_abs, a.tol_rel, a.ignorar, a.vigilar, a.conservar)
        if a.informe:
            Path(a.informe).write_text(json.dumps(inf, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return codigo_de(inf)
    if a.cmd == "manifiesto":
        return correr_manifiestos(Path(a.ruta), a.tol_abs, a.tol_rel, a.conservar)
    try:
        if a.cmd == "valor":
            v = leer_valor(cargar_json(a.archivo), a.ruta)
            print(json.dumps(v, ensure_ascii=False, indent=1) if isinstance(v, (dict, list)) else repr(v))
            return 0
        va = leer_valor(cargar_json(a.a), a.ruta_a)
        vb = leer_valor(cargar_json(a.b), a.ruta_b)
    except (KeyError, OSError, ValueError) as e:
        print(f"FALTA · {e}")
        return 2
    ok = numeros_iguales(float(va), float(vb), a.tol_abs, a.tol_rel) if es_numero(va) and es_numero(vb) else va == vb
    dif = f" · dif {vb - va:+,.6g}" if es_numero(va) and es_numero(vb) else ""
    print(f"{'COINCIDE' if ok else 'NO COINCIDE'} · {Path(a.a).name}:{a.ruta_a} = {va!r} · "
          f"{Path(a.b).name}:{a.ruta_b} = {vb!r}{dif}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
