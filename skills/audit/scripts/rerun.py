#!/usr/bin/env python3
"""rerun.py · the mechanical half of the audit (stage 4 of ingeniero-agentico).

Runs a method's script again with the SAME input, but leaving the output in a temporary folder,
and compares figure by figure what came out now against the archived output. It writes nothing
outside the temporary folder and warns if the script touched the case folders.

    python rerun.py run --script <method.py> --original <output.json> [--original <output.xlsx>]
                        [--tol-abs 0.01] [--tol-rel 1e-9] [--ignore <path>] -- <args with {OUT}>
    python rerun.py manifest <context/runs/ | one-run.json>
    python rerun.py value <file.json> <path>               (e.g. indicators.npv)
    python rerun.py cross <a.json> <path_a> <b.json> <path_b> [--tol-abs 1]
    python rerun.py --test

{OUT} in the arguments is replaced by the temporary folder ({SALIDA}, the pre-v1.1 marker, too).
Each --original is compared with the file of the same name in that folder (or another name:
--original path=new_name). Paths like  alternatives["Leasing 36 payments"].cae  or  rows[0].value.

Pre-v1.1 names still work: subcommands correr, manifiesto, valor, cruzar; flags --ignorar,
--vigilar, --informe, --conservar, --probar; Spanish manifest keys (metodo, salidas,
cifra_decisiva...); and a manifest that names an old script path (skills/evaluacion-economica/
scripts/evaluar.py) runs the renamed one.

Exit: 0 = all the same · 1 = differences or changes in the case · 2 = the script failed or
something is missing. Standard library only; if openpyxl is installed, .xlsx files are compared
cell by cell too.
"""
from __future__ import annotations

import argparse
import fnmatch
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.dont_write_bytecode = True
REPO = Path(__file__).resolve().parents[3]
MARK = "{OUT}"
MARKS = (MARK, "{SALIDA}")

# Script paths before v1.1 -> after (a manifest written before the rename still re-runs).
SCRIPT_RENAMES = {
    "skills/adquisiciones/scripts/adquisiciones.py": "skills/procurement/scripts/procurement.py",
    "skills/costo-y-precio/scripts/costo_precio.py": "skills/cost-and-pricing/scripts/cost_pricing.py",
    "skills/evaluacion-economica/scripts/evaluar.py": "skills/economic-evaluation/scripts/evaluate.py",
    "skills/evaluacion-financiera/scripts/evaluacion_financiera.py":
        "skills/financial-evaluation/scripts/financial_evaluation.py",
    "skills/cartografiar/scripts/mapa_contexto.py": "skills/map-context/scripts/map_context.py",
    "skills/investigar/scripts/validar_investigacion.py": "skills/research/scripts/validate_research.py",
    "skills/auditar/scripts/recorrer.py": "skills/audit/scripts/rerun.py",
}
SUBCOMMANDS_ES = {"correr": "run", "manifiesto": "manifest", "valor": "value", "cruzar": "cross"}
MANIFEST_KEYS = {"metodo": "method", "fecha": "date", "carpeta_salida": "out_dir", "salidas": "outputs",
                 "cifra_decisiva": "decisive_figure", "supuestos": "assumptions", "ignorar": "ignore"}
FIGURE_KEYS = {"nombre": "name", "archivo": "file", "ruta": "path", "valor": "value"}


# ------------------------------------------------------------------ paths inside a JSON
def split_path(path: str) -> list:
    """'a.b[0]["c d"].e' -> ['a', 'b', 0, 'c d', 'e']"""
    steps, i, n = [], 0, len(path)
    current = ""
    while i < n:
        ch = path[i]
        if ch == ".":
            if current:
                steps.append(current)
            current = ""
            i += 1
        elif ch == "[":
            if current:
                steps.append(current)
            current = ""
            end = path.index("]", i) if path[i + 1] not in "\"'" else path.index(path[i + 1] + "]", i + 2) + 1
            inside = path[i + 1:end]
            if inside[:1] in "\"'":
                steps.append(inside[1:-1])
            else:
                steps.append(int(inside))
            i = end + 1
        else:
            current += ch
            i += 1
    if current:
        steps.append(current)
    return steps


def read_value(data, path: str):
    v = data
    for p in split_path(path):
        if isinstance(p, int):
            if not isinstance(v, list) or p >= len(v):
                raise KeyError(f"{path}: there is no index [{p}]")
            v = v[p]
        else:
            if not isinstance(v, dict) or p not in v:
                raise KeyError(f"{path}: the key «{p}» does not exist")
            v = v[p]
    return v


def load_json(path) -> object:
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


# ------------------------------------------------------------------ compare
def numbers_equal(a: float, b: float, tol_abs: float, tol_rel: float) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if math.isnan(a) and math.isnan(b):
        return True
    return abs(a - b) <= tol_abs or abs(a - b) <= tol_rel * max(abs(a), abs(b))


def is_number(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


# ------------------------------------------------------------------ outputs archived before v1.1
ALIASES_FILE = Path(__file__).resolve().parent / "output_aliases.json"
RE_IDENT_PATH = re.compile(r"^[A-Za-z_áéíóúñ][\w\[\]\.áéíóúñ]*$")


class Legacy:
    """Reads an output archived before v1.1 (Spanish JSON keys) against a re-run with English keys.

    The map (output_aliases.json) was derived by running each new script on old inputs and
    pairing its output with the archived one. Only key names and enum words are translated: every
    number and every other text is still compared exactly. Fields a re-run carries that the old
    format never had are listed as "added since v1.1", not compared (there is nothing to compare)."""

    def __init__(self, data):
        self.keys = data.get("keys", {})
        self.scoped = data.get("scoped", {})
        self.values = data.get("values", {})
        self.tokens = {**self.keys, **data.get("path_tokens", {})}
        self.renamed = 0
        self.added = []

    @classmethod
    def load(cls):
        try:
            return cls(json.loads(ALIASES_FILE.read_text(encoding="utf-8")))
        except (OSError, ValueError):
            return None

    def key(self, parent, k):
        return self.scoped.get(f"{parent}.{k}") or self.keys.get(k)

    def same_text(self, a, b):
        if self.values.get(a) == b:
            return True
        if RE_IDENT_PATH.match(a):                 # a path written as text: ingresos[0].ocupacion
            return re.sub(r"[A-Za-z_áéíóúñ][A-Za-z_áéíóúñ]*", lambda m: self.tokens.get(m.group(0), m.group(0)), a) == b
        return False

    def applies_to(self, tree, new_tree) -> bool:
        """True if the archived tree uses old key names that the re-run does not."""
        old, new = set(), set()
        _collect_keys(tree, old)
        _collect_keys(new_tree, new)
        return any(k in self.keys and k not in new and self.keys[k] in new for k in old)


def _collect_keys(x, out):
    if isinstance(x, dict):
        for k, v in x.items():
            out.add(k)
            _collect_keys(v, out)
    elif isinstance(x, list):
        for v in x:
            _collect_keys(v, out)


def _signature(x, legacy):
    """What identifies a list element regardless of key names: its leaf values."""
    leaves = []

    def go(v):
        if isinstance(v, dict):
            for w in v.values():
                go(w)
        elif isinstance(v, list):
            for w in v:
                go(w)
        elif is_number(v):
            leaves.append(round(float(v), 2))
        elif isinstance(v, str) and legacy is not None:
            leaves.append(legacy.values.get(v, v))
        else:
            leaves.append(v)

    go(x)
    return json.dumps(sorted(map(str, leaves)))


def compare_values(a, b, path, diffs, count, tol_abs, tol_rel, ignore, replacements, legacy=None, parent=""):
    if any(fnmatch.fnmatchcase(path, pattern) for pattern in ignore):
        return
    if isinstance(a, dict) and isinstance(b, dict):
        used = set()
        for k in a:
            target = k if k in b else (legacy.key(parent, k) if legacy else None)
            if target is not None and (target not in b or target in used or (target != k and target in a)):
                target = None
            sub = f"{path}.{target or k}" if path else str(target or k)
            if target is None:
                diffs.append((sub, a[k], "<missing>"))
                continue
            if target != k:
                legacy.renamed += 1
            used.add(target)
            compare_values(a[k], b[target], sub, diffs, count, tol_abs, tol_rel, ignore, replacements,
                           legacy, target)
        for k in b:
            if k not in used and k not in a:
                sub = f"{path}.{k}" if path else str(k)
                if legacy is not None:
                    legacy.added.append(sub)
                else:
                    diffs.append((sub, "<missing>", b[k]))
        return
    if isinstance(a, list) and isinstance(b, list):
        if len(a) != len(b) and legacy is not None and len(a) < len(b):
            # a re-run may carry rows the old format did not have: align by leaf values
            import difflib  # noqa: PLC0415
            sm = difflib.SequenceMatcher(a=[_signature(x, legacy) for x in a],
                                         b=[_signature(y, legacy) for y in b], autojunk=False)
            pairs, extra_b, lost_a = [], [], []
            for tag, i1, i2, j1, j2 in sm.get_opcodes():
                if tag == "equal" or (tag == "replace" and i2 - i1 == j2 - j1):
                    pairs += list(zip(range(i1, i2), range(j1, j2)))
                elif tag == "insert":
                    extra_b += list(range(j1, j2))
                else:
                    lost_a += list(range(i1, i2))
                    extra_b += list(range(j1, j2))
            if not lost_a:
                for j in extra_b:
                    legacy.added.append(f"{path}[{j}]")
                for i, j in pairs:
                    compare_values(a[i], b[j], f"{path}[{j}]", diffs, count, tol_abs, tol_rel, ignore,
                                   replacements, legacy, parent)
                return
        if len(a) != len(b):
            diffs.append((f"{path} (length)", len(a), len(b)))
        for i, (x, y) in enumerate(zip(a, b)):
            compare_values(x, y, f"{path}[{i}]", diffs, count, tol_abs, tol_rel, ignore, replacements,
                           legacy, parent)
        return
    count[0] += 1
    if is_number(a) and is_number(b):
        if not numbers_equal(float(a), float(b), tol_abs, tol_rel):
            diffs.append((path, a, b))
        return
    if isinstance(a, str) and isinstance(b, str):
        for old, new in replacements:
            b = b.replace(old, new)
        if a != b and not (legacy is not None and legacy.same_text(a, b)):
            diffs.append((path, a, b))
        return
    if a != b:
        diffs.append((path, a, b))


def compare_json(orig: Path, new: Path, tol_abs, tol_rel, ignore, replacements, info=None):
    diffs, count = [], [0]
    ta, tb = load_json(orig), load_json(new)
    legacy = Legacy.load()
    if legacy is None or not legacy.applies_to(ta, tb):
        legacy = None
    compare_values(ta, tb, "", diffs, count, tol_abs, tol_rel, ignore, replacements, legacy)
    if info is not None and legacy is not None:
        info.update(legacy=True, renamed=legacy.renamed, added=legacy.added)
    return diffs, count[0]


def compare_xlsx(orig: Path, new: Path, tol_abs, tol_rel, replacements, ignore=()):
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        return None, 0
    wa = openpyxl.load_workbook(orig)
    wb = openpyxl.load_workbook(new)
    diffs, count = [], 0
    if wa.sheetnames != wb.sheetnames:
        diffs.append(("sheets", wa.sheetnames, wb.sheetnames))
    for name in wa.sheetnames:
        if name not in wb.sheetnames:
            continue
        ha, hb = wa[name], wb[name]
        rows = max(ha.max_row, hb.max_row)
        cols = max(ha.max_column, hb.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                ca, cb = ha.cell(row=r, column=c), hb.cell(row=r, column=c)
                va, vb = ca.value, cb.value
                if va is None and vb is None:
                    continue
                where = f"{name}!{ca.coordinate}"
                if any(fnmatch.fnmatchcase(where, pattern) for pattern in ignore):
                    continue
                count += 1
                if is_number(va) and is_number(vb):
                    if not numbers_equal(float(va), float(vb), tol_abs, tol_rel):
                        diffs.append((where, va, vb))
                else:
                    if isinstance(vb, str):
                        for old, new_txt in replacements:
                            vb = vb.replace(old, new_txt)
                    if va != vb:
                        diffs.append((where, va, vb))
                fa = ca.fill.fgColor.rgb if ca.fill and ca.fill.fill_type else None
                fb = cb.fill.fgColor.rgb if cb.fill and cb.fill.fill_type else None
                if fa != fb:
                    diffs.append((where + " (color)", fa, fb))
    return diffs, count


# ------------------------------------------------------------------ watch the case folders
def snapshot(folders) -> dict:
    state = {}
    for c in folders:
        c = Path(c)
        if not c.is_dir():
            continue
        for root, dirs, files in os.walk(c):
            dirs[:] = [d for d in dirs if d != "__pycache__"]
            for a in files:
                p = Path(root) / a
                try:
                    st = p.stat()
                except OSError:
                    continue
                state[str(p)] = (st.st_size, st.st_mtime_ns)
    return state


def changes(before: dict, after: dict) -> list:
    out = []
    for p in sorted(set(before) | set(after)):
        if p not in before:
            out.append(f"NEW        {p}")
        elif p not in after:
            out.append(f"DELETED    {p}")
        elif before[p] != after[p]:
            out.append(f"MODIFIED   {p}")
    return out


# ------------------------------------------------------------------ run and compare
def resolve_script(script: str) -> Path:
    p = Path(script)
    if not p.is_absolute():
        p = (REPO / p) if (REPO / p).exists() else p.resolve()
    if not p.exists():
        posix = str(p).replace("\\", "/")
        for old, new in SCRIPT_RENAMES.items():
            if posix.endswith(old):
                renamed = Path(posix[: -len(old)] + new)
                p = renamed if renamed.exists() else (REPO / new)
                break
    return p


def has_mark(a: str) -> bool:
    return any(m in a for m in MARKS)


def prepare_args(args, tmp: Path) -> list:
    out = []
    for a in args:
        for m in MARKS:
            a = a.replace(m, str(tmp))
        if not os.path.isabs(a) and not a.startswith("-") and os.path.exists(a):
            a = os.path.abspath(a)
        out.append(a)
    return out


def rerun(script, originals, args, tol_abs=0.01, tol_rel=1e-9, ignore=(), watch=(),
          keep=False, max_diffs=25, say=print) -> dict:
    """Runs the script in a temporary folder and compares. Returns a report (dict)."""
    script_path = resolve_script(script)
    pairs = []
    for o in originals:
        if "=" in o and not Path(o).exists():
            o, name = o.rsplit("=", 1)
        else:
            name = Path(o).name
        pairs.append((Path(o).resolve(), name))
    rep = {"script": str(script_path), "args": list(args), "files": [], "case_changes": [],
           "code": None, "verdict": None}

    missing = [str(o) for o, _ in pairs if not o.exists()]
    if not script_path.exists():
        missing.append(str(script_path))
    if missing:
        rep["verdict"] = "MISSING"
        rep["error"] = "does not exist: " + " · ".join(missing)
        say(f"MISSING · does not exist: {' · '.join(missing)}")
        return rep
    if not any(has_mark(a) for a in args):
        rep["verdict"] = "MISSING"
        rep["error"] = f"no argument carries {MARK}: the output would go to the case folder"
        say(f"MISSING · {rep['error']}. Not run.")
        return rep

    folders = set(watch) or {str(o.parent) for o, _ in pairs}
    for a in args:
        if not a.startswith("-") and not has_mark(a) and os.path.isfile(a):
            folders.add(str(Path(a).resolve().parent))
    tmp = Path(tempfile.mkdtemp(prefix="audit-"))
    try:
        argv = [sys.executable, "-B", str(script_path)] + prepare_args(args, tmp)
        before = snapshot(folders)
        env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")
        proc = subprocess.run(argv, cwd=tmp, capture_output=True, env=env)
        after = snapshot(folders)
        rep["code"] = proc.returncode
        rep["case_changes"] = changes(before, after)
        out_txt = proc.stdout.decode("utf-8", "replace")
        err_txt = proc.stderr.decode("utf-8", "replace")
        say(f"RERUN · {script_path.name}")
        say("Command  : python " + " ".join(f'"{x}"' if " " in x else x for x in argv[2:]))
        say(f"Temporary: {tmp}{'  (kept)' if keep else ''}")
        say(f"Exit code: {proc.returncode}")
        if proc.returncode != 0:
            rep["verdict"] = "FAIL"
            rep["stdout"], rep["stderr"] = out_txt, err_txt
            say("The script FAILED. Output as it came:")
            say(out_txt.rstrip())
            say(err_txt.rstrip())
            return rep
        replacements = [(str(tmp), MARK)]
        any_diff = False
        for orig, name in pairs:
            new = tmp / name
            reg = {"original": str(orig), "new": name, "compared": 0, "differences": []}
            if not new.exists():
                reg["status"] = "NOT GENERATED"
                any_diff = True
                say(f"[NOT GENERATED] {name}: the script did not leave this file")
                rep["files"].append(reg)
                continue
            suf = orig.suffix.lower()
            info = {}
            if suf == ".json":
                diffs, n = compare_json(orig, new, tol_abs, tol_rel, ignore, replacements, info)
            elif suf == ".xlsx":
                diffs, n = compare_xlsx(orig, new, tol_abs, tol_rel, replacements, ignore)
                if diffs is None:
                    reg["status"] = "NOT COMPARED"
                    say(f"[NOT COMPARED] {name}: openpyxl is not installed")
                    rep["files"].append(reg)
                    continue
            else:
                a, b = orig.read_bytes(), new.read_bytes()
                diffs, n = ([] if a == b else [("bytes", len(a), len(b))]), 1
            reg["compared"] = n
            reg["differences"] = [{"where": d, "original": x, "new": y} for d, x, y in diffs]
            reg["status"] = "SAME" if not diffs else "DIFFERENT"
            any_diff = any_diff or bool(diffs)
            say(f"[{reg['status']}] {orig.name}: {n} values compared, {len(diffs)} differences")
            if info.get("legacy"):
                reg["archived_before_v1_1"] = {"renamed_keys": info["renamed"], "added_since_v1_1": info["added"]}
                fields = sorted({re.sub(r"\[\d+\]", "[]", x) for x in info["added"]})
                say(f"      archived before v1.1: {info['renamed']} Spanish keys read through output_aliases.json; "
                    f"{len(info['added'])} fields added since v1.1 not compared"
                    + (f" ({', '.join(fields[:6])}{' …' if len(fields) > 6 else ''})" if fields else ""))
            for d, x, y in diffs[:max_diffs]:
                dif = ""
                if is_number(x) and is_number(y):
                    dif = f" · diff {y - x:+,.6g}"
                say(f"      {d}: original {x!r} · now {y!r}{dif}")
            if len(diffs) > max_diffs:
                say(f"      … and {len(diffs) - max_diffs} more")
            rep["files"].append(reg)
        if rep["case_changes"]:
            say("CHANGES IN THE CASE FOLDERS (the script wrote outside the temporary folder):")
            for c in rep["case_changes"]:
                say(f"      {c}")
        else:
            say(f"Watch: no changes in {len(folders)} case folder(s)")
        rep["verdict"] = "DIFFERENT" if any_diff or rep["case_changes"] else "SAME"
        say(f"MECHANICAL VERDICT: {rep['verdict']}")
        return rep
    finally:
        if not keep:
            shutil.rmtree(tmp, ignore_errors=True)


def exit_code(rep) -> int:
    return {"SAME": 0, "DIFFERENT": 1}.get(rep["verdict"], 2)


# ------------------------------------------------------------------ the executor's manifests
def normalize_manifest(m: dict) -> dict:
    """English keys; a manifest written before v1.1 (Spanish keys) is accepted."""
    m = {MANIFEST_KEYS.get(k, k): v for k, v in m.items()}
    if isinstance(m.get("decisive_figure"), dict):
        m["decisive_figure"] = {FIGURE_KEYS.get(k, k): v for k, v in m["decisive_figure"].items()}
    return m


def read_manifests(path: Path) -> list:
    if path.is_dir():
        return [(p, normalize_manifest(load_json(p))) for p in sorted(path.glob("*.json"))]
    return [(path, normalize_manifest(load_json(path)))]


def run_manifests(path: Path, tol_abs, tol_rel, keep) -> int:
    worst = 0
    for p, m in read_manifests(path):
        print(f"\n=== {p.name} · {m.get('id', '')} · {m.get('method', '')}")
        rep = rerun(m["script"], m["outputs"], m["args"], m.get("tol_abs", tol_abs),
                    m.get("tol_rel", tol_rel), m.get("ignore", ()), keep=keep)
        c = exit_code(rep)
        fig = m.get("decisive_figure")
        if c == 0 and fig and fig.get("file") and fig.get("path"):
            try:
                real = read_value(load_json(fig["file"]), fig["path"])
                ok = (numbers_equal(float(real), float(fig["value"]), max(tol_abs, 0.5), tol_rel)
                      if is_number(real) and is_number(fig.get("value")) else real == fig.get("value"))
                print(f"Declared decisive figure {fig.get('name', fig['path'])} = {fig.get('value')!r} · "
                      f"in the output = {real!r} · {'MATCHES' if ok else 'DOES NOT MATCH'}")
                if not ok:
                    c = max(c, 1)
            except (KeyError, OSError, ValueError) as e:
                print(f"Decisive figure: could not be read ({e})")
                c = max(c, 1)
        worst = max(worst, c)
    return worst


# ------------------------------------------------------------------ test
def run_test() -> int:
    failures = []

    def check(cond, text):
        print(("  ok   " if cond else "  FAIL ") + text)
        if not cond:
            failures.append(text)

    quiet = lambda *_a, **_k: None  # noqa: E731
    base = Path(tempfile.mkdtemp(prefix="rerun-test-"))
    try:
        case = base / "case"
        case.mkdir()
        script = base / "fictitious_method.py"
        script.write_text(
            "import json, sys, pathlib\n"
            "e = json.loads(pathlib.Path(sys.argv[1]).read_text(encoding='utf-8'))\n"
            "if e.get('break'): print('ERROR: the price field is missing', file=sys.stderr); sys.exit(2)\n"
            "npv = sum(f / (1 + e['rate']) ** t for t, f in enumerate(e['flows']))\n"
            "s = pathlib.Path(sys.argv[2]); s.mkdir(parents=True, exist_ok=True)\n"
            "out = {'company': 'Ferretería Andes', 'indicators': {'npv': round(npv, 2)},\n"
            "       'alternatives': {'Leasing 36 payments': {'cae': 0.2046}},\n"
            "       'assumptions': [{'parameter': 'rate', 'mark': 'ASSUMPTION'}]}\n"
            "(s / 'output.json').write_text(json.dumps(out, ensure_ascii=False), encoding='utf-8')\n"
            "if e.get('dirty'): (pathlib.Path(sys.argv[1]).parent / 'junk.txt').write_text('x')\n",
            encoding="utf-8")
        inp = case / "input.json"
        inp.write_text(json.dumps({"rate": 0.12, "flows": [-1000, 400, 400, 400]}), encoding="utf-8")
        subprocess.run([sys.executable, "-B", str(script), str(inp), str(case)], check=True)
        archived = case / "output.json"

        rep = rerun(str(script), [str(archived)], [str(inp), MARK], say=quiet)
        check(rep["verdict"] == "SAME", "same input, same script: SAME")
        check(rep["files"][0]["compared"] == 5, "compares the 5 leaf values of the JSON")

        d = load_json(archived)
        d["indicators"]["npv"] += 37.0
        archived.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        rep = rerun(str(script), [str(archived)], [str(inp), MARK], say=quiet)
        diffs = rep["files"][0]["differences"]
        check(rep["verdict"] == "DIFFERENT" and diffs and diffs[0]["where"] == "indicators.npv",
              "an archived output altered by $37: DIFFERENT at indicators.npv")
        rep = rerun(str(script), [str(archived)], [str(inp), MARK], tol_abs=50, say=quiet)
        check(rep["verdict"] == "SAME", "with a tolerance of 50 the same difference passes")

        broken = case / "input-broken.json"
        broken.write_text(json.dumps({"break": True}), encoding="utf-8")
        rep = rerun(str(script), [str(archived)], [str(broken), MARK], say=quiet)
        check(rep["verdict"] == "FAIL" and "the price field is missing" in rep.get("stderr", ""),
              "if the script fails, it returns FAIL with the error as it came")

        intact = archived.read_text(encoding="utf-8")
        rep = rerun(str(script), [str(archived)], [str(inp), str(case)], say=quiet)
        check(rep["verdict"] == "MISSING" and archived.read_text(encoding="utf-8") == intact,
              f"without {MARK} in the arguments it does not run (it does not overwrite the archived output)")

        dirty = case / "input-dirty.json"
        dirty.write_text(json.dumps({"rate": 0.12, "flows": [-1000, 400, 400, 400], "dirty": True}),
                         encoding="utf-8")
        rep = rerun(str(script), [str(archived) + "=output.json"], [str(dirty), MARK],
                    tol_abs=50, say=quiet)
        check(rep["verdict"] == "DIFFERENT" and any("junk.txt" in c for c in rep["case_changes"]),
              "if the script writes in the case folder, the watch detects it")

        # an output archived before v1.1: Spanish keys and enum words, one field fewer
        old_dir = base / "archived-v1"
        old_dir.mkdir()
        good = load_json(case / "output.json")
        good["indicators"]["npv"] -= 37.0          # undo the alteration made above
        spanish = {"empresa": good["company"], "indicadores": {"van": good["indicators"]["npv"]},
                   "alternativas": {"Leasing 36 payments": {"cae": 0.2046}}}
        old_out = old_dir / "output.json"
        old_out.write_text(json.dumps(spanish, ensure_ascii=False), encoding="utf-8")
        rep = rerun(str(script), [str(old_out)], [str(inp), MARK], say=quiet)
        leg = rep["files"][0].get("archived_before_v1_1", {})
        check(rep["verdict"] == "SAME" and leg.get("renamed_keys") == 4 and "assumptions" in leg.get("added_since_v1_1", []),
              "an output archived before v1.1 (Spanish keys) compares SAME through output_aliases.json")
        spanish["indicadores"]["van"] += 1.0
        old_out.write_text(json.dumps(spanish, ensure_ascii=False), encoding="utf-8")
        rep = rerun(str(script), [str(old_out)], [str(inp), MARK], say=quiet)
        diffs = rep["files"][0]["differences"]
        check(rep["verdict"] == "DIFFERENT" and [x["where"] for x in diffs] == ["indicators.npv"],
              "…and a changed number in it is still DIFFERENT (only names are translated)")
        lg = Legacy.load()
        check(lg is not None and lg.same_text("ingresos[0].ocupacion", "revenues[0].occupancy")
              and lg.same_text("lineal", "straight_line") and not lg.same_text("lineal", "accelerated"),
              "aliases translate paths written as text and enum words, nothing else")

        data = {"alternatives": {"Leasing 36 payments": {"cae": 0.2046}}, "rows": [{"v": 1}, {"v": 2}]}
        check(read_value(data, 'alternatives["Leasing 36 payments"].cae') == 0.2046,
              "path with quotes and spaces")
        check(read_value(data, "rows[1].v") == 2, "path with an index")

        runs = base / "runs"
        runs.mkdir()
        (runs / "fictitious-method.json").write_text(json.dumps({
            "id": "test", "method": "fictitious", "script": str(script),
            "args": [str(inp), MARK], "outputs": [str(archived)], "tol_abs": 50,
            "decisive_figure": {"name": "NPV", "file": str(archived), "path": "indicators.npv",
                                "value": d["indicators"]["npv"]}}), encoding="utf-8")
        legacy = base / "ejecuciones"
        legacy.mkdir()
        (legacy / "metodo-ficticio.json").write_text(json.dumps({
            "id": "prueba", "metodo": "ficticio", "script": str(script),
            "args": [str(inp), "{SALIDA}"], "salidas": [str(archived)], "tol_abs": 50,
            "cifra_decisiva": {"nombre": "VAN", "archivo": str(archived), "ruta": "indicators.npv",
                               "valor": d["indicators"]["npv"]}}), encoding="utf-8")
        old = sys.stdout
        sys.stdout = open(os.devnull, "w", encoding="utf-8")
        try:
            c = run_manifests(runs, 0.01, 1e-9, False)
            c_legacy = run_manifests(legacy, 0.01, 1e-9, False)
            c_value = main(["valor", str(archived), "indicators.npv"])
            c_cross = main(["cruzar", str(archived), "indicators.npv", str(archived), "indicators.npv"])
        finally:
            sys.stdout.close()
            sys.stdout = old
        check(c == 0, "executor manifest: re-runs and the decisive figure matches")
        check(c_legacy == 0, "a pre-v1.1 manifest (Spanish keys, {SALIDA}) re-runs the same way")
        renamed = resolve_script("skills/evaluacion-economica/scripts/evaluar.py")
        check(renamed.as_posix().endswith("skills/economic-evaluation/scripts/evaluate.py"),
              "an old script path resolves to the renamed script")
        check(c_value == 0 and c_cross == 0, "the old subcommands «valor» and «cruzar» still work")
    finally:
        shutil.rmtree(base, ignore_errors=True)

    print("OK" if not failures else f"FAIL ({len(failures)})")
    return 0 if not failures else 1


# ------------------------------------------------------------------ main
def _alias(ap, old, dest, **kw):
    """A pre-v1.1 Spanish flag that still works and does not show in --help."""
    ap.add_argument(old, dest=dest, help=argparse.SUPPRESS, default=argparse.SUPPRESS, **kw)


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    argv = list(sys.argv[1:] if argv is None else argv)
    if "--test" in argv or "--probar" in argv:
        return run_test()
    extra = []
    if "--" in argv:
        i = argv.index("--")
        argv, extra = argv[:i], argv[i + 1:]
    if argv and argv[0] in SUBCOMMANDS_ES:
        argv[0] = SUBCOMMANDS_ES[argv[0]]
    ap = argparse.ArgumentParser(description="Re-runs a method in a temporary folder and compares the output "
                                             "figure by figure.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    pc = sub.add_parser("run", help="re-run one script and compare its outputs")
    pc.add_argument("--script", required=True, help="the method's script (absolute or relative to the repo)")
    pc.add_argument("--original", action="append", required=True, help="archived output (repeatable)")
    pc.add_argument("--tol-abs", type=float, default=0.01)
    pc.add_argument("--tol-rel", type=float, default=1e-9)
    pc.add_argument("--ignore", action="append", default=[],
                    help="JSON path or Excel SHEET!CELL to ignore (accepts * and ?)")
    pc.add_argument("--watch", action="append", default=[], help="case folder to watch")
    pc.add_argument("--report", help="saves the report as JSON (outside the case)")
    pc.add_argument("--keep", action="store_true", help="does not delete the temporary folder")
    _alias(pc, "--ignorar", "ignore", action="append")
    _alias(pc, "--vigilar", "watch", action="append")
    _alias(pc, "--informe", "report")
    _alias(pc, "--conservar", "keep", action="store_true")
    pm = sub.add_parser("manifest", help="re-run every manifest of context/runs/")
    pm.add_argument("path")
    pm.add_argument("--tol-abs", type=float, default=0.01)
    pm.add_argument("--tol-rel", type=float, default=1e-9)
    pm.add_argument("--keep", action="store_true")
    _alias(pm, "--conservar", "keep", action="store_true")
    pv = sub.add_parser("value", help="read one value by path")
    pv.add_argument("file")
    pv.add_argument("path")
    px = sub.add_parser("cross", help="compare one value of two chained outputs")
    px.add_argument("a")
    px.add_argument("path_a")
    px.add_argument("b")
    px.add_argument("path_b")
    px.add_argument("--tol-abs", type=float, default=0.01)
    px.add_argument("--tol-rel", type=float, default=1e-9)
    a = ap.parse_args(argv)

    if a.cmd == "run":
        rep = rerun(a.script, a.original, extra, a.tol_abs, a.tol_rel, a.ignore, a.watch, a.keep)
        if a.report:
            Path(a.report).write_text(json.dumps(rep, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return exit_code(rep)
    if a.cmd == "manifest":
        return run_manifests(Path(a.path), a.tol_abs, a.tol_rel, a.keep)
    try:
        if a.cmd == "value":
            v = read_value(load_json(a.file), a.path)
            print(json.dumps(v, ensure_ascii=False, indent=1) if isinstance(v, (dict, list)) else repr(v))
            return 0
        va = read_value(load_json(a.a), a.path_a)
        vb = read_value(load_json(a.b), a.path_b)
    except (KeyError, OSError, ValueError) as e:
        print(f"MISSING · {e}")
        return 2
    ok = numbers_equal(float(va), float(vb), a.tol_abs, a.tol_rel) if is_number(va) and is_number(vb) else va == vb
    dif = f" · diff {vb - va:+,.6g}" if is_number(va) and is_number(vb) else ""
    print(f"{'MATCHES' if ok else 'DOES NOT MATCH'} · {Path(a.a).name}:{a.path_a} = {va!r} · "
          f"{Path(a.b).name}:{a.path_b} = {vb!r}{dif}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
