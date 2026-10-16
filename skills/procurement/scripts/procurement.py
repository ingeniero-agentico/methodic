# -*- coding: utf-8 -*-
"""
procurement · compare supplier quotes and allocate the purchase in cascade.

    python procurement.py <input.json> --out <PROCUREMENT.xlsx>
                          [--profile <profile.yaml|name>] [--company <name>] [--json <result.json>]
    python procurement.py --test

What it computes (the judgement is set by whoever fills the input; the script only computes):

1. Checks every quote against its printed subtotal. If the lines do not add up to the printed
   figure, it aborts: a line is missing, one is extra, or a total was copied wrong.
2. Derives the effective unit price: line_total / quantity. The printed unit price is a
   reference, never data (some suppliers print list prices and discount only in the total).
3. Converts to the profile currency with the exchange rate the quote declares.
4. Cascade allocation: the cheapest supplier sells what it has, the balance goes to the next one,
   and whatever nobody covers stays TO QUOTE. Price ties are settled by the profile
   (`purchasing.tie_break`).
5. Buys in the supplier's unit: if it sells 6 m lengths, whole lengths are bought and the
   leftover is declared and paid for (it stays in the cost).

Profile search order: --profile · METHODIC_DATA, then the legacy INGENIERO_DATOS (profiles/, then
perfiles/) · sibling folder ../methodic-data/, or the legacy ../ingeniero-agentico-datos/ when the new
one does not exist (profiles/, then perfiles/) · the repo's profiles/_template.yaml, then
perfiles/_plantilla.yaml.

Spanish input keys, Spanish profile keys and the old Spanish flags are still accepted.
Output labels (Excel, console, the SUPUESTO / ASSUMPTION marker) follow the profile `language`
(es by default).

Standard library + openpyxl only.
"""
import copy
import json
import math
import os
import shutil
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

EPS = 1e-9
HERE = Path(__file__).resolve()
SKILL = HERE.parents[1]
REPO = HERE.parents[3] if len(HERE.parents) > 3 else SKILL

LANG = "es"   # output language; set from the profile


def set_language(lang):
    global LANG
    LANG = "en" if str(lang or "es").strip().lower() == "en" else "es"


def tr(es, en):
    """Pick the text for the output language."""
    return en if LANG == "en" else es


def mark():
    return tr("SUPUESTO", "ASSUMPTION")


class InputError(Exception):
    """An error that prevents computing. Fix the input and run again."""


# ---------------------------------------------------------------- key aliases
# Spanish keys are accepted at load time and renamed to the English canonical ones.
INPUT_ALIASES = {
    "titulo": "title", "requerimiento": "requirement", "cotizaciones": "quotes",
    "cantidad": "quantity", "unidad": "unit", "proveedor": "supplier", "identificador": "tax_id",
    "vendedor": "salesperson", "fecha": "date", "vence": "expires", "condiciones": "terms",
    "plazo_dias": "lead_time_days", "moneda": "currency", "tipo_cambio": "exchange_rate",
    "fecha_cambio": "exchange_rate_date", "subtotal_impreso": "printed_subtotal", "archivo": "file",
    "lineas": "lines", "linea": "line", "descripcion": "description",
    "unitario_impreso": "printed_unit_price", "descuento": "discount", "total_linea": "line_total",
    "disponible": "available", "unidad_venta": "sale_unit", "nombre": "name", "contenido": "content",
    "largo": "length", "textual": "verbatim", "descarte": "rejected", "nota": "note",
    "supuestos": "assumptions", "parametro": "parameter", "valor": "value", "fuente": "source",
    "texto": "text", "_ayuda": "_help",
}
PROFILE_ALIASES = {
    "idioma": "language", "empresa": "company", "nombre": "name", "rut": "tax_id", "pais": "country",
    "moneda": "currency", "decimales_moneda": "currency_decimals", "iva": "vat",
    "impuesto_renta": "income_tax", "compras": "purchasing", "proveedores": "suppliers",
    "plazo_entrega_dias": "lead_time_days", "desempate": "tie_break", "precios": "pricing",
    "recargo_por_defecto": "default_markup", "gastos_generales": "overhead", "finanzas": "finance",
    "tasa_descuento": "discount_rate", "horizonte_anios": "horizon_years", "depreciacion": "depreciation",
}
PROFILE_VALUES = {
    "mayor_disponibilidad": "highest_availability", "menor_plazo": "shortest_lead_time",
    "lineal": "straight_line", "acelerada": "accelerated",
}
# Spanish name of each field, to speak the user's language in messages.
_ES_FIELD = {}
for _es, _en in INPUT_ALIASES.items():
    _ES_FIELD.setdefault(_en, _es)
_ES_FIELD.update({"content": "contenido", "currency_decimals": "decimales_moneda", "vat": "iva",
                  "tie_break": "desempate"})
ES_TIE = {"highest_availability": "mayor_disponibilidad", "shortest_lead_time": "menor_plazo"}


def F(field):
    """Field name in the output language (English key, or its Spanish alias in es)."""
    return _ES_FIELD.get(field, field) if LANG == "es" else field


def rename_keys(obj, aliases):
    """Rename Spanish keys to English, recursively. When both are present, the English one wins."""
    if isinstance(obj, list):
        return [rename_keys(x, aliases) for x in obj]
    if not isinstance(obj, dict):
        return obj
    out = {}
    for k, v in obj.items():
        if k not in aliases:
            out[k] = rename_keys(v, aliases)
    for k, v in obj.items():
        if k in aliases and aliases[k] not in out:
            out[aliases[k]] = rename_keys(v, aliases)
    return out


def normalize_input(raw):
    return rename_keys(raw, INPUT_ALIASES)


def normalize_profile(raw):
    prof = rename_keys(raw or {}, PROFILE_ALIASES)
    pur = prof.get("purchasing")
    if isinstance(pur, dict) and isinstance(pur.get("tie_break"), str):
        pur["tie_break"] = PROFILE_VALUES.get(pur["tie_break"], pur["tie_break"])
    fin = prof.get("finance")
    if isinstance(fin, dict) and isinstance(fin.get("depreciation"), str):
        fin["depreciation"] = PROFILE_VALUES.get(fin["depreciation"], fin["depreciation"])
    return prof


# ---------------------------------------------------------------- minimal YAML
def _scalar(s):
    s = s.strip()
    if s in ("", "~", "null", "Null", "NULL"):
        return None
    if s in ("true", "True"):
        return True
    if s in ("false", "False"):
        return False
    if s == "[]":
        return []
    if s == "{}":
        return {}
    if len(s) >= 2 and s[0] == s[-1] and s[0] in "\"'":
        return s[1:-1]
    if s.startswith("[") and s.endswith("]"):
        return [_scalar(x) for x in s[1:-1].split(",") if x.strip()]
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return s


def _strip_comment(line):
    inside, q = False, ""
    for i, ch in enumerate(line):
        if ch in "\"'":
            if not inside:
                inside, q = True, ch
            elif ch == q:
                inside = False
        elif ch == "#" and not inside and (i == 0 or line[i - 1] in " \t"):
            return line[:i].rstrip()
    return line.rstrip()


def read_yaml(text):
    """The YAML subset profiles use: nested maps, lists and scalars."""
    rows = []
    for raw in text.splitlines():
        ln = _strip_comment(raw)
        if ln.strip():
            rows.append((len(ln) - len(ln.lstrip(" ")), ln.strip()))

    def block(i, ind):
        if i >= len(rows):
            return None, i
        if rows[i][1].startswith("- ") or rows[i][1] == "-":
            items = []
            while i < len(rows) and rows[i][0] == ind and rows[i][1].startswith("-"):
                rest = rows[i][1][1:].strip()
                if not rest:
                    val, i = block(i + 1, rows[i + 1][0]) if i + 1 < len(rows) else (None, i + 1)
                    items.append(val)
                    continue
                if ":" in rest and not rest.startswith(("\"", "'")):
                    # map item: «- key: value» and its continuation lines further in
                    sub_ind = ind + 2
                    rows[i] = (sub_ind, rest)
                    val, i = block(i, sub_ind)
                    items.append(val)
                else:
                    items.append(_scalar(rest))
                    i += 1
            return items, i
        mapping = {}
        while i < len(rows) and rows[i][0] == ind and not rows[i][1].startswith("- "):
            key, _, value = rows[i][1].partition(":")
            key = key.strip().strip("\"'")
            if value.strip():
                mapping[key] = _scalar(value)
                i += 1
            elif i + 1 < len(rows) and rows[i + 1][0] > ind:
                mapping[key], i = block(i + 1, rows[i + 1][0])
            elif i + 1 < len(rows) and rows[i + 1][0] == ind and rows[i + 1][1].startswith("- "):
                mapping[key], i = block(i + 1, ind)
            else:
                mapping[key] = None
                i += 1
        return mapping, i

    data, _ = block(0, rows[0][0]) if rows else ({}, 0)
    return data or {}


# ---------------------------------------------------------------- profile
DEFAULTS = {"currency": "CLP", "currency_decimals": 0, "vat": 0.19, "tie_break": "highest_availability"}
TIE_BREAKS = ("highest_availability", "shortest_lead_time")


def show_tie(v):
    return ES_TIE.get(v, v) if LANG == "es" else v


def template_path(root=None):
    """The repo's profile template: profiles/_template.yaml, then perfiles/_plantilla.yaml."""
    root = Path(root) if root else REPO
    for p in (root / "profiles" / "_template.yaml", root / "perfiles" / "_plantilla.yaml"):
        if p.is_file():
            return p
    return None


def _is_template(p):
    return p.name.startswith(("_template", "_plantilla"))


ENV_DATA = ("METHODIC_DATA", "INGENIERO_DATOS")              # the second one is the legacy name
SIBLING_DATA = ("methodic-data", "ingeniero-agentico-datos")  # the second one is the legacy name


def env_data():
    """Values of $METHODIC_DATA and then the legacy $INGENIERO_DATOS, the ones that are set."""
    return [os.environ[v] for v in ENV_DATA if os.environ.get(v)]


def sibling_data():
    """../methodic-data, or the legacy ../ingeniero-agentico-datos when the new one does not exist."""
    new, old = (REPO.parent / n for n in SIBLING_DATA)
    return old if not new.is_dir() and old.is_dir() else new


def data_dirs():
    """Folders where company profiles live, in search order."""
    dirs = []
    for env in env_data():
        if not Path(env).is_file():
            dirs += [Path(env) / "profiles", Path(env) / "perfiles", Path(env)]
    sibling = sibling_data()
    dirs += [sibling / "profiles", sibling / "perfiles"]
    return dirs


def _pick(dirs, company):
    seen = []
    for d in dirs:
        if not d.is_dir():
            continue
        yamls = sorted(p for p in d.glob("*.yaml") if not p.name.startswith("_"))
        if not yamls:
            continue
        if company:
            for p in yamls:
                if p.stem.lower() == company.lower():
                    return p
            seen += [p.stem for p in yamls]
            continue
        if len(yamls) == 1:
            return yamls[0]
        raise InputError(tr("Hay varios perfiles; elige uno con --company: ",
                            "There are several profiles; pick one with --company: ")
                         + ", ".join(p.stem for p in yamls))
    if company:
        raise InputError(tr(f"No hay perfil «{company}». Hay: {', '.join(seen) or 'ninguno'}",
                            f"No profile «{company}». Found: {', '.join(seen) or 'none'}"))
    return None


def find_profile(arg=None, company=None):
    """Return (path, is_template). Order: --profile · METHODIC_DATA · INGENIERO_DATOS · sibling · template."""
    if arg:
        p = Path(arg)
        if p.is_file():
            return p, _is_template(p)
        name = arg if arg.endswith((".yaml", ".yml")) else arg + ".yaml"
        for d in data_dirs():
            if (d / name).is_file():
                return d / name, False
        raise InputError(tr(f"--profile no existe: {arg}", f"--profile not found: {arg}"))
    env = (env_data() or [None])[0]  # the first variable set may point straight at a profile file
    if env and Path(env).is_file():
        return Path(env), False
    found = _pick(data_dirs(), company)
    if found:
        return found, False
    tpl = template_path()
    return tpl, True


def assumption(parameter, value, origin, body):
    """One assumption: what, which value, where it came from, and the line shown to people."""
    return {"parameter": parameter, "value": value, "origin": origin, "text": f"{mark()} · {body}"}


def load_profile(path=None, company=None, language=None):
    p, is_template = find_profile(path, company)
    raw = normalize_profile(read_yaml(p.read_text(encoding="utf-8-sig"))) if p else {}
    set_language(language or raw.get("language") or "es")
    comp = raw.get("company") or {}
    pur = raw.get("purchasing") or {}
    profile = {"source": p.name if p else tr("(sin archivo: valores por defecto)", "(no file: built-in defaults)"),
               "is_template": is_template or p is None, "language": LANG, "assumptions": [],
               "suppliers": pur.get("suppliers") or []}
    fallback = []
    for field, src in (("currency", comp), ("currency_decimals", comp), ("vat", comp), ("tie_break", pur)):
        v = src.get(field)
        if v in (None, ""):
            if field in ("currency", "vat") and not profile["is_template"]:
                raise InputError(tr(f"El perfil {profile['source']} no trae `{F(field)}`. "
                                    f"Se pregunta a la empresa; no se inventa.",
                                    f"Profile {profile['source']} has no `{field}`. "
                                    f"Ask the company; it is not invented."))
            v = DEFAULTS[field]
            profile["assumptions"].append(assumption(
                field, v, "default", tr(f"{F(field)} = {show_tie(v)} (el perfil no lo trae)",
                                        f"{field} = {v} (the profile does not have it)")))
        elif profile["is_template"]:
            fallback.append(assumption(field, v, "template",
                                       tr(f"{F(field)} = {show_tie(v)} (de la plantilla del perfil)",
                                          f"{field} = {v} (from the profile template)")))
        profile[field] = v
    if profile["is_template"]:
        profile["assumptions"][0:0] = [assumption(
            "profile", profile["source"], "template",
            tr("se usó la plantilla del perfil, no el de la empresa: moneda, IVA y desempate son los de ejemplo",
               "the profile template was used, not the company's: currency, VAT and tie-break are sample "
               "values"))] + fallback
    profile["currency"] = str(profile["currency"]).upper()
    profile["currency_decimals"] = int(profile["currency_decimals"])
    profile["vat"] = float(profile["vat"])
    if profile["tie_break"] not in TIE_BREAKS:
        raise InputError(tr(f"desempate «{profile['tie_break']}» no existe. Usa: {', '.join(ES_TIE.values())}",
                            f"tie_break «{profile['tie_break']}» does not exist. Use: {', '.join(TIE_BREAKS)}"))
    return profile


# ---------------------------------------------------------------- helpers
def num(v, field, where):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        raise InputError(tr(f"{where}: `{F(field)}` no es un número", f"{where}: `{field}` is not a number"))
    if isinstance(v, (int, float)):
        return float(v)
    raise InputError(tr(f"{where}: `{F(field)}` debe ser número (vino «{v}»). "
                        f"Escríbelo sin separador de miles: 1234567.89",
                        f"{where}: `{field}` must be a number (got «{v}»). "
                        f"Write it without thousands separators: 1234567.89"))


def fmt(x, dec=0):
    """Number for text notes: 1.234.567 · 1.234,56 in es; 1,234,567 · 1,234.56 in en."""
    if x is None:
        return "—"
    s = f"{x:,.{dec}f}"
    if LANG == "en":
        return s
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def g(x):
    """Quantity without trailing zeros."""
    if abs(x - round(x)) < 1e-6:
        return fmt(x, 0)
    return fmt(x, 2).rstrip("0").rstrip("." if LANG == "en" else ",")


def declared_assumptions(inp):
    """Assumptions the user wrote in the input (`assumptions`, alias `supuestos`)."""
    out = []
    items = inp.get("assumptions") or []
    if not isinstance(items, list):
        raise InputError(tr("`supuestos` debe ser una lista", "`assumptions` must be a list"))
    for a in items:
        if isinstance(a, str):
            out.append({"parameter": "", "value": None, "origin": "input", "text": f"{mark()} · {a}"})
        elif isinstance(a, dict):
            par, val = a.get("parameter") or "", a.get("value")
            src = a.get("source") or a.get("note") or a.get("text") or ""
            head = f"{par} = {val}" if par else ("" if val is None else str(val))
            body = " · ".join(x for x in (head, str(src)) if x)
            e = {"parameter": par, "value": val, "origin": "input", "text": f"{mark()} · {body}"}
            if src:
                e["source"] = src
            out.append(e)
        else:
            raise InputError(tr(f"`supuestos`: no se entiende «{a}»", f"`assumptions`: cannot read «{a}»"))
    return out


# ---------------------------------------------------------------- prepare
def prepare(inp, profile):
    """Validate the input and turn it into a list of comparable offers."""
    errors, warnings, assumptions = [], [], []
    req = inp.get("requirement") or []
    quotes = inp.get("quotes") or []
    if not req:
        errors.append(tr(f"La entrada no trae `{F('requirement')}`.", "The input has no `requirement`."))
    if not quotes:
        errors.append(tr(f"La entrada no trae `{F('quotes')}`.", "The input has no `quotes`."))
    ns = [r.get("n") for r in req]
    for n in sorted({n for n in ns if ns.count(n) > 1}, key=str):
        errors.append(tr(f"N° de ítem repetido en el requerimiento: {n}", f"Item number repeated in the requirement: {n}"))
    for r in req:
        if num(r.get("quantity"), "quantity", tr(f"ítem {r.get('n')}", f"item {r.get('n')}")) in (None, 0.0):
            errors.append(tr(f"Ítem {r.get('n')}: sin cantidad requerida", f"Item {r.get('n')}: no required quantity"))
    valid = set(ns)

    ids = [c.get("id") for c in quotes]
    for i in sorted({i for i in ids if ids.count(i) > 1}, key=str):
        errors.append(tr(f"Cotización con id repetido: {i}", f"Quote id repeated: {i}"))

    base = profile["currency"]
    offers, reconciliation, printed, no_item = [], [], [], []
    no_avail = 0
    for c in quotes:
        cid, sup = c.get("id"), c.get("supplier")
        where = tr(f"Cotización {cid}", f"Quote {cid}")
        if not cid or not sup:
            errors.append(tr(f"{where}: falta `id` o `{F('supplier')}`", f"{where}: missing `id` or `supplier`"))
            continue
        currency = (c.get("currency") or "").upper()
        if not currency:
            currency = base
            assumptions.append(assumption("currency", base, "script",
                                          tr(f"{cid}: no declara moneda; se asume {base}",
                                             f"{cid}: no currency declared; {base} is assumed")))
        rate = 1.0
        if currency != base:
            rate = num(c.get("exchange_rate"), "exchange_rate", where)
            if not rate:
                errors.append(tr(f"{where}: cotiza en {currency} y no declara `{F('exchange_rate')}`. El cambio "
                                 f"lo pone quien paga (el que cobra su banco), no una web.",
                                 f"{where}: quotes in {currency} and declares no `exchange_rate`. The rate is set "
                                 f"by whoever pays (what their bank charges), not a website."))
                rate = None
            elif not c.get("exchange_rate_date"):
                warnings.append(tr(f"{where}: tipo de cambio sin `{F('exchange_rate_date')}`; no se podrá reconstruir.",
                                   f"{where}: exchange rate without `exchange_rate_date`; it cannot be rebuilt."))
        if not c.get("tax_id"):
            warnings.append(tr(f"{where} ({sup}): sin identificador tributario. Sin él no se emite la orden.",
                               f"{where} ({sup}): no tax ID. Without it the purchase order is not issued."))
        dec_q = 2 if currency != base else profile["currency_decimals"]
        total_sum, seen, desc_seen = 0.0, set(), defaultdict(list)
        for k, ln in enumerate(c.get("lines") or [], 1):
            lid = str(ln.get("line", k))
            wl = tr(f"{where} línea {lid}", f"{where} line {lid}")
            if lid in seen:
                errors.append(tr(f"{wl}: número de línea repetido en la misma cotización",
                                 f"{wl}: line number repeated in the same quote"))
            seen.add(lid)
            qty = num(ln.get("quantity"), "quantity", wl)
            tot = num(ln.get("line_total"), "line_total", wl)
            prt = num(ln.get("printed_unit_price"), "printed_unit_price", wl)
            if not qty or qty <= 0:
                errors.append(tr(f"{wl}: sin cantidad. Una línea sin cantidad no se compara ni se compra.",
                                 f"{wl}: no quantity. A line without quantity is neither compared nor bought."))
                continue
            tot_assumed = False
            if tot is None:
                if prt is None:
                    errors.append(tr(f"{wl}: sin {F('line_total')} ni {F('printed_unit_price')}",
                                     f"{wl}: neither line_total nor printed_unit_price"))
                    continue
                tot, tot_assumed = prt * qty, True
                assumptions.append(assumption("line_total", tot, "script",
                                              tr(f"{wl}: no imprime total; se usó unitario impreso × cantidad",
                                                 f"{wl}: no printed total; printed unit price × quantity was used")))
            total_sum += tot
            desc_seen[(ln.get("description") or "").strip().upper()].append(lid)
            u_eff = tot / qty
            if prt and abs(prt - u_eff) > max(abs(prt) * 0.005, 0.01):
                printed.append({"quote": cid, "supplier": sup, "line": lid, "printed": prt,
                                "effective": u_eff, "currency": currency, "diff": u_eff / prt - 1})
            item = ln.get("item")
            if item in (None, ""):
                no_item.append({"quote": cid, "supplier": sup, "line": lid,
                                "description": ln.get("description") or "", "total": tot, "currency": currency})
                continue
            if item not in valid:
                errors.append(tr(f"{wl} apunta al ítem {item}, que no existe en el requerimiento",
                                 f"{wl} points to item {item}, which is not in the requirement"))
                continue
            su = ln.get("sale_unit") or None
            content = 1.0
            if su:
                content = num(su.get("content", su.get("length")), "content", wl)
                if not content or content <= 0:
                    errors.append(tr(f"{wl}: `{F('sale_unit')}` sin `{F('content')}` (cuánto del ítem trae "
                                     f"cada unidad que vende el proveedor)",
                                     f"{wl}: `sale_unit` without `content` (how much of the item each unit "
                                     f"the supplier sells holds)"))
                    continue
                if not su.get("verbatim"):
                    warnings.append(tr(f"{wl}: `{F('sale_unit')}` sin `{F('verbatim')}`; es la prueba de dónde salió "
                                       f"el {g(content)}.",
                                       f"{wl}: `sale_unit` without `verbatim`; it is the proof of where the "
                                       f"{g(content)} came from."))
            avail = num(ln.get("available"), "available", wl)
            avail_assumed = False
            if avail is None:
                avail, avail_assumed = qty, True
                no_avail += 1
            q_req = round(qty * content, 6)
            offers.append({
                "item": item, "supplier": sup, "salesperson": (c.get("salesperson") or "").strip(),
                "quote": cid, "line": lid, "date": c.get("date") or "",
                "lead_time_days": c.get("lead_time_days"), "tax_id": c.get("tax_id") or "",
                "currency": currency, "exchange_rate": rate, "quantity": qty, "content": content,
                "sale_unit": su, "q_req": q_req, "avail_units": avail, "avail_req": round(avail * content, 6),
                "avail_assumed": avail_assumed, "total_src": tot, "total_assumed": tot_assumed,
                "u_src": tot / q_req, "u": (tot * rate / q_req) if rate else None,
                "rejected": (ln.get("rejected") or "").strip(), "note": ln.get("note") or "",
                "description": ln.get("description") or "",
            })
        for d, lids in desc_seen.items():
            if d and len(lids) > 1:
                warnings.append(tr(f"{where}: «{d[:40]}» aparece en las líneas {', '.join(lids)}. No es error "
                                   f"si venía en dos partidas: van las dos.",
                                   f"{where}: «{d[:40]}» appears in lines {', '.join(lids)}. Not an error "
                                   f"if it came in two batches: both go."))
        sub = num(c.get("printed_subtotal"), "printed_subtotal", where)
        tol = max(2 * 10 ** -dec_q, 0.05) if dec_q else 2.0
        status = "no_subtotal"
        if sub is None:
            warnings.append(tr(f"{where}: sin `{F('printed_subtotal')}`; la extracción no se pudo verificar completa.",
                               f"{where}: no `printed_subtotal`; the extraction could not be verified complete."))
        elif abs(total_sum - sub) > tol:
            errors.append(tr(f"{where}: NO CUADRA. Las líneas suman {fmt(total_sum, dec_q)} y el documento "
                             f"imprime {fmt(sub, dec_q)} (diferencia {fmt(total_sum - sub, dec_q)}). "
                             f"Falta una línea, sobra una o un total está mal copiado.",
                             f"{where}: MISMATCH. The lines add up to {fmt(total_sum, dec_q)} and the document "
                             f"prints {fmt(sub, dec_q)} (difference {fmt(total_sum - sub, dec_q)}). "
                             f"A line is missing, one is extra, or a total was copied wrong."))
            status = "mismatch"
        else:
            status = "matches"
        reconciliation.append({"quote": cid, "supplier": sup, "currency": currency,
                               "exchange_rate": rate if currency != base else None,
                               "lines": len(c.get("lines") or []), "sum": round(total_sum, 2),
                               "printed_subtotal": sub, "status": status,
                               "expires": c.get("expires") or "", "terms": c.get("terms") or "",
                               "lead_time_days": c.get("lead_time_days"), "file": c.get("file") or "",
                               "tax_id": c.get("tax_id") or "",
                               "salesperson": (c.get("salesperson") or "").strip(), "date": c.get("date") or ""})
    if no_avail:
        assumptions.append(assumption("available", no_avail, "script",
                                      tr(f"{no_avail} línea(s) sin `{F('available')}`: se asume que el proveedor "
                                         f"entrega todo lo que cotizó",
                                         f"{no_avail} line(s) without `available`: the supplier is assumed to "
                                         f"deliver everything it quoted")))

    # Counteroffer: same item + same supplier + same salesperson in TWO quotes.
    # The newest replaces the old one; they do not compete with each other.
    replacements = []
    groups = defaultdict(set)
    for o in offers:
        groups[(o["item"], o["supplier"], o["salesperson"])].add((o["date"], o["quote"]))
    out = set()
    for (item, sup, sp), qs in groups.items():
        if len(qs) < 2:
            continue
        dates = [d for d, _ in qs]
        if any(not d for d in dates) or len(set(dates)) != len(dates):
            errors.append(tr(f"Ítem {item} · {sup}{' · ' + sp if sp else ''}: aparece en "
                             f"{len(qs)} cotizaciones sin fecha que las ordene. Si es contraoferta, "
                             f"pon `{F('date')}`; si son vendedores distintos, pon `{F('salesperson')}`.",
                             f"Item {item} · {sup}{' · ' + sp if sp else ''}: appears in "
                             f"{len(qs)} quotes with no date to order them. If it is a counteroffer, "
                             f"set `date`; if they are different salespeople, set `salesperson`."))
            continue
        newest = max(qs)[1]
        for d, cid in qs:
            if cid != newest:
                out.add((item, cid))
        before = [o["u"] for o in offers if o["item"] == item and (item, o["quote"]) in out
                  and o["supplier"] == sup and o["u"] is not None]
        now = [o["u"] for o in offers if o["item"] == item and o["quote"] == newest and o["u"] is not None]
        replacements.append({"item": item, "supplier": sup, "salesperson": sp, "valid": newest,
                             "before": min(before) if before else None, "after": min(now) if now else None})
    for o in offers:
        o["replaced"] = (o["item"], o["quote"]) in out

    used = {o["supplier"] for o in offers}
    for c in quotes:
        if c.get("supplier") and c["supplier"] not in used:
            warnings.append(tr(f"{c['supplier']} ({c.get('id')}): ninguna línea quedó asignada a un ítem. "
                               f"¿Falta emparejar su cotización?",
                               f"{c['supplier']} ({c.get('id')}): no line was matched to an item. "
                               f"Is its quote still unmatched?"))
    covered = {o["item"] for o in offers if not o["rejected"] and not o["replaced"]}
    for r in req:
        if r.get("n") not in covered:
            warnings.append(tr(f"Ítem {r.get('n')} «{r.get('item')}»: nadie lo cubre",
                               f"Item {r.get('n')} «{r.get('item')}»: nobody covers it"))
    if errors:
        raise InputError("\n".join(errors))
    return offers, reconciliation, printed, replacements, warnings, assumptions, no_item


# ---------------------------------------------------------------- allocation
def lead_time_of(o, profile):
    if o.get("lead_time_days") not in (None, ""):
        return float(o["lead_time_days"])
    for p in profile.get("suppliers") or []:
        if isinstance(p, dict) and o["tax_id"] and str(p.get("tax_id")) == o["tax_id"]:
            if p.get("lead_time_days") not in (None, ""):
                return float(p["lead_time_days"])
    return math.inf


def sale_unit_label(o):
    su = o["sale_unit"]
    if not su:
        return ""
    name = su.get("name") or tr("unidad", "unit")
    return tr(f"{name} de {g(o['content'])} {su.get('unit', '')}",
              f"{name} of {g(o['content'])} {su.get('unit', '')}").strip()


def cascade(requirement, offers, profile):
    alive = [o for o in offers if not o["rejected"] and not o["replaced"]]
    by_item = defaultdict(list)
    for o in alive:
        by_item[o["item"]].append(o)
    rule = profile["tie_break"]

    def key(o):
        if rule == "shortest_lead_time":
            return (o["u"], lead_time_of(o, profile), -o["avail_req"])
        return (o["u"], -o["avail_req"])

    gap_label = tr("POR COTIZAR", "TO QUOTE")
    plan = []
    for r in requirement:
        n, missing = r["n"], float(r["quantity"])
        cands = sorted(by_item.get(n, []), key=key)   # stable sort: equal keys keep input order
        allocs, order = [], 0
        for o in cands:
            if missing <= EPS:
                break
            if o["avail_req"] <= EPS:
                continue
            take, leftover, units = min(missing, o["avail_req"]), 0.0, None
            notes = []
            if o["sale_unit"]:
                # Buy in the supplier's unit: whole lengths, bars, boxes.
                units = math.ceil(round(take / o["content"], 6) - EPS)
                cap = math.floor(round(o["avail_units"], 6) + EPS)
                if units > cap:
                    units = cap
                if units <= 0:
                    continue
                take_whole = round(units * o["content"], 6)
                leftover = max(0.0, round(take_whole - take, 6))
                take = take_whole
                if leftover > EPS:
                    notes.append(tr(f"sobran {g(leftover)} {r.get('unit', '')}".rstrip()
                                    + " por comprar unidades enteras (se pagan)",
                                    f"{g(leftover)} {r.get('unit', '')}".rstrip()
                                    + " left over from buying whole units (paid for)"))
            else:
                units = take / o["content"]
            tied = [x for x in cands if x is not o and abs(x["u"] - o["u"]) < 1e-9]
            if tied:
                others = ", ".join(f"{x['supplier']} ({x['quote']})" for x in tied)
                notes.append(tr(f"empate de precio con {others}; desempate: {show_tie(rule).replace('_', ' ')}",
                                f"price tie with {others}; tie-break: {rule.replace('_', ' ')}"))
            if o["avail_req"] + EPS < float(r["quantity"]):
                notes.append(tr(f"disponible {g(o['avail_req'])} de {g(float(r['quantity']))}",
                                f"available {g(o['avail_req'])} of {g(float(r['quantity']))}"))
            if o["avail_assumed"]:
                notes.append(tr("disponible SUPUESTO = lo cotizado", "available ASSUMPTION = quoted"))
            if o["note"]:
                notes.append(o["note"])
            order += 1
            subtotal = round(take * o["u"], 2)
            foreign = o["currency"] != profile["currency"]
            trace = (tr(f"{o['quote']} línea {o['line']}: ", f"{o['quote']} line {o['line']}: ")
                     + f"{fmt(o['total_src'], 2 if foreign else profile['currency_decimals'])} {o['currency']}"
                     + f" ÷ {g(o['q_req'])}"
                     + (f" ({g(o['quantity'])} × {g(o['content'])})" if o["sale_unit"] else "")
                     + (f" × {fmt(o['exchange_rate'], 2)}" if foreign else "")
                     + tr(f" = {fmt(o['u'], 2)} por {r.get('unit') or 'unidad'}",
                          f" = {fmt(o['u'], 2)} per {r.get('unit') or 'unit'}"))
            allocs.append({
                "order": order, "gap": False, "supplier": o["supplier"], "salesperson": o["salesperson"],
                "quote": o["quote"], "line": o["line"],
                "units": round(units, 6), "sale_unit": sale_unit_label(o),
                "quantity": round(take, 6), "leftover": leftover,
                "unit_price": round(o["u"], 2), "subtotal": subtotal,
                "source_currency": o["currency"],
                "source_unit_price": round(o["u_src"] * (o["content"] if o["sale_unit"] else 1), 4),
                "source_subtotal": round(take * o["u_src"], 2) if foreign else None,
                "trace": trace, "note": " · ".join(notes),
            })
            missing = max(0.0, round(missing - take, 6))
        reason = ""
        if missing > EPS:
            rejected = [o for o in offers if o["item"] == n and o["rejected"]]
            if cands:
                reason = tr("nadie cubre el saldo", "nobody covers the balance")
            elif rejected:
                reason = tr(f"solo hubo ofertas descartadas ({len(rejected)})",
                            f"only rejected offers ({len(rejected)})")
            else:
                reason = tr("nadie lo cotizó", "nobody quoted it")
            allocs.append({"order": order + 1, "gap": True, "supplier": gap_label, "quote": "",
                           "line": "", "quantity": missing, "unit_price": None, "subtotal": None,
                           "trace": "", "note": reason})
        plan.append({"req": r, "allocations": allocs, "missing": missing, "reason": reason,
                     "candidates": cands})
    return plan


def summarize(plan, profile):
    net = 0.0
    by_sup = {}
    for p in plan:
        for a in p["allocations"]:
            if a["subtotal"] is None:
                continue
            net = round(net + a["subtotal"], 2)
            k = a["supplier"] + (f" · {a['salesperson']}" if a.get("salesperson") else "")
            e = by_sup.setdefault(k, {"supplier": k, "quotes": [], "lines": 0, "net": 0.0})
            e["lines"] += 1
            e["net"] = round(e["net"] + a["subtotal"], 2)
            if a["quote"] not in e["quotes"]:
                e["quotes"].append(a["quote"])
    vat = round(net * profile["vat"], 2)
    complete = sum(1 for p in plan if p["missing"] <= EPS)
    return {"net": net, "vat": vat, "total": round(net + vat, 2), "items": len(plan),
            "complete": complete, "incomplete": len(plan) - complete}, \
        sorted(by_sup.values(), key=lambda e: -e["net"])


def calculate(inp, profile):
    set_language(profile.get("language"))
    inp = normalize_input(inp)
    declared = declared_assumptions(inp)
    offers, reconciliation, printed, replacements, warnings, assumptions, no_item = prepare(inp, profile)
    plan = cascade(inp["requirement"], offers, profile)
    totals, by_sup = summarize(plan, profile)
    return {
        "title": inp.get("title", ""),
        "profile": {k: profile[k] for k in ("source", "language", "currency", "currency_decimals", "vat",
                                            "tie_break")},
        "totals": totals,
        "by_supplier": by_sup,
        "items": [{"n": p["req"]["n"], "item": p["req"].get("item", ""),
                   "quantity": float(p["req"]["quantity"]), "unit": p["req"].get("unit", ""),
                   "missing": p["missing"], "allocations": p["allocations"]} for p in plan],
        "missing": [{"n": p["req"]["n"], "item": p["req"].get("item", ""),
                     "quantity": p["missing"], "reason": p["reason"]}
                    for p in plan if p["missing"] > EPS],
        "rejected": [{"item": o["item"], "supplier": o["supplier"], "quote": o["quote"],
                      "line": o["line"], "reason": o["rejected"]} for o in offers if o["rejected"]],
        "counteroffers": replacements,
        "lines_without_item": no_item,
        "printed_vs_effective": printed,
        "reconciliation": reconciliation,
        "assumptions": profile["assumptions"] + declared + assumptions,
        "warnings": warnings,
        "_offers": offers, "_plan": plan,
    }


# ---------------------------------------------------------------- Excel
SHEETS = {"es": ("RESUMEN", "ASIGNACION", "COMPARATIVA", "PROVEEDORES", "VERIFICACION"),
          "en": ("SUMMARY", "ALLOCATION", "COMPARISON", "SUPPLIERS", "VERIFICATION")}
STATUS = {"matches": ("CUADRA", "MATCHES"), "mismatch": ("NO CUADRA", "MISMATCH"),
          "no_subtotal": ("SIN ANCLA", "NO ANCHOR")}


def write_excel(res, dest):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    set_language(res["profile"]["language"])
    S_SUM, S_ALLOC, S_COMP, S_SUP, S_VER = SHEETS[LANG]
    dec = res["profile"]["currency_decimals"]
    MONEY = "#,##0" + ("." + "0" * dec if dec else "")
    MONEY2 = "#,##0.00"
    BLUE, GREY, GREEN, RED, YELLOW = "1F3A5F", "E7E6E6", "E2F0D9", "C00000", "FFF2CC"
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_title, f_head, f_bold = Font(bold=True, size=13, color="FFFFFF"), Font(bold=True, color="FFFFFF"), Font(bold=True)
    fill = lambda c: PatternFill("solid", fgColor=c)  # noqa: E731
    status_label = lambda s: STATUS[s][1 if LANG == "en" else 0]  # noqa: E731

    wb = openpyxl.Workbook()

    def sheet(name, title, width):
        ws = wb.create_sheet(name)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        c = ws.cell(row=1, column=1, value=title)
        c.font, c.fill = f_title, fill(BLUE)
        return ws

    def header(ws, row, cols, widths=None):
        for j, t in enumerate(cols, 1):
            c = ws.cell(row=row, column=j, value=t)
            c.font, c.fill, c.border = f_head, fill(BLUE), box
            c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            for j, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    cur = res["profile"]["currency"]
    vat = res["profile"]["vat"]
    t = res["totals"]

    # SUMMARY
    ws = sheet(S_SUM, tr(f"ADQUISICIONES · {res['title']}", f"PROCUREMENT · {res['title']}"), 6)
    ws.cell(row=2, column=1, value=tr(f"Perfil: {res['profile']['source']} · moneda {cur} · IVA "
                                      f"{vat:.0%} · desempate {show_tie(res['profile']['tie_break'])}",
                                      f"Profile: {res['profile']['source']} · currency {cur} · VAT "
                                      f"{vat:.0%} · tie-break {res['profile']['tie_break']}"))
    rows = [(tr("TOTAL NETO", "NET TOTAL"), t["net"]), (tr(f"IVA {vat:.0%}", f"VAT {vat:.0%}"), t["vat"]),
            (tr("TOTAL CON IVA", "TOTAL WITH VAT"), t["total"])]
    for i, (k, v) in enumerate(rows):
        ws.cell(row=4 + i, column=1, value=k).font = f_bold
        c = ws.cell(row=4 + i, column=2, value=v)
        c.number_format = MONEY
    ws.cell(row=7, column=1, value=tr("Ítems cubiertos al 100 %", "Items 100 % covered"))
    ws.cell(row=7, column=2, value=tr(f"{t['complete']} de {t['items']}", f"{t['complete']} of {t['items']}"))
    r = 9
    header(ws, r, tr(["PROVEEDOR", "COTIZACIONES", "LÍNEAS", f"NETO {cur}", "IVA", "% DE LA COMPRA"],
                     ["SUPPLIER", "QUOTES", "LINES", f"NET {cur}", "VAT", "% OF PURCHASE"]),
           [34, 22, 9, 16, 14, 14])
    for e in res["by_supplier"]:
        r += 1
        ws.cell(row=r, column=1, value=e["supplier"])
        ws.cell(row=r, column=2, value=", ".join(e["quotes"]))
        ws.cell(row=r, column=3, value=e["lines"])
        ws.cell(row=r, column=4, value=e["net"]).number_format = MONEY
        ws.cell(row=r, column=5, value=round(e["net"] * vat, 2)).number_format = MONEY
        ws.cell(row=r, column=6, value=e["net"] / t["net"] if t["net"] else None).number_format = "0.0%"
    r += 2
    ws.cell(row=r, column=1, value=tr("LO QUE NADIE CUBRE (POR COTIZAR)", "WHAT NOBODY COVERS (TO QUOTE)")
            ).font = Font(bold=True, color=RED)
    if not res["missing"]:
        ws.cell(row=r + 1, column=1, value=tr("Nada: todo el requerimiento quedó asignado.",
                                              "Nothing: the whole requirement was allocated."))
        r += 1
    for f in res["missing"]:
        r += 1
        ws.cell(row=r, column=1, value=tr(f"Ítem {f['n']} · {f['item']}", f"Item {f['n']} · {f['item']}"))
        ws.cell(row=r, column=2, value=tr(f"falta {g(f['quantity'])}", f"missing {g(f['quantity'])}"))
        ws.cell(row=r, column=3, value=f["reason"])
    r += 2
    ws.cell(row=r, column=1, value=tr("SUPUESTOS (no vinieron del usuario)",
                                      "ASSUMPTIONS (not given by the user)")).font = f_bold
    for s in [a["text"] for a in res["assumptions"]] or [tr("Ninguno.", "None.")]:
        r += 1
        ws.cell(row=r, column=1, value=s)

    # ALLOCATION
    ws = sheet(S_ALLOC, tr("ASIGNACIÓN EN CASCADA · al más barato lo que tenga, el saldo al siguiente",
                           "CASCADE ALLOCATION · the cheapest sells what it has, the balance goes to the next"), 14)
    header(ws, 3, tr(["N°", "ÍTEM", "CANT. REQ.", "UN.", "ORD.", "PROVEEDOR", "COTIZACIÓN", "LÍNEA",
                      "UNIDADES DEL PROVEEDOR", "CANT. ASIGNADA", "SOBRA", f"UNITARIO {cur}",
                      f"SUBTOTAL {cur}", "DE DÓNDE SALE (traza) · NOTA"],
                     ["No.", "ITEM", "REQ. QTY", "UNIT", "ORD.", "SUPPLIER", "QUOTE", "LINE",
                      "SUPPLIER UNITS", "ALLOCATED QTY", "LEFTOVER", f"UNIT PRICE {cur}",
                      f"SUBTOTAL {cur}", "WHERE IT COMES FROM (trace) · NOTE"]),
           [5, 38, 10, 6, 5, 20, 12, 7, 22, 12, 8, 13, 15, 80])
    r = 3
    for it in res["items"]:
        for i, a in enumerate(it["allocations"]):
            r += 1
            if i == 0:
                ws.cell(row=r, column=1, value=it["n"])
                ws.cell(row=r, column=2, value=it["item"])
                ws.cell(row=r, column=3, value=it["quantity"])
                ws.cell(row=r, column=4, value=it["unit"])
            ws.cell(row=r, column=5, value=tr(f"{a['order']}°", f"{a['order']}"))
            ws.cell(row=r, column=6, value=a["supplier"])
            ws.cell(row=r, column=7, value=a["quote"])
            ws.cell(row=r, column=8, value=a["line"])
            if a.get("sale_unit"):
                ws.cell(row=r, column=9, value=f"{g(a['units'])} × {a['sale_unit']}")
            ws.cell(row=r, column=10, value=a["quantity"])
            if a.get("leftover"):
                ws.cell(row=r, column=11, value=a["leftover"])
            if a["unit_price"] is not None:
                ws.cell(row=r, column=12, value=a["unit_price"]).number_format = MONEY2
            if a["subtotal"] is None:
                c = ws.cell(row=r, column=13, value=tr("FALTA", "MISSING"))
                c.font = Font(bold=True, color=RED)
                for j in range(1, 15):
                    ws.cell(row=r, column=j).fill = fill("FBE5E5")
            else:
                ws.cell(row=r, column=13, value=a["subtotal"]).number_format = MONEY
            ws.cell(row=r, column=14, value=" · ".join(x for x in (a["trace"], a["note"]) if x))
            if len(it["allocations"]) > 1 and a["subtotal"] is not None:
                for j in range(5, 14):
                    ws.cell(row=r, column=j).fill = fill(YELLOW)
    r += 1
    ws.cell(row=r, column=2, value=tr("TOTAL NETO", "NET TOTAL")).font = f_bold
    c = ws.cell(row=r, column=13, value=t["net"])
    c.font, c.number_format = f_bold, MONEY
    ws.cell(row=r + 1, column=2, value=tr("Amarillo = ítem partido entre proveedores. Rojo = lo que falta.",
                                          "Yellow = item split between suppliers. Red = what is missing."))

    # COMPARISON
    sups = []
    for o in res["_offers"]:
        k = (o["supplier"], o["quote"])
        if k not in sups:
            sups.append(k)
    width = 3 + len(sups) + 3
    ws = sheet(S_COMP, tr(f"UNITARIO EFECTIVO EN {cur} (total ÷ cantidad; nunca el impreso)",
                          f"EFFECTIVE UNIT PRICE IN {cur} (total ÷ quantity; never the printed one)"), width)
    header(ws, 3, tr(["N°", "ÍTEM", "CANT."], ["No.", "ITEM", "QTY"]) + [f"{p}\n{c}" for p, c in sups]
           + tr(["MEJOR", "PEOR", "DIF. %"], ["BEST", "WORST", "DIFF. %"]),
           [5, 38, 8] + [16] * len(sups) + [13, 13, 9])
    ws.row_dimensions[3].height = 32
    r = 3
    won = {(it["n"], a["quote"]) for it in res["items"] for a in it["allocations"] if a["subtotal"] is not None}
    for it in res["items"]:
        r += 1
        ws.cell(row=r, column=1, value=it["n"])
        ws.cell(row=r, column=2, value=it["item"])
        ws.cell(row=r, column=3, value=it["quantity"])
        vals = []
        for j, (p, cid) in enumerate(sups, 4):
            mine = [o for o in res["_offers"] if o["item"] == it["n"] and o["quote"] == cid]
            alive = [o for o in mine if not o["rejected"] and not o["replaced"]]
            c = ws.cell(row=r, column=j)
            if alive:
                u = min(o["u"] for o in alive)
                vals.append(u)
                c.value, c.number_format = round(u, 2), MONEY2
                if (it["n"], cid) in won:
                    c.fill = fill(GREEN)
            elif mine:
                c.value = (tr("DESCARTADA", "REJECTED") if any(o["rejected"] for o in mine)
                           else tr("REEMPLAZADA", "REPLACED"))
                c.fill = fill(GREY)
            else:
                c.value = "—"
            c.border = box
        if vals:
            best, worst = min(vals), max(vals)
            ws.cell(row=r, column=width - 2, value=round(best, 2)).number_format = MONEY2
            ws.cell(row=r, column=width - 1, value=round(worst, 2)).number_format = MONEY2
            ws.cell(row=r, column=width, value=(worst / best - 1) if len(vals) > 1 and best else None
                    ).number_format = "0.0%"
        else:
            ws.cell(row=r, column=width - 2, value=tr("NADIE", "NOBODY")).font = Font(bold=True, color=RED)
    ws.cell(row=r + 2, column=2, value=tr("Verde = recibió asignación. Gris = línea fuera del reparto "
                                          "(motivo en VERIFICACION).",
                                          "Green = got an allocation. Grey = line out of the allocation "
                                          "(reason in VERIFICATION)."))

    # SUPPLIERS
    ws = sheet(S_SUP, tr("COTIZACIONES · cuadratura contra el subtotal impreso",
                         "QUOTES · reconciliation against the printed subtotal"), 13)
    header(ws, 3, tr(["COTIZACIÓN", "PROVEEDOR", "VENDEDOR", "IDENTIFICADOR", "FECHA", "MONEDA",
                      "TIPO CAMBIO", "LÍNEAS", "SUMA LÍNEAS", "SUBTOTAL IMPRESO", "ESTADO",
                      f"ASIGNADO NETO {cur}", "VENCE · CONDICIONES · PLAZO · ARCHIVO"],
                     ["QUOTE", "SUPPLIER", "SALESPERSON", "TAX ID", "DATE", "CURRENCY",
                      "EXCH. RATE", "LINES", "LINES SUM", "PRINTED SUBTOTAL", "STATUS",
                      f"ALLOCATED NET {cur}", "EXPIRES · TERMS · LEAD TIME · FILE"]),
           [12, 22, 14, 15, 11, 8, 11, 7, 15, 15, 11, 15, 60])
    alloc = defaultdict(float)
    for it in res["items"]:
        for a in it["allocations"]:
            if a["subtotal"] is not None:
                alloc[a["quote"]] += a["subtotal"]
    r = 3
    for q in res["reconciliation"]:
        r += 1
        lead = (tr(f"{q['lead_time_days']} días", f"{q['lead_time_days']} days")
                if q["lead_time_days"] not in (None, "") else "")
        vals = [q["quote"], q["supplier"], q["salesperson"], q["tax_id"] or tr("FALTA", "MISSING"),
                q["date"], q["currency"], q["exchange_rate"], q["lines"], q["sum"],
                q["printed_subtotal"], status_label(q["status"]), round(alloc.get(q["quote"], 0.0), 2),
                " · ".join(str(x) for x in (q["expires"], q["terms"], lead, q["file"]) if x)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            if j in (9, 10):
                c.number_format = MONEY2
            if j == 12:
                c.number_format = MONEY
        st = ws.cell(row=r, column=11)
        st.fill = fill(GREEN if q["status"] == "matches" else ("FBE5E5" if q["status"] == "mismatch" else YELLOW))

    # VERIFICATION
    ws = sheet(S_VER, tr("VERIFICACIÓN · leer antes de emitir órdenes de compra",
                         "VERIFICATION · read before issuing purchase orders"), 6)
    for j, w in enumerate([5, 22, 14, 10, 14, 90], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    r = 3

    def section(title, rows, cols):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = f_bold
        r += 1
        if not rows:
            ws.cell(row=r, column=2, value=tr("Nada que reportar.", "Nothing to report."))
            r += 2
            return
        for j, h in enumerate(cols, 1):
            c = ws.cell(row=r, column=j, value=h)
            c.font, c.fill = f_head, fill(BLUE)
        for f in rows:
            r += 1
            for j, v in enumerate(f, 1):
                ws.cell(row=r, column=j, value=v)
        r += 2

    asked = sum(it["quantity"] for it in res["items"])
    covered = sum(a["quantity"] for it in res["items"] for a in it["allocations"] if a["subtotal"] is not None)
    leftovers = sum(a.get("leftover") or 0 for it in res["items"] for a in it["allocations"])
    missing = sum(f["quantity"] for f in res["missing"])
    ok = abs(asked - (covered - leftovers + missing)) < 1e-4
    section(tr("1 · CUADRATURA DE CANTIDADES (pedido = asignado − sobra + falta)",
               "1 · QUANTITY CHECK (asked = allocated − leftover + missing)"),
            [["", tr("pedido", "asked"), g(asked), "", "", ""],
             ["", tr("asignado", "allocated"), g(covered), "", "", ""],
             ["", tr("sobra por unidad entera", "leftover from whole units"), g(leftovers), "", "", ""],
             ["", tr("falta", "missing"), g(missing), "", "", ""],
             ["", tr("cuadra", "balances"), tr("SÍ", "YES") if ok else "NO", "", "", ""]],
            ["", tr("CONCEPTO", "CONCEPT"), tr("CANTIDAD", "QUANTITY"), "", "", ""])
    section(tr("2 · LÍNEAS DESCARTADAS POR ESPECIFICACIÓN (fuera del reparto)",
               "2 · LINES REJECTED ON SPECIFICATION (out of the allocation)"),
            [[d["item"], d["supplier"], d["quote"], d["line"], "", d["reason"]] for d in res["rejected"]],
            tr(["ÍTEM", "PROVEEDOR", "COTIZACIÓN", "LÍNEA", "", "MOTIVO"],
               ["ITEM", "SUPPLIER", "QUOTE", "LINE", "", "REASON"]))
    pct = (lambda d: f"{d * 100:+.1f} %".replace(".", ",")) if LANG == "es" else (lambda d: f"{d * 100:+.1f} %")
    section(tr("3 · UNITARIO IMPRESO ≠ EFECTIVO (descuento aplicado en el total; se usó el efectivo)",
               "3 · PRINTED UNIT PRICE ≠ EFFECTIVE (discount applied in the total; the effective one was used)"),
            [["", x["supplier"], x["quote"], x["line"], pct(x["diff"]),
              tr(f"impreso {fmt(x['printed'], 2)} · efectivo {fmt(x['effective'], 2)} {x['currency']}",
                 f"printed {fmt(x['printed'], 2)} · effective {fmt(x['effective'], 2)} {x['currency']}")]
             for x in res["printed_vs_effective"]],
            tr(["", "PROVEEDOR", "COTIZACIÓN", "LÍNEA", "DIF.", "DETALLE"],
               ["", "SUPPLIER", "QUOTE", "LINE", "DIFF.", "DETAIL"]))
    section(tr("4 · CONTRAOFERTAS (la cotización más nueva reemplaza a la vieja)",
               "4 · COUNTEROFFERS (the newest quote replaces the old one)"),
            [[x["item"], x["supplier"], x["valid"], "", "",
              tr(f"antes {fmt(x['before'], 2)} → ahora {fmt(x['after'], 2)}",
                 f"before {fmt(x['before'], 2)} → now {fmt(x['after'], 2)}")] for x in res["counteroffers"]],
            tr(["ÍTEM", "PROVEEDOR", "VALE", "", "", "DETALLE"], ["ITEM", "SUPPLIER", "VALID", "", "", "DETAIL"]))
    section(tr("5 · LÍNEAS COTIZADAS SIN ÍTEM (flete, embalaje, lo no pedido): no entran al reparto; decidir aparte",
               "5 · QUOTED LINES WITHOUT ITEM (freight, packing, not requested): out of the allocation; "
               "decide separately"),
            [["", x["supplier"], x["quote"], x["line"], fmt(x["total"], 2), f"{x['description']} ({x['currency']})"]
             for x in res["lines_without_item"]],
            tr(["", "PROVEEDOR", "COTIZACIÓN", "LÍNEA", "TOTAL", "DESCRIPCIÓN"],
               ["", "SUPPLIER", "QUOTE", "LINE", "TOTAL", "DESCRIPTION"]))
    section(tr("6 · SUPUESTOS", "6 · ASSUMPTIONS"), [["", "", "", "", "", a["text"]] for a in res["assumptions"]],
            ["", "", "", "", "", mark()])
    section(tr("7 · AVISOS", "7 · WARNINGS"), [["", "", "", "", "", w] for w in res["warnings"]],
            ["", "", "", "", "", tr("AVISO", "WARNING")])

    del wb["Sheet"]
    wb.save(dest)


# ---------------------------------------------------------------- JSON and console
def to_json(res):
    return {k: v for k, v in res.items() if not k.startswith("_")}


def print_result(res):
    set_language(res["profile"]["language"])
    cur, dec = res["profile"]["currency"], res["profile"]["currency_decimals"]
    t = res["totals"]
    tie = show_tie(res["profile"]["tie_break"])
    print(tr(f"Perfil       : {res['profile']['source']} ({cur}, IVA {res['profile']['vat']:.0%}, desempate {tie})",
             f"Profile      : {res['profile']['source']} ({cur}, VAT {res['profile']['vat']:.0%}, tie-break {tie})"))
    print(tr(f"TOTAL NETO   : {fmt(t['net'], 2)} {cur}   IVA {fmt(t['vat'], 2)}   TOTAL {fmt(t['total'], 2)}",
             f"NET TOTAL    : {fmt(t['net'], 2)} {cur}   VAT {fmt(t['vat'], 2)}   TOTAL {fmt(t['total'], 2)}"))
    print(tr(f"Ítems 100 %  : {t['complete']} de {t['items']}   incompletos: {t['incomplete']}",
             f"Items 100 %  : {t['complete']} of {t['items']}   incomplete: {t['incomplete']}"))
    print(tr(f"Órdenes      : {len(res['by_supplier'])}", f"Orders       : {len(res['by_supplier'])}"))
    for e in res["by_supplier"]:
        print(tr(f"   {e['supplier'][:28]:<28} {e['lines']:>3} líneas  {fmt(e['net'], dec):>16}",
                 f"   {e['supplier'][:28]:<28} {e['lines']:>3} lines   {fmt(e['net'], dec):>16}"))
    for f in res["missing"]:
        print(tr(f"POR COTIZAR  : ítem {f['n']} «{f['item'][:40]}» falta {g(f['quantity'])} ({f['reason']})",
                 f"TO QUOTE     : item {f['n']} «{f['item'][:40]}» missing {g(f['quantity'])} ({f['reason']})"))
    for a in res["assumptions"]:
        print(f"   {a['text']}")
    if res["warnings"]:
        print(tr(f"Avisos ({len(res['warnings'])}), detalle en VERIFICACION:",
                 f"Warnings ({len(res['warnings'])}), detail in VERIFICATION:"))
        for w in res["warnings"][:10]:
            print(f"   [!] {w}")


# ---------------------------------------------------------------- test
def _to_spanish(obj):
    """The sample case written with Spanish keys, to prove the alias map."""
    back = {}
    for es, en in INPUT_ALIASES.items():
        back.setdefault(en, es)
    if isinstance(obj, list):
        return [_to_spanish(x) for x in obj]
    if isinstance(obj, dict):
        return {back.get(k, k): _to_spanish(v) for k, v in obj.items()}
    return obj


def run_test():
    case = SKILL / "tests" / "sample-case.json"
    expected = json.loads((SKILL / "tests" / "expected.json").read_text(encoding="utf-8"))
    tpl = template_path()
    inp = json.loads(case.read_text(encoding="utf-8"))
    failures = []

    def same(a, b, what):
        if isinstance(b, (int, float)) and not isinstance(b, bool):
            if a is None or abs(float(a) - float(b)) > 0.01:
                failures.append(f"{what}: got {a}, expected {b}")
        elif a != b:
            failures.append(f"{what}: got {a!r}, expected {b!r}")

    def check(res, tag):
        for k, v in expected["totals"].items():
            same(res["totals"].get(k), v, f"[{tag}] totals.{k}")
        got = {e["supplier"]: e for e in res["by_supplier"]}
        same(sorted(got), sorted(e["supplier"] for e in expected["by_supplier"]), f"[{tag}] suppliers with an order")
        for e in expected["by_supplier"]:
            if e["supplier"] in got:
                same(got[e["supplier"]]["net"], e["net"], f"[{tag}] net of {e['supplier']}")
                same(got[e["supplier"]]["lines"], e["lines"], f"[{tag}] lines of {e['supplier']}")
        items = {it["n"]: it for it in res["items"]}
        for it in expected["items"]:
            al = items[it["n"]]["allocations"]
            same(len(al), len(it["allocations"]), f"[{tag}] item {it['n']} number of allocations")
            for a, b in zip(al, it["allocations"]):
                for k, v in b.items():
                    same(a.get(k), v, f"[{tag}] item {it['n']} allocation {b.get('order')} {k}")
        same([f["n"] for f in res["missing"]], [f["n"] for f in expected["missing"]], f"[{tag}] missing")
        same(len(res["printed_vs_effective"]), expected["printed_differs_from_effective"],
             f"[{tag}] lines with printed ≠ effective")
        same(len(res["rejected"]), expected["rejected"], f"[{tag}] rejected")

    def sheets_of(res):
        tmp = Path(tempfile.mkdtemp(prefix="procurement-"))
        try:
            write_excel(res, tmp / "test.xlsx")
            import openpyxl
            wb = openpyxl.load_workbook(tmp / "test.xlsx")
            names = wb.sheetnames
            cells = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
            wb.close()
            return names, cells
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    # 1 · es (default): the fictitious case against expected.json.
    profile = load_profile(str(tpl) if tpl else None, language="es")
    res = calculate(inp, profile)
    check(res, "es")
    if res["items"][4]["allocations"][0]["supplier"] != "POR COTIZAR":
        failures.append("[es] the gap is not labelled POR COTIZAR")
    names, _ = sheets_of(res)
    for s in SHEETS["es"]:
        if s not in names:
            failures.append(f"[es] sheet {s} missing in the workbook")
    if not any(a["origin"] == "template" for a in res["assumptions"]):
        failures.append("[es] profile fields taken from the template are not listed as assumptions")

    # 2 · the same case with Spanish input keys gives the same result.
    res_es_keys = calculate(_to_spanish(inp), profile)
    check(res_es_keys, "Spanish keys")

    # 3 · Spanish profile keys and values are normalized.
    tmp = Path(tempfile.mkdtemp(prefix="procurement-"))
    try:
        (tmp / "perfil.yaml").write_text("idioma: es\nempresa:\n  moneda: clp\n  decimales_moneda: 0\n"
                                         "  iva: 0.19\ncompras:\n  desempate: menor_plazo\n", encoding="utf-8")
        p_es = load_profile(str(tmp / "perfil.yaml"))
        same((p_es["currency"], p_es["vat"], p_es["tie_break"], p_es["is_template"]),
             ("CLP", 0.19, "shortest_lead_time", False), "Spanish profile keys")
        # template fallback: profiles/_template.yaml, then perfiles/_plantilla.yaml
        (tmp / "perfiles").mkdir()
        (tmp / "perfiles" / "_plantilla.yaml").write_text("empresa:\n  moneda: CLP\n", encoding="utf-8")
        same(template_path(tmp).name if template_path(tmp) else None, "_plantilla.yaml", "template fallback")
        (tmp / "profiles").mkdir()
        (tmp / "profiles" / "_template.yaml").write_text("company:\n  currency: CLP\n", encoding="utf-8")
        same(template_path(tmp).name, "_template.yaml", "template preference")
        # data folder: $METHODIC_DATA · legacy $INGENIERO_DATOS · ../methodic-data · legacy sibling
        saved_env = {v: os.environ.pop(v) for v in ENV_DATA if v in os.environ}
        saved_repo = REPO
        try:
            globals()["REPO"] = tmp / "lookup" / "repo"
            steps = (("ingeniero-agentico-datos", None, "legacy sibling ../ingeniero-agentico-datos"),
                     ("methodic-data", None, "../methodic-data over the legacy sibling"),
                     ("env-old", "INGENIERO_DATOS", "legacy $INGENIERO_DATOS"),
                     ("env-new", "METHODIC_DATA", "$METHODIC_DATA over $INGENIERO_DATOS"))
            for folder, var, what in steps:
                d = tmp / "lookup" / folder / "profiles"
                d.mkdir(parents=True)
                (d / f"{folder}.yaml").write_text("company:\n  currency: CLP\n", encoding="utf-8")
                if var:
                    os.environ[var] = str(tmp / "lookup" / folder)
                found = find_profile()[0]
                same(found.name if found else None, f"{folder}.yaml", what)
        finally:
            globals()["REPO"] = saved_repo
            for v in ENV_DATA:
                os.environ.pop(v, None)
            os.environ.update(saved_env)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # 4 · declared assumptions travel to the output with their origin.
    with_a = copy.deepcopy(inp)
    with_a["assumptions"] = ["freight is paid by the client", {"parameter": "exchange_rate", "value": 950,
                                                               "source": "bank quote"}]
    got = [a for a in calculate(with_a, profile)["assumptions"] if a["origin"] == "input"]
    same(len(got), 2, "declared assumptions in the output")

    # 5 · the traps: validation must stop these errors.
    bad = copy.deepcopy(inp)
    bad["quotes"][0]["lines"][1]["line_total"] += 1000
    try:
        calculate(bad, profile)
        failures.append("a total copied wrong did not abort the check against the printed subtotal")
    except InputError as e:
        if "NO CUADRA" not in str(e):
            failures.append(f"it aborted for another reason: {e}")
    bad = copy.deepcopy(inp)
    bad["quotes"][0]["lines"][0]["item"] = 99
    try:
        calculate(bad, profile)
        failures.append("a line pointing to a missing item did not abort")
    except InputError:
        pass
    bad = copy.deepcopy(inp)
    bad["quotes"][0]["currency"] = "USD"
    try:
        calculate(bad, profile)
        failures.append("a quote in another currency without exchange rate did not abort")
    except InputError:
        pass

    print_result(res)

    # 6 · en: same numbers, English sheets and marker.
    profile_en = load_profile(str(tpl) if tpl else None, language="en")
    res_en = calculate(inp, profile_en)
    check(res_en, "en")
    if res_en["items"][4]["allocations"][0]["supplier"] != "TO QUOTE":
        failures.append("[en] the gap is not labelled TO QUOTE")
    names, cells = sheets_of(res_en)
    for s in SHEETS["en"]:
        if s not in names:
            failures.append(f"[en] sheet {s} missing in the workbook")
    if any("SUPUESTO" in c for c in cells) or not any(c.startswith("ASSUMPTION") for c in cells):
        failures.append("[en] the assumption marker is not ASSUMPTION")
    print()
    print_result(res_en)

    print()
    if failures:
        print("FAIL")
        for f in failures:
            print(f"   [X] {f}")
        return 1
    print("OK · the fictitious case matches tests/expected.json in es and en, Spanish keys give the same "
          "result, and the three traps abort")
    return 0


# ---------------------------------------------------------------- main
def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    import argparse
    hide = argparse.SUPPRESS
    ap = argparse.ArgumentParser(description="Compare supplier quotes and allocate the purchase in cascade.")
    ap.add_argument("input", nargs="?", help="input JSON (see templates/input.json)")
    ap.add_argument("--out", dest="out", help="output Excel workbook")
    ap.add_argument("--salida", dest="out", help=hide)
    ap.add_argument("--json", help="result as JSON (default: next to the workbook)")
    ap.add_argument("--profile", dest="profile", help="profile YAML path or name")
    ap.add_argument("--perfil", dest="profile", help=hide)
    ap.add_argument("--company", dest="company", help="profile name when there are several")
    ap.add_argument("--empresa", dest="company", help=hide)
    ap.add_argument("--test", dest="test", action="store_true", help="run the fictitious case")
    ap.add_argument("--probar", dest="test", action="store_true", help=hide)
    a = ap.parse_args(argv)
    try:
        if a.test:
            return run_test()
        if not a.input or not a.out:
            ap.print_help()
            return 1
        profile = load_profile(a.profile, a.company)
        inp = json.loads(Path(a.input).read_text(encoding="utf-8-sig"))
        res = calculate(inp, profile)
    except InputError as e:
        print(tr("ERRORES · no se genera nada hasta corregirlos:", "ERRORS · nothing is produced until they are fixed:"))
        for line in str(e).splitlines():
            print(f"   [X] {line}")
        return 2
    write_excel(res, a.out)
    dest_json = a.json or str(Path(a.out).with_suffix(".json"))
    Path(dest_json).write_text(json.dumps(to_json(res), ensure_ascii=False, indent=1), encoding="utf-8")
    print_result(res)
    print(tr(f"\nLibro: {a.out}\nJSON : {dest_json}", f"\nWorkbook: {a.out}\nJSON    : {dest_json}"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
