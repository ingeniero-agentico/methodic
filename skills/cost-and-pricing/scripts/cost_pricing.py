# -*- coding: utf-8 -*-
"""
cost-and-pricing · real cost per item, price by markup, and break-even.

    python cost_pricing.py <case.json> [--profile <path|name>] [--company <name>] [--out cost.xlsx] [--json res.json]
    python cost_pricing.py --test

What it computes (all net, without value added tax):

  (a) Direct cost per item from purchase order lines or awarded quotes. If the item was bought
      from several suppliers, the cost is the quantity-WEIGHTED AVERAGE. If it was bought in the
      supplier's unit (length, roll, box) and some is left over, the leftover is charged to the
      cost of the unit sold: total cost / units sold.
  (b) Overhead as a share of direct cost (simplified absorption costing).
  (c) Price by MARKUP on cost (0.30 = cost x 1.30), per item or global, and next to it the REAL
      MARGIN on sales: (price - cost) / price. The two names are never mixed.
  (d) Contribution margin and break-even of a service or product; if units produced and sold
      are declared, profit under variable costing and under absorption costing.

Profile search order: --profile (path or name) · the case's `profile` field · INGENIERO_DATOS
(profiles/, then perfiles/) · sibling folder ../ingeniero-agentico-datos/ (profiles/, then
perfiles/) · the repo's profiles/_template.yaml, then perfiles/_plantilla.yaml.

Spanish input keys, Spanish profile keys and the old Spanish flags are still accepted.
Output labels (Excel, console, the SUPUESTO / ASSUMPTION marker) follow the profile `language`
(es by default).

Standard library + openpyxl only.
"""
import argparse
import copy
import decimal
import json
import math
import os
import shutil
import sys
import tempfile

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is missing: pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
REPO = os.path.dirname(os.path.dirname(SKILL))

LANG = "es"   # output language; set from the profile


def set_language(lang):
    global LANG
    LANG = "en" if str(lang or "es").strip().lower() == "en" else "es"


def tr(es, en):
    """Pick the text for the output language."""
    return en if LANG == "en" else es


def mark():
    return tr("SUPUESTO", "ASSUMPTION")


class CaseError(Exception):
    """A missing or inconsistent datum: it is asked, never invented."""


# --------------------------------------------------------------------------- key aliases
INPUT_ALIASES = {
    "caso": "case", "perfil": "profile", "moneda": "currency", "decimales_moneda": "currency_decimals",
    "recargo_global": "global_markup", "gastos_generales": "overhead",
    "margen_minimo_alerta": "min_margin_alert", "tipos_de_cambio": "exchange_rates", "valor": "value",
    "fecha": "date", "fuente": "source", "descripcion": "description", "unidad_venta": "sale_unit",
    "cantidad_venta": "quantity_sold", "recargo": "markup", "precio_manual": "manual_price",
    "compras": "purchases", "linea": "line", "proveedor": "supplier", "cantidad": "quantity",
    "unidad": "unit", "factor_unidad": "unit_factor", "precio_unit": "unit_price", "equilibrio": "breakeven",
    "nombre": "name", "costo_variable_unit": "unit_variable_cost",
    "costos_fijos_produccion": "fixed_production_costs", "costos_fijos_operacion": "fixed_operating_costs",
    "volumen_esperado": "expected_volume", "unidades_producidas": "units_produced",
    "unidades_vendidas": "units_sold", "supuestos": "assumptions", "parametro": "parameter",
    "nota": "note", "texto": "text", "_ayuda": "_help",
}
PROFILE_ALIASES = {
    "idioma": "language", "empresa": "company", "nombre": "name", "rut": "tax_id", "pais": "country",
    "moneda": "currency", "decimales_moneda": "currency_decimals", "iva": "vat",
    "impuesto_renta": "income_tax", "compras": "purchasing", "proveedores": "suppliers",
    "plazo_entrega_dias": "lead_time_days", "desempate": "tie_break", "precios": "pricing",
    "recargo_por_defecto": "default_markup", "gastos_generales": "overhead", "finanzas": "finance",
    "tasa_descuento": "discount_rate", "horizonte_anios": "horizon_years", "depreciacion": "depreciation",
}
PROFILE_VALUES = {
    "mayor_disponibilidad": "highest_availability", "menor_plazo": "shortest_lead_time",
    "lineal": "straight_line", "acelerada": "accelerated",
}
_ES_FIELD = {}
for _es, _en in INPUT_ALIASES.items():
    _ES_FIELD.setdefault(_en, _es)


def F(field):
    """Field name in the output language (English key, or its Spanish alias in es)."""
    return _ES_FIELD.get(field, field) if LANG == "es" else field


def rename_keys(obj, aliases):
    """Rename Spanish keys to English, recursively. When both are present, the English one wins."""
    if isinstance(obj, list):
        return [rename_keys(x, aliases) for x in obj]
    if not isinstance(obj, dict):
        return obj
    out = {}
    for k, v in obj.items():
        if k not in aliases:
            out[k] = rename_keys(v, aliases)
    for k, v in obj.items():
        if k in aliases and aliases[k] not in out:
            out[aliases[k]] = rename_keys(v, aliases)
    return out


def normalize_profile(raw):
    prof = rename_keys(raw or {}, PROFILE_ALIASES)
    for sec, key in (("purchasing", "tie_break"), ("finance", "depreciation")):
        d = prof.get(sec)
        if isinstance(d, dict) and isinstance(d.get(key), str):
            d[key] = PROFILE_VALUES.get(d[key], d[key])
    return prof


# --------------------------------------------------------------------------- profile (YAML)

def _scalar(txt):
    t = txt.strip()
    if t == "" or t in ("~", "null", "Null", "NULL"):
        return None
    if t in ("[]",):
        return []
    if t in ("{}",):
        return {}
    if t in ("true", "True"):
        return True
    if t in ("false", "False"):
        return False
    if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
        return t[1:-1]
    try:
        return int(t)
    except ValueError:
        pass
    try:
        return float(t)
    except ValueError:
        return t


def _strip_comment(line):
    """Remove the comment (# preceded by a space), respecting quotes."""
    inside = None
    for i, ch in enumerate(line):
        if ch in "\"'":
            inside = None if inside == ch else (ch if inside is None else inside)
        elif ch == "#" and inside is None and (i == 0 or line[i - 1] in " \t"):
            return line[:i].rstrip()
    return line.rstrip()


def read_yaml(text):
    """
    Minimal reader for the YAML subset profiles use: maps nested by indentation, scalars,
    lists with «- » (of scalars or of maps) and empty [] / {}. PyYAML is not needed.
    """
    lines = []
    for raw in text.splitlines():
        s = _strip_comment(raw)
        if s.strip():
            lines.append((len(s) - len(s.lstrip(" ")), s.strip()))

    def block(i, indent):
        if i < len(lines) and lines[i][1].startswith("- "):
            items = []
            while i < len(lines) and lines[i][0] == indent and lines[i][1].startswith("- "):
                rest = lines[i][1][2:]
                if ":" in rest and not rest.startswith(("'", '"')):
                    # a map inside the list: the first key sits on the same line
                    sub = indent + 2
                    lines[i] = (sub, rest)
                    value, i = block(i, sub)
                    items.append(value)
                else:
                    items.append(_scalar(rest))
                    i += 1
            return items, i
        mapping = {}
        while i < len(lines) and lines[i][0] == indent:
            key, _, value = lines[i][1].partition(":")
            key = key.strip().strip("\"'")
            i += 1
            if value.strip() == "":
                if i < len(lines) and lines[i][0] > indent:
                    mapping[key], i = block(i, lines[i][0])
                elif i < len(lines) and lines[i][0] == indent and lines[i][1].startswith("- "):
                    mapping[key], i = block(i, indent)
                else:
                    mapping[key] = None
            else:
                mapping[key] = _scalar(value)
        return mapping, i

    if not lines:
        return {}
    data, _ = block(0, lines[0][0])
    return data


def template_path(root=None):
    """The repo's profile template: profiles/_template.yaml, then perfiles/_plantilla.yaml."""
    root = root or REPO
    for p in (os.path.join(root, "profiles", "_template.yaml"), os.path.join(root, "perfiles", "_plantilla.yaml")):
        if os.path.isfile(p):
            return p
    return None


def _is_template(path):
    return os.path.basename(path).startswith(("_template", "_plantilla"))


def data_dirs():
    """Where company profiles live, in order: INGENIERO_DATOS and the repo's sibling folder."""
    dirs = []
    env = os.environ.get("INGENIERO_DATOS")
    if env and not os.path.isfile(env):
        dirs += [os.path.join(env, "profiles"), os.path.join(env, "perfiles"), env]
    sibling = os.path.join(os.path.dirname(REPO), "ingeniero-agentico-datos")
    dirs += [os.path.join(sibling, "profiles"), os.path.join(sibling, "perfiles")]
    return dirs


def _pick(dirs, company):
    seen = []
    for d in dirs:
        if not os.path.isdir(d):
            continue
        yamls = sorted(os.path.join(d, n) for n in os.listdir(d) if n.endswith(".yaml") and not n.startswith("_"))
        if not yamls:
            continue
        stems = [os.path.splitext(os.path.basename(p))[0] for p in yamls]
        if company:
            for p, s in zip(yamls, stems):
                if s.lower() == company.lower():
                    return p
            seen += stems
            continue
        if len(yamls) == 1:
            return yamls[0]
        raise CaseError(tr("Hay varios perfiles; elige uno con --company: ",
                           "There are several profiles; pick one with --company: ") + ", ".join(stems))
    if company:
        raise CaseError(tr(f"No hay perfil «{company}». Hay: {', '.join(seen) or 'ninguno'}",
                           f"No profile «{company}». Found: {', '.join(seen) or 'none'}"))
    return None


def find_profile(arg=None, company=None):
    """Return (path, is_template). Order: arg (path or name) · INGENIERO_DATOS · sibling · template."""
    if arg:
        if os.path.isfile(arg):
            return arg, _is_template(arg)
        name = arg if arg.endswith((".yaml", ".yml")) else arg + ".yaml"
        tried = [os.path.join(d, name) for d in data_dirs()]
        for c in tried:
            if os.path.isfile(c):
                return c, False
        raise CaseError(tr(f"No encontré el perfil «{arg}» (busqué en: {', '.join(tried)})",
                           f"Profile «{arg}» not found (looked in: {', '.join(tried)})"))
    env = os.environ.get("INGENIERO_DATOS")
    if env and os.path.isfile(env):
        return env, False
    found = _pick(data_dirs(), company)
    if found:
        return found, False
    tpl = template_path()
    if not tpl:
        raise CaseError(tr("No hay perfil ni plantilla de perfil en el repo", "No profile and no profile template in the repo"))
    return tpl, True


def load_profile(arg=None, company=None, language=None):
    """
    Return (profile, source). The profile comes normalized to English keys; source is the file
    name, or "template" when the repo template was used (and then everything taken from it is
    an assumption).
    """
    path, is_template = find_profile(arg, company)
    with open(path, encoding="utf-8-sig") as fh:
        prof = normalize_profile(read_yaml(fh.read()))
    prof["language"] = language or prof.get("language") or "es"
    set_language(prof["language"])
    return prof, ("template" if is_template else os.path.basename(path))


def _g(d, *path):
    for k in path:
        if not isinstance(d, dict):
            return None
        d = d.get(k)
    return d


# --------------------------------------------------------------------------- calculation

def round_half_up(x, dec):
    """
    Commercial rounding (half up), the same as a spreadsheet's ROUND; Python's round() is not
    (it rounds half to even and drags binary noise: 18,932.355 would come out 18,932.35).
    """
    q = decimal.Decimal(1).scaleb(-dec)
    return float(decimal.Decimal(repr(round(x, 9))).quantize(q, rounding=decimal.ROUND_HALF_UP))


def _num(x, field, where, minimum=None, strict=False):
    if x is None or x == "":
        raise CaseError(tr(f"{where}: falta «{F(field)}»", f"{where}: missing «{field}»"))
    try:
        v = float(x)
    except (TypeError, ValueError):
        raise CaseError(tr(f"{where}: «{F(field)}» no es número ({x!r})", f"{where}: «{field}» is not a number ({x!r})"))
    if minimum is not None and (v <= minimum if strict else v < minimum):
        raise CaseError(tr(f"{where}: «{F(field)}» = {v} fuera de rango", f"{where}: «{field}» = {v} out of range"))
    return v


def declared_assumptions(case):
    """Assumptions the user wrote in the case (`assumptions`, alias `supuestos`)."""
    out = []
    items = case.get("assumptions") or []
    if not isinstance(items, list):
        raise CaseError(tr("«supuestos» debe ser una lista", "«assumptions» must be a list"))
    for a in items:
        if isinstance(a, str):
            out.append({"parameter": "", "value": None, "origin": "input", "text": f"{mark()} · {a}"})
        elif isinstance(a, dict):
            par, val = a.get("parameter") or "", a.get("value")
            src = a.get("source") or a.get("note") or a.get("text") or ""
            head = f"{par} = {val}" if par else ("" if val is None else str(val))
            e = {"parameter": par, "value": val, "origin": "input",
                 "text": f"{mark()} · " + " · ".join(x for x in (head, str(src)) if x)}
            if src:
                e["source"] = src
            out.append(e)
        else:
            raise CaseError(tr(f"«supuestos»: no se entiende «{a}»", f"«assumptions»: cannot read «{a}»"))
    return out


def calculate(case, profile, profile_source="profile"):
    set_language(profile.get("language"))
    case = rename_keys(case, INPUT_ALIASES)
    warnings, params, extra = [], [], []
    from_template = profile_source == "template"
    prof_src = "template" if from_template else "profile"
    profile_label = (tr(f"perfil {profile_source}", f"profile {profile_source}") if not from_template
                     else tr(f"plantilla del repo · {mark()}", f"repo template · {mark()}"))

    currency = case.get("currency") or _g(profile, "company", "currency")
    if not currency:
        raise CaseError(tr("No hay moneda: ni el caso ni el perfil la declaran",
                           "No currency: neither the case nor the profile declares it"))
    dec = case.get("currency_decimals", _g(profile, "company", "currency_decimals"))
    if dec is None:
        dec = 0 if currency == "CLP" else 2
        extra.append({"parameter": "currency_decimals", "value": dec, "origin": "default",
                      "text": tr(f"{mark()} · decimales de la moneda = {dec} (nadie los declara)",
                                 f"{mark()} · currency decimals = {dec} (nobody declares them)")})
    dec = int(dec)
    params.append({"parameter": tr("moneda", "currency"),
                   "value": tr(f"{currency} · {dec} decimales", f"{currency} · {dec} decimals"),
                   "origin": tr("caso", "case") if case.get("currency") else profile_label,
                   "source": "input" if case.get("currency") else prof_src})

    # Global markup: case > profile. Without one, every item must bring its own.
    if case.get("global_markup") is not None:
        global_markup = _num(case["global_markup"], "global_markup", tr("caso", "case"))
        origin_gm, src_gm = tr("caso", "case"), "input"
    elif _g(profile, "pricing", "default_markup") is not None:
        global_markup = float(_g(profile, "pricing", "default_markup"))
        origin_gm, src_gm = profile_label, prof_src
    else:
        global_markup, origin_gm, src_gm = None, None, None
    if global_markup is not None:
        params.append({"parameter": tr("recargo global (sobre costo)", "global markup (on cost)"),
                       "value": global_markup, "origin": origin_gm, "format": "pct", "source": src_gm})

    # Overhead: case > profile > 0 marked as an assumption.
    if case.get("overhead") is not None:
        overhead, origin_oh, src_oh = _num(case["overhead"], "overhead", tr("caso", "case"), 0), tr("caso", "case"), "input"
    elif _g(profile, "pricing", "overhead") is not None:
        overhead, origin_oh, src_oh = float(_g(profile, "pricing", "overhead")), profile_label, prof_src
    else:
        overhead, src_oh = 0.0, "default"
        origin_oh = tr(f"{mark()}: ni el caso ni el perfil los declaran; se usa 0 %",
                       f"{mark()}: neither the case nor the profile declares it; 0 % is used")
        warnings.append(tr("Gastos generales no declarados: se calculó con 0 % (SUPUESTO). Pregúntalos.",
                           "Overhead not declared: computed with 0 % (ASSUMPTION). Ask for it."))
    params.append({"parameter": tr("gastos generales (fracción del costo directo)", "overhead (share of direct cost)"),
                   "value": overhead, "origin": origin_oh, "format": "pct", "source": src_oh})

    # Declared exchange rates: never invented.
    rates = {}
    for cur, d in (case.get("exchange_rates") or {}).items():
        if cur.startswith("_"):
            continue
        if not isinstance(d, dict):
            d = {"value": d}
        where = tr(f"tipo de cambio {cur}", f"exchange rate {cur}")
        v = _num(d.get("value"), "value", where, 0, strict=True)
        source = d.get("source") or ""
        date = d.get("date") or ""
        rates[cur] = v
        origin = (tr(f"caso · {source or 'fuente no declarada'}", f"case · {source or 'source not declared'}")
                  + (f" · {date}" if date else ""))
        if not source or not date:
            warnings.append(tr(f"Tipo de cambio {cur}: falta {'fuente' if not source else 'fecha'}; decláralo.",
                               f"Exchange rate {cur}: missing {'source' if not source else 'date'}; declare it."))
        params.append({"parameter": tr(f"tipo de cambio {cur}/{currency}", f"exchange rate {cur}/{currency}"),
                       "value": round_half_up(v, 2), "origin": origin, "format": "rate", "source": "input"})

    purchases = case.get("purchases") or []
    if not purchases:
        raise CaseError(tr("El caso no trae «compras» (líneas de OC o cotizaciones adjudicadas)",
                           "The case has no «purchases» (purchase order lines or awarded quotes)"))

    # Items: if the case does not declare them, one is created per «item» in the purchases.
    declared = case.get("items")
    if declared:
        items = {}
        for k, it in enumerate(declared, start=1):
            if not it.get("id"):
                raise CaseError(tr(f"items[{k}]: falta «id»", f"items[{k}]: missing «id»"))
            if it["id"] in items:
                raise CaseError(tr(f"items: «{it['id']}» repetido", f"items: «{it['id']}» repeated"))
            items[it["id"]] = dict(it)
    else:
        items = {}
        for ln in purchases:
            iid = ln.get("item")
            if iid and iid not in items:
                items[iid] = {"id": iid, "description": ln.get("description") or iid,
                              "sale_unit": ln.get("sale_unit") or ln.get("unit") or "un"}

    lines = []
    for k, ln in enumerate(purchases, start=1):
        where = tr(f"compras[{k}] ({ln.get('doc', '?')} línea {ln.get('line', '?')})",
                   f"purchases[{k}] ({ln.get('doc', '?')} line {ln.get('line', '?')})")
        iid = ln.get("item")
        if not iid:
            raise CaseError(tr(f"{where}: falta «item» (a qué ítem de venta se carga)",
                               f"{where}: missing «item» (which sale item it is charged to)"))
        if iid not in items:
            raise CaseError(tr(f"{where}: el ítem «{iid}» no está en «items»", f"{where}: item «{iid}» is not in «items»"))
        if not ln.get("doc"):
            raise CaseError(tr(f"{where}: falta «doc» (OC o cotización de origen)",
                               f"{where}: missing «doc» (source purchase order or quote)"))
        qty = _num(ln.get("quantity"), "quantity", where, 0, strict=True)
        price = _num(ln.get("unit_price"), "unit_price", where, 0)
        factor = _num(ln.get("unit_factor", 1), "unit_factor", where, 0, strict=True)
        cur = ln.get("currency") or currency
        if cur == currency:
            rate = 1.0
        elif cur in rates:
            rate = rates[cur]
        else:
            raise CaseError(tr(f"{where}: viene en {cur} y el caso no declara tipo de cambio {cur}",
                               f"{where}: comes in {cur} and the case declares no exchange rate for {cur}"))
        total_src = qty * price
        total_local = total_src * rate
        lines.append({
            "doc": str(ln["doc"]), "line": ln.get("line", ""), "supplier": ln.get("supplier", ""),
            "item": iid, "description": ln.get("description") or items[iid].get("description", iid),
            "quantity": qty, "unit": ln.get("unit", "un"), "unit_factor": factor,
            "quantity_su": qty * factor, "unit_price": price, "currency": cur, "rate": rate,
            "local_price": price * rate, "local_cost_su": price * rate / factor,
            "total_src": total_src, "total_local": total_local,
        })

    margin_alert = case.get("min_margin_alert")
    rows = []
    for iid, it in items.items():
        ls = [x for x in lines if x["item"] == iid]
        where = tr(f"ítem «{iid}»", f"item «{iid}»")
        if not ls:
            raise CaseError(tr(f"{where}: no tiene ninguna línea de compra; sin costo no hay precio",
                               f"{where}: has no purchase line; without a cost there is no price"))
        total = sum(x["total_local"] for x in ls)
        bought = sum(x["quantity_su"] for x in ls)
        if it.get("quantity_sold") not in (None, ""):
            sold = _num(it["quantity_sold"], "quantity_sold", where, 0, strict=True)
            origin_q = tr("caso", "case")
        else:
            sold, origin_q = bought, tr("= lo comprado (sin sobrante)", "= what was bought (no leftover)")
        if sold > bought + 1e-9:
            raise CaseError(tr(f"{where}: se venden {sold:g} y se compraron {bought:g} "
                               f"{it.get('sale_unit', '')}: la compra no cubre la venta",
                               f"{where}: {sold:g} are sold and {bought:g} {it.get('sale_unit', '')} "
                               f"were bought: the purchase does not cover the sale"))
        leftover = bought - sold
        for x in ls:
            x["weight"] = x["total_local"] / total if total else 0.0
        weighted = total / bought
        dcu = total / sold
        tcu = dcu * (1 + overhead)
        manual = it.get("manual_price")
        if manual not in (None, ""):
            price = _num(manual, "manual_price", where, 0, strict=True)
            markup = price / tcu - 1
            origin_m, src_m = tr("precio manual (el recargo se recalcula)", "manual price (the markup is recomputed)"), "manual"
        else:
            if it.get("markup") not in (None, ""):
                markup, origin_m, src_m = _num(it["markup"], "markup", where, -1, strict=True), tr("ítem", "item"), "item"
            elif global_markup is not None:
                markup, origin_m, src_m = global_markup, f"global ({origin_gm})", "global"
            else:
                raise CaseError(tr(f"{where}: no tiene recargo y no hay recargo global en el caso ni en el perfil",
                                   f"{where}: has no markup and there is no global markup in the case or the profile"))
            price = round_half_up(tcu * (1 + markup), dec)
        margin = (price - tcu) / price if price else None
        unit_su = [x["local_cost_su"] for x in ls]
        sups = sorted({x["supplier"] for x in ls if x["supplier"]})
        spread = (max(unit_su) - min(unit_su)) / min(unit_su) if min(unit_su) else 0.0
        if len(sups) > 1 and spread > 0.15:
            warnings.append(tr(f"{where}: entre proveedores el costo por unidad difiere {spread:.1%}; "
                               "confirma que sea el mismo producto",
                               f"{where}: the unit cost differs {spread:.1%} between suppliers; "
                               "confirm it is the same product"))
        if leftover > 1e-9:
            warnings.append(tr(f"{where}: sobran {leftover:g} {it.get('sale_unit', '')} "
                               f"(se compró en la unidad del proveedor); su costo va en la unidad vendida",
                               f"{where}: {leftover:g} {it.get('sale_unit', '')} left over "
                               f"(bought in the supplier's unit); its cost goes into the unit sold"))
        if margin_alert is not None and margin is not None and margin < float(margin_alert):
            warnings.append(tr(f"{where}: margen real {margin:.1%} bajo el mínimo de alerta {float(margin_alert):.1%}",
                               f"{where}: real margin {margin:.1%} below the alert minimum {float(margin_alert):.1%}"))
        rows.append({
            "id": iid, "description": it.get("description", iid), "sale_unit": it.get("sale_unit", "un"),
            "quantity_sold": sold, "quantity_origin": origin_q, "quantity_bought": bought,
            "leftover": leftover, "leftover_cost": leftover * weighted, "direct_cost_total": total,
            "weighted_unit_cost": weighted, "direct_unit_cost": dcu, "overhead": overhead, "total_unit_cost": tcu,
            "markup": markup, "markup_source": src_m, "markup_origin": origin_m, "unit_price": price,
            "real_margin": margin, "sales_total": price * sold, "cost_total": tcu * sold,
            "profit": price * sold - tcu * sold, "suppliers": sups,
            "n_suppliers": len(sups), "cost_spread": spread,
            "documents": sorted({f"{x['doc']}:{x['line']}" for x in ls}),
        })

    dc = sum(f["direct_cost_total"] for f in rows)
    tc = sum(f["cost_total"] for f in rows)
    st = sum(f["sales_total"] for f in rows)
    totals = {"direct_cost": dc, "overhead": tc - dc, "total_cost": tc,
              "sales_total": st, "profit": st - tc,
              "overall_margin": (st - tc) / st if st else None,
              "effective_markup": st / tc - 1 if tc else None}

    breakeven = [break_even(e) for e in (case.get("breakeven") or [])]

    if from_template:
        warnings.insert(0, tr("Se usó la PLANTILLA de perfil: moneda, recargo y gastos son SUPUESTOS. "
                              "Pasa --profile con el de la empresa.",
                              "The profile TEMPLATE was used: currency, markup and overhead are ASSUMPTIONS. "
                              "Pass --profile with the company's."))
    declared_a = declared_assumptions(case)

    def as_text(p):
        origin = p["origin"].replace(f" · {mark()}", "").replace(f"{mark()}: ", "")
        return f"{mark()} · {p['parameter']} = {p['value']} ({origin})"

    assumptions = ([{"parameter": p["parameter"], "value": p["value"], "origin": p["source"], "text": as_text(p)}
                    for p in params if p["source"] in ("template", "default")] + extra + declared_a)
    return {"case": case.get("case", ""), "profile": profile_source, "language": LANG, "currency": currency,
            "decimals": dec, "items": rows, "lines": lines, "totals": totals, "breakeven": breakeven,
            "overhead_assumed": src_oh == "default", "parameters": params, "declared": declared_a,
            "assumptions": assumptions, "warnings": warnings}


def break_even(e):
    """
    Contribution margin and break-even. With units produced and sold, also profit under
    variable costing and under absorption costing (no opening inventory).
    """
    name = e.get("name", tr("servicio", "service"))
    where = tr(f"equilibrio «{name}»", f"break-even «{name}»")
    p = _num(e.get("unit_price"), "unit_price", where, 0, strict=True)
    cv = _num(e.get("unit_variable_cost"), "unit_variable_cost", where, 0)
    fpc = _num(e.get("fixed_production_costs", 0), "fixed_production_costs", where, 0)
    foc = _num(e.get("fixed_operating_costs", 0), "fixed_operating_costs", where, 0)
    cm = p - cv
    if cm <= 0:
        raise CaseError(tr(f"{where}: el precio no cubre el costo variable (margen de contribución {cm:g}); no hay equilibrio",
                           f"{where}: the price does not cover the variable cost (contribution margin {cm:g}); "
                           f"there is no break-even"))
    fc = fpc + foc
    be = fc / cm
    r = {"name": name, "unit": e.get("unit", "un"), "unit_price": p, "unit_variable_cost": cv,
         "fixed_production_costs": fpc, "fixed_operating_costs": foc, "fixed_costs": fc,
         "unit_contribution_margin": cm, "contribution_ratio": cm / p,
         "breakeven_units": be, "breakeven_units_whole": math.ceil(be - 1e-9), "breakeven_sales": be * p}
    v = e.get("expected_volume")
    if v not in (None, ""):
        v = _num(v, "expected_volume", where, 0, strict=True)
        r.update({"expected_volume": v, "expected_profit": cm * v - fc,
                  "safety_margin_units": v - be, "safety_margin": (v - be) / v,
                  "absorption_unit_cost": cv + fpc / v})
    qp, qs = e.get("units_produced"), e.get("units_sold")
    if qp not in (None, "") and qs not in (None, ""):
        qp = _num(qp, "units_produced", where, 0, strict=True)
        qs = _num(qs, "units_sold", where, 0)
        if qs > qp:
            raise CaseError(tr(f"{where}: se venden {qs:g} y se producen {qp:g}; sin inventario inicial no cuadra",
                               f"{where}: {qs:g} are sold and {qp:g} produced; without opening inventory it does not add up"))
        cu_abs = cv + fpc / qp
        p_var = cm * qs - fc
        p_abs = (p - cu_abs) * qs - foc
        inv = qp - qs
        r.update({"units_produced": qp, "units_sold": qs, "absorption_unit_cost": cu_abs,
                  "variable_costing_profit": p_var, "absorption_costing_profit": p_abs,
                  "profit_difference": p_abs - p_var, "fixed_cost_in_inventory": fpc / qp * inv,
                  "ending_inventory_variable": cv * inv, "ending_inventory_absorption": cu_abs * inv})
    return r


# --------------------------------------------------------------------------- Excel

BLUE = "1F3A5F"
F_TITLE = Font(bold=True, size=13, color="FFFFFF")
F_HEAD = Font(bold=True, size=10, color="FFFFFF")
F_BOLD = Font(bold=True)
F_NOTE = Font(italic=True, size=9)
FILL_TITLE = PatternFill("solid", fgColor=BLUE)
FILL_ASSUMED = PatternFill("solid", fgColor="FCE4D6")   # orange = assumption
FILL_SPLIT = PatternFill("solid", fgColor="DCE6F1")     # light blue = bought from several suppliers
PCT = "0.0%"
SHEETS = {"es": ("COSTO Y PRECIO", "TRAZABILIDAD", "EQUILIBRIO", "SUPUESTOS"),
          "en": ("COST AND PRICE", "TRACEABILITY", "BREAK-EVEN", "ASSUMPTIONS")}


def _money_fmt(dec):
    return "#,##0" if dec == 0 else "#,##0." + "0" * dec


def _header(ws, row, cols):
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill = F_HEAD, FILL_TITLE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 42


def _title(ws, text, ncols, note):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font, c.fill = F_TITLE, FILL_TITLE
    ws.row_dimensions[1].height = 24
    ws.cell(row=2, column=1, value=note).font = F_NOTE


def write_excel(res, path):
    set_language(res["language"])
    S_COST, S_TRACE, S_BE, S_ASSUME = SHEETS[LANG]
    dec = res["decimals"]
    MONEY = _money_fmt(dec)
    MONEY2 = "#,##0.00"
    cur = res["currency"]
    wb = openpyxl.Workbook()

    # ---- COST AND PRICE
    ws = wb.active
    ws.title = S_COST
    cols = tr(["ÍTEM", "DESCRIPCIÓN", "UN. VENTA", "CANT. VENDIDA", "CANT. COMPRADA\n(un. venta)",
               "SOBRANTE", "PROV.", "COSTO DIRECTO\nTOTAL", "COSTO PONDERADO\n$/un comprada",
               "COSTO DIRECTO\n$/un vendida", "GASTOS\nGENERALES %", "COSTO TOTAL\n$/un",
               "RECARGO %\n(sobre costo)", "ORIGEN DEL\nRECARGO", "PRECIO UNIT.\nVENTA",
               "MARGEN REAL\n(sobre venta)", "VENTA TOTAL", "UTILIDAD", "DE DÓNDE VIENE\n(doc:línea)",
               "COSTO DEL\nSOBRANTE"],
              ["ITEM", "DESCRIPTION", "SALE UNIT", "QTY SOLD", "QTY BOUGHT\n(sale units)",
               "LEFTOVER", "SUPP.", "DIRECT COST\nTOTAL", "WEIGHTED COST\n$/unit bought",
               "DIRECT COST\n$/unit sold", "OVERHEAD\n%", "TOTAL COST\n$/unit",
               "MARKUP %\n(on cost)", "MARKUP\nORIGIN", "UNIT SALE\nPRICE",
               "REAL MARGIN\n(on sales)", "SALES TOTAL", "PROFIT", "WHERE IT COMES FROM\n(doc:line)",
               "LEFTOVER\nCOST"])
    _title(ws, tr(f"COSTO Y PRECIO · {res['case']}", f"COST AND PRICE · {res['case']}"), len(cols),
           tr(f"Interno: lleva costos. Montos netos en {cur}. RECARGO = sobre el costo (0,30 → costo × 1,30); "
              "MARGEN REAL = (precio − costo) ÷ precio. Fondo azul: comprado a más de un proveedor (costo ponderado). "
              "Naranjo: SUPUESTO.",
              f"Internal: it carries costs. Net amounts in {cur}. MARKUP = on cost (0.30 → cost × 1.30); "
              "REAL MARGIN = (price − cost) ÷ price. Blue background: bought from more than one supplier (weighted cost). "
              "Orange: ASSUMPTION."))
    _header(ws, 4, cols)
    r = 5
    for f in res["items"]:
        vals = [f["id"], f["description"], f["sale_unit"], f["quantity_sold"], f["quantity_bought"],
                f["leftover"], f["n_suppliers"], round_half_up(f["direct_cost_total"], 2),
                round_half_up(f["weighted_unit_cost"], 2), round_half_up(f["direct_unit_cost"], 2), f["overhead"],
                round_half_up(f["total_unit_cost"], 2), f["markup"], f["markup_origin"], f["unit_price"],
                f["real_margin"], round_half_up(f["sales_total"], 2), round_half_up(f["profit"], 2),
                " · ".join(f["documents"]), round_half_up(f["leftover_cost"], 2)]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)
        for c in (8, 17, 18, 20):
            ws.cell(row=r, column=c).number_format = MONEY
        for c in (9, 10, 12):
            ws.cell(row=r, column=c).number_format = MONEY2
        ws.cell(row=r, column=15).number_format = MONEY
        for c in (11, 13, 16):
            ws.cell(row=r, column=c).number_format = PCT
        if f["n_suppliers"] > 1:
            for c in (1, 2, 9):
                ws.cell(row=r, column=c).fill = FILL_SPLIT
        if res["overhead_assumed"]:
            ws.cell(row=r, column=11).fill = FILL_ASSUMED
        ws.cell(row=r, column=15).font = F_BOLD
        r += 1
    t = res["totals"]
    r += 1
    for label, col, value, fmt in ((tr("COSTO DIRECTO", "DIRECT COST"), 8, t["direct_cost"], MONEY),
                                   (tr("COSTO TOTAL (con gastos generales)", "TOTAL COST (with overhead)"), 12,
                                    t["total_cost"], MONEY),
                                   (tr("VENTA TOTAL", "SALES TOTAL"), 17, t["sales_total"], MONEY),
                                   (tr("UTILIDAD", "PROFIT"), 18, t["profit"], MONEY)):
        ws.cell(row=r, column=2, value=label).font = F_BOLD
        c = ws.cell(row=r, column=col, value=round_half_up(value, 2))
        c.number_format, c.font = fmt, F_BOLD
        r += 1
    ws.cell(row=r, column=2, value=tr("RECARGO EFECTIVO (venta ÷ costo total − 1)",
                                      "EFFECTIVE MARKUP (sales ÷ total cost − 1)")).font = F_BOLD
    c = ws.cell(row=r, column=13, value=t["effective_markup"])
    c.number_format, c.font = PCT, F_BOLD
    r += 1
    ws.cell(row=r, column=2, value=tr("MARGEN REAL GLOBAL (utilidad ÷ venta)",
                                      "OVERALL REAL MARGIN (profit ÷ sales)")).font = F_BOLD
    c = ws.cell(row=r, column=16, value=t["overall_margin"])
    c.number_format, c.font = PCT, F_BOLD
    for col, w in zip("ABCDEFGHIJKLMNOPQRST",
                      [12, 40, 9, 10, 12, 9, 6, 14, 15, 14, 10, 13, 10, 22, 13, 12, 14, 13, 28, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"

    # ---- TRACEABILITY
    wt = wb.create_sheet(S_TRACE)
    cols = tr(["DOC", "LÍNEA", "PROVEEDOR", "ÍTEM", "DESCRIPCIÓN EN EL DOC", "CANT.", "UN. PROVEEDOR",
               "FACTOR\n(un. venta por un.)", "CANT.\n(un. venta)", "PRECIO UNIT.\nORIGINAL", "MONEDA",
               "TIPO DE\nCAMBIO", f"PRECIO UNIT.\n{cur}", "TOTAL\nORIGINAL", f"TOTAL\n{cur}", "PESO EN\nEL ÍTEM"],
              ["DOC", "LINE", "SUPPLIER", "ITEM", "DESCRIPTION IN THE DOC", "QTY", "SUPPLIER UNIT",
               "FACTOR\n(sale units per unit)", "QTY\n(sale units)", "ORIGINAL\nUNIT PRICE", "CURRENCY",
               "EXCHANGE\nRATE", f"UNIT PRICE\n{cur}", "ORIGINAL\nTOTAL", f"TOTAL\n{cur}", "WEIGHT IN\nTHE ITEM"])
    _title(wt, tr("TRAZABILIDAD · de qué documento y línea viene cada costo",
                  "TRACEABILITY · which document and line each cost comes from"), len(cols),
           tr("Una fila por línea de OC o cotización adjudicada. Moneda extranjera con 2 decimales y el tipo de "
              "cambio declarado en el caso (hoja SUPUESTOS).",
              "One row per purchase order line or awarded quote line. Foreign currency with 2 decimals and the "
              "exchange rate declared in the case (ASSUMPTIONS sheet)."))
    _header(wt, 4, cols)
    r = 5
    for x in res["lines"]:
        foreign = x["currency"] != cur
        vals = [x["doc"], x["line"], x["supplier"], x["item"], x["description"], x["quantity"], x["unit"],
                x["unit_factor"], x["quantity_su"], x["unit_price"], x["currency"], x["rate"],
                round_half_up(x["local_price"], 2), round_half_up(x["total_src"], 2),
                round_half_up(x["total_local"], 2), x.get("weight")]
        for i, v in enumerate(vals, start=1):
            wt.cell(row=r, column=i, value=v)
        wt.cell(row=r, column=10).number_format = MONEY2
        wt.cell(row=r, column=12).number_format = MONEY2
        wt.cell(row=r, column=13).number_format = MONEY2
        wt.cell(row=r, column=14).number_format = MONEY2 if foreign else MONEY
        wt.cell(row=r, column=15).number_format = MONEY2
        wt.cell(row=r, column=16).number_format = PCT
        r += 1
    for col, w in zip("ABCDEFGHIJKLMNOP", [12, 7, 22, 12, 40, 8, 12, 11, 10, 13, 8, 10, 13, 13, 14, 9]):
        wt.column_dimensions[col].width = w
    wt.freeze_panes = "A5"

    # ---- BREAK-EVEN
    if res["breakeven"]:
        we = wb.create_sheet(S_BE)
        _title(we, tr("MARGEN DE CONTRIBUCIÓN Y PUNTO DE EQUILIBRIO", "CONTRIBUTION MARGIN AND BREAK-EVEN POINT"),
               2 + len(res["breakeven"]),
               tr("Costeo variable: el fijo es gasto del período. Costeo por absorción: el fijo de producción se "
                  "reparte en las unidades producidas y lo no vendido queda en inventario.",
                  "Variable costing: the fixed cost is an expense of the period. Absorption costing: the fixed "
                  "production cost is spread over the units produced and what is not sold stays in inventory."))
        _header(we, 4, tr(["CONCEPTO", "UNIDAD"], ["CONCEPT", "UNIT"]) + [e["name"] for e in res["breakeven"]])
        rows = [(tr("Precio unitario", "Unit price"), "$/un", "unit_price", MONEY),
                (tr("Costo variable unitario", "Unit variable cost"), "$/un", "unit_variable_cost", MONEY),
                (tr("Margen de contribución unitario", "Unit contribution margin"), "$/un", "unit_contribution_margin", MONEY),
                (tr("Razón de contribución", "Contribution ratio"), "%", "contribution_ratio", PCT),
                (tr("Costos fijos de producción (período)", "Fixed production costs (period)"), "$",
                 "fixed_production_costs", MONEY),
                (tr("Costos fijos de operación (período)", "Fixed operating costs (period)"), "$",
                 "fixed_operating_costs", MONEY),
                (tr("Punto de equilibrio", "Break-even point"), "un", "breakeven_units", "#,##0.00"),
                (tr("Punto de equilibrio (unidades enteras)", "Break-even point (whole units)"), "un",
                 "breakeven_units_whole", "#,##0"),
                (tr("Ventas de equilibrio", "Break-even sales"), "$", "breakeven_sales", MONEY),
                (tr("Volumen esperado", "Expected volume"), "un", "expected_volume", "#,##0.##"),
                (tr("Utilidad al volumen esperado", "Profit at the expected volume"), "$", "expected_profit", MONEY),
                (tr("Margen de seguridad", "Safety margin"), "%", "safety_margin", PCT),
                (tr("Costo unitario por absorción", "Absorption unit cost"), "$/un", "absorption_unit_cost", "#,##0.00"),
                (tr("Unidades producidas", "Units produced"), "un", "units_produced", "#,##0.##"),
                (tr("Unidades vendidas", "Units sold"), "un", "units_sold", "#,##0.##"),
                (tr("Utilidad · costeo variable", "Profit · variable costing"), "$", "variable_costing_profit", MONEY),
                (tr("Utilidad · costeo por absorción", "Profit · absorption costing"), "$",
                 "absorption_costing_profit", MONEY),
                (tr("Diferencia (absorción − variable)", "Difference (absorption − variable)"), "$",
                 "profit_difference", MONEY),
                (tr("Fijo de producción que quedó en inventario", "Fixed production cost left in inventory"), "$",
                 "fixed_cost_in_inventory", MONEY)]
        r = 5
        for label, unit, key, fmt in rows:
            if not any(key in e for e in res["breakeven"]):
                continue
            we.cell(row=r, column=1, value=label)
            we.cell(row=r, column=2, value=unit)
            for j, e in enumerate(res["breakeven"], start=3):
                if key in e:
                    v = e[key]
                    c = we.cell(row=r, column=j, value=round(v, 4) if isinstance(v, float) else v)
                    c.number_format = fmt
            r += 1
        we.column_dimensions["A"].width = 44
        we.column_dimensions["B"].width = 8
        for j in range(len(res["breakeven"])):
            we.column_dimensions[chr(ord("C") + j)].width = 22

    # ---- ASSUMPTIONS
    wsu = wb.create_sheet(S_ASSUME)
    _title(wsu, tr("SUPUESTOS Y PARÁMETROS · de dónde salió cada uno", "ASSUMPTIONS AND PARAMETERS · where each came from"), 3,
           tr(f"Perfil usado: {profile_name(res)}. Lo marcado SUPUESTO no vino del usuario: confírmalo antes de cotizar.",
              f"Profile used: {profile_name(res)}. What is marked ASSUMPTION did not come from the user: confirm it "
              f"before quoting."))
    _header(wsu, 4, tr(["PARÁMETRO", "VALOR", "ORIGEN"], ["PARAMETER", "VALUE", "ORIGIN"]))
    r = 5
    for s in res["parameters"]:
        wsu.cell(row=r, column=1, value=s["parameter"])
        c = wsu.cell(row=r, column=2, value=s["value"])
        if s.get("format") == "pct":
            c.number_format = PCT
        elif s.get("format") == "rate":
            c.number_format = "#,##0.00"
        wsu.cell(row=r, column=3, value=s["origin"])
        if s["source"] in ("template", "default"):
            for cc in (1, 2, 3):
                wsu.cell(row=r, column=cc).fill = FILL_ASSUMED
        r += 1
    for a in res["declared"]:
        wsu.cell(row=r, column=1, value=a["parameter"] or a["text"])
        wsu.cell(row=r, column=2, value=a["value"])
        wsu.cell(row=r, column=3, value=tr(f"entrada · supuestos · {mark()}", f"input · assumptions · {mark()}")
                 + (f" · {a['source']}" if a.get("source") else ""))
        for cc in (1, 2, 3):
            wsu.cell(row=r, column=cc).fill = FILL_ASSUMED
        r += 1
    if res["warnings"]:
        r += 1
        wsu.cell(row=r, column=1, value=tr("AVISOS", "WARNINGS")).font = F_BOLD
        for a in res["warnings"]:
            r += 1
            wsu.cell(row=r, column=1, value=a)
    wsu.column_dimensions["A"].width = 48
    wsu.column_dimensions["B"].width = 18
    wsu.column_dimensions["C"].width = 60

    wb.save(path)


def profile_name(res):
    return tr("plantilla", "template") if res["profile"] == "template" else res["profile"]


# --------------------------------------------------------------------------- output and test

def summary_json(res):
    """The comparable part: rounded as shown."""
    dec = res["decimals"]
    out = {"currency": res["currency"], "profile": res["profile"], "language": res["language"],
           "items": {}, "totals": {}, "breakeven": {}}
    for f in res["items"]:
        out["items"][f["id"]] = {
            "quantity_sold": f["quantity_sold"], "quantity_bought": f["quantity_bought"],
            "leftover": round(f["leftover"], 6), "leftover_cost": round_half_up(f["leftover_cost"], 2),
            "n_suppliers": f["n_suppliers"],
            "direct_cost_total": round_half_up(f["direct_cost_total"], 2),
            "weighted_unit_cost": round_half_up(f["weighted_unit_cost"], 2),
            "direct_unit_cost": round_half_up(f["direct_unit_cost"], 2),
            "total_unit_cost": round_half_up(f["total_unit_cost"], 2),
            "markup": round(f["markup"], 6), "markup_source": f["markup_source"], "markup_origin": f["markup_origin"],
            "unit_price": round_half_up(f["unit_price"], dec), "real_margin": round(f["real_margin"], 6),
            "sales_total": round_half_up(f["sales_total"], 2), "profit": round_half_up(f["profit"], 2)}
    out["totals"] = {k: (round(v, 6) if k in ("overall_margin", "effective_markup") else round_half_up(v, 2))
                     for k, v in res["totals"].items() if v is not None}
    for e in res["breakeven"]:
        out["breakeven"][e["name"]] = {k: (round(v, 6) if isinstance(v, float) else v)
                                       for k, v in e.items() if k not in ("name", "unit")}
    out["parameters"] = [{k: v for k, v in p.items() if k != "format"} for p in res["parameters"]]
    out["assumptions"] = res["assumptions"]
    out["warnings"] = res["warnings"]
    return out


def compare(got, expected, path="", failures=None):
    failures = [] if failures is None else failures
    if isinstance(expected, dict):
        for k, v in expected.items():
            if k.startswith("_"):
                continue
            if not isinstance(got, dict) or k not in got:
                failures.append(f"{path}/{k}: missing in the result")
            else:
                compare(got[k], v, f"{path}/{k}", failures)
    elif isinstance(expected, (int, float)) and not isinstance(expected, bool):
        tol = 1e-4 if abs(expected) < 10 else 0.01
        if not isinstance(got, (int, float)) or abs(got - expected) > tol:
            failures.append(f"{path}: expected {expected}, got {got}")
    elif got != expected:
        failures.append(f"{path}: expected {expected!r}, got {got!r}")
    return failures


def print_result(res):
    set_language(res["language"])
    dec = res["decimals"]
    print(tr(f"\nCaso: {res['case']} · perfil: {profile_name(res)} · moneda {res['currency']}",
             f"\nCase: {res['case']} · profile: {profile_name(res)} · currency {res['currency']}"))
    w = min(max([4] + [len(str(f["id"])) for f in res["items"]]), 34)
    print(tr(f"{'ÍTEM':<{w}} {'CANT':>8} {'SOBRA':>6} {'COSTO/u':>12} {'RECARGO':>8} {'PRECIO':>12} {'MARGEN':>7}",
             f"{'ITEM':<{w}} {'QTY':>8} {'LEFT':>6} {'COST/u':>12} {'MARKUP':>8} {'PRICE':>12} {'MARGIN':>7}"))
    for f in res["items"]:
        print(f"{str(f['id'])[:w]:<{w}} {f['quantity_sold']:>8g} {f['leftover']:>6g} {f['total_unit_cost']:>12,.2f} "
              f"{f['markup']:>8.1%} {f['unit_price']:>12,.{dec}f} {f['real_margin']:>7.1%}")
    t = res["totals"]
    print(tr(f"Costo total {t['total_cost']:,.2f} · venta {t['sales_total']:,.{dec}f} · "
             f"recargo efectivo {t['effective_markup']:.1%} · margen real {t['overall_margin']:.1%}",
             f"Total cost {t['total_cost']:,.2f} · sales {t['sales_total']:,.{dec}f} · "
             f"effective markup {t['effective_markup']:.1%} · real margin {t['overall_margin']:.1%}"))
    for e in res["breakeven"]:
        print(tr(f"Equilibrio {e['name']}: MC {e['unit_contribution_margin']:,.2f}/{e['unit']} · "
                 f"PE {e['breakeven_units']:,.2f} {e['unit']} ({e['breakeven_units_whole']} enteras)",
                 f"Break-even {e['name']}: CM {e['unit_contribution_margin']:,.2f}/{e['unit']} · "
                 f"BE {e['breakeven_units']:,.2f} {e['unit']} ({e['breakeven_units_whole']} whole)"))
    for a in res["assumptions"]:
        print(f"  {a['text']}")
    for a in res["warnings"]:
        print(f"  [!] {a}")


def _to_spanish(obj):
    """The sample case written with Spanish keys, to prove the alias map."""
    back = {}
    for es, en in INPUT_ALIASES.items():
        back.setdefault(en, es)
    if isinstance(obj, list):
        return [_to_spanish(x) for x in obj]
    if isinstance(obj, dict):
        return {back.get(k, k): _to_spanish(v) for k, v in obj.items()}
    return obj


SPANISH_PROFILE = """# the sample profile written with the old Spanish keys
idioma: es
empresa:
  nombre: "Ferretería Andes"
  pais: CL
  moneda: CLP
  decimales_moneda: 0
  iva: 0.19
compras:
  proveedores: []
  desempate: mayor_disponibilidad
precios:
  recargo_por_defecto: 0.30
  gastos_generales: 0.05
finanzas:
  depreciacion: lineal
"""


def run_test():
    tests = os.path.join(SKILL, "tests")
    with open(os.path.join(tests, "sample-case.json"), encoding="utf-8") as fh:
        case = json.load(fh)
    with open(os.path.join(tests, "expected.json"), encoding="utf-8") as fh:
        expected = json.load(fh)
    sample_profile = os.path.join(tests, "sample-profile.yaml")
    profile, source = load_profile(sample_profile, language="es")
    res = calculate(case, profile, source)
    failures = ["[es] " + f for f in compare(summary_json(res), expected)]
    gm = [p for p in res["parameters"] if p["parameter"] == "recargo global (sobre costo)"]
    if not gm or gm[0]["source"] != "profile":
        failures.append("[es] the global markup did not come from the profile")

    # The same case with Spanish input keys, and the profile with Spanish keys.
    failures += ["[Spanish keys] " + f for f in compare(summary_json(calculate(_to_spanish(case), profile, source)), expected)]
    tmp = tempfile.mkdtemp(prefix="cost-pricing-")
    try:
        p_es = os.path.join(tmp, "perfil-ficticio.yaml")
        with open(p_es, "w", encoding="utf-8") as fh:
            fh.write(SPANISH_PROFILE)
        prof_es, src_es = load_profile(p_es)
        exp_es = copy.deepcopy(expected)
        failures += ["[Spanish profile] " + f for f in compare(summary_json(calculate(case, prof_es, src_es)), exp_es)]
        if _g(prof_es, "purchasing", "tie_break") != "highest_availability" or \
                _g(prof_es, "finance", "depreciation") != "straight_line":
            failures.append("[Spanish profile] values were not normalized")
        # template fallback: profiles/_template.yaml, then perfiles/_plantilla.yaml
        os.makedirs(os.path.join(tmp, "perfiles"))
        with open(os.path.join(tmp, "perfiles", "_plantilla.yaml"), "w", encoding="utf-8") as fh:
            fh.write("precios:\n  recargo_por_defecto: 0.30\n")
        if os.path.basename(template_path(tmp) or "") != "_plantilla.yaml":
            failures.append("template fallback to perfiles/_plantilla.yaml does not work")
        os.makedirs(os.path.join(tmp, "profiles"))
        with open(os.path.join(tmp, "profiles", "_template.yaml"), "w", encoding="utf-8") as fh:
            fh.write("pricing:\n  default_markup: 0.30\n")
        if os.path.basename(template_path(tmp) or "") != "_template.yaml":
            failures.append("profiles/_template.yaml is not preferred")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # The repo template must be readable, and what comes from it is an assumption.
    tpl = template_path()
    if not tpl:
        failures.append("no profile template in the repo")
    else:
        tprof, tsrc = load_profile(tpl, language="es")
        if _g(tprof, "pricing", "default_markup") != 0.30:
            failures.append("profile template: pricing.default_markup was not read")
        tres = calculate(case, tprof, tsrc)
        if not any(a["origin"] == "template" for a in tres["assumptions"]):
            failures.append("profile fields taken from the template are not listed as assumptions")

    # Declared assumptions travel to the output with their origin.
    with_a = copy.deepcopy(case)
    with_a["assumptions"] = ["the client pays freight", {"parameter": "exchange rate USD", "value": 945.30,
                                                         "source": "bank, 01-10-2026"}]
    got = [a for a in summary_json(calculate(with_a, profile, source))["assumptions"] if a["origin"] == "input"]
    if len(got) != 2:
        failures.append(f"declared assumptions in the output: {len(got)}, expected 2")

    # The errors that must not be swallowed silently.
    for name, mod in (("no exchange rate", lambda c: c.pop("exchange_rates")),
                      ("purchase does not cover the sale", lambda c: c["items"][1].update(quantity_sold=43)),
                      ("line without document", lambda c: c["purchases"][0].pop("doc"))):
        c2 = json.loads(json.dumps(case))
        mod(c2)
        try:
            calculate(c2, profile, source)
            failures.append(f"did not abort: {name}")
        except CaseError:
            pass

    def excel_check(r_, tag):
        tmpd = tempfile.mkdtemp(prefix="cost-pricing-")
        try:
            path = os.path.join(tmpd, "test.xlsx")
            write_excel(r_, path)
            wb = openpyxl.load_workbook(path)
            names = set(wb.sheetnames)
            for h in SHEETS[r_["language"]]:
                if h not in names:
                    failures.append(f"[{tag}] Excel: sheet {h} missing")
            ws = wb[SHEETS[r_["language"]][0]]
            prices = {ws.cell(i, 1).value: ws.cell(i, 15).value for i in range(5, 5 + len(r_["items"]))}
            for iid, e in expected["items"].items():
                if prices.get(iid) != e["unit_price"]:
                    failures.append(f"[{tag}] Excel: price of {iid} = {prices.get(iid)}, expected {e['unit_price']}")
            cells = [str(c.value) for w in wb for row in w.iter_rows() for c in row if c.value is not None]
            wb.close()
            return cells
        finally:
            shutil.rmtree(tmpd, ignore_errors=True)

    excel_check(res, "es")
    print_result(res)

    # en: same numbers, English sheets and marker.
    profile_en, source_en = load_profile(sample_profile, language="en")
    res_en = calculate(case, profile_en, source_en)
    failures += ["[en] " + f for f in compare(summary_json(res_en), expected)]
    cells = excel_check(res_en, "en")
    tres_en = calculate(case, load_profile(tpl, language="en")[0], "template") if tpl else None
    if any("SUPUESTO" in c for c in cells) or not tres_en or not tres_en["assumptions"] or \
            not all(a["text"].startswith("ASSUMPTION") for a in tres_en["assumptions"]):
        failures.append("[en] the assumption marker is not ASSUMPTION")
    print_result(res_en)

    if failures:
        print("\nFAIL")
        for f in failures:
            print("  -", f)
        return 1
    print(f"\nOK · {len(expected['items'])} items, totals, {len(expected['breakeven'])} break-evens, "
          "3 provoked errors and the Excel match expected.json in es and en; Spanish keys give the same result")
    return 0


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    hide = argparse.SUPPRESS
    ap = argparse.ArgumentParser(description="Real cost per item, price by markup, and break-even")
    ap.add_argument("case", nargs="?", help="case.json (see templates/case.json)")
    ap.add_argument("--profile", dest="profile", help="profile YAML path or company name (looked up in INGENIERO_DATOS "
                                                      "and in ../ingeniero-agentico-datos, profiles/ then perfiles/)")
    ap.add_argument("--perfil", dest="profile", help=hide)
    ap.add_argument("--company", dest="company", help="profile name when there are several")
    ap.add_argument("--empresa", dest="company", help=hide)
    ap.add_argument("--out", dest="out", help="output Excel (.xlsx)")
    ap.add_argument("--salida", dest="out", help=hide)
    ap.add_argument("--json", help="write the comparable summary to this JSON")
    ap.add_argument("--test", dest="test", action="store_true", help="run the fictitious case")
    ap.add_argument("--probar", dest="test", action="store_true", help=hide)
    a = ap.parse_args(argv)
    if a.test:
        return run_test()
    if not a.case:
        ap.error("the case is missing (or use --test)")
    try:
        with open(a.case, encoding="utf-8-sig") as fh:
            case = json.load(fh)
        case_profile = case.get("profile") or case.get("perfil")
        profile, source = load_profile(a.profile or case_profile, a.company)
        res = calculate(case, profile, source)
    except CaseError as e:
        print(f"ERROR: {e}")
        return 2
    print_result(res)
    out = a.out or os.path.splitext(a.case)[0] + tr("_costo-y-precio.xlsx", "_cost-and-pricing.xlsx")
    write_excel(res, out)
    print(f"\nExcel: {out}")
    if a.json:
        with open(a.json, "w", encoding="utf-8") as fh:
            json.dump(summary_json(res), fh, ensure_ascii=False, indent=2)
        print(f"JSON:  {a.json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
